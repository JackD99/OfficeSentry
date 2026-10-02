"""Cyber-insurance evidence pack: the controls insurers ask about, answered from collected data.

Each control maps to existing checks. The pack states what the data shows and
when it was collected; anything Office Sentry can't see (backups, training)
is listed as "Not assessed" rather than guessed.
"""

from __future__ import annotations

from dataclasses import dataclass, field
from datetime import datetime

from .. import coverage
from ..rules import RuleResult
from ..snapshot import TenantSnapshot
from .monthly_review import build as build_review

IN_PLACE, PARTLY, NOT_IN_PLACE, NOT_ASSESSED = "In place", "Partly in place", "Not in place", "Not assessed"


@dataclass(frozen=True)
class Control:
    key: str
    question: str
    rules: tuple[str, ...]
    outside_scope: str = ""  # set when the data can't answer it at all


CONTROLS = (
    Control("mfa_all", "Is multi-factor authentication required for all users accessing email and cloud apps?",
            ("mfa_enforced", "users_mfa_registered")),
    Control("mfa_admins", "Is MFA required for all privileged and administrator accounts?",
            ("admin_mfa", "admins_ca_mfa")),
    Control("legacy_auth", "Are legacy authentication protocols (which can't do MFA) blocked?",
            ("legacy_auth_blocked", "smtp_auth")),
    Control("privileged", "Is administrative access limited to named, internal accounts?",
            ("global_admin_count", "guest_admins", "service_principal_admins")),
    Control("email_auth", "Are SPF, DKIM and DMARC configured for your email domains?",
            ("spf_record", "dkim_enabled", "dmarc_policy")),
    Control("forwarding", "Is automatic forwarding of email to external addresses blocked or monitored?",
            ("auto_forward_policy", "external_forwarding", "suspicious_inbox_rules")),
    Control("logging", "Is audit logging enabled for email and user activity?", ("mailbox_auditing",)),
    Control("accounts", "Are inactive and leavers' accounts disabled promptly?",
            ("stale_users", "licensed_disabled", "pending_guests")),
    Control("apps", "Is access by third-party applications controlled?",
            ("user_consent", "risky_app_permissions")),
    Control("sharing", "Is external file sharing restricted?", ("sharepoint_sharing", "guest_invites")),
    Control("backups", "Are backups of Microsoft 365 data taken and kept separately?", (),
            "Backups are outside what Microsoft 365 reports. Provide evidence from the backup product."),
    Control("training", "Do staff receive security awareness and phishing training?", (),
            "Training isn't visible in Microsoft 365 data. Provide records from the training provider."),
)


@dataclass
class ControlResult:
    control: Control
    status: str
    evidence: list[str] = field(default_factory=list)
    gaps: list[str] = field(default_factory=list)


@dataclass
class EvidencePack:
    tenant: dict
    generated_at: datetime
    controls: list[ControlResult]
    secure_score: float | None
    oldest_data: str | None

    @property
    def counts(self) -> dict[str, int]:
        out = {s: 0 for s in (IN_PLACE, PARTLY, NOT_IN_PLACE, NOT_ASSESSED)}
        for c in self.controls:
            out[c.status] += 1
        return out


def _status(results: list[RuleResult]) -> str:
    ran = [r for r in results if r.status != "unavailable"]
    if not ran:
        return NOT_ASSESSED
    if all(r.status == "pass" for r in ran):
        return IN_PLACE if len(ran) == len(results) else PARTLY
    severe = any(f.severity in ("critical", "high") for r in ran if r.status == "fail" for f in r.findings)
    if severe and not any(r.status == "pass" for r in ran):
        return NOT_IN_PLACE
    return PARTLY


def _extra_evidence(key: str, snap: TenantSnapshot) -> list[str]:
    if key == "mfa_all":
        out = []
        pct = snap.metric("mfa.registered_pct")
        if pct is not None:
            out.append(f"{pct:g}% of checked users have registered an MFA method.")
        cov = coverage.user_coverage(snap)
        defaults = snap.objects("conditional_access", "security_defaults")
        if cov and not (defaults and defaults[0].get("isEnabled")):
            out.append(f"{sum(c.mfa for c in cov)} of {len(cov)} enabled accounts are covered by an enforced "
                       "Conditional Access policy requiring MFA.")
        return out
    return []


def build(conn, tenant, now: datetime | None = None) -> EvidencePack:
    review = build_review(conn, tenant, now)
    snap = TenantSnapshot(conn, tenant)
    by_key = {r.key: r for r in review.results}
    accepted_keys = {r.key for r, _f, _l in review.accepted}
    controls = []
    for c in CONTROLS:
        if c.outside_scope:
            controls.append(ControlResult(c, NOT_ASSESSED, [c.outside_scope]))
            continue
        results = [by_key[k] for k in c.rules if k in by_key]
        status = _status([r if r.status != "accepted" else _as_pass(r) for r in results])
        evidence, gaps = [], []
        for r in results:
            if r.status == "pass":
                evidence.append(_sentence(r.summary))
            elif r.status == "unavailable":
                evidence.append(f"Not checked: {r.title[0].lower() + r.title[1:]} ({r.summary.rstrip('.')}).")
            elif r.status == "accepted":
                evidence.append(f"{r.title}: exceptions accepted as a documented risk.")
            else:
                gaps.append(f"{r.title}: {r.summary}.")
                gaps += [f"  {f.subject}: {f.detail}" for f in r.findings[:5]]
                if len(r.findings) > 5:
                    gaps.append(f"  …and {len(r.findings) - 5} more")
            if r.key in accepted_keys and r.status != "accepted":
                evidence.append(f"{r.title}: some exceptions accepted as documented risks.")
        controls.append(ControlResult(c, status, evidence + _extra_evidence(c.key, snap), gaps))

    finished = [r["finished_at"] for r in snap.runs.values() if r["finished_at"]]
    return EvidencePack(dict(tenant), review.generated_at, controls, snap.metric("secure_score.pct"),
                        min(finished) if finished else None)


def _sentence(text: str) -> str:
    text = text.strip()
    return text[:1].upper() + text[1:] + ("" if text.endswith(".") else ".")


def _as_pass(r: RuleResult) -> RuleResult:
    return RuleResult(r.key, r.title, r.category, "pass", r.summary, r.recommendation, [])


def to_xlsx(pack: EvidencePack) -> bytes:
    from io import BytesIO

    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font

    wb = Workbook()
    ws = wb.active
    ws.title = "Controls"
    ws.append(["Question", "Status", "Evidence", "Gaps"])
    for cell in ws[1]:
        cell.font = Font(bold=True)
    for c in pack.controls:
        ws.append([c.control.question, c.status, "\n".join(c.evidence), "\n".join(c.gaps)])
    for col, width in zip("ABCD", (60, 16, 70, 70), strict=True):
        ws.column_dimensions[col].width = width
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(wrap_text=True, vertical="top")
    ws.freeze_panes = "A2"
    info = wb.create_sheet("About")
    info.append(["Tenant", pack.tenant["name"]])
    info.append(["Tenant ID", pack.tenant["tenant_ref"]])
    info.append(["Prepared (UTC)", pack.generated_at.strftime("%Y-%m-%d %H:%M")])
    info.append(["Oldest data used (UTC)", (pack.oldest_data or "")[:16].replace("T", " ")])
    info.append(["Secure Score", f"{pack.secure_score:g}%" if pack.secure_score is not None else "n/a"])
    out = BytesIO()
    wb.save(out)
    return out.getvalue()
