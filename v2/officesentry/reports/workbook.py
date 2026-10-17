"""The shared look of every XLSX download: a summary sheet, then one styled sheet per table.

Every table sheet has its column names in row 1 (bold, on the MSP's accent colour),
frozen, with a filter, sensible widths, real dates and numbers (``cells.xlsx_value``),
and prints landscape on one page width with the header repeated. The first sheet
says what the workbook is, who it's for, who prepared it and when, and lists the
other sheets. Values are never stored as formulas (``safe_xlsx``).
"""

from __future__ import annotations

from dataclasses import dataclass, field
from datetime import date, datetime
from io import BytesIO

from .cells import printable, xlsx_value
from .safe_xlsx import no_formulas

DEFAULT_ACCENT = "1F2937"
INK = "111827"
MUTED = "6B7280"
BAND = "F3F4F6"
DATE_FORMAT = "d mmm yyyy"
DATETIME_FORMAT = "d mmm yyyy hh:mm"
MIN_WIDTH, MAX_WIDTH = 8, 60
# Severity flag cells: (background, text), matching the colours on the page.
FLAG_FILLS = {
    "critical": ("FDECEA", "B42318"),
    "high": ("FDF0E6", "B54708"),
    "medium": ("FDF6E1", "8A5A00"),
    "low": ("F0F2F5", "475467"),
}


@dataclass
class Cover:
    """What the summary sheet says about the workbook."""

    title: str
    client: str = ""  # client name, or "All clients"
    generated_at: datetime | None = None
    prepared_by: str = ""
    tenant_ref: str = ""
    accent: str = ""  # "#0d7a5f", the MSP's brand colour
    intro: str = ""
    facts: list[tuple] = field(default_factory=list)  # (measure, value, note) rows
    notes: list[tuple[str, str]] = field(default_factory=list)  # (label, text) rows under the figures


def _hex(colour: str) -> str:
    colour = (colour or "").lstrip("#").upper()
    return colour if len(colour) == 6 and all(c in "0123456789ABCDEF" for c in colour) else DEFAULT_ACCENT


def _on(fill: str) -> str:
    r, g, b = (int(fill[i : i + 2], 16) / 255 for i in (0, 2, 4))
    lum = 0.2126 * r + 0.7152 * g + 0.0722 * b
    return "FFFFFF" if lum < 0.55 else INK


def _row(values: list) -> list:
    """A summary sheet row with its text cleaned of characters a workbook can't hold."""
    return [printable(v) if isinstance(v, str) else v for v in values]


def _number_format(value, source) -> str | None:
    """How a cell read from ``source`` shows: dates as dates, percentages as percentages, decimals to two
    places, and a change ("+4", "-1.5", "+4.6%") with its sign."""
    if isinstance(value, datetime):
        return DATETIME_FORMAT
    if isinstance(value, date):
        return DATE_FORMAT
    text = source if isinstance(source, str) else ""
    signed = text[:1] in ("+", "-")
    if isinstance(value, float) and text.endswith("%"):
        shown = "0%" if round(value * 100, 6).is_integer() else "0.0%"  # 0.29 * 100 is 28.999...
    elif isinstance(value, (int, float)) and signed:
        decimals = len(text.partition(".")[2])
        shown = "0." + "0" * decimals if decimals else "0"
    elif isinstance(value, int) and abs(value) >= 10_000:
        return "#,##0"  # 113280 items reads as 113,280; smaller counts need no separator
    elif isinstance(value, float):
        return "#,##0.00"
    else:
        return None
    return f"+{shown};-{shown};{shown}" if signed else shown


def _width(value) -> int:
    if isinstance(value, datetime):
        return 17
    if isinstance(value, date):
        return 12
    if isinstance(value, float):
        return len(f"{value:,.2f}")
    text = "" if value is None else str(value)
    return max((len(line) for line in text.split("\n")), default=0)


class Book:
    """Build a workbook: ``table()`` for each sheet, then ``save(cover)``."""

    def __init__(self, accent: str = ""):
        from openpyxl import Workbook

        self.wb = Workbook()
        self.wb.remove(self.wb.active)
        self.accent = _hex(accent)
        self.used: set[str] = set()
        self.sheets: list[tuple[str, int, str]] = []  # (name, rows, description)

    def _name(self, title: str) -> str:
        from .tables import sheet_name

        return sheet_name(title, self.used)

    def header(self, ws, row: int, count: int) -> None:
        from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

        fill = PatternFill("solid", fgColor=self.accent)
        font = Font(bold=True, color=_on(self.accent))
        edge = Border(bottom=Side(style="thin", color=self.accent))
        for col in range(1, count + 1):
            cell = ws.cell(row, col)
            cell.fill, cell.font, cell.border = fill, font, edge
            cell.alignment = Alignment(vertical="center", wrap_text=True)
        ws.row_dimensions[row].height = 20

    def table(
        self,
        title: str,
        columns: list[str],
        rows: list[list],
        *,
        flags: list | None = None,
        links: dict[tuple[int, int], str] | None = None,
        description: str = "",
        wrap: bool = True,
        flag_name: str = "Flag",
    ):
        """A sheet of ``rows`` under ``columns``. ``flags`` (a severity per row) add a coloured first column
        (called ``flag_name``) reading as the page does ("High");
        ``links`` maps (row index, column index) to an https address."""
        from openpyxl.styles import Alignment, Font, PatternFill
        from openpyxl.utils import get_column_letter

        ws = self.wb.create_sheet(self._name(title))
        flagged = bool(flags) and any(flags)
        first = 2 if flagged else 1  # the column the row's own values start in
        header = ([flag_name] if flagged else []) + [printable(str(c)) for c in columns]
        ws.append(header)
        self.header(ws, 1, len(header))
        widths = [len(str(c)) + 3 for c in header]  # room for the filter button
        long_cols: set[int] = set()
        by_row: dict[int, list[tuple[int, str]]] = {}  # each row's links, so a row looks only at its own
        for (ri, ci), href in (links or {}).items():
            by_row.setdefault(ri, []).append((ci, href))
        for i, row in enumerate(rows):
            values = [xlsx_value(v) for v in row]
            flag = (flags[i] if flags and i < len(flags) else None) if flagged else None
            ws.append(([flag.capitalize() if flag else None] if flagged else []) + values)
            r = ws.max_row
            for col, v in enumerate(values, first):
                cell = ws.cell(r, col)
                if number_format := _number_format(v, row[col - first]):
                    cell.number_format = number_format
                w = _width(v)
                if col <= len(widths):
                    widths[col - 1] = max(widths[col - 1], min(w, MAX_WIDTH))
                    if w > MAX_WIDTH or (isinstance(v, str) and "\n" in v):
                        long_cols.add(col)
            if flag in FLAG_FILLS:
                cell = ws.cell(r, 1)
                cell.fill = PatternFill("solid", fgColor=FLAG_FILLS[flag][0])
                cell.font = Font(bold=True, color=FLAG_FILLS[flag][1])
                widths[0] = max(widths[0], len(flag) + 1)
            for ci, href in by_row.get(i, ()):
                cell = ws.cell(r, ci + first)
                cell.hyperlink = href
                cell.font = Font(color="1D4ED8", underline="single")
        for col, w in enumerate(widths, 1):
            ws.column_dimensions[get_column_letter(col)].width = max(MIN_WIDTH, w + 2)
        if wrap and long_cols:
            for col in long_cols:
                for (cell,) in ws.iter_rows(min_row=2, min_col=col, max_col=col):
                    cell.alignment = Alignment(wrap_text=True, vertical="top")
        ws.freeze_panes = "A2"
        if rows:
            ws.auto_filter.ref = f"A1:{get_column_letter(len(header))}{ws.max_row}"
        self._print_setup(ws, title)
        self.sheets.append((ws.title, len(rows), description))
        return ws

    def _print_setup(self, ws, title: str) -> None:
        ws.print_title_rows = "1:1"
        ws.page_setup.orientation = "landscape"
        ws.page_setup.paperSize = ws.PAPERSIZE_A4
        ws.page_setup.fitToWidth, ws.page_setup.fitToHeight = 1, 0
        ws.sheet_properties.pageSetUpPr.fitToPage = True
        ws.oddFooter.left.text = printable(title).replace("&", "&&")
        ws.oddFooter.right.text = "Page &P of &N"

    def save(self, cover: Cover | None = None) -> bytes:
        if cover is not None:
            self._cover(cover)
        if not self.wb.worksheets:
            self.table("No data", ["Note"], [["Nothing to export."]])
        no_formulas(self.wb)
        out = BytesIO()
        self.wb.save(out)
        return out.getvalue()

    def _cover(self, cover: Cover) -> None:
        from openpyxl.styles import Alignment, Font

        name = "Summary"
        if name.lower() in self.used:  # a report table already called Summary: keep it, and put this first
            name = self._name("About this workbook")
        else:
            self.used.add(name.lower())
        ws = self.wb.create_sheet(name, 0)
        ws.sheet_view.showGridLines = False
        ws.column_dimensions["A"].width = 34
        ws.column_dimensions["B"].width = 22
        ws.column_dimensions["C"].width = 70
        ws.append(_row([cover.title]))
        ws["A1"].font = Font(bold=True, size=16, color=self.accent if self.accent != DEFAULT_ACCENT else INK)
        ws.row_dimensions[1].height = 24
        if cover.client:
            ws.append(_row([cover.client]))
            ws.cell(ws.max_row, 1).font = Font(bold=True, size=13, color=INK)
        line = []
        if cover.prepared_by:
            line.append(f"Prepared by {cover.prepared_by}")
        if cover.generated_at:
            line.append(f"{cover.generated_at.day} {cover.generated_at:%b %Y, %H:%M} UTC")
        if cover.tenant_ref:
            line.append(f"Microsoft 365 tenant {cover.tenant_ref}")
        if line:
            ws.append(_row([" · ".join(line)]))
            ws.cell(ws.max_row, 1).font = Font(color=MUTED)
        if cover.intro:
            ws.append(_row([cover.intro]))
            ws.cell(ws.max_row, 1).font = Font(color=INK)
        if cover.facts:
            ws.append([])
            ws.append(["Measure", "Value", "Note"])
            self.header(ws, ws.max_row, 3)
            for fact in cover.facts:
                label, value, note = (list(fact) + ["", ""])[:3]
                ws.append(_row([label, xlsx_value(value), note or None]))
                cell = ws.cell(ws.max_row, 2)
                cell.alignment = Alignment(horizontal="left")
                if number_format := _number_format(cell.value, value):
                    cell.number_format = number_format
                ws.cell(ws.max_row, 3).alignment = Alignment(wrap_text=True, vertical="top")
        for label, text in cover.notes:
            ws.append(_row([label, text]))
            ws.cell(ws.max_row, 1).font = Font(bold=True)
            ws.merge_cells(start_row=ws.max_row, start_column=2, end_row=ws.max_row, end_column=3)
            ws.cell(ws.max_row, 2).alignment = Alignment(wrap_text=True, vertical="top")
        if self.sheets:
            ws.append([])
            ws.append(["Sheet", "Rows", "What it lists"])
            self.header(ws, ws.max_row, 3)
            for sheet, count, description in self.sheets:
                ws.append(_row([sheet, count, description or None]))
                cell = ws.cell(ws.max_row, 1)
                cell.hyperlink = "#'" + sheet.replace("'", "''") + "'!A1"  # Excel doubles a quote in a sheet name
                cell.font = Font(color="1D4ED8", underline="single")
                ws.cell(ws.max_row, 2).alignment = Alignment(horizontal="left")
                ws.cell(ws.max_row, 3).alignment = Alignment(wrap_text=True, vertical="top")
        self._print_setup(ws, cover.title)
        ws.page_setup.orientation = "portrait"
