"""Report exports: XLSX (always) and PDF (when WeasyPrint is installed)."""

from __future__ import annotations

from io import BytesIO

from .monthly_review import Report


def pdf_available() -> bool:
    try:
        import weasyprint  # noqa: F401
    except (ImportError, OSError):  # OSError: missing pango/cairo system libraries
        return False
    return True


def to_pdf(html: str, base_url: str) -> bytes:
    import weasyprint

    return weasyprint.HTML(string=html, base_url=base_url).write_pdf()


def to_xlsx(report: Report) -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Font
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    bold = Font(bold=True)

    def sheet(title, header, rows, first=False):
        ws = wb.active if first else wb.create_sheet()
        ws.title = title
        ws.append(header)
        for cell in ws[1]:
            cell.font = bold
        for row in rows:
            ws.append(list(row))
        ws.freeze_panes = "A2"
        for i, _ in enumerate(header, 1):
            width = max((len(str(c.value or "")) for c in ws[get_column_letter(i)]), default=10)
            ws.column_dimensions[get_column_letter(i)].width = min(max(width + 2, 10), 60)
        return ws

    sheet("Summary", ["Measure", "Value", "4 weeks ago", "Note"], [
        (k.label, k.value, k.previous, k.note) for k in report.kpis
    ] + [("Tenant", report.tenant["name"], None, report.tenant["tenant_ref"]),
         ("Generated (UTC)", report.generated_at.strftime("%Y-%m-%d %H:%M"), None, "")], first=True)
    sheet("Checks", ["Category", "Check", "Result", "Summary", "Recommendation"], [
        (r.category, r.title, r.status, r.summary, r.recommendation) for r in report.results])
    sheet("Findings", ["Severity", "Check", "Subject", "Detail"], [
        (f.severity, r.title, f.subject, f.detail) for r in report.failing for f in r.findings])
    sheet("Licences", ["Licence", "Part number", "Purchased", "Assigned", "Unassigned", "On disabled users",
                       "On inactive users", f"Monthly price ({report.currency})",
                       f"Monthly waste ({report.currency})", "Free/trial"], [
        (ln.name, ln.part_number, ln.purchased, ln.assigned, ln.unassigned, ln.on_disabled_users,
         ln.on_stale_users, ln.monthly_price, ln.monthly_waste, "yes" if ln.free else "")
        for ln in report.licences.lines])
    sheet("Data freshness", ["Collector", "Last successful run (UTC)", "Status"], [
        (f.title, f.finished_at, f.status or "never run") for f in report.freshness])

    out = BytesIO()
    wb.save(out)
    return out.getvalue()
