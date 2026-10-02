"""Detail reports Microsoft doesn't give you in one place.

Each builder reads stored snapshots and returns titled tables; the page and
the XLSX export only format them. Nothing here calls Microsoft.
"""

from __future__ import annotations

from collections.abc import Callable
from dataclasses import dataclass, field
from datetime import UTC, datetime

from .. import coverage, dnscheck
from ..collectors import REGISTRY
from ..licences import last_sign_in
from ..mailflow import inbox_rule_risk, is_external, mailbox_forwarding, tenant_domains
from ..rules import SHARING_LEVELS, app_risks
from ..snapshot import TenantSnapshot


@dataclass
class Section:
    title: str
    columns: list[str]
    rows: list[list] = field(default_factory=list)
    flags: list[str | None] = field(default_factory=list)  # severity per row, for highlighting
    intro: str = ""
    empty: str = "Nothing to show."

    def add(self, row: list, flag: str | None = None) -> None:
        self.rows.append([_cell(v) for v in row])
        self.flags.append(flag)

    @property
    def flagged(self) -> int:
        return sum(1 for f in self.flags if f in ("critical", "high", "medium"))


@dataclass
class Highlight:
    """A key figure shown above the tables, so the page leads with the answer."""
    label: str
    value: object
    foot: str = ""
    tone: str | None = None  # bad | warn | ok, for the dot next to the label


@dataclass
class DataReport:
    key: str
    title: str
    description: str
    tenant: dict
    generated_at: datetime
    sections: list[Section] = field(default_factory=list)
    unavailable: str | None = None
    sources: list[tuple[str, str | None]] = field(default_factory=list)  # (collector title, finished_at)
    highlights: list[Highlight] = field(default_factory=list)
    notes: list[str] = field(default_factory=list)  # what couldn't be checked, and why
    chart: str = ""  # server-side SVG from reports.charts


def _cell(value):
    if value is None:
        return ""
    if isinstance(value, bool):
        return "Yes" if value else "No"
    if isinstance(value, (list, tuple, set)):
        return ", ".join(str(v) for v in value)
    return value


def _date(value: str | None) -> str:
    return value[:10] if value else ""


# -- Mailboxes ------------------------------------------------------------------


def mailboxes(snap: TenantSnapshot, now: datetime, report: DataReport) -> None:
    boxes = snap.objects("mailboxes", "mailbox")
    rules = snap.objects("inbox_rules", "inbox_rule")
    if boxes is None and rules is None:
        report.unavailable = ("Mailbox data hasn't been collected yet. "
                              "Refresh Mailboxes and Inbox rules on the tenant page.")
        return
    status = snap.objects("mailboxes", "exchange_status")
    if status and not status[0].get("available"):
        report.unavailable = "This tenant has no Exchange Online licence."
        return
    domains = tenant_domains(snap) or set()

    fwd = Section("Forwarding", ["Mailbox", "Forwards to", "Set by", "Outside the organisation", "Keeps a copy"],
                  intro="Every place mail leaves a mailbox automatically: admin-set forwarding and inbox rules.",
                  empty="No mailbox forwards mail.")
    for m in boxes or []:
        target = mailbox_forwarding(m)
        if target:
            external = is_external(target, domains)
            fwd.add([m.get("PrimarySmtpAddress"), target, "Admin (mailbox setting)", external,
                     m.get("DeliverToMailboxAndForward")], "high" if external else None)
    suspicious = Section("Inbox rules worth checking", ["Mailbox", "Rule", "Enabled", "Why it stands out"],
                         intro="Rules that forward outside, delete, or bury mail. These are the classic signs of a "
                               "compromised mailbox.", empty="No inbox rule looks suspicious.")
    every = Section("All inbox rules", ["Mailbox", "Rule", "Enabled", "Conditions", "Actions"],
                    empty="No inbox rules found.")
    for r in sorted(rules or [], key=lambda r: ((r.get("mailbox") or "").lower(), r.get("sequence") or 0)):
        risk = inbox_rule_risk(r, domains)
        for target in risk.forwards_to:
            external = target in risk.external
            fwd.add([r.get("mailbox"), target, f"Inbox rule \"{r.get('displayName')}\"", external, True],
                    ("high" if r.get("isEnabled", True) else "low") if external else None)
        if risk.severity:
            suspicious.add([r.get("mailbox"), r.get("displayName"), r.get("isEnabled", True),
                            "; ".join(risk.reasons)], risk.severity)
        every.add([r.get("mailbox"), r.get("displayName"), r.get("isEnabled", True),
                   _describe_conditions(r.get("conditions") or {}), _describe_actions(r)], risk.severity)

    sections = [fwd, suspicious]
    if boxes is not None:
        access = Section("Delegated access", ["Mailbox", "Mailbox type", "Who", "Access"],
                         intro="People who can open or send as someone else's mailbox.",
                         empty="Nobody has access to another mailbox.")
        types = {m.get("PrimarySmtpAddress"): _type(m.get("RecipientTypeDetails")) for m in boxes}
        for p in sorted(snap.objects("mailboxes", "mailbox_permission") or [], key=lambda p: p.get("mailbox") or ""):
            access.add([p.get("mailbox"), _type(p.get("mailboxType")), p.get("User"), "Full access"])
        for p in snap.objects("mailboxes", "send_as") or []:
            access.add([p.get("Identity"), "", p.get("Trustee"), "Send as"])
        for m in boxes:
            for who in m.get("GrantSendOnBehalfTo") or []:
                access.add([m.get("PrimarySmtpAddress"), types.get(m.get("PrimarySmtpAddress")), who, "Send on behalf"])

        users = {u["_id"]: u for u in snap.objects("users", "user") or []}
        perms = snap.objects("mailboxes", "mailbox_permission") or []
        shared = Section("Shared mailboxes", ["Mailbox", "Sign-in allowed", "Licensed", "People with full access"],
                         intro="A shared mailbox's account doesn't need to sign in. If it can, it usually has no MFA.",
                         empty="No shared mailboxes.")
        for m in sorted(boxes, key=lambda m: m.get("PrimarySmtpAddress") or ""):
            if m.get("RecipientTypeDetails") != "SharedMailbox":
                continue
            user = users.get(m.get("ExternalDirectoryObjectId")) or {}
            enabled = user.get("accountEnabled")
            shared.add([m.get("PrimarySmtpAddress"), enabled if user else "Unknown", user.get("licensed"),
                        sum(1 for p in perms if p.get("mailboxId") == m["_id"])], "medium" if enabled else None)
        sections += [access, shared]
    report.sections = sections + [every]


def _type(value: str | None) -> str:
    return {"UserMailbox": "User", "SharedMailbox": "Shared", "RoomMailbox": "Room",
            "EquipmentMailbox": "Equipment"}.get(value or "", value or "")


def _describe_conditions(c: dict) -> str:
    parts = []
    for key, label in (("fromAddresses", "from"), ("sentToAddresses", "sent to")):
        addrs = [((a or {}).get("emailAddress") or {}).get("address") for a in c.get(key) or []]
        if addrs:
            parts.append(f"{label} {', '.join(a for a in addrs if a)}")
    for key, label in (("subjectContains", "subject has"), ("bodyContains", "body has"),
                       ("bodyOrSubjectContains", "subject or body has"), ("senderContains", "sender has")):
        if c.get(key):
            parts.append(f"{label} {', '.join(c[key])}")
    return "; ".join(parts) or "All messages"


def _describe_actions(r: dict) -> str:
    a = r.get("actions") or {}
    parts = []
    for key, label in (("forwardTo", "forward to"), ("redirectTo", "redirect to"),
                       ("forwardAsAttachmentTo", "forward as attachment to")):
        addrs = [((x or {}).get("emailAddress") or {}).get("address") for x in a.get(key) or []]
        if addrs:
            parts.append(f"{label} {', '.join(x for x in addrs if x)}")
    if a.get("moveToFolder"):
        parts.append(f"move to {r.get('moveToFolderName') or 'a folder'}")
    if a.get("delete") or a.get("permanentDelete"):
        parts.append("delete")
    if a.get("markAsRead"):
        parts.append("mark as read")
    if a.get("stopProcessingRules"):
        parts.append("stop processing rules")
    return "; ".join(parts) or "Other"


# -- Apps ----------------------------------------------------------------------------


def apps(snap: TenantSnapshot, now: datetime, report: DataReport) -> None:
    risks = app_risks(snap)
    if risks is None:
        report.unavailable = "App data hasn't been collected yet. Refresh Apps and consents on the tenant page."
        return
    third = Section("Third-party apps with access", ["App", "Publisher", "Application permissions",
                                                     "Consented for everyone", "Consented by users", "Risk"],
                    intro="Apps from outside Microsoft that hold permissions in this tenant. Application permissions "
                          "work without anyone signed in.", empty="No third-party app holds permissions.")
    for a in risks:
        publisher = a["publisher"] + (" (verified)" if a["verified"] else (" (unverified)" if a["publisher"]
                                                                           else "Unverified"))
        users = f"{a['users']} user(s): {', '.join(a['user_scopes'])}" if a["user_scopes"] else ""
        third.add([a["name"], publisher, a["app_permissions"], a["tenant_scopes"], users,
                   ", ".join(a["risky"]) or "Low"], a["severity"])

    regs = Section("App registrations and their credentials", ["App", "Created", "Credentials", "Next expiry",
                                                                "Status"],
                   intro="Apps registered in this tenant. Expired secrets usually mean an integration has stopped "
                         "working, or the app is no longer used.", empty="No app registrations.")
    soon = now.timestamp() + 30 * 86400
    registrations = snap.objects("app_consents", "application") or []
    for a in sorted(registrations, key=lambda a: (a.get("displayName") or "").lower()):
        creds = a.get("credentials") or []
        ends = sorted(c["endDateTime"] for c in creds if c.get("endDateTime"))
        future = [e for e in ends if _ts(e) >= now.timestamp()]
        if not creds:
            status, flag = "No credentials", None
        elif not future:
            status, flag = "All expired", "low"
        elif _ts(future[-1]) < soon:
            status, flag = "Expiring soon", "medium"
        else:
            status, flag = "OK", None
        kinds = ", ".join(f"{sum(c['type'] == k for c in creds)} {k}(s)" for k in ("secret", "certificate")
                          if any(c["type"] == k for c in creds))
        regs.add([a.get("displayName"), _date(a.get("createdDateTime")), kinds or "None",
                  _date(future[0]) if future else (_date(ends[-1]) + " (expired)" if ends else ""), status], flag)
    report.sections = [third, regs]


def _ts(iso: str) -> float:
    return datetime.fromisoformat(iso.replace("Z", "+00:00")).timestamp()


# -- Conditional Access coverage -------------------------------------------------------


def conditional_access(snap: TenantSnapshot, now: datetime, report: DataReport) -> None:
    defaults = snap.objects("conditional_access", "security_defaults")
    d = coverage.directory(snap)
    cov = coverage.user_coverage(snap)
    if defaults is None or d is None or cov is None:
        report.unavailable = "Conditional Access data hasn't been collected yet."
        return
    policies = snap.objects("conditional_access", "ca_policy") or []
    if defaults and defaults[0].get("isEnabled"):
        report.sections.append(Section("Security Defaults", ["Setting", "Value"], rows=[[
            "Security Defaults", "On: MFA registration and sign-in prompts apply to everyone"]], flags=[None]))
    order = {"enabled": 0, "enabledForReportingButNotEnforced": 1, "disabled": 2}
    pol = Section("Policies", ["Policy", "State", "Applies to", "Except", "Apps", "Conditions", "Requires"],
                  intro="Every policy in plain words, enforced ones first.",
                  empty="No Conditional Access policies." + (" The tenant has no Entra ID P1."
                                                              if snap.capabilities.get("entra_p1") is False else ""))
    for p in sorted(policies, key=lambda p: (order.get(p.get("state"), 3), (p.get("displayName") or "").lower())):
        row = coverage.describe(p, d)
        pol.add([row.name, row.state, row.users, row.excluded, row.apps, row.conditions, row.grant])
    report.sections.append(pol)

    covered = sum(c.mfa for c in cov)
    sd = bool(defaults and defaults[0].get("isEnabled"))
    accounts = Section("Coverage by account", ["Account", "Type", "Admin roles", "MFA required by",
                                               "Legacy sign-in blocked", "Excluded from"],
                       intro=(f"{covered} of {len(cov)} enabled account(s) are covered by an enforced policy that "
                              "requires MFA for all apps." + (" Security Defaults also apply." if sd else "")),
                       empty="No enabled accounts.")
    for c in cov:
        if c.mfa or sd:
            flag = None
        else:
            flag = "high" if c.roles else ("low" if c.user_type == "Guest" else "medium")
        accounts.add([c.label, c.user_type, c.roles, c.mfa_policies or ("Security Defaults" if sd else "Nothing"),
                      c.legacy_blocked or sd, c.excluded_from], flag)
    report.sections.append(accounts)


# -- Guests and sharing -----------------------------------------------------------


INVITES = {"none": "Nobody", "adminsAndGuestInviters": "Admins and Guest Inviters",
           "adminsGuestInvitersAndAllMembers": "Admins and members", "everyone": "Everyone, including guests"}


def guests(snap: TenantSnapshot, now: datetime, report: DataReport) -> None:
    users = snap.objects("users", "user")
    policy = (snap.objects("tenant_policies", "authorization_policy") or [None])[0]
    sharepoint = (snap.objects("tenant_policies", "sharepoint_settings") or [None])[0]
    if users is None and policy is None:
        report.unavailable = "User and sharing data hasn't been collected yet."
        return
    settings = Section("External access settings", ["Setting", "Value"])
    if policy:
        settings.add(["Who can invite guests", INVITES.get(policy.get("allowInvitesFrom"),
                                                             policy.get("allowInvitesFrom"))],
                     "low" if policy.get("allowInvitesFrom") == "everyone" else None)
        grants = policy.get("permissionGrantPolicies") or []
        legacy = any(p.endswith("microsoft-user-default-legacy") for p in grants)
        settings.add(["User consent to apps", "Any app" if legacy else (
            "Verified publishers, low-risk permissions" if grants else "Off")], "medium" if legacy else None)
        settings.add(["Users can register apps", policy.get("allowedToCreateApps")])
    if sharepoint:
        level = sharepoint.get("sharingCapability")
        settings.add(["SharePoint and OneDrive sharing", SHARING_LEVELS.get(level, level)],
                     "medium" if level == "externalUserAndGuestSharing" else None)
        mode = sharepoint.get("sharingDomainRestrictionMode")
        if mode == "allowList":
            settings.add(["Sharing limited to domains", sharepoint.get("sharingAllowedDomainList") or []])
        elif mode == "blockList":
            settings.add(["Sharing blocked for domains", sharepoint.get("sharingBlockedDomainList") or []])
        settings.add(["Guests can reshare", sharepoint.get("isResharingByExternalUsersEnabled")])
    sections = [settings]
    if users is not None:
        activity = any(u.get("signInActivityAvailable") for u in users)
        guest = Section("Guest accounts", ["Guest", "Invited", "Invitation", "Last sign-in", "Sign-in allowed"],
                        intro="Every external account, oldest activity first." + (
                            "" if activity else " Sign-in dates need Entra ID P1."), empty="No guest accounts.")
        rows = [u for u in users if u.get("userType") == "Guest"]
        rows.sort(key=lambda u: (last_sign_in(u) or datetime.min.replace(tzinfo=UTC), u.get("createdDateTime") or ""))
        cutoff = now.timestamp() - 90 * 86400
        for u in rows:
            last = last_sign_in(u)
            pending = u.get("externalUserState") == "PendingAcceptance"
            stale = activity and u.get("accountEnabled") and (last is None or last.timestamp() < cutoff)
            guest.add([u.get("mail") or u.get("userPrincipalName"), _date(u.get("createdDateTime")),
                       "Not accepted" if pending else "Accepted",
                       last.date().isoformat() if last else ("Never" if activity else "Unknown"),
                       u.get("accountEnabled")], "low" if stale or pending else None)
        sections.append(guest)
    report.sections = sections


# -- Email domains --------------------------------------------------------------------


def domains(snap: TenantSnapshot, now: datetime, report: DataReport) -> None:
    rows = snap.objects("domain_health", "domain")
    if not rows:
        report.unavailable = "Domain checks haven't run yet, or the tenant has no custom domains."
        return
    overview = Section("Domains", ["Domain", "Receives mail via", "SPF", "DMARC", "DKIM", "What to fix"],
                       intro="SPF, DMARC and DKIM stop other people sending email as the client's domains.")
    records = Section("Published records", ["Domain", "SPF record", "DNS lookups", "DMARC record", "MX"])
    for row in sorted(rows, key=lambda r: (not r.get("mx"), r["domain"])):
        h = dnscheck.DomainHealth.from_dict(row)
        found = dnscheck.issues(h)
        words = dnscheck.summary(h)
        via = "Microsoft 365" if h.microsoft_365 else (h.mx[0] if h.mx else "Doesn't receive mail")
        overview.add([h.domain, via, words["SPF"], words["DMARC"], words["DKIM"],
                      "; ".join(i.detail for i in found) or "Nothing"], found[0].severity if found else None)
        records.add([h.domain, h.spf or "None", h.spf_lookups if h.spf else "", h.dmarc or "None", h.mx])
    report.sections = [overview, records]


# -- Registry ----------------------------------------------------------------------------

Builder = Callable[[TenantSnapshot, datetime, DataReport], None]

REPORTS: dict[str, tuple[str, str, tuple[str, ...], Builder]] = {
    "mailboxes": ("Mailbox access and forwarding",
                  "Forwarding, suspicious inbox rules, delegated access and shared mailboxes.",
                  ("mailboxes", "inbox_rules", "users"), mailboxes),
    "apps": ("Apps and consents", "Third-party apps with access to the tenant, and app credential expiry.",
             ("app_consents",), apps),
    "conditional-access": ("Conditional Access coverage",
                           "Every policy in plain words, and which accounts each one actually covers.",
                           ("conditional_access", "admin_roles", "users"), conditional_access),
    "domains": ("Email domains", "SPF, DMARC and DKIM for each of the tenant's domains, from public DNS.",
                ("domain_health",), domains),
    "guests": ("Guests and sharing", "Who can invite and share externally, and every guest account.",
               ("tenant_policies", "users"), guests),
}


# Identity and licence reports live in their own module; it needs Section and
# DataReport from above, so it's imported here rather than at the top.
from .identity_reports import REPORTS as _IDENTITY_REPORTS  # noqa: E402

REPORTS.update(_IDENTITY_REPORTS)


def build(conn, tenant, key: str, now: datetime | None = None) -> DataReport:
    title, description, collectors, builder = REPORTS[key]
    now = now or datetime.now(UTC)
    snap = TenantSnapshot(conn, tenant)
    report = DataReport(key, title, description, dict(tenant), now)
    report.sources = [(REGISTRY[c].title, (snap.runs.get(c) or {"finished_at": None})["finished_at"])
                      for c in collectors]
    builder(snap, now, report)
    return report


def to_xlsx(report: DataReport) -> bytes:
    from io import BytesIO

    from openpyxl import Workbook
    from openpyxl.styles import Font
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    wb.remove(wb.active)
    used = set()
    if report.highlights or report.notes:
        ws = wb.create_sheet("Summary")
        used.add("Summary")
        ws.append([report.title, report.tenant.get("name", "")])
        ws["A1"].font = Font(bold=True)
        for h in report.highlights:
            ws.append([h.label, h.value, h.foot])
        for note in report.notes:
            ws.append(["Not checked", note])
        ws.column_dimensions["A"].width = 32
        ws.column_dimensions["B"].width = 24
        ws.column_dimensions["C"].width = 60
    for section in report.sections or [Section("No data", ["Note"], rows=[[report.unavailable or ""]])]:
        name = section.title[:31]
        while name in used:
            name = name[:29] + str(len(used))
        used.add(name)
        ws = wb.create_sheet(name)
        ws.append(section.columns + (["Flag"] if any(section.flags) else []))
        for cell in ws[1]:
            cell.font = Font(bold=True)
        for row, flag in zip(section.rows, section.flags or [None] * len(section.rows), strict=False):
            ws.append(list(row) + ([flag or ""] if any(section.flags) else []))
        ws.freeze_panes = "A2"
        for i in range(1, ws.max_column + 1):
            width = max((len(str(c.value or "")) for c in ws[get_column_letter(i)]), default=10)
            ws.column_dimensions[get_column_letter(i)].width = min(max(width + 2, 10), 60)
    out = BytesIO()
    wb.save(out)
    return out.getvalue()
