"""Plain tables and their CSV and XLSX exports, shared by every report and rollup.

Collected values come from client tenants (display names, rule names, app
names), so exports must never let one become a spreadsheet formula. CSV cells
that start with a formula character are prefixed with ``'``; XLSX cells are
always written as text or numbers, never formulas.
"""

from __future__ import annotations

import csv
import re
from dataclasses import dataclass, field
from io import BytesIO, StringIO

FORMULA_START = ("=", "+", "-", "@", "\t", "\r", "\n")


@dataclass
class Table:
    title: str
    columns: list[str]
    rows: list[list] = field(default_factory=list)


def csv_cell(value) -> str | int | float:
    if value is None:
        return ""
    if isinstance(value, bool):
        return "Yes" if value else "No"
    if isinstance(value, (int, float)):
        return value
    if isinstance(value, (list, tuple, set)):
        value = ", ".join(str(v) for v in value)
    text = str(value)
    if text.startswith(FORMULA_START):
        return "'" + text
    return text


def to_csv(table: Table) -> bytes:
    out = StringIO()
    writer = csv.writer(out, lineterminator="\r\n")
    writer.writerow([csv_cell(c) for c in table.columns])
    for row in table.rows:
        writer.writerow([csv_cell(v) for v in row])
    return ("﻿" + out.getvalue()).encode("utf-8")  # BOM so Excel reads UTF-8


def no_formulas(wb) -> None:
    """Store every cell openpyxl would treat as a formula (text starting with '=') as text."""
    for ws in wb.worksheets:
        for row in ws.iter_rows():
            for cell in row:
                if cell.data_type == "f":
                    cell.data_type = "s"


def sheet_name(title: str, used: set[str]) -> str:
    name = re.sub(r"[\[\]:*?/\\]", " ", title).strip()[:31] or "Sheet"
    base, n = name, 2
    while name.lower() in used:
        suffix = f" ({n})"
        name = base[:31 - len(suffix)] + suffix
        n += 1
    used.add(name.lower())
    return name


def to_xlsx(tables: list[Table]) -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Font
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    wb.remove(wb.active)
    used: set[str] = set()
    for table in tables or [Table("No data", ["Note"], [["Nothing to export."]])]:
        ws = wb.create_sheet(sheet_name(table.title, used))
        ws.append(list(table.columns))
        for cell in ws[1]:
            cell.font = Font(bold=True)
        for row in table.rows:
            ws.append([_xlsx_cell(v) for v in row])
        ws.freeze_panes = "A2"
        for i in range(1, ws.max_column + 1):
            letter = get_column_letter(i)
            width = max((len(str(c.value or "")) for c in ws[letter]), default=10)
            ws.column_dimensions[letter].width = min(max(width + 2, 10), 60)
    no_formulas(wb)
    out = BytesIO()
    wb.save(out)
    return out.getvalue()


def _xlsx_cell(value):
    if value is None:
        return ""
    if isinstance(value, bool):
        return "Yes" if value else "No"
    if isinstance(value, (list, tuple, set)):
        return ", ".join(str(v) for v in value)
    return value


def slug(text: str) -> str:
    return re.sub(r"[^a-z0-9]+", "-", str(text).lower()).strip("-") or "report"
