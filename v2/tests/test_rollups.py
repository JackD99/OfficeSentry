"""Cross-tenant rollups, CSV/XLSX exports and the month-over-month change report."""

from __future__ import annotations

import csv
import io
from datetime import UTC, datetime, timedelta

import pytest
from openpyxl import load_workbook

from officesentry import store
from officesentry.collectors import site_permissions, teams
from officesentry.reports import data_reports, month_changes, rollups, tables
from officesentry.reports.data_reports import Section
from officesentry.reports.identity_reports import Money
from officesentry.snapshot import TenantSnapshot
from tests.conftest import login, make_user
from tests.seed import seed_tenant
from tests.test_collectors import USERS

ISO = "%Y-%m-%dT%H:%M:%SZ"


def _csv(body: bytes) -> list[list[str]]:
    assert body.startswith("﻿".encode())
    return list(csv.reader(io.StringIO(body.decode("utf-8-sig"))))


def _backdate(db, tenant_id: int, days: int) -> None:
    """Make everything collected so far look as if it was collected ``days`` ago."""
    when = (datetime.now(UTC) - timedelta(days=days)).strftime(ISO)
    db.execute("UPDATE runs SET finished_at = ?, created_at = ? WHERE tenant_id = ?", (when, when, tenant_id))
    db.execute("UPDATE metrics SET captured_at = ? WHERE tenant_id = ?", (when, tenant_id))


# -- Exports --------------------------------------------------------------------


def test_csv_neutralises_formulas_and_keeps_numbers():
    table = tables.Table(
        "T",
        ["Name", "Count"],
        [['=HYPERLINK("http://x")', 3], ["+cmd", -2], ["@SUM(A1)", 1.5], ["plain", None], [True, ["a", "b"]]],
    )
    rows = _csv(tables.to_csv(table))
    assert rows[1] == ['\'=HYPERLINK("http://x")', "3"]
    assert rows[2] == ["'+cmd", "-2"]
    assert rows[3][0] == "'@SUM(A1)"
    assert rows[4] == ["plain", ""]
    assert rows[5] == ["Yes", "a, b"]


def test_csv_leaves_a_signed_number_alone_but_not_a_formula_that_starts_like_one():
    for text in ("+2", "-1", "+4.6%", "-1.5", "+£23.00", "-$1,234.50"):
        assert tables.csv_cell(text) == text
    for text in ("-1+cmd|' /C calc'!A0", "+2\n", "-2+3", "+A$1", "-1e3", "+SUM(1)", "-", "+£"):
        assert tables.csv_cell(text) == "'" + text


def test_xlsx_never_contains_formulas_and_names_sheets_safely():
    body = tables.to_xlsx(
        [
            tables.Table("Mail: forwarding/rules [all]", ["A"], [["=1+1"]]),
            tables.Table("Mail: forwarding/rules [all]", ["A"], [["x"]]),
        ]
    )
    wb = load_workbook(io.BytesIO(body))
    assert len(wb.sheetnames) == 2 and len(set(wb.sheetnames)) == 2
    assert all(len(n) <= 31 and not set(n) & set("[]:*?/\\") for n in wb.sheetnames)
    cell = wb.worksheets[0]["A2"]
    assert cell.value == "=1+1" and cell.data_type == "s"


def test_review_download_says_which_figures_are_percentages_and_writes_counts_whole():
    columns = [
        rollups.Column("secure_score.pct", "Secure Score", "pct", "up"),
        rollups.Column("admins", "Privileged accounts", "num", "down"),
        rollups.Column("users", "Licensed users"),
        rollups.Column("waste", "Licence waste / month", "money"),
    ]
    values = {"secure_score.pct": 50.9, "admins": 4.0, "users": 54.0, "waste": Money(120, "GBP")}
    row = rollups.RollupRow({"name": "Alpha", "tenant_ref": "a"}, values, {"secure_score.pct": 2.0, "admins": 0.0})
    table = rollups.Rollup("review", "Monthly review", "", columns, [row], datetime.now(UTC)).to_table()
    assert table.columns[2:8] == [
        "Secure Score %",
        "Secure Score: change",
        "Privileged accounts",
        "Privileged accounts: change",
        "Licensed users",
        "Licence waste / month (GBP)",
    ]
    assert table.rows[0][2:8] == [50.9, 2, 4, 0, 54, 120] and [type(v) for v in table.rows[0][3:7]] == [int] * 4
    assert _csv(tables.to_csv(table))[1][2:7] == ["50.9", "2", "4", "0", "54"]
    ws = load_workbook(io.BytesIO(tables.to_xlsx([table]))).active
    assert [(c.value, c.number_format) for c in ws[2][4:7]] == [(4, "General"), (0, "General"), (54, "General")]


# -- Rollups ----------------------------------------------------------------------


def test_every_report_has_a_rollup_that_survives_tenants_without_data(db, two_tenants):
    alpha, bravo = two_tenants
    seed_tenant(db, alpha)
    tenants = db.execute("SELECT * FROM tenants ORDER BY id").fetchall()
    keys = [k for k, _ in rollups.available()]
    assert keys[:2] == ["review", "evidence"] and set(data_reports.REPORTS) <= set(keys)
    for key in keys:
        r = rollups.build(db, tenants, key)
        assert {row.tenant["id"] for row in r.rows} == {alpha, bravo}, key
        table = r.to_table()
        assert len(table.rows) == 2 and all(len(row) == len(table.columns) for row in table.rows), key
    review = rollups.build(db, tenants, "review")
    by_id = {row.tenant["id"]: row for row in review.rows}
    assert by_id[bravo].note and not by_id[alpha].note
    assert review.rows[-1].tenant["id"] == bravo  # tenants without data sort last


def test_new_detail_reports_get_a_rollup_automatically(db, two_tenants, monkeypatch):
    alpha, _ = two_tenants

    def builder(snap, now, report):
        s = Section("Widgets", ["Widget", "State"])
        s.add(["w1", "ok"])
        s.add(["w2", "bad"], "high")
        report.sections = [s, Section("Settings", ["Setting", "Value"], rows=[["x", "y"]], flags=[None])]

    monkeypatch.setitem(data_reports.REPORTS, "widgets", ("Widgets", "Test", (), builder))
    assert ("widgets", "Widgets") in rollups.available()
    r = rollups.build(db, db.execute("SELECT * FROM tenants").fetchall(), "widgets")
    assert [c.label for c in r.columns] == ["To review", "Widgets"]  # settings tables aren't counted
    row = next(row for row in r.rows if row.tenant["id"] == alpha)
    assert row.values == {"_flagged": 1, "c:Widgets": 2}


def test_rollup_pages_and_downloads(client, db, two_tenants):
    alpha, _ = two_tenants
    seed_tenant(db, alpha)
    make_user(db, "admin")
    login(client, "admin")
    for key, _title in rollups.available():
        page = client.get(f"/across/{key}")
        assert page.status_code == 200, key
        assert "Alpha Ltd" in page.get_data(as_text=True)
    assert client.get("/across/nope").status_code == 404
    assert client.get("/across?key=mailboxes").headers["Location"].endswith("/across/mailboxes")
    resp = client.get("/across/review.csv")
    assert resp.headers["Content-Type"] == "text/csv; charset=utf-8"  # the charset once
    assert "attachment" in resp.headers["Content-Disposition"]
    rows = _csv(resp.data)
    assert rows[0][:2] == ["Client", "Tenant ID"] and {r[0] for r in rows[1:]} == {"Alpha Ltd", "Bravo Inc"}
    resp = client.get("/across/mailboxes.xlsx")
    assert resp.status_code == 200
    wb = load_workbook(io.BytesIO(resp.data))
    assert wb.sheetnames[0] == "Summary" and wb["Summary"]["A2"].value == "All clients"
    assert wb.worksheets[1]["A1"].value == "Client"
    assert db.execute("SELECT COUNT(*) FROM audit_log WHERE action = 'rollup_downloaded'").fetchone()[0] == 2


def test_rollups_only_show_granted_tenants(client, db, two_tenants):
    alpha, bravo = two_tenants
    make_user(db, "analyst", role="analyst", tenants=[alpha])
    login(client, "analyst")
    page = client.get("/across/review").get_data(as_text=True)
    assert "Alpha Ltd" in page and "Bravo Inc" not in page
    assert b"Bravo" not in client.get("/across/apps.csv").data
    assert b"Bravo" not in client.get("/changes.csv").data
    assert client.get(f"/tenants/{bravo}/changes").status_code == 404
    assert client.get(f"/tenants/{bravo}/data/apps.csv").status_code == 404


def test_clients_get_their_own_downloads_but_no_cross_tenant_pages(client, db, two_tenants):
    alpha, bravo = two_tenants
    seed_tenant(db, alpha)
    make_user(db, "client", role="client", tenants=[alpha])
    login(client, "client")
    assert client.get("/across/review").status_code == 403
    assert client.get("/changes").status_code == 403
    assert client.get(f"/tenants/{alpha}/changes").status_code == 200
    for url in (
        f"/tenants/{alpha}/report.csv",
        f"/tenants/{alpha}/evidence/insurance.csv",
        f"/tenants/{alpha}/data/mailboxes.csv",
        f"/tenants/{alpha}/changes.csv",
    ):
        resp = client.get(url)
        assert resp.status_code == 200 and resp.mimetype == "text/csv", url
    assert client.get(f"/tenants/{bravo}/report.csv").status_code == 404


def test_detail_report_csv_downloads_one_table(client, db, two_tenants):
    alpha, _ = two_tenants
    seed_tenant(db, alpha)
    make_user(db, "admin")
    login(client, "admin")
    report = data_reports.build(db, db.execute("SELECT * FROM tenants WHERE id = ?", (alpha,)).fetchone(), "guests")
    for i, section in enumerate(report.sections, 1):
        rows = _csv(client.get(f"/tenants/{alpha}/data/guests.csv?section={i}").data)
        assert rows[0][: len(section.columns)] == section.columns
        assert len(rows) == len(section.rows) + 1
    assert client.get(f"/tenants/{alpha}/data/guests.csv?section=99").status_code == 404
    assert client.get(f"/tenants/{alpha}/data/nope.csv").status_code == 404
    page = client.get(f"/tenants/{alpha}/data/guests").get_data(as_text=True)
    assert "guests.csv?section=1" in page and "/across/guests" in page


# -- Month-over-month changes ---------------------------------------------------------


def test_compare_sections_matches_rows_and_ignores_time_columns():
    old = Section("Guests", ["Guest", "Invitation", "Last sign-in"])
    new = Section("Guests", ["Guest", "Invitation", "Last sign-in"])
    for s, rows in (
        (old, [["a@x", "Accepted", "2026-08-01"], ["b@x", "Not accepted", "Never"]]),
        (new, [["a@x", "Accepted", "2026-09-30"], ["b@x", "Accepted", "Never"], ["c@x", "Accepted", ""]]),
    ):
        for row in rows:
            s.add(row)
    diff = month_changes.compare_sections(old, new)
    kinds = {(c.kind, c.subject) for c in diff.changes}
    assert kinds == {("changed", "b@x"), ("added", "c@x")}  # a@x only signed in again
    changed = next(c for c in diff.changes if c.kind == "changed")
    assert changed.detail == "Invitation: Not accepted → Accepted"
    assert month_changes.compare_sections(Section("Log", ["When"], compare=False), None) is None


def test_day_counts_that_only_move_with_the_clock_are_not_changes():
    # The same data collected 40 days apart: the "then" report counts days to the period start.
    old = Section("Inactive members", ["Account", "Inactive for", "Days", "Licences"])
    new = Section("Inactive members", ["Account", "Inactive for", "Days", "Licences"])
    old.add(["ann@x", "180+ days", 170, "E3"])
    new.add(["ann@x", "180+ days", 210, "E5"])
    diff = month_changes.compare_sections(old, new)
    assert [c.detail for c in diff.changes] == ["Licences: E3 → E5"]
    devices_then, devices_now = (Section("Needs attention", ["Device", "Problems"]) for _ in range(2))
    devices_then.add(["PC-1", "Not compliant; No check-in for 35 days"])
    devices_now.add(["PC-1", "Not compliant; No check-in for 75 days"])
    assert month_changes.compare_sections(devices_then, devices_now).changes == []
    devices_now.rows[0][1] = "Not encrypted; No check-in for 75 days"
    assert len(month_changes.compare_sections(devices_then, devices_now).changes) == 1


def test_reasons_and_buckets_that_only_move_with_the_clock_are_not_changes():
    phones_then, phones_now = (Section("Devices that need attention", ["Mailbox", "Device", "Why"]) for _ in range(2))
    phones_then.add(["ann@x", "iPhone", ["Not synced for 201 days: the partnership can be removed"]])
    phones_now.add(["ann@x", "iPhone", ["Not synced for 231 days: the partnership can be removed"]])
    assert month_changes.compare_sections(phones_then, phones_now).changes == []
    lists_then, lists_now = (Section("Lists that need attention", ["List", "Why"]) for _ in range(2))
    lists_then.add(["Sales", ["No mail in 40 days (last 1 Aug 2026)"]])
    lists_now.add(["Sales", ["No mail in 70 days (last 1 Aug 2026)"]])
    assert month_changes.compare_sections(lists_then, lists_now).changes == []

    # The same last sign-in, a month later: only the clock moved the bucket.
    columns = ["Account", "Inactive for", "Last sign-in", "Days"]
    then, now = Section("Inactive members", columns), Section("Inactive members", columns)
    then.add(["ann@x", "90 days to 6 months", "2026-04-20", 170])
    now.add(["ann@x", "6 to 12 months", "2026-04-20", 200])
    then.add(["bob@x", "Never signed in (created 100 days ago)", "Never", 100])
    now.add(["bob@x", "Never signed in (created 130 days ago)", "Never", 130])
    assert month_changes.compare_sections(then, now).changes == []
    apps_then, apps_now = (Section("Enterprise apps", ["App", "Last sign-in", "Activity"]) for _ in range(2))
    apps_then.add(["Zoom", "2026-09-20", "Last 30 days"])
    apps_now.add(["Zoom", "2026-09-20", "30 to 90 days ago"])
    apps_then.add(["Dropbox", "2026-01-05", "Over 90 days ago"])
    apps_now.add(["Dropbox", "2026-10-10", "Last 30 days"])  # used again: that is a change
    diff = month_changes.compare_sections(apps_then, apps_now)
    assert [(c.subject, c.detail) for c in diff.changes] == [("Dropbox", "Activity: Over 90 days ago → Last 30 days")]


def test_a_shared_name_only_compares_the_rows_that_share_it_whole():
    # Two groups called Marketing: Clients is still matched on its name, so its member count is a change.
    old, new = Section("Groups", ["Group", "Teams", "Members"]), Section("Groups", ["Group", "Teams", "Members"])
    for row in (["Clients", "Yes", 11], ["Marketing", "No", 5], ["Marketing", "No", 7]):
        old.add(row)
    for row in (["Clients", "Yes", 10], ["Marketing", "No", 5], ["Marketing", "No", 8]):
        new.add(row)
    diff = month_changes.compare_sections(old, new)
    assert [(c.kind, c.subject, c.detail) for c in diff.changes] == [
        ("changed", "Clients", "Members: 11 → 10"),
        ("added", "Marketing", "Teams: No; Members: 8"),
        ("removed", "Marketing", "Teams: No; Members: 7"),
    ]


def test_compare_sections_without_a_unique_key_reports_added_and_removed():
    old, new = Section("Access", ["Mailbox", "Who"]), Section("Access", ["Mailbox", "Who"])
    for row in (["box", "ann"], ["box", "bob"]):
        old.add(row)
    for row in (["box", "ann"], ["box", "cat"], ["box", "cat"]):
        new.add(row)
    diff = month_changes.compare_sections(old, new)
    assert sorted((c.kind, c.row[1]) for c in diff.changes) == [("added", "cat"), ("added", "cat"), ("removed", "bob")]


def test_change_report_compares_every_report_with_the_start_of_the_period(client, db, two_tenants):
    alpha, bravo = two_tenants
    seed_tenant(db, alpha)
    _backdate(db, alpha, 40)
    newcomer = {
        "id": "u9",
        "displayName": "Visitor",
        "userPrincipalName": "v#EXT#@y.com",
        "mail": "v@y.com",
        "userType": "Guest",
        "accountEnabled": True,
        "assignedLicenses": [],
        "createdDateTime": "2026-09-20T00:00:00Z",
    }
    seed_tenant(db, alpha, users=[u for u in USERS if u["id"] != "u3"] + [newcomer])

    tenant = db.execute("SELECT * FROM tenants WHERE id = ?", (alpha,)).fetchone()
    report = month_changes.build(db, tenant, 30)
    assert report.has_baseline
    accounts = {(c.kind, c.subject) for c in report.key_changes if c.area == "Accounts"}
    assert ("added", "v#EXT#@y.com") in accounts and ("removed", "g#EXT#@x.com") in accounts
    guests = next(r for r in report.reports if r.key == "guests")
    rows = {(c.kind, c.subject) for s in guests.sections for c in s.changes}
    assert ("added", "v@y.com") in rows
    assert all(not r.count for r in report.reports if r.key == "admin-activity")  # event logs aren't diffed
    assert month_changes.build(db, tenant, 7).has_baseline  # the second collection is the 7-day baseline

    make_user(db, "admin")
    login(client, "admin")
    page = client.get(f"/tenants/{alpha}/changes?days=30").get_data(as_text=True)
    assert "v@y.com" in page and "Guests and sharing" in page
    across = client.get("/changes?days=30").get_data(as_text=True)
    assert "Alpha Ltd" in across and "No data from 30 days ago yet" in across  # Bravo has no history
    rows = _csv(client.get(f"/tenants/{alpha}/changes.csv?days=30").data)
    assert any(r[3] == "v@y.com" for r in rows[1:])
    wb = load_workbook(io.BytesIO(client.get(f"/tenants/{alpha}/changes.xlsx").data))
    assert {"Key figures", "Key changes", "Changes by report"} <= set(wb.sheetnames)
    assert client.get("/changes.xlsx?days=90").status_code == 200
    assert client.get(f"/tenants/{bravo}/changes?days=999").status_code == 200  # bad periods fall back


def test_a_report_missing_one_of_its_data_sources_at_the_start_is_not_compared(db, two_tenants):
    # MFA registration failed before the period began: every account's registration would look new.
    alpha, _ = two_tenants
    seed_tenant(db, alpha)
    _backdate(db, alpha, 40)
    db.execute("UPDATE runs SET status = 'failed' WHERE tenant_id = ? AND collector = 'mfa_registration'", (alpha,))
    seed_tenant(db, alpha)
    report = month_changes.build(db, store.get_tenant(db, alpha), 30)
    mfa = next(r for r in report.reports if r.key == "mfa")
    assert mfa.count == 0 and mfa.note.startswith("MFA registration hadn't been collected by the start of the period")
    assert not next(r for r in report.reports if r.key == "guests").note  # it had all its data then


def test_a_permission_approved_or_still_awaited_is_not_compared_as_change(db, two_tenants):
    # Teams channels and site permissions waiting for the client's approval on one of the two days.
    waiting = {
        "teams_kw": {"status": 403},
        "sharepoint_kw": {"rest": "consent"},
        "our_app": {"missing": (teams.CHANNELS, teams.APPS, site_permissions.SITES_READ)},
    }
    alpha, bravo = two_tenants
    for tenant_id, then, now in ((alpha, waiting, {}), (bravo, {}, waiting)):
        seed_tenant(db, tenant_id, **then)
        _backdate(db, tenant_id, 10)
        seed_tenant(db, tenant_id, **now)
    approved = {r.key: r for r in month_changes.build(db, store.get_tenant(db, alpha), 7).reports}
    assert approved["teams"].count == 0 and approved["teams"].note.startswith("The client approved the permission")
    assert (approved["site-permissions"].count, approved["site-permissions"].note) == (0, month_changes.NO_START)
    awaited = {r.key: r for r in month_changes.build(db, store.get_tenant(db, bravo), 7).reports}
    assert awaited["teams"].count == 0 and awaited["teams"].note.startswith("Parts of it are waiting for the client")
    assert awaited["site-permissions"].note.startswith("Waiting for the client to approve the SharePoint permission")


def test_old_snapshots_skip_runs_whose_objects_were_pruned(db, two_tenants):
    alpha, _ = two_tenants
    tenant = db.execute("SELECT * FROM tenants WHERE id = ?", (alpha,)).fetchone()
    ids = []
    for name in ("kept", "pruned"):
        run_id = store.enqueue_run(db, alpha, "users")
        store.claim_next_run(db)
        store.add_objects(db, run_id, "user", [("u1", {"displayName": name})])
        store.finish_run(db, run_id, "succeeded")
        ids.append(run_id)
    db.execute("UPDATE runs SET finished_at = ? WHERE id = ?", ("2026-08-01T00:00:00Z", ids[0]))
    db.execute("UPDATE runs SET finished_at = ? WHERE id = ?", ("2026-08-02T00:00:00Z", ids[1]))
    db.execute("DELETE FROM objects WHERE run_id = ?", (ids[1],))
    snap = TenantSnapshot(db, tenant, as_of=datetime(2026, 8, 3, tzinfo=UTC))
    assert snap.objects("users", "user")[0]["displayName"] == "kept"
    assert TenantSnapshot(db, tenant, as_of=datetime(2026, 7, 1, tzinfo=UTC)).objects("users", "user") is None


@pytest.mark.parametrize("days", month_changes.PERIODS)
def test_change_report_without_history_explains_itself(db, two_tenants, days):
    _, bravo = two_tenants
    report = month_changes.build(db, db.execute("SELECT * FROM tenants WHERE id = ?", (bravo,)).fetchone(), days)
    assert not report.has_baseline and report.total == 0
    assert all(r.note for r in report.reports)
