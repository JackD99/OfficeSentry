"""Devices: collector, checks, Windows support dates and the Devices report."""

from datetime import UTC, date, datetime, timedelta
from io import BytesIO

import pytest
import responses
from openpyxl import load_workbook

from officesentry import store
from officesentry.changes import all_changes
from officesentry.devices import WINDOWS_RELEASES, WindowsRelease, windows_support
from officesentry.graph import GraphClient
from officesentry.reports import data_reports, evidence
from officesentry.rules import evaluate
from officesentry.snapshot import TenantSnapshot
from tests import devicedata as dd
from tests.conftest import login, make_user
from tests.seed import seed_tenant
from tests.test_collectors import SKUS, V1, _metrics, _run


@pytest.fixture
def factory():
    return lambda tenant, connection: GraphClient(lambda: "t", sleep=lambda s: None)


def _results(db, tenant_id):
    return {r.key: r for r in evaluate(TenantSnapshot(db, store.get_tenant(db, tenant_id)))}


# -- Windows lifecycle ------------------------------------------------------------------


def test_windows_support_by_build():
    today = date(2026, 10, 2)
    assert windows_support("10.0.19045.5011", today).status == "ended"  # Windows 10
    assert windows_support("10.0.22621.4317", today).status == "ended"  # 22H2, all editions
    assert windows_support("10.0.22631.4317", today).status == "pro_ended"  # 23H2: Enterprise until Nov 2026
    assert windows_support("10.0.26100.1", today).status == "ending"  # 24H2 Home/Pro: 13 Oct 2026
    assert windows_support("10.0.26200.1", today).status == "supported"  # 25H2
    assert windows_support("10.0.26100.1", date(2026, 6, 1)).status == "supported"
    assert windows_support("6.3.9600", today).status == "unknown"
    assert windows_support(None, today).status == "unknown"
    assert "Extended Security Updates" in windows_support("10.0.19045.1", today).detail


# -- Collector ---------------------------------------------------------------------------


@responses.activate
def test_collects_entra_and_intune_devices(db, two_tenants, factory):
    alpha, _ = two_tenants
    responses.get(f"{V1}/devices", json={"value": dd.ENTRA_DEVICES})
    responses.get(f"{V1}/informationProtection/bitlocker/recoveryKeys", json={"value": dd.RECOVERY_KEYS})
    responses.get(f"{V1}/directory/deviceLocalCredentials", json={"value": dd.LAPS_BACKUPS})
    responses.get(f"{V1}/deviceManagement/managedDevices", json={"value": dd.MANAGED_DEVICES})
    run_id, status = _run(db, alpha, "devices", factory)
    assert status == "succeeded"
    m = _metrics(db, run_id)
    assert (m["devices.entra_total"], m["devices.entra_stale"]) == (5, 1)
    assert (m["devices.managed"], m["devices.compliant"], m["devices.noncompliant"]) == (3, 2, 1)
    assert (m["devices.unencrypted_computers"], m["devices.not_synced"], m["devices.unsupported_windows"]) == (1, 1, 1)
    assert m["devices.recovery_keys_saved"] == 2  # LAPTOP-ANN and OLD-PC
    managed = next(c for c in responses.calls if "managedDevices" in c.request.url)
    assert "select=id%2CdeviceName" in managed.request.url  # only the fields we use


@responses.activate
def test_tenant_without_intune_collects_entra_only(db, two_tenants, factory):
    alpha, _ = two_tenants
    no_intune = [
        {**s, "servicePlans": [p for p in s["servicePlans"] if p["servicePlanName"] != "INTUNE_A"]} for s in SKUS
    ]
    responses.get(f"{V1}/organization", json={"value": [{"id": "org1", "displayName": "Alpha"}]})
    responses.get(f"{V1}/subscribedSkus", json={"value": no_intune})
    responses.get(f"{V1}/directory/subscriptions", json={"value": []})
    # Hybrid-joined servers: Windows Server 2022 and 2025, and a 2022 whose record just says "Windows".
    servers = [
        {
            **dd.ENTRA_DEVICES[3],
            "id": f"srv{n}",
            "deviceId": f"d-srv{n}",
            "displayName": f"SRV-{n}",
            "operatingSystem": os_name,
            "operatingSystemVersion": version,
        }
        for n, (os_name, version) in enumerate(
            (
                ("Windows Server", "10.0.20348.2700"),
                ("Windows Server 2025", "10.0.26100.1"),
                ("Windows", "10.0.20348.2700"),
            )
        )
    ]
    responses.get(f"{V1}/devices", json={"value": dd.ENTRA_DEVICES + servers})
    responses.get(f"{V1}/informationProtection/bitlocker/recoveryKeys", json={"value": []})
    responses.get(f"{V1}/directory/deviceLocalCredentials", json={"value": dd.LAPS_BACKUPS})
    assert _run(db, alpha, "tenant_profile", factory)[1] == "succeeded"
    assert _run(db, alpha, "devices", factory)[1] == "succeeded"
    assert not any("managedDevices" in c.request.url for c in responses.calls)
    r = _results(db, alpha)
    assert r["device_compliance"].status == r["device_encryption"].status == "unavailable"
    assert "Intune" in r["device_compliance"].summary
    # Without Intune, Windows support comes from devices active in Entra (OLD-PC is stale, so skipped).
    assert r["supported_windows"].status == "fail"
    assert [f.subject for f in r["supported_windows"].findings] == ["LAPTOP-BOB"]  # no server is "Windows 10"
    report = data_reports.build(db, store.get_tenant(db, alpha), "devices")
    systems = {row[0]: row[1] for row in report.sections[0].rows}
    assert systems["Windows Server"] == 1 and systems["Windows Server 2025"] == 1
    assert "Windows 10" in systems and systems["Windows 10"] == 1  # LAPTOP-BOB only


@responses.activate
def test_intune_forbidden_is_partial_not_failed(db, two_tenants, factory):
    alpha, _ = two_tenants
    responses.get(f"{V1}/devices", json={"value": dd.ENTRA_DEVICES})
    responses.get(f"{V1}/informationProtection/bitlocker/recoveryKeys", json={"value": dd.RECOVERY_KEYS})
    responses.get(f"{V1}/directory/deviceLocalCredentials", json={"value": dd.LAPS_BACKUPS})
    responses.get(
        f"{V1}/deviceManagement/managedDevices",
        status=403,
        json={"error": {"code": "Forbidden", "message": "Application is not authorized"}},
    )
    assert _run(db, alpha, "devices", factory)[1] == "partial"
    r = _results(db, alpha)
    assert r["device_compliance"].status == "unavailable"
    assert r["stale_devices"].status == "fail"


# -- Checks ------------------------------------------------------------------------------


def test_device_checks(db, two_tenants):
    alpha, _ = two_tenants
    seed_tenant(db, alpha)
    r = _results(db, alpha)
    assert r["device_compliance"].status == "fail"
    assert [(f.severity, f.subject) for f in r["device_compliance"].findings] == [("medium", "LAPTOP-BOB (bob@x.com)")]
    assert [f.severity for f in r["device_encryption"].findings] == ["high"]
    assert r["device_encryption"].summary == "1 of 2 computers not encrypted"  # the iPhone isn't a computer
    assert [(f.severity, f.subject) for f in r["supported_windows"].findings] == [("high", "LAPTOP-BOB (bob@x.com)")]
    assert [f.subject for f in r["intune_enrolled"].findings] == ["DESK-RECEPTION"]
    assert "Hybrid joined" in r["intune_enrolled"].findings[0].detail
    assert [f.subject for f in r["device_checkin"].findings] == ["LAPTOP-BOB (bob@x.com)"]
    assert [f.subject for f in r["stale_devices"].findings] == ["OLD-PC"]
    assert all(r[k].category == "Devices" for k in ("device_compliance", "stale_devices"))


def test_windows_summary_names_the_computers_reaching_end_of_support_too(client, db, two_tenants, monkeypatch):
    """Overview shows the summary beside the number of computers found, so it counts every one of them."""
    alpha, _ = two_tenants
    today = datetime.now(UTC).date()
    # A release whose Home and Pro updates end in a month, whatever day this runs.
    soon = WindowsRelease("Windows 11 (test)", 26300, today + timedelta(days=30), today + timedelta(days=395))
    monkeypatch.setattr("officesentry.devices.WINDOWS_RELEASES", (soon, *WINDOWS_RELEASES))
    _, bob, phone = dd.MANAGED_DEVICES  # Bob's laptop is on Windows 10
    desks = [{**bob, "id": f"md-desk{n}", "deviceName": f"DESK-{n}", "osVersion": "10.0.26300.1"} for n in (1, 2)]
    monkeypatch.setattr(dd, "MANAGED_DEVICES", [bob, phone, *desks])
    seed_tenant(db, alpha)
    r = _results(db, alpha)["supported_windows"]
    assert r.summary == "1 of 3 Windows computers past end of support, 2 more reaching it within 90 days"
    assert [f.severity for f in r.findings] == ["high", "low", "low"]
    make_user(db, "admin")
    login(client, "admin")
    page = client.get(f"/tenants/{alpha}").get_data(as_text=True)
    row = page.split(f'<div class="title">{r.summary}</div>')[1].split("</li>")[0]
    assert '<span class="pill" title="Accounts or items affected">3</span>' in row
    # Past it, nothing to come: the summary says only that.
    monkeypatch.setattr(dd, "MANAGED_DEVICES", [bob, phone])
    seed_tenant(db, alpha)
    assert _results(db, alpha)["supported_windows"].summary == "1 of 1 Windows computer past end of support"


def test_jailbroken_device_is_high(db, two_tenants, monkeypatch):
    alpha, _ = two_tenants
    phone = {**dd.MANAGED_DEVICES[2], "jailBroken": "True"}
    monkeypatch.setattr(dd, "MANAGED_DEVICES", dd.MANAGED_DEVICES[:2] + [phone])
    seed_tenant(db, alpha)
    findings = _results(db, alpha)["device_compliance"].findings
    assert findings[0].severity == "high" and findings[0].detail == "Jailbroken or rooted"


def test_clean_estate_passes(db, two_tenants, monkeypatch):
    alpha, _ = two_tenants
    monkeypatch.setattr(dd, "MANAGED_DEVICES", [dd.MANAGED_DEVICES[0], dd.MANAGED_DEVICES[2]])
    monkeypatch.setattr(dd, "ENTRA_DEVICES", [dd.ENTRA_DEVICES[0], dd.ENTRA_DEVICES[4]])
    seed_tenant(db, alpha)
    r = _results(db, alpha)
    for key in (
        "device_compliance",
        "device_encryption",
        "supported_windows",
        "intune_enrolled",
        "device_checkin",
        "stale_devices",
    ):
        assert r[key].status == "pass", (key, r[key].summary)


def test_never_collected_is_unavailable(db, two_tenants):
    _, bravo = two_tenants
    r = _results(db, bravo)
    assert all(r[k].status == "unavailable" for k in ("device_compliance", "supported_windows", "stale_devices"))


def test_insurance_pack_has_devices_control(db, two_tenants):
    alpha, _ = two_tenants
    seed_tenant(db, alpha)
    control = {c.control.key: c for c in evidence.build(db, store.get_tenant(db, alpha)).controls}["devices"]
    assert control.status == evidence.NOT_IN_PLACE  # unencrypted laptop on Windows 10, nothing passes
    assert any("not encrypted" in g for g in control.gaps)


def test_device_changes_since_last_month(db, two_tenants, monkeypatch):
    alpha, _ = two_tenants
    seed_tenant(db, alpha)
    old = (datetime.now(UTC) - timedelta(days=30)).strftime("%Y-%m-%dT%H:%M:%SZ")
    db.execute("UPDATE runs SET finished_at = ?", (old,))
    ann = {**dd.MANAGED_DEVICES[0], "complianceState": "noncompliant", "isEncrypted": False}
    new = {**dd.MANAGED_DEVICES[0], "id": "md9", "deviceName": "LAPTOP-NEW", "userPrincipalName": None}
    bob = {**dd.MANAGED_DEVICES[1], "complianceState": "compliant"}
    monkeypatch.setattr(dd, "MANAGED_DEVICES", [ann, bob, new])  # the iPhone was retired
    seed_tenant(db, alpha)
    since = (datetime.now(UTC) - timedelta(days=28)).strftime("%Y-%m-%dT%H:%M:%SZ")
    got = {(c.kind, c.subject, c.detail) for c in all_changes(db, alpha, since) if c.area == "Devices"}
    assert got == {
        ("added", "LAPTOP-NEW", "Enrolled Windows"),
        ("changed", "LAPTOP-ANN (ann@x.com)", "Now not compliant"),
        ("changed", "LAPTOP-ANN (ann@x.com)", "Disk no longer encrypted"),
        ("changed", "LAPTOP-BOB (bob@x.com)", "Compliant again"),
        ("removed", "Ann's iPhone", "Removed from Intune"),
    }


# -- Report ------------------------------------------------------------------------------


def test_devices_report(db, two_tenants):
    alpha, _ = two_tenants
    seed_tenant(db, alpha)
    report = data_reports.build(db, store.get_tenant(db, alpha), "devices", datetime.now(UTC))
    sections = {s.title: s for s in report.sections}
    assert list(sections) == [
        "Operating systems",
        "Needs attention",
        "Managed devices",
        "BitLocker recovery keys",
        "Local admin passwords",
        "Entra ID device records",
    ]
    systems = {row[0]: row for row in sections["Operating systems"].rows}
    assert systems["Windows 10"][1:] == [1, "Ended"]
    attention = sections["Needs attention"]
    assert [row[0] for row in attention.rows] == ["LAPTOP-BOB"]
    assert attention.flags == ["high"]
    assert "Not encrypted" in attention.rows[0][3] and "No check-in for 45 days" in attention.rows[0][3]
    records = sections["Entra ID device records"]
    assert records.rows[0][0] == "OLD-PC" and records.flags[0] == "low"  # least recently seen first
    assert dict(zip(records.columns, records.rows[-1], strict=True))["Join type"] == "Entra registered"


def test_devices_report_without_data(db, two_tenants):
    _, bravo = two_tenants
    report = data_reports.build(db, store.get_tenant(db, bravo), "devices")
    assert report.unavailable and not report.sections


def test_devices_page_and_xlsx(client, db, two_tenants):
    alpha, _ = two_tenants
    seed_tenant(db, alpha)
    make_user(db, "admin")
    login(client, "admin")
    page = client.get(f"/tenants/{alpha}/data/devices")
    assert page.status_code == 200
    html = page.get_data(as_text=True)
    assert "LAPTOP-BOB" in html and "Needs attention" in html
    resp = client.get(f"/tenants/{alpha}/data/devices.xlsx")
    assert resp.status_code == 200
    wb = load_workbook(BytesIO(resp.data))
    assert wb.sheetnames == [
        "Summary",
        "Operating systems",
        "Needs attention",
        "Managed devices",
        "BitLocker recovery keys",
        "Local admin passwords",
        "Entra ID device records",
    ]
