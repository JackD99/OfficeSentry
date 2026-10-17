"""Coming up: renewals, expiring credentials, mailbox quotas, DMARC and accepted risks, on the page and by email."""

import copy
import json
from datetime import UTC, datetime, timedelta
from io import BytesIO

import pytest
from openpyxl import load_workbook

from officesentry import delivery, lifecycle, store, upcoming
from officesentry.db import utcnow
from tests import graphdata as gd
from tests import groupdata as grp
from tests import licencedata as ld
from tests.conftest import login, make_user
from tests.seed import seed_tenant
from tests.test_email import _configure, fake_smtp, sends  # noqa: F401  (fake_smtp: autouse, nothing is sent)

GB = 1024**3
ISO = "%Y-%m-%dT%H:%M:%SZ"


def _in(days: float) -> str:
    return (datetime.now(UTC) + timedelta(days=days)).strftime(ISO)


def _exchange(used_gb: float = 48):
    """Ann's mailbox at ``used_gb`` of a 49.5 GB send limit; the shared mailbox has no limit."""
    data = copy.deepcopy(gd.EXCHANGE)
    ann, info = data["Get-Mailbox"]
    ann["ProhibitSendQuota"] = "49.5 GB (53,150,220,288 bytes)"
    info["ProhibitSendQuota"] = "Unlimited"
    data["Get-MailboxStatistics"] = {
        "g-u1": [{"TotalItemSize": f"{used_gb} GB ({int(used_gb * GB):,} bytes)", "ItemCount": 9000}],
        "g-u7": [{"TotalItemSize": "1 GB (1,073,741,824 bytes)", "ItemCount": 13}],
    }
    return data


def _apps() -> list[dict]:
    """Dated when the test reads them (a run may pass midnight UTC after this module loads)."""
    return [
        {
            "id": "app1",
            "appId": "a1",
            "displayName": "Payroll sync",
            "createdDateTime": "2024-01-01T00:00:00Z",
            "passwordCredentials": [{"displayName": "prod", "endDateTime": _in(10)}],
            "keyCredentials": [],
        },
        {
            "id": "app2",
            "appId": "a2",
            "displayName": "Backup agent",
            "createdDateTime": "2024-01-01T00:00:00Z",
            "passwordCredentials": [],
            "keyCredentials": [{"displayName": "cert", "endDateTime": _in(45)}],
        },
        {
            "id": "app3",
            "appId": "a3",
            "displayName": "Old connector",
            "createdDateTime": "2023-01-01T00:00:00Z",
            "passwordCredentials": [{"displayName": "s1", "endDateTime": _in(-5)}],
            "keyCredentials": [],
        },
        {
            "id": "app4",
            "appId": "a4",
            "displayName": "Rotated app",
            "createdDateTime": "2024-01-01T00:00:00Z",
            "passwordCredentials": [
                {"displayName": "old", "endDateTime": _in(20)},
                {"displayName": "new", "endDateTime": _in(700)},
            ],
            "keyCredentials": [],
        },
        {
            "id": "app5",
            "appId": "a5",
            "displayName": "Long gone",
            "createdDateTime": "2020-01-01T00:00:00Z",
            "passwordCredentials": [{"displayName": "s1", "endDateTime": "2021-01-01T00:00:00Z"}],
            "keyCredentials": [],
        },
    ]


@pytest.fixture
def tenant(db, two_tenants, monkeypatch):
    alpha, _ = two_tenants
    monkeypatch.setattr(gd, "APPLICATIONS", _apps())
    seed_tenant(db, alpha, skus=ld.SKUS, subscriptions=ld.TENANT_SUBSCRIPTIONS, exchange=_exchange())
    return store.get_tenant(db, alpha)


def _items(db, tenants, **kw) -> dict[str, upcoming.Item]:
    up = upcoming.collect(db, tenants, **kw)
    assert not up.without
    return {i.title: i for i in up.items}


def test_subscriptions_trials_and_lapsing(db, tenant):
    items = _items(db, [tenant])
    trial = items["Microsoft 365 Copilot trial ends"]
    assert (trial.type, trial.window, trial.severity, trial.days) == ("Trial", "Next 30 days", "high", 12)
    assert "2 people use it" in trial.detail and "Buy it before then" in trial.detail
    visio = items["Visio Plan 2 (grace period)"]
    assert (visio.window, visio.severity, visio.now) == ("Now", "high", True)
    exo = items["Exchange Online (Plan 1) (suspended)"]
    assert exo.window == "Now" and exo.days == 70  # suspended: needs attention now, deleted in 70 days
    # Business Premium renews in 300 days and Flow is free, so neither is listed.
    assert not any("Business Premium" in t or "FLOW" in t or "Power Automate" in t for t in items)


def test_app_secrets_and_certificates(db, tenant):
    items = _items(db, [tenant])
    secret = items["Payroll sync: “prod” secret"]
    assert (secret.type, secret.window, secret.severity) == ("App secret", "Next 30 days", "high")
    cert = items["Backup agent: “cert” certificate"]
    assert (cert.type, cert.window, cert.severity) == ("App certificate", "31 to 60 days", "medium")
    expired = items["Old connector: “s1” secret"]
    assert (expired.window, expired.severity, expired.when) == ("Now", "high", "5 days ago")
    rotated = items["Rotated app: “old” secret"]
    assert rotated.severity == "low" and "probably been replaced" in rotated.detail
    assert not any(t.startswith("Long gone") for t in items)  # expired years ago: the Apps report has it


def test_teams_and_groups_expiring(db, two_tenants):
    alpha, _ = two_tenants
    expiry = {"g4": 20, "g2": 50, "g1": 200, "g3": -2}
    groups = [{**g, "expirationDateTime": _in(expiry[g["id"]])} if g["id"] in expiry else g for g in grp.GROUPS]
    seed_tenant(db, alpha, groups=groups)
    items = _items(db, [store.get_tenant(db, alpha)])
    board = items["Board expires"]
    assert (board.type, board.window, board.severity, board.days) == (
        "Team or group expiry",
        "Next 30 days",
        "medium",
        20,
    )
    assert (
        board.detail
        == "Team with 2 members. Ann can renew it; otherwise it's deleted, and can be restored for 30 days."
    )
    assert board.href.endswith("/data/group-structure")
    phoenix = items["Project Phoenix expires"]
    assert (phoenix.window, phoenix.severity) == ("31 to 60 days", "medium")  # its only owner is blocked
    assert phoenix.detail.startswith("Team with 3 members. No owner can sign in to renew it;")
    marketing = items["Marketing expires"]
    assert (marketing.window, marketing.now, marketing.severity) == ("Now", True, "high")  # no owner, past its date
    assert "Sales expires" not in items  # 200 days away


def test_mailbox_nearly_full_now(db, two_tenants):
    alpha, _ = two_tenants
    seed_tenant(db, alpha, exchange=_exchange())  # the default licences include Exchange Online
    items = _items(db, [store.get_tenant(db, alpha)])
    ann = items["ann@x.com is nearly full"]
    assert (ann.type, ann.window, ann.due) == ("Mailbox quota", "Now", None)  # no history to forecast from
    assert "48.0 GB of 49.5 GB (97%)" in ann.detail
    assert not any(t.startswith("info@x.com") for t in items)  # no limit


def _older_sizes(db, tenant_id: int, days_ago: int, sizes: dict[str, int]) -> None:
    """A mailbox size collection from ``days_ago`` days back, for the growth forecast."""
    run = store.enqueue_run(db, tenant_id, "mailbox_usage")
    finished = (datetime.now(UTC) - timedelta(days=days_ago)).strftime(ISO)
    db.execute(
        "UPDATE runs SET status = 'succeeded', started_at = ?, finished_at = ? WHERE id = ?", (finished, finished, run)
    )
    store.add_objects(db, run, "mailbox_size", [(k, {"totalBytes": v}) for k, v in sizes.items()])


def test_mailbox_growth_forecast(db, two_tenants):
    alpha, bravo = two_tenants
    _older_sizes(db, alpha, 30, {"u1": 30 * GB})  # 10 GB in 30 days; 9.5 GB left: full in about 28 days
    seed_tenant(db, alpha, exchange=_exchange(used_gb=40))
    seed_tenant(db, bravo, exchange=_exchange(used_gb=40))
    assert "ann@x.com fills up" not in _items(db, [store.get_tenant(db, bravo)])  # 81%, no history: nothing yet
    ann = _items(db, [store.get_tenant(db, alpha)])["ann@x.com fills up"]
    assert (ann.window, ann.severity) == ("Next 30 days", "medium") and 27 <= ann.days <= 29
    assert "growing about 10.0 GB a month" in ann.detail


def test_dmarc_monitor_only(db, tenant):
    items = _items(db, [tenant])
    dmarc = items["x.com DMARC is p=none"]
    assert (dmarc.window, dmarc.due) == ("No date", None)
    assert "Monitor-only since" in dmarc.detail and "p=quarantine" in dmarc.detail
    # Accepting the DMARC finding as it is takes the domain off the list.
    db.execute(
        "UPDATE findings SET accepted_at = ?, accepted_by = 'jack' WHERE tenant_id = ?"
        " AND rule_key = 'dmarc_policy' AND object_id = 'x.com:DMARC'",
        (utcnow(), tenant["id"]),
    )
    assert "x.com DMARC is p=none" not in _items(db, [tenant])


def test_accepted_risks_due_for_review(db, tenant):
    f = db.execute(
        "SELECT * FROM findings WHERE tenant_id = ? AND resolved_at IS NULL"
        " AND rule_key != 'dmarc_policy' ORDER BY id LIMIT 2",
        (tenant["id"],),
    ).fetchall()
    db.execute(
        "UPDATE findings SET accepted_at = ?, accepted_by = 'jack', accepted_until = ? WHERE id = ?",
        (utcnow(), _in(20), f[0]["id"]),
    )
    db.execute(
        "UPDATE findings SET accepted_at = ?, accepted_by = 'jack', accepted_until = ? WHERE id = ?",
        (utcnow(), _in(200), f[1]["id"]),
    )  # too far off
    risks = [i for i in _items(db, [tenant]).values() if i.type == "Accepted risk"]
    assert len(risks) == 1
    risk = risks[0]
    assert (risk.window, risk.severity, risk.days) == ("Next 30 days", "medium", 20)
    assert risk.title.endswith(f": {f[0]['subject']}") and "Accepted by jack" in risk.detail


def test_our_own_certificates(db, tenant, connection_id):
    db.execute("UPDATE app_connections SET not_after = ? WHERE id = ?", (_in(20), connection_id))
    assert not any(i.type == "Office Sentry certificate" for i in _items(db, [tenant]).values())
    ours = _items(db, [tenant], ours=True)["The Test app certificate"]
    assert (ours.tenant, ours.window, ours.severity, ours.days) == (None, "Next 30 days", "high", 20)
    assert "Collection from every client stops" in ours.detail
    assert ours.href == f"/admin/connections#conn-{connection_id}"


def test_soonest_first_and_paused_tenants_left_out(db, tenant):
    up = upcoming.collect(db, [tenant])
    windows = [i.window for i in up.items]
    assert windows == sorted(windows, key=upcoming.WINDOWS.index)
    assert up.soonest.window == "Now"
    dated = [i.due for i in up.in_window("Next 30 days")]
    assert dated == sorted(dated)
    lifecycle.pause(db, tenant, "jack")
    paused = upcoming.collect(db, [store.get_tenant(db, tenant["id"])])
    assert paused.items == [] and paused.tenants == 0


def test_odd_data_and_a_failing_part(db, tenant, monkeypatch):
    run = store.latest_run(db, tenant["id"], "app_consents")["id"]
    db.execute("DELETE FROM objects WHERE run_id = ? AND kind = 'application'", (run,))
    store.add_objects(
        db,
        run,
        "application",
        [
            (
                "x1",
                {
                    "displayName": None,
                    "credentials": [
                        {"type": "secret", "endDateTime": "not-a-date"},
                        {"type": "secret", "endDateTime": None},
                        "junk",
                    ],
                },
            ),
            ("x2", {"displayName": "No creds", "credentials": None}),
            ("x3", {"credentials": "0001-01-01T00:00:00Z"}),
        ],
    )
    up = upcoming.collect(db, [tenant])
    assert not up.without and not any(i.type == "App secret" for i in up.items)
    assert any(i.type == "Trial" for i in up.items)

    def boom(*a):
        raise ValueError("bad data")

    monkeypatch.setattr(upcoming, "_mailboxes", boom)
    up = upcoming.collect(db, [tenant])
    assert [part for _, part in up.without] == ["Mailbox quotas"]  # left out for this tenant, the rest still there
    assert any(i.type == "Trial" for i in up.items)
    assert "Mailbox quotas couldn't be read" in upcoming.report(up).notes[0]


# -- The page -------------------------------------------------------------------------------------


def test_page_leads_with_the_answer_and_links_rows(client, db, tenant, connection_id):
    db.execute("UPDATE app_connections SET not_after = ? WHERE id = ?", (_in(40), connection_id))
    make_user(db, "admin")
    login(client, "admin")
    page = client.get("/upcoming").get_data(as_text=True)
    assert "items need attention now." in page  # the answer first
    assert "Needs attention now" in page and "Next 30 days" in page and "61 to 90 days" in page
    assert "Copilot trial ends" in page and "Payroll sync" in page and "The Test app certificate" in page
    assert f'href="/tenants/{tenant["id"]}/data/subscriptions"' in page
    # A tile opens its rows: the table filtered to that window.
    filtered = client.get("/upcoming?f=coming-up/when:Next+30+days").get_data(as_text=True)
    assert "Copilot trial ends" in filtered and "Backup agent" not in filtered
    for ext in ("xlsx", "csv"):
        assert client.get(f"/upcoming.{ext}").status_code == 200
    book = load_workbook(BytesIO(client.get("/upcoming.xlsx").data))
    assert "Coming up" in book.sheetnames
    csv = client.get("/upcoming.csv?f=coming-up/type:Trial").get_data(as_text=True)
    assert "Copilot" in csv and "Payroll" not in csv
    assert client.get("/upcoming.pdf").status_code == 200


def test_analyst_sees_only_granted_tenants_and_not_our_certificates(
    client, db, two_tenants, monkeypatch, connection_id
):
    alpha, bravo = two_tenants
    seed_tenant(db, alpha)
    seed_tenant(db, bravo)
    db.execute("UPDATE app_connections SET not_after = ? WHERE id = ?", (_in(40), connection_id))
    make_user(db, "analyst", role="analyst", tenants=[alpha])
    login(client, "analyst")
    page = client.get("/upcoming").get_data(as_text=True)
    assert "Alpha Ltd" in page and "Bravo Inc" not in page and "certificate expires" not in page
    assert "The Test app certificate" not in page
    assert b"Bravo" not in client.get("/upcoming.csv").data
    make_user(db, "client", role="client", tenants=[alpha])
    client.post("/logout")
    login(client, "client")
    assert client.get("/upcoming").status_code == 403


def test_empty_state(client, db):
    make_user(db, "admin")
    login(client, "admin")
    page = client.get("/upcoming").get_data(as_text=True)
    assert "No clients yet" in page  # an admin can add one; anyone else sees "No clients to show"


# -- The weekly email -----------------------------------------------------------------------------


def test_weekly_email_to_the_team(client, db, tenant, app, cfg, vault):
    make_user(db, "admin")
    login(client, "admin")
    _configure(client)
    team = delivery.add_contact(db, None, "Me", "me@msp.com")
    outsider = delivery.add_contact(db, tenant["id"], "Client", "it@client.com")
    assert client.get("/emails/coming-up/new").status_code == 200
    resp = client.post(
        "/emails/coming-up/new",
        data={
            "frequency": "weekly",
            "day_of_week": "0",
            "day_of_month": "1",
            "hour": "8",
            "include_xlsx": "1",
            "contacts": [str(team), str(outsider)],
        },
    )
    assert resp.status_code == 302
    s = db.execute("SELECT * FROM email_schedules WHERE kind = 'digest'").fetchone()
    assert json.loads(s["reports"]) == ["upcoming"] and delivery.is_upcoming(s)
    assert [c["email"] for c in delivery.recipients(db, "schedule", s["id"])] == ["me@msp.com"]
    overview = client.get("/emails").get_data(as_text=True)
    assert "Coming up" in overview and "Set up the weekly email" not in overview
    assert client.get(f"/schedules/{s['id']}/preview").status_code == 200
    assert "Copilot trial ends" in client.get(f"/schedules/{s['id']}/preview/body").get_data(as_text=True)

    client.post(f"/schedules/{s['id']}/send")
    ((_, _, msg, to_addrs),) = sends()
    assert to_addrs == ["me@msp.com"]
    assert msg["Subject"].startswith("Coming up: ") and "need attention now" in msg["Subject"]
    html = msg.get_body(("html",)).get_content()
    assert "Needs attention now" in html and "Copilot trial ends" in html and "Alpha Ltd" in html
    assert f"/tenants/{tenant['id']}/data/subscriptions" in html
    text = msg.get_body(("plain",)).get_content()
    assert "Next 30 days:" in text and "Copilot trial ends" in text
    (attachment,) = [p.get_filename() for p in msg.iter_attachments() if p.get_filename().endswith(".xlsx")]
    assert attachment.startswith("coming-up-")
    row = db.execute("SELECT * FROM email_deliveries ORDER BY id DESC").fetchone()
    assert row["kind"] == "digest" and row["status"] == "sent"

    # Editing keeps it a Coming up email.
    client.post(
        f"/schedules/{s['id']}/edit",
        data={"frequency": "weekly", "day_of_week": "2", "day_of_month": "1", "hour": "9", "contacts": [str(team)]},
    )
    assert delivery.is_upcoming(db.execute("SELECT * FROM email_schedules WHERE id = ?", (s["id"],)).fetchone())


def test_weekly_email_goes_out_on_its_slot(db, tenant, app, cfg, vault):
    for key, value in {"smtp_host": "smtp.example.com", "smtp_from": "r@msp.com", "smtp_security": "none"}.items():
        db.execute("INSERT INTO settings (key, value) VALUES (?, ?)", (key, value))
    team = delivery.add_contact(db, None, "Me", "me@msp.com")
    sid = db.execute(
        "INSERT INTO email_schedules (kind, reports, include_xlsx, frequency, day_of_week, hour, timezone,"
        " last_period, created_at) VALUES ('digest', '[\"upcoming\"]', 0, 'weekly', 0, 8, 'UTC', '2026-W40', ?)",
        (utcnow(),),
    ).lastrowid
    delivery.set_recipients(db, "schedule", sid, [team])
    monday = datetime(2026, 10, 5, 8, 30, tzinfo=UTC)
    assert delivery.tick(app, db, vault, cfg, monday) == 1
    ((_, _, msg, _),) = sends()
    assert msg["Subject"].startswith("Coming up")
    assert not [p for p in msg.iter_attachments() if (p.get_filename() or "").endswith(".xlsx")]
    assert db.execute("SELECT last_period FROM email_schedules WHERE id = ?", (sid,)).fetchone()[0] == "2026-W41"


def test_subject_leads_with_the_answer():
    counts = dict.fromkeys(upcoming.WINDOWS, 0)
    assert delivery.upcoming_subject(counts) == "Coming up: nothing due in the next 90 days"
    assert delivery.upcoming_subject({**counts, "61 to 90 days": 2}) == "Coming up: 2 items in the next 90 days"
    assert (
        delivery.upcoming_subject({**counts, "Now": 1, "Next 30 days": 4})
        == "Coming up: 1 needs attention now, 4 due within 30 days"
    )


def test_analyst_cannot_set_up_the_email(client, db, two_tenants):
    alpha, _ = two_tenants
    make_user(db, "analyst", role="analyst", tenants=[alpha])
    login(client, "analyst")
    assert client.get("/emails/coming-up/new").status_code == 403
