"""The Group membership report: who is in which team and Microsoft 365 group."""

from datetime import UTC, datetime
from io import BytesIO

from openpyxl import load_workbook

from officesentry.collectors.groups import GROUP_LIMIT, MEMBER_IDS_KEPT
from officesentry.reports import data_reports, group_membership, library, views
from officesentry.reports.data_reports import DataReport
from tests.conftest import login, make_user
from tests.seed import seed_tenant

NOW = datetime(2026, 10, 3, tzinfo=UTC)


def _tenant(db, tenant_id):
    return db.execute("SELECT * FROM tenants WHERE id = ?", (tenant_id,)).fetchone()


def _table(report, title):
    section = next(s for s in report.sections if s.title == title)
    return [dict(zip(section.columns, row, strict=True)) for row in section.rows], section


def test_registered_in_files_and_teams_and_the_library():
    assert "group-membership" in data_reports.REPORTS
    keys = list(data_reports.REPORTS)
    assert keys.index("group-membership") == keys.index("groups") + 1
    files = next(g for g in data_reports.grouped() if g.key == "files")
    assert "group-membership" in [k for k, _, _ in files.reports]
    entry = next(e for a in library.catalogue() for e in a.entries if e.key == "group-membership")
    assert "who is in" in entry.search and "offboarding" in entry.uses


def test_seeded_report(db, two_tenants):
    alpha, _ = two_tenants
    seed_tenant(db, alpha)
    report = data_reports.build(db, _tenant(db, alpha), "group-membership")
    assert report.unavailable is None
    tiles = {t.label: t.value for t in report.tiles}
    assert tiles == {
        "Teams and groups checked": 5,
        "People in at least one": 4,
        "People in none": 1,
        "Guests in teams": 2,
    }

    people, section = _table(report, "By person")
    by_account = {str(p["Account"]): p for p in people}
    ann = by_account["ann@x.com"]
    assert ann["Teams and groups"] == 4 and list(ann["Member of"].parts) == [
        "Board",
        "Marketing",
        "Provisioned",
        "Sales",
    ]
    assert ann["Type"] == "Member" and ann["Account status"] == "Enabled"
    # A disabled leaver still in a team is flagged.
    bob = people.index(by_account["bob@x.com"])
    assert by_account["bob@x.com"]["Account status"] == "Disabled" and section.flags[bob] == "low"
    assert by_account["g#EXT#@x.com"]["Teams and groups"] == 0  # the guest account is in nothing
    assert by_account["pat@supplier.com"]["Type"] == "Guest"
    assert by_account["pat@supplier.com"]["Teams and groups"] == 2

    groups, _ = _table(report, "By team or group")
    sales = next(g for g in groups if g["Name"] == "Sales")
    assert sales["Team"] == "Yes" and sales["Members"] == 3
    assert "ann@x.com" in sales["Member list"].parts
    assert "1 other member (groups, devices or contacts)" in sales["Member list"].parts
    assert list(sales["Guests"].parts) == ["pat@supplier.com"] and sales["Note"] == ""

    every, section = _table(report, "Every membership")
    assert section.compare is False
    assert len(every) == 4 + 1 + 2 + 1  # Ann, Bob, Pat, Lee
    ann_sales = next(r for r in every if str(r["Account"]) == "ann@x.com" and r["Team or group"] == "Sales")
    assert ann_sales["Owner"] == "Yes" and ann_sales["Team"] == "Yes"
    assert any("Security groups and distribution lists" in n for n in report.notes)


def test_headline_links_open_rows(client, db, two_tenants):
    alpha, _ = two_tenants
    seed_tenant(db, alpha)
    make_user(db, "admin")
    login(client, "admin")
    report = data_reports.build(db, _tenant(db, alpha), "group-membership")
    path = f"/tenants/{alpha}/data/group-membership"
    links = views.headline_links(report, path)["tiles"]
    assert all(links)
    none = views.apply(report, views.View.parse(links[2].split("?")[1].split("#")[0]))
    assert [str(r[0]) for r in none.sections[0].rows] == ["g#EXT#@x.com"]
    page = client.get(path)
    assert page.status_code == 200 and "Every membership" in page.get_data(as_text=True)
    xlsx = client.get(f"{path}.xlsx")
    assert xlsx.status_code == 200
    assert any(name.startswith("Every membership") for name in load_workbook(BytesIO(xlsx.data)).sheetnames)


class _Snap:
    def __init__(self, groups, users):
        self._data = {("groups", "group"): groups, ("users", "user"): users}
        self.capabilities = {}

    def objects(self, collector, kind):
        return self._data.get((collector, kind))


def _build(groups, users):
    report = DataReport("group-membership", "Group membership", "", {"id": 1, "name": "T"}, NOW)
    group_membership.membership(_Snap(groups, users), NOW, report)
    return report


def test_truncated_and_unexamined_groups_are_noted():
    users = [
        {
            "_id": "u1",
            "userPrincipalName": "ann@x.com",
            "displayName": "Ann",
            "userType": "Member",
            "accountEnabled": True,
        }
    ]
    groups = [
        {
            "_id": "big",
            "displayName": "Everyone",
            "kind": "microsoft365",
            "isTeam": True,
            "visibility": "Public",
            "examined": True,
            "memberIds": ["u1"],
            "memberIdsComplete": False,
            "memberCount": MEMBER_IDS_KEPT + 10,
            "owners": [],
            "ownerCount": 0,
            "guests": [],
            "guestCount": 0,
        },
        {"_id": "late", "displayName": "Zeta", "kind": "microsoft365", "isTeam": False, "examined": False},
        {"_id": "sec", "displayName": "VPN", "kind": "security"},
    ]
    report = _build(groups, users)
    rows, _ = _table(report, "By team or group")
    assert [r["Name"] for r in rows] == ["Everyone", "Zeta"]  # security groups aren't included
    big, late = rows
    assert (
        "Member list cut" in big["Note"]
        and f"{MEMBER_IDS_KEPT + 9} more members not listed" in big["Member list"].parts
    )
    assert late["Members"] == "Not checked" and str(GROUP_LIMIT) in late["Note"]
    notes = " ".join(report.notes)
    assert "Everyone" in notes and "1 of 2" in notes
    tiles = {t.label: (t.value, t.tone) for t in report.tiles}
    assert tiles["Teams and groups checked"] == (1, "warn")


def test_guest_in_many_groups_is_flagged():
    users = [
        {
            "_id": "gx",
            "userPrincipalName": "pat_s.com#EXT#@x.com",
            "mail": "pat@s.com",
            "userType": "Guest",
            "accountEnabled": True,
        }
    ]
    groups = [
        {
            "_id": f"g{i}",
            "displayName": f"G{i}",
            "kind": "microsoft365",
            "examined": True,
            "memberIds": ["gx"],
            "memberIdsComplete": True,
            "memberCount": 1,
            "guests": ["pat@s.com"],
            "guestCount": 1,
            "owners": [],
            "ownerCount": 0,
        }
        for i in range(group_membership.GUEST_MANY)
    ]
    report = _build(groups, users)
    rows, section = _table(report, "By person")
    assert len(rows) == 1 and rows[0]["Teams and groups"] == group_membership.GUEST_MANY
    assert section.flags == ["low"]


def test_unavailable_without_groups_or_users():
    assert "Group data" in _build(None, []).unavailable
    assert "User data" in _build([], None).unavailable
