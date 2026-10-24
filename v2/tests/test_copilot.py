"""The Microsoft 365 Copilot report: collected from Copilot's usage reports, matched to licences and accounts."""

import io
from datetime import UTC, datetime

import pytest
from openpyxl import load_workbook

from officesentry import store, usage
from officesentry.reports import data_reports, rollups
from officesentry.snapshot import TenantSnapshot
from tests import licencedata as ld
from tests import usagedata as ud
from tests.conftest import login, make_user
from tests.seed import seed_tenant
from tests.test_collectors import SKUS

COPILOT_PLANS = [
    {"servicePlanName": name, "provisioningStatus": "Success"}
    for name in ("M365_COPILOT_APPS", "M365_COPILOT_TEAMS", "M365_COPILOT_BUSINESS_CHAT", "GRAPH_CONNECTORS_COPILOT")
]
# Five Copilot add-on licences bought, four assigned: Ann uses hers, Bob's account is blocked, Carol was given hers
# three days ago and Dan's is in Copilot's report but has never been used.
COPILOT_SKU = {
    "skuId": "sku-cop",
    "skuPartNumber": "Microsoft_365_Copilot",
    "appliesTo": "User",
    "capabilityStatus": "Enabled",
    "prepaidUnits": {"enabled": 5},
    "consumedUnits": 4,
    "servicePlans": COPILOT_PLANS,
}
# E7 includes Copilot along with a mailbox, so its Copilot has no price of its own.
E7_SKU = {
    "skuId": "sku-e7",
    "skuPartNumber": "MICROSOFT_365_E7",
    "appliesTo": "User",
    "capabilityStatus": "Enabled",
    "prepaidUnits": {"enabled": 1},
    "consumedUnits": 1,
    "servicePlans": [*COPILOT_PLANS, {"servicePlanName": "EXCHANGE_S_ENTERPRISE", "provisioningStatus": "Success"}],
}
NEW = {**ld.state("sku-cop"), "lastUpdatedDateTime": ld.days_from_now(-3)}
USERS = [
    ld.user("u1", "Ann", [ld.state("sku-bp"), ld.state("sku-cop")]),
    ld.user("u2", "Bob", [ld.state("sku-bp"), ld.state("sku-cop")], enabled=False),
    ld.user("u4", "Carol", [ld.state("sku-bp"), NEW]),
    ld.user("u6", "Dan", [ld.state("sku-bp"), ld.state("sku-cop")]),
    ld.user("u7", "Erin", [ld.state("sku-bp")]),
]
COPILOT = [
    ud.copilot_user("ann@x.com", 1, word=1, teams=4, chat=45),  # Copilot Chat last used over 30 days ago
    ud.copilot_user("bob@x.com"),
    ud.copilot_user("dan@x.com"),
]
TREND = ud.copilot_trend([0, 1, 1, 2, 1])


def _seed(db, tenant_id, **usage_kw):
    seed_tenant(
        db,
        tenant_id,
        users=USERS,
        skus=[*SKUS, COPILOT_SKU],
        usage_kw={"copilot": COPILOT, "trend": TREND, **usage_kw},
    )


def _report(db, tenant_id):
    return data_reports.build(db, store.get_tenant(db, tenant_id), "copilot")


def _rows(report, title):
    section = next(s for s in report.sections if s.title == title)
    return {r[0]: r for r in section.rows}, dict(zip([r[0] for r in section.rows], section.flags, strict=True))


def _price(db, part, price):
    db.execute(
        "INSERT INTO sku_prices (sku_part_number, monthly_price, updated_at) VALUES (?, ?, '2026-10-01')", (part, price)
    )


def test_copilot_licences_are_told_by_their_service_plan():
    suite = dict(SKUS[0])
    assert set(usage.copilot_skus([suite, COPILOT_SKU, E7_SKU, {"skuId": "x"}, "junk"])) == {"sku-cop", "sku-e7"}
    assert usage.copilot_addon(COPILOT_SKU) and not usage.copilot_addon(E7_SKU)


def test_collector_keeps_each_app_and_the_daily_counts(db, two_tenants):
    alpha, _ = two_tenants
    _seed(db, alpha)
    snap = TenantSnapshot(db, store.get_tenant(db, alpha))
    ann = usage.person(snap, "ann@x.com")
    assert ann["copilotLicensed"] and ann["copilotActive"]
    assert ann["copilotApps"]["word"] == ud.days_ago(1) and ann["copilotApps"]["excel"] is None
    assert usage.person(snap, "erin@x.com")["copilotLicensed"] is False
    status = usage.status(snap)
    assert status["read"]["copilot_trend"] is True and status["copilotReportDate"] == ud.REFRESHED
    days = snap.objects("m365_usage", "copilot_day")
    assert [d["active"] for d in days] == [0, 1, 1, 2, 1] and days[-1]["apps"]["word"] == 1
    # Bob is blocked from signing in, so of the two others with Copilot in the report, Ann used it.
    assert dict(snap.history("copilot.licensed"))
    assert [v for _, v in snap.history("copilot.licensed")] == [2]
    assert [v for _, v in snap.history("copilot.active")] == [1]


def test_daily_counts_are_only_asked_for_when_someone_holds_copilot(db, two_tenants):
    alpha, bravo = two_tenants
    seed_tenant(db, alpha, usage_kw={"trend_status": 500})
    status = usage.status(TenantSnapshot(db, store.get_tenant(db, alpha)))
    assert status["read"]["copilot"] is True and "copilot_trend" not in status["read"]
    # With Copilot, a failure is noted and the report leaves the chart out.
    _seed(db, bravo, trend_status=500)
    status = usage.status(TenantSnapshot(db, store.get_tenant(db, bravo)))
    assert status["read"]["copilot_trend"] is False
    assert not _report(db, bravo).chart


def test_copilot_report(db, two_tenants):
    alpha, _ = two_tenants
    _seed(db, alpha)
    _price(db, "Microsoft_365_Copilot", 25.0)
    report = _report(db, alpha)
    assert not report.unavailable and report.parts == {"copilot": "read"}
    tiles = {t.label: t for t in report.tiles}
    assert (tiles["Copilot licences"].value, tiles["Copilot licences"].foot) == (4, "5 bought, 1 not assigned")
    # Of the three in Copilot's report, Ann used it.
    assert (tiles["Using Copilot"].value, tiles["Using Copilot"].foot) == ("33%", "1 of 3 used it in the last 30 days")
    # Bob's and Dan's, and the one nobody has; Carol's is too new to count.
    assert tiles["Unused licences"].value == 3
    assert float(tiles["Unused cost / month"].value) == 75.0
    assert report.notes == ["1 Copilot licence given in the last 14 days isn't counted as unused yet."]

    apps, _ = _rows(report, "Copilot in each app")
    assert apps["Word"] == ["Word", 1, "33%"] and apps["Teams"][1] == 1 and apps["Copilot Chat"][1] == 0
    people, flags = _rows(report, "People with Copilot")
    assert list(people) == ["ann@x.com", "bob@x.com", "carol@x.com", "dan@x.com"]
    assert people["ann@x.com"][2].parts == ("Teams", "Word")
    assert people["carol@x.com"][1] == "Not in Microsoft's report yet" and people["dan@x.com"][1] == "Never"
    assert flags == {"ann@x.com": None, "bob@x.com": "medium", "carol@x.com": None, "dan@x.com": "medium"}

    unused, _ = _rows(report, "Unused Copilot licences")
    assert list(unused) == ["bob@x.com", "dan@x.com"]
    assert unused["bob@x.com"][3] == "Account blocked from sign-in" and unused["dan@x.com"][3] == "Never used"
    assert float(unused["dan@x.com"][4]) == 25.0 and unused["dan@x.com"][2]  # when Dan last signed in

    licences, licence_flags = _rows(report, "Copilot licences")
    assert licences["Microsoft 365 Copilot"][1:5] == [5, 4, 1, 1]
    assert licence_flags["Microsoft 365 Copilot"] == "medium"

    # Erin is busiest without Copilot (Teams meetings first); Ann, Bob, Carol and Dan hold it.
    nxt, _ = _rows(report, "Who could use Copilot next")
    assert list(nxt) == ["erin@x.com"]
    assert nxt["erin@x.com"][1:4] == [9, 60, 70] and nxt["erin@x.com"][4].parts == ("Outlook", "Word", "Excel")
    assert report.chart and "People who used Copilot" in report.chart_title


def test_a_suite_that_includes_copilot_isnt_costed(db, two_tenants):
    alpha, _ = two_tenants
    users = [ld.user("u1", "Ann", [ld.state("sku-e7")])]
    seed_tenant(db, alpha, users=users, skus=[E7_SKU], usage_kw={"copilot": [ud.copilot_user("ann@x.com")]})
    _price(db, "MICROSOFT_365_E7", 99.0)
    report = _report(db, alpha)
    unused, _ = _rows(report, "Unused Copilot licences")
    assert unused["ann@x.com"][4] == "Part of Microsoft 365 E7" or unused["ann@x.com"][4].startswith("Part of")
    tiles = {t.label: t for t in report.tiles}
    assert float(tiles["Unused cost / month"].value) == 0
    licences, _ = _rows(report, "Copilot licences")
    assert next(iter(licences.values()))[5] == "In the suite's price"


def test_without_copilot_the_report_still_says_who_could_use_it(db, two_tenants):
    alpha, _ = two_tenants
    seed_tenant(db, alpha)  # nobody holds Copilot
    report = _report(db, alpha)
    tiles = {t.label: t for t in report.tiles}
    assert tiles["Copilot licences"].value == 0
    assert tiles["Copilot licences"].foot == "Nobody here has Microsoft 365 Copilot"
    assert tiles["Using Copilot"].value == "No one licensed"
    assert "Copilot licences" not in [s.title for s in report.sections]
    nxt, _ = _rows(report, "Who could use Copilot next")
    assert list(nxt)[0] == "erin@x.com" and "free@x.com" not in nxt and "bob@x.com" not in nxt
    assert not report.chart


def test_when_copilots_report_fails_nobody_is_called_unused(db, two_tenants):
    alpha, _ = two_tenants
    _seed(db, alpha, copilot_status=400)
    report = _report(db, alpha)
    assert report.parts == {"copilot": "failed"} and report.notes[0].startswith("Microsoft's Copilot usage report")
    tiles = {t.label: t for t in report.tiles}
    assert tiles["Using Copilot"].value == "Not known"
    unused, _ = _rows(report, "Unused Copilot licences")
    assert not unused
    people, _ = _rows(report, "People with Copilot")
    assert {r[1] for r in people.values()} == {"Not known"}


def test_hidden_names_use_the_report_alone(db, two_tenants):
    alpha, _ = two_tenants
    code = "A" * 32
    active = [ud.active_user(code, ["MICROSOFT 365 BUSINESS PREMIUM", "MICROSOFT 365 COPILOT"], exchange=1)]
    seed_tenant(
        db,
        alpha,
        users=USERS,
        skus=[*SKUS, COPILOT_SKU],
        usage_kw={"active": active, "copilot": [ud.copilot_user(code)], "concealed": True},
    )
    report = _report(db, alpha)
    people, _ = _rows(report, "People with Copilot")
    assert list(people) == [f"Hidden name ({code[:6]})"]
    assert any("hides names" in n for n in report.notes)


@pytest.mark.parametrize("cloud", ["gcc_high", "dod", "china"])
def test_clouds_without_copilot(db, two_tenants, cloud):
    alpha, _ = two_tenants
    db.execute("UPDATE tenants SET cloud = ? WHERE id = ?", (cloud, alpha))
    report = _report(db, alpha)
    assert report.unavailable and "Copilot" in report.unavailable


def test_pages_downloads_and_rollup(client, db, two_tenants):
    alpha, bravo = two_tenants
    _seed(db, alpha)
    make_user(db, "admin")
    login(client, "admin")
    page = client.get(f"/tenants/{alpha}/data/copilot")
    assert page.status_code == 200 and "Who could use Copilot next" in page.get_data(as_text=True)
    wb = load_workbook(io.BytesIO(client.get(f"/tenants/{alpha}/data/copilot.xlsx").data))
    assert "Summary" in wb.sheetnames
    assert client.get(f"/tenants/{bravo}/data/copilot").status_code == 200  # nothing collected yet
    assert client.get("/across/copilot").status_code == 200
    rollup = rollups.build(db, [store.get_tenant(db, t) for t in (alpha, bravo)], "copilot", datetime.now(UTC))
    assert next(r for r in rollup.rows if r.tenant["id"] == alpha).values["c:Copilot licences"] == 4
