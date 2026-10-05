"""Mailbox holds and archiving: litigation hold, archives, leavers without a hold, large mailboxes."""

import copy
from io import BytesIO

import pytest
from openpyxl import load_workbook

from officesentry import store
from officesentry.reports import library
from officesentry.reports.data_reports import GROUPS, REPORTS, build
from tests import graphdata as gd
from tests.conftest import login, make_user
from tests.seed import seed_tenant
from tests.test_collectors import SKUS, USERS

KEY = "mailbox-holds"
# A shared mailbox converted from a person who has left: their account has signed in before.
OLD_USER = {"id": "u8", "displayName": "Old", "userPrincipalName": "old@x.com", "userType": "Member",
            "accountEnabled": True, "assignedLicenses": [], "licenseAssignmentStates": [],
            "signInActivity": {"lastSignInDateTime": "2026-03-01T09:00:00Z"}}


def _exchange():
    data = copy.deepcopy(gd.EXCHANGE)
    ann, info = data["Get-Mailbox"]
    # Ann: 48 GB of a 49.5 GB quota and no archive.
    ann.update({"ProhibitSendQuota": "49.5 GB (53,150,220,288 bytes)", "ArchiveStatus": "None",
                "LitigationHoldEnabled": False})
    info.update({"ProhibitSendQuota": "Unlimited", "ArchiveStatus": "None", "LitigationHoldEnabled": True})
    data["Get-Mailbox"] += [
        # Bob has left (sign-in blocked); his mailbox has an archive but no hold.
        {"ExternalDirectoryObjectId": "u2", "Guid": "g-u2", "UserPrincipalName": "bob@x.com",
         "PrimarySmtpAddress": "bob@x.com", "DisplayName": "Bob", "RecipientTypeDetails": "UserMailbox",
         "GrantSendOnBehalfTo": [], "ArchiveStatus": "Active", "LitigationHoldEnabled": False},
        {"ExternalDirectoryObjectId": "u8", "Guid": "g-u8", "UserPrincipalName": "old@x.com",
         "PrimarySmtpAddress": "old@x.com", "DisplayName": "Old", "RecipientTypeDetails": "SharedMailbox",
         "GrantSendOnBehalfTo": [], "ArchiveStatus": "None", "LitigationHoldEnabled": False},
        {"ExternalDirectoryObjectId": "", "Guid": "g-disc", "PrimarySmtpAddress": "DiscoverySearchMailbox@x.com",
         "Name": "DiscoverySearchMailbox", "RecipientTypeDetails": "DiscoveryMailbox", "GrantSendOnBehalfTo": []},
    ]
    data["Get-MailboxStatistics"] = {
        "g-u1": [{"TotalItemSize": "48 GB (51,539,607,552 bytes)", "ItemCount": 90000}],
        "g-u7": [{"TotalItemSize": "586.6 KB (600,640 bytes)", "ItemCount": 13}],
        "g-u2": [{"TotalItemSize": "2 GB (2,147,483,648 bytes)", "ItemCount": 500}],
        "g-u8": [{"TotalItemSize": "45 GB (48,318,382,080 bytes)", "ItemCount": 800}],
    }
    return data


@pytest.fixture
def tenant(db, two_tenants):
    alpha, _ = two_tenants
    seed_tenant(db, alpha, users=USERS + [gd.INFO_USER, OLD_USER], exchange=_exchange())
    return store.get_tenant(db, alpha)


def _report(db, tenant):
    report = build(db, tenant, KEY)
    assert not report.unavailable, report.unavailable
    return report, {s.title: s for s in report.sections}


def test_every_mailbox(db, tenant):
    report, s = _report(db, tenant)
    every = s["Every mailbox"]
    assert every.columns == ["Mailbox", "Type", "Account", "Litigation hold", "Archive", "Mailbox size"]
    rows = {r[0]: r for r in every.rows}
    assert set(rows) == {"ann@x.com", "bob@x.com", "info@x.com", "old@x.com"}  # no discovery mailbox
    assert rows["ann@x.com"][1:] == ["User", "Active", "No", "No", "48.0 GB"]
    assert rows["bob@x.com"][1:] == ["User", "Sign-in blocked", "No", "Yes", "2.0 GB"]
    assert rows["info@x.com"][1:5] == ["Shared", "Active", "Yes", "No"]
    assert "1 of 4 mailboxes is on litigation hold and 1 has an archive" in every.intro
    tiles = {t.label: t for t in report.tiles}
    assert tiles["Mailboxes on litigation hold"].value == "1 of 4"
    assert tiles["Archive turned on"].value == "1 of 4"
    assert tiles["Leavers' mailboxes without a hold"].value == 2
    assert tiles["Leavers' mailboxes without a hold"].tone == "warn"


def test_leavers_without_a_hold(db, tenant):
    _, s = _report(db, tenant)
    leavers = s["Leavers without a hold"]
    assert [r[0] for r in leavers.rows] == ["bob@x.com", "old@x.com"]
    assert leavers.rows[0][2] == "Sign-in blocked"
    assert "converted" in leavers.rows[1][2]
    assert leavers.flags == ["medium", "medium"]
    assert "licence is removed" in leavers.intro and leavers.intro.startswith("2 mailboxes")


def test_large_mailboxes_without_an_archive(db, tenant):
    _, s = _report(db, tenant)
    large = s["Large mailboxes without an archive"]
    # Ann is at 97% of her quota; old@ has no quota but is over 40 GB. Bob is small and has an archive.
    assert [r[0] for r in large.rows] == ["ann@x.com", "old@x.com"]
    assert large.rows[0][2:] == ["48.0 GB", "49.5 GB", "97%"]
    assert large.rows[1][3:] == ["No limit", ""]
    assert large.flags == ["low", "low"]


def test_seeded_tenant_without_sizes(db, two_tenants):
    alpha, _ = two_tenants
    seed_tenant(db, alpha)
    report, s = _report(db, store.get_tenant(db, alpha))
    assert s["Large mailboxes without an archive"].rows == []
    assert report.tiles[2].value == 0 and report.tiles[2].tone == "ok"


def test_unavailable_before_collection_and_without_exchange(db, two_tenants):
    alpha, beta = two_tenants
    assert "hasn't been collected" in build(db, store.get_tenant(db, beta), KEY).unavailable
    no_exchange = [{**SKUS[0], "servicePlans": [{"servicePlanName": "AAD_PREMIUM", "provisioningStatus": "Success"}]}]
    seed_tenant(db, alpha, skus=no_exchange)
    assert build(db, store.get_tenant(db, alpha), KEY).unavailable == "This tenant has no Exchange Online licence."


def test_registered_grouped_and_searchable():
    assert REPORTS[KEY][0] == "Mailbox holds and archiving"
    assert KEY in dict((k, keys) for k, _, keys in GROUPS)["email"]
    assert "ediscovery" in library.KEYWORDS[KEY]
    assert any(KEY in use.reports for use in library.USES)


def test_page_and_xlsx(client, db, tenant):
    make_user(db, "admin")
    login(client, "admin")
    page = client.get(f"/tenants/{tenant['id']}/data/{KEY}")
    assert page.status_code == 200
    text = page.get_data(as_text=True)
    assert "Mailbox holds and archiving" in text and "Leavers without a hold" in text
    resp = client.get(f"/tenants/{tenant['id']}/data/{KEY}.xlsx")
    assert resp.status_code == 200
    assert load_workbook(BytesIO(resp.data)).sheetnames
