"""Cyber-insurance evidence pack."""

from io import BytesIO

from openpyxl import load_workbook

from officesentry import findings, store
from officesentry.reports import evidence
from officesentry.reports.export import pdf_available
from tests import graphdata as gd
from tests.conftest import login, make_user
from tests.seed import seed_tenant
from tests.test_collectors import USERS


def _by_key(pack):
    return {c.control.key: c for c in pack.controls}


def test_controls_answered_from_checks(db, two_tenants):
    alpha, _ = two_tenants
    seed_tenant(db, alpha, users=USERS + [gd.INFO_USER])
    pack = evidence.build(db, store.get_tenant(db, alpha))
    c = _by_key(pack)
    assert c["mfa_admins"].status == evidence.NOT_IN_PLACE  # breakglass without MFA (critical)
    assert c["email_auth"].status == evidence.PARTLY
    assert c["logging"].status == evidence.IN_PLACE
    assert c["backups"].status == c["training"].status == evidence.NOT_ASSESSED
    assert any("enabled accounts are covered" in e for e in c["mfa_all"].evidence)
    assert sum(pack.counts.values()) == len(evidence.CONTROLS)


def test_a_control_whose_only_check_fails_for_the_organisation_is_not_in_place(db, two_tenants, monkeypatch):
    alpha, _ = two_tenants
    monkeypatch.setitem(gd.EXCHANGE, "Get-OrganizationConfig", [{"AuditDisabled": True}])  # a medium finding
    seed_tenant(db, alpha)
    pack = evidence.build(db, store.get_tenant(db, alpha))
    assert _by_key(pack)["logging"].status == evidence.NOT_IN_PLACE
    answer = next(a for a in pack.answers if a.key == "audit_logs")
    assert answer.answer == "No"


def test_findings_about_a_few_accounts_are_gaps_not_a_missing_control():
    from officesentry.checks import Finding, RuleResult

    excluded = RuleResult("mfa_enforced", "MFA", "Identity", "fail", "", "", [Finding("medium", "bg@x.com", "x")])
    off = RuleResult("mailbox_auditing", "Audit", "Email", "fail", "", "", [Finding("medium", "Organisation", "x")])
    assert evidence._status([excluded]) == evidence.PARTLY
    assert evidence._status([off]) == evidence.NOT_IN_PLACE


def test_never_collected_is_not_assessed(db, two_tenants):
    _, bravo = two_tenants
    pack = evidence.build(db, store.get_tenant(db, bravo))
    assert all(c.status == evidence.NOT_ASSESSED for c in pack.controls)


def test_accepted_risk_counts_as_documented(db, two_tenants):
    alpha, _ = two_tenants
    seed_tenant(db, alpha)
    for row in db.execute("SELECT id FROM findings WHERE rule_key = 'smtp_auth'").fetchall():
        findings.accept(db, row["id"], "jack", "Scanner needs SMTP", None)
    c = _by_key(evidence.build(db, store.get_tenant(db, alpha)))
    assert any("accepted" in e for e in c["legacy_auth"].evidence)


def test_pages(client, db, two_tenants):
    alpha, bravo = two_tenants
    seed_tenant(db, alpha)
    make_user(db, "client", role="client", tenants=[alpha])
    login(client, "client")
    assert client.get(f"/tenants/{bravo}/evidence").status_code == 404
    page = client.get(f"/tenants/{alpha}/evidence?set=insurance").get_data(as_text=True)
    assert "Is multi-factor authentication (MFA) enforced for all users accessing email" in page
    controls = client.get(f"/tenants/{alpha}/evidence?set=controls").get_data(as_text=True)
    assert "Is multi-factor authentication required" in controls
    wb = load_workbook(BytesIO(client.get(f"/tenants/{alpha}/evidence.xlsx").data))
    assert wb["Controls"].max_row == len(evidence.CONTROLS) + 1
    if pdf_available():
        resp = client.get(f"/tenants/{alpha}/evidence.pdf")
        assert resp.data[:4] == b"%PDF"
