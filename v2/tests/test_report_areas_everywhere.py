"""What a role can't open on a client stays closed wherever that client's data goes: scheduled emails and
their previews, the report history, report designs, month-end notes, the home pages and What changed.

A role holds report areas (``reports.identity``, ``reports.security``, ...), and an admin can take one away
from a person on one client. ``tests/test_roles.py`` checks each page's own gate; these are the places that
show or send a report from inside another page, each with the request that used to get through. The client
cases run for a client account with a password and for one signed in with its own Microsoft account (granted
another organisation's client too, which must count for nothing).
"""

from __future__ import annotations

import json
import re
from datetime import UTC, datetime

import pytest

from officesentry import delivery, month_end, roles
from officesentry.reports import designs
from tests import test_client_sign_in as csi
from tests.conftest import login, make_user
from tests.seed import seed_tenant
from tests.test_email import _smtp_ready, fake_smtp, sends  # noqa: F401 (autouse fixture)
from tests.test_isolation_walk import microsoft_client, sign_in_with_microsoft
from tests.test_roles import role_user, store_user

MARK = "Zebedee"  # in every name of the seeded tenant's people
NOW = datetime.now(UTC)
PERIOD = f"{NOW.year}-{NOW.month:02d}"
FORM = {"frequency": "monthly", "day_of_month": "1", "hour": "8", "timezone": "UTC", "include_xlsx": "1"}


@pytest.fixture
def alpha(db, two_tenants):
    alpha, _ = two_tenants
    db.execute("UPDATE tenants SET consented_at = '2026-01-01T00:00:00Z'")
    seed_tenant(db, alpha, mark=MARK)
    _smtp_ready(db)
    return alpha


def _deny(db, uid: int, capability: str, tenant_id: int) -> None:
    roles.set_overrides(db, "admin", store_user(db, uid), [roles.Override(capability, tenant_id, False)])


def _switch(client, username: str, how: str = "password") -> None:
    client.post("/logout")
    if how == "microsoft":
        sign_in_with_microsoft(client, username)
    else:
        login(client, username)


# A client's person signs in with a password, or with their own work account from Alpha's tenant.
HOW = ["password", "microsoft"]


def _client_account(db, vault, username: str, alpha: int, bravo: int, how: str) -> int:
    """Alpha's client account. One signing in with Microsoft is granted Bravo too (Bravo's Directory ID known,
    another organisation's), which must count for nothing."""
    if how == "password":
        return make_user(db, username, role="client", tenants=[alpha])
    db.execute("UPDATE tenants SET tenant_ref = ? WHERE id = ?", (csi.BRAVO, bravo))
    return microsoft_client(db, vault, username, alpha, [alpha, bravo])


def _schedule(db, tenant_id: int, reports: list[str], email: str = "it@alpha.example") -> int:
    contact = delivery.add_contact(db, tenant_id, "IT", email)
    sid = db.execute(
        "INSERT INTO email_schedules (kind, tenant_id, reports, include_xlsx, frequency, day_of_month, hour,"
        " timezone, created_by, created_at) VALUES ('report', ?, ?, 1, 'monthly', 1, 8, 'UTC', 'admin',"
        " '2026-01-01T00:00:00Z')",
        (tenant_id, json.dumps(reports)),
    ).lastrowid
    delivery.set_recipients(db, "schedule", sid, [contact])
    return sid


def test_a_schedule_only_sends_reports_its_maker_can_open(client, db, alpha):
    """An account manager doesn't hold identity reports: the form doesn't offer them, and naming one, or a
    saved view of one, sends nothing of it."""
    role_user(db, "am", "account_manager", tenants=[alpha])
    login(client, "am")
    assert client.get(f"/tenants/{alpha}/data/all-users").status_code == 403
    form = client.get(f"/tenants/{alpha}/schedules/new").data.decode()
    assert 'value="data:all-users"' not in form and 'value="data:licences"' in form
    client.post("/reports/all-users/views", data={"tenant_id": str(alpha), "name": "Everyone", "query": "hide=x/y"})
    view = db.execute("SELECT id FROM saved_views WHERE report_key = 'all-users'").fetchone()[0]
    client.post(
        f"/tenants/{alpha}/schedules/new",
        data={
            **FORM,
            "reports": ["data:all-users", "data:licences"],
            "views": f"s{view}",
            f"view_format_s{view}": "xlsx",
            "new_email": "me@home.example",
        },
    )
    s = db.execute("SELECT * FROM email_schedules WHERE tenant_id = ?", (alpha,)).fetchone()
    assert json.loads(s["reports"]) == ["data:licences"]
    assert delivery.schedule_views(db, s["id"]) == []


def test_a_schedule_sending_more_than_you_open_only_changes_for_someone_who_can(client, db, alpha):
    """Changing who it goes to would hand over the rest, so the schedule is left as it is."""
    sid = _schedule(db, alpha, ["review", "data:all-users"])
    role_user(db, "am", "account_manager", tenants=[alpha])
    login(client, "am")
    resp = client.get(f"/schedules/{sid}/edit", follow_redirects=True)
    assert "Only someone who can open every report this email sends can change it." in resp.data.decode()
    client.post(f"/schedules/{sid}/edit", data={**FORM, "reports": ["review"], "new_email": "me@home.example"})
    assert [c["email"] for c in delivery.recipients(db, "schedule", sid)] == ["it@alpha.example"]
    assert json.loads(db.execute("SELECT reports FROM email_schedules").fetchone()[0]) == [
        "review",
        "data:all-users",
    ]
    # Someone who opens every report it sends still changes it.
    make_user(db, "admin")
    _switch(client, "admin")
    assert client.get(f"/schedules/{sid}/edit").status_code == 200


def _email_plan(db, tenant_id: int, reports: list[str]) -> int:
    """``tenant_id`` on a send-on-approval plan that emails ``reports`` to its contacts."""
    plan = month_end.save_plan(
        db,
        {
            "name": "Monthly",
            "reports": reports,
            "include_xlsx": 0,
            "day_of_month": 1,
            "hour": 8,
            "timezone": "UTC",
            "subject": "",
            "message": "",
            "mode": "approval",
            "delivery": "email",
            "bcc_team": 0,
        },
        "admin",
    )
    month_end.apply_plan(db, month_end.get_plan(db, plan), [tenant_id], "admin", "UTC")
    return plan


def _plan_recipients(db, tenant_id: int) -> list[str]:
    (schedule_id,) = db.execute(
        "SELECT schedule_id FROM report_plan_tenants WHERE tenant_id = ?", (tenant_id,)
    ).fetchone()
    return [c["email"] for c in delivery.recipients(db, "schedule", schedule_id)]


def test_only_someone_who_opens_every_plan_report_adds_contacts(client, db, alpha):
    """A new contact gets the client's plan emails. While the plan sends an identity report, an account manager
    (no identity reports) adds nobody, by any of the forms that take a new address, so they can't send it to
    themselves; once the plan sends only what they open, they can."""
    plan = _email_plan(db, alpha, ["review", "data:all-users"])
    role_user(db, "am", "account_manager", tenants=[alpha])
    login(client, "am")
    page = client.get(f"/tenants/{alpha}/schedules").data.decode()
    assert f"/tenants/{alpha}/contacts" not in page and "someone who can open all of them adds contacts" in page
    assert 'name="new_email"' not in client.get(f"/tenants/{alpha}/schedules/new").data.decode()
    resp = client.post(
        f"/tenants/{alpha}/contacts", data={"name": "Me", "email": "me@home.example"}, follow_redirects=True
    )
    assert "only someone who can open all of them can add a contact" in resp.data.decode()
    client.post(
        f"/tenants/{alpha}/schedules/new", data={**FORM, "reports": ["data:licences"], "new_email": "me@home.example"}
    )
    client.post(f"/tenants/{alpha}/alert", data={"min_severity": "critical", "new_email": "me@home.example"})
    assert db.execute("SELECT COUNT(*) FROM contacts").fetchone()[0] == 0
    assert db.execute("SELECT COUNT(*) FROM email_schedules WHERE plan_id IS NULL").fetchone()[0] == 0

    # Someone who opens everything it sends adds contacts, and they join the plan.
    make_user(db, "admin")
    _switch(client, "admin")
    client.post(f"/tenants/{alpha}/contacts", data={"name": "IT", "email": "it@alpha.example"})
    assert _plan_recipients(db, alpha) == ["it@alpha.example"]

    # A plan of reports the account manager opens takes their contacts too.
    db.execute("UPDATE report_plans SET reports = ? WHERE id = ?", (json.dumps(["review", "data:licences"]), plan))
    month_end.apply_plan(db, month_end.get_plan(db, plan), [alpha], "admin", "UTC")
    _switch(client, "am")
    assert f"/tenants/{alpha}/contacts" in client.get(f"/tenants/{alpha}/schedules").data.decode()
    client.post(f"/tenants/{alpha}/contacts", data={"name": "Finance", "email": "finance@alpha.example"})
    assert _plan_recipients(db, alpha) == ["finance@alpha.example", "it@alpha.example"]


def test_choosing_who_gets_alerts_needs_security_findings(client, db, alpha):
    """An alert names the client's new findings, so someone with those taken away doesn't choose who gets it."""
    uid = role_user(db, "sl", "second_line", tenants=[alpha])
    _deny(db, uid, "reports.security", alpha)
    login(client, "sl")
    page = client.get(f"/tenants/{alpha}/schedules").data.decode()
    assert f"/tenants/{alpha}/alert" not in page and "someone who can see security findings chooses" in page
    resp = client.post(f"/tenants/{alpha}/alert", data={"min_severity": "critical", "new_email": "me@home.example"})
    assert resp.status_code == 403
    assert db.execute("SELECT COUNT(*) FROM email_alerts").fetchone()[0] == 0
    assert db.execute("SELECT COUNT(*) FROM contacts").fetchone()[0] == 0


@pytest.mark.parametrize("how", HOW)
def test_report_history_needs_the_report_area_and_downloads(client, db, vault, alpha, two_tenants, how):
    sid = _schedule(db, alpha, ["data:all-users", "data:licences"])
    make_user(db, "admin")
    login(client, "admin")
    assert client.post(f"/schedules/{sid}/send").status_code == 302
    files = {
        r["kind"]: r["id"] for r in db.execute("SELECT kind, id FROM issued_reports WHERE tenant_id = ?", (alpha,))
    }
    assert set(files) == {"data:all-users", "data:licences"}
    role_user(db, "am", "account_manager", tenants=[alpha])
    _switch(client, "am")
    page = client.get(f"/tenants/{alpha}/report-history").data.decode()
    assert (
        f"report-history/{files['data:licences']}" in page and f"report-history/{files['data:all-users']}" not in page
    )
    assert client.get(f"/tenants/{alpha}/report-history/{files['data:all-users']}").status_code == 403
    assert client.get(f"/tenants/{alpha}/report-history/{files['data:licences']}").status_code == 200
    # Downloads taken away on this client: the files are listed, not downloadable.
    viewer = _client_account(db, vault, "it@alpha.example", alpha, two_tenants[1], how)
    _deny(db, viewer, "export", alpha)
    _switch(client, "it@alpha.example", how)
    page = client.get(f"/tenants/{alpha}/report-history").data.decode()
    assert "report-history/" not in page and "All users" in page and "Bravo Inc" not in page
    assert client.get(f"/tenants/{alpha}/report-history/{files['data:licences']}").status_code == 403
    assert client.get(f"/tenants/{two_tenants[1]}/report-history").status_code == 404


def test_email_previews_need_security_findings(client, db, alpha):
    """Every report email opens with the client's security summary, and an alert lists findings."""
    sid = _schedule(db, alpha, ["data:mailboxes"])
    aid = db.execute(
        "INSERT INTO email_alerts (tenant_id, min_severity, created_by, created_at) VALUES (?, 'high', 'admin',"
        " '2026-01-01T00:00:00Z')",
        (alpha,),
    ).lastrowid
    role_user(db, "fl", "first_line", tenants=[alpha])
    login(client, "fl")
    page = client.get(f"/tenants/{alpha}/schedules").data.decode()
    assert f"/schedules/{sid}/preview" not in page and f"/alerts/{aid}/preview" not in page
    for url in (f"/schedules/{sid}/preview", f"/schedules/{sid}/preview/body", f"/alerts/{aid}/preview"):
        assert client.get(url).status_code == 403, url


def test_draft_commentary_stays_out_of_emails_while_reports_wait_for_approval(client, db, alpha, monkeypatch):
    """A client on a send-on-approval plan gets the month's commentary once it's approved, also in the
    review attached to any other scheduled email."""
    pages = []
    monkeypatch.setattr(delivery.export, "to_pdf", lambda html, base: pages.append(html) or html.encode())
    plan = month_end.save_plan(
        db,
        {
            "name": "Monthly",
            "reports": ["review"],
            "include_xlsx": 0,
            "day_of_month": 1,
            "hour": 8,
            "timezone": "UTC",
            "subject": "",
            "message": "",
            "mode": "approval",
            "delivery": "email",
            "bcc_team": 0,
        },
        "admin",
    )
    month_end.apply_plan(db, month_end.get_plan(db, plan), [alpha], "admin", "UTC")
    month_end.update_review(db, alpha, PERIOD, "sl", state="in_review", commentary="Still being written.")
    sid = _schedule(db, alpha, ["review"])
    make_user(db, "admin")
    login(client, "admin")
    client.post(f"/schedules/{sid}/send")
    assert pages and not any("Still being written." in p for p in pages)
    month_end.update_review(db, alpha, PERIOD, "admin", state="approved")
    pages.clear()
    db.execute("DELETE FROM email_deliveries")
    client.post(f"/schedules/{sid}/send")
    assert any("Still being written." in p for p in pages)


@pytest.mark.parametrize("how", HOW)
def test_clients_never_see_who_on_the_team_sent_an_email(client, db, vault, alpha, two_tenants, how):
    sid = _schedule(db, alpha, ["data:mailboxes"])
    make_user(db, "jane.admin", role="admin")
    login(client, "jane.admin")
    client.post(f"/schedules/{sid}/send")
    assert b"sent by jane.admin" in client.get(f"/tenants/{alpha}/schedules").data
    _client_account(db, vault, "it@alpha.example", alpha, two_tenants[1], how)
    _switch(client, "it@alpha.example", how)
    page = client.get(f"/tenants/{alpha}/schedules").data.decode()
    assert "it@alpha.example" in page and "jane.admin" not in page and "Bravo Inc" not in page
    assert client.get(f"/tenants/{two_tenants[1]}/schedules").status_code == 404


def test_staff_notes_follow_month_end_review_on_each_client(client, db, two_tenants):
    alpha, bravo = two_tenants
    seed_tenant(db, alpha)
    db.execute(
        "INSERT INTO finding_notes (tenant_id, rule_key, body, visibility, created_at)"
        " VALUES (?, 'admin_mfa', 'Only for the team.', 'staff', '2026-01-01T00:00:00Z')",
        (alpha,),
    )
    month_end.update_review(db, alpha, PERIOD, "x", state="in_review", commentary="A draft.")
    uid = role_user(db, "sl", "second_line", tenants=[alpha, bravo])
    _deny(db, uid, "month_end.review", alpha)
    login(client, "sl")
    page = client.get(f"/tenants/{alpha}/report").data.decode()
    assert "Only for the team." not in page and "A draft." not in page


def _design(db, *blocks) -> int:
    return designs.create(db, "Board pack", "test", {"settings": {}, "blocks": list(blocks)})


def test_designs_leave_out_what_the_viewer_cant_open(client, db, alpha, two_tenants):
    table = {"kind": "data_table", "settings": {"report": "all-users", "max_rows": 50}}
    did = _design(db, {"kind": "findings", "settings": {"severities": ["critical", "high"]}}, table)
    # An account manager holds security findings but not identity reports.
    role_user(db, "am", "account_manager", tenants=[alpha])
    login(client, "am")
    assert MARK.encode() not in client.get(f"/tenants/{alpha}/designs/{did}").data
    # Second line with identity taken away on Alpha previews it there without the table, and with security
    # taken away on Bravo can't preview it on Bravo.
    _, bravo = two_tenants
    seed_tenant(db, bravo)
    uid = role_user(db, "sl", "second_line", tenants=[alpha, bravo])
    roles.set_overrides(
        db,
        "admin",
        store_user(db, uid),
        [roles.Override("reports.identity", alpha, False), roles.Override("reports.security", bravo, False)],
    )
    _switch(client, "sl")
    only_table = _design(db, table)
    page = client.get(f"/reporting/designs/{only_table}/preview?tenant={alpha}").data.decode()
    assert MARK not in page and "Left out: you can&#39;t open this report for this client." in page
    page = client.get(f"/reporting/designs/{did}/preview?tenant={bravo}").data.decode()
    assert "Bravo Inc" not in page


def _tile(body: str, label: str) -> str:
    m = re.search(re.escape(label) + r".*?<div class=\"value[^\"]*\">\s*(?:<span>)?([^<]*)", body, re.S)
    return m.group(1).strip() if m else ""


def test_home_pages_count_only_what_the_person_sees(client, db, two_tenants):
    alpha, bravo = two_tenants
    seed_tenant(db, bravo)
    uid = role_user(db, "sl", "second_line", tenants=[alpha, bravo])
    roles.set_overrides(
        db,
        "admin",
        store_user(db, uid),
        [roles.Override("reports.security", bravo, False), roles.Override("projects.view", bravo, False)],
    )
    login(client, "sl")
    assert _tile(client.get("/").data.decode(), "Critical and high findings") == "0"
    accounts = client.get("/clients").data.decode()
    assert "Bravo Inc" in accounts and f"/tenants/{bravo}/projects" not in accounts
    # First line holds no security findings: Today doesn't count them.
    role_user(db, "fl", "first_line", tenants=[bravo])
    _switch(client, "fl")
    today = client.get("/today").data.decode()
    assert _tile(today, "New critical and high") == "0" and "Failing:" not in today


@pytest.mark.parametrize("how", HOW)
def test_a_client_home_leaves_out_areas_taken_away(client, db, vault, alpha, two_tenants, how):
    bravo = two_tenants[1]
    seed_tenant(db, bravo)
    uid = _client_account(db, vault, "it@alpha.example", alpha, bravo, how)
    _switch(client, "it@alpha.example", how)
    page = client.get("/").data.decode()
    assert 'class="card client-recs"' in page and "Bravo Inc" not in page
    _deny(db, uid, "reports.security", alpha)  # which signs them out
    _switch(client, "it@alpha.example", how)
    page = client.get("/").data.decode()
    assert 'class="card client-recs"' not in page and "Checks needing attention" not in page
    assert f"/tenants/{alpha}/findings" not in page and "Users with MFA" in page and "Bravo Inc" not in page
    assert client.get(f"/tenants/{bravo}/findings").status_code == 404


# -- What changed ------------------------------------------------------------------------------------------


def _a_month_of_changes(db, tenant_id: int) -> None:
    """Collected again 40 days on, with the guest account deleted: an Accounts change (identity) among
    others, such as admin activity (tenant)."""
    from tests.test_collectors import USERS
    from tests.test_rollups import _backdate

    _backdate(db, tenant_id, 40)
    seed_tenant(db, tenant_id, mark=MARK, users=[u for u in USERS if u["id"] != "u3"])


GUEST = "g#EXT#@x.com"  # the deleted account
ON_PAGE = f'<span class="strong">Accounts</span> {GUEST}'  # its key change on the Changes tab
IN_CSV = f"Accounts,,Removed,{GUEST}"
ON_OVERVIEW = '<span class="title">Accounts</span>'


def _sheets(client, url: str) -> list[str]:
    from io import BytesIO

    from openpyxl import load_workbook

    return load_workbook(BytesIO(client.get(url).data)).sheetnames


def test_what_changed_shows_only_the_areas_the_person_opens(client, db, alpha):
    _a_month_of_changes(db, alpha)
    make_user(db, "admin")
    login(client, "admin")
    assert ON_PAGE in client.get(f"/tenants/{alpha}/changes").data.decode()
    assert IN_CSV in client.get(f"/tenants/{alpha}/changes.csv").data.decode()
    assert ON_OVERVIEW in client.get(f"/tenants/{alpha}").data.decode()

    # An account manager holds findings and tenant reports, but no identity, email, files or devices reports.
    role_user(db, "am", "account_manager", tenants=[alpha])
    _switch(client, "am")
    page = client.get(f"/tenants/{alpha}/changes").data.decode()
    assert ON_PAGE not in page and "Users with MFA" not in page and '<div class="label">New findings</div>' in page
    assert "can&#39;t open them for this client: Identity and access, Email" in page
    csv = client.get(f"/tenants/{alpha}/changes.csv").data.decode()
    assert IN_CSV not in csv and "Admin activity," in csv  # their own areas stay
    overview = client.get(f"/tenants/{alpha}").data.decode()
    assert ON_OVERVIEW not in overview and '<span class="title">Admin activity</span>' in overview

    # First line opens identity reports but not findings.
    role_user(db, "fl", "first_line", tenants=[alpha])
    _switch(client, "fl")
    page = client.get(f"/tenants/{alpha}/changes").data.decode()
    assert ON_PAGE in page and "Left out: you can't open findings for this client." in page
    assert '<div class="label">New findings</div>' not in page
    assert '<div class="label">Secure Score</div>' not in page  # the Secure Score trend is a tenant report
    assert "can&#39;t open them for this client: Findings, Files and Teams, Tenant and licences." in page
    assert "New findings" not in _sheets(client, f"/tenants/{alpha}/changes.xlsx")
    assert "new findings" not in client.get(f"/tenants/{alpha}").data.decode()

    # Taken away on one client only.
    uid = role_user(db, "sl", "second_line", tenants=[alpha])
    _deny(db, uid, "reports.identity", alpha)
    _switch(client, "sl")
    assert ON_PAGE not in client.get(f"/tenants/{alpha}/changes").data.decode()


def test_what_changed_across_clients_counts_only_the_areas_the_person_opens(client, db, alpha):
    _a_month_of_changes(db, alpha)
    make_user(db, "admin")
    login(client, "admin")
    (everything,) = [r for r in client.get("/changes.csv").data.decode().splitlines() if r.startswith("Alpha Ltd")]
    assert "Accounts" in everything
    role_user(db, "am", "account_manager", tenants=[alpha])
    _switch(client, "am")
    for _ in range(2):  # built, then from the stored row, which holds every area's counts
        (row,) = [r for r in client.get("/changes.csv").data.decode().splitlines() if r.startswith("Alpha Ltd")]
        cells = row.split(",")
        assert "Accounts" not in row and "Admin activity" in row
        assert int(cells[2]) < int(everything.split(",")[2])  # fewer changes
        assert cells[8] == ""  # no Users with MFA figure without identity reports
    assert "Accounts" not in client.get("/changes").data.decode()


def test_every_change_feed_and_figure_names_the_report_it_comes_from():
    """What changed leaves out a key change or figure whose report isn't named, so each one must be."""
    from pathlib import Path

    from officesentry import changes
    from officesentry.reports import data_reports, month_changes

    source = Path(changes.__file__).read_text(encoding="utf-8")
    areas = set(re.findall(r'Change\(\s*"([^"]+)"', source))
    assert areas and areas <= set(month_changes.CHANGE_REPORTS)
    assert {key for key, *_ in month_changes.FIGURES} <= set(month_changes.FIGURE_REPORTS)
    named = {*month_changes.CHANGE_REPORTS.values(), *month_changes.FIGURE_REPORTS.values()}
    assert named <= set(data_reports.REPORTS)
