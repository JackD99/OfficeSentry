"""A tenant as it was on a past day (``?on=``) and what changed between any two days (``?from=&to=``), read
from the daily history once retention has removed the snapshots (web/as_of.py, state_history.objects_of)."""

from __future__ import annotations

import io
from datetime import UTC, date, datetime, timedelta

import pytest
from openpyxl import load_workbook

from officesentry import changes, retention, state_history, store
from officesentry.reports import data_reports, month_changes
from officesentry.snapshot import TenantSnapshot
from officesentry.web import as_of
from tests.seed import seed_tenant
from tests.test_collectors import USERS

from .conftest import login, make_user

NOW = datetime.now(UTC)
TODAY = NOW.date()


@pytest.fixture(autouse=True)
def _now():
    """NOW and TODAY as each test starts, not as the module loaded: a run may pass midnight UTC in between."""
    global NOW, TODAY
    NOW = datetime.now(UTC)
    TODAY = NOW.date()


def day(days_ago: int) -> str:
    return (TODAY - timedelta(days=days_ago)).isoformat()


def _with_guest(mail: str, bob=True) -> list[dict]:
    return [u if u["id"] != "u3" else {**u, "mail": mail} for u in USERS if bob or u["id"] != "u2"]


@pytest.fixture
def then_and_now(db, two_tenants):
    """Alpha collected 60 days ago, with retention since removing that collection's snapshot so only the daily
    history holds it, and again today with its guest's address changed and Bob gone."""
    alpha = two_tenants[0]
    seed_tenant(db, alpha, users=_with_guest("gwen@then.example"))
    old = [r[0] for r in db.execute("SELECT id FROM runs WHERE tenant_id = ?", (alpha,))]
    marks = ",".join("?" * len(old))
    then = (NOW - timedelta(days=60)).strftime("%Y-%m-%dT%H:%M:%SZ")
    db.execute(
        f"UPDATE runs SET created_at = ?, started_at = ?, finished_at = ? WHERE id IN ({marks})",
        (then, then, then, *old),
    )
    db.execute(f"UPDATE state_runs SET finished_at = ? WHERE run_id IN ({marks})", (then, *old))
    seed_tenant(db, alpha, users=_with_guest("gwen@now.example", bob=False))
    db.execute(f"DELETE FROM objects WHERE run_id IN ({marks})", old)  # as the retention policy does
    assert db.execute(f"SELECT COUNT(*) FROM state_runs WHERE run_id IN ({marks})", old).fetchone()[0] > 20
    return alpha


def _collected_again(db, tenant_id: int, days_ago: int) -> None:
    """Move today's collection back ``days_ago`` days."""
    when = (NOW - timedelta(days=days_ago)).strftime("%Y-%m-%dT%H:%M:%SZ")
    new = [
        r[0] for r in db.execute("SELECT id FROM runs WHERE tenant_id = ? AND finished_at > ?", (tenant_id, day(30)))
    ]
    marks = ",".join("?" * len(new))
    db.execute(
        f"UPDATE runs SET created_at = ?, started_at = ?, finished_at = ? WHERE id IN ({marks})", (when,) * 3 + (*new,)
    )
    db.execute(f"UPDATE state_runs SET finished_at = ? WHERE run_id IN ({marks})", (when, *new))


@pytest.fixture
def admin(client, db):
    make_user(db, "admin")
    login(client, "admin")
    return client


# -- The days asked for ----------------------------------------------------------------------------


def test_a_day_is_moved_into_the_data_held(db, then_and_now):
    first = TODAY - timedelta(days=60)
    assert as_of.earliest(db, then_and_now) == first
    assert as_of.day(db, then_and_now, day(30)).on == TODAY - timedelta(days=30)
    assert as_of.day(db, then_and_now, day(30)).as_of == datetime.combine(
        TODAY - timedelta(days=30), datetime.max.time().replace(microsecond=0), UTC
    )
    early = as_of.day(db, then_and_now, day(400))
    assert early.on == first and "starts on" in early.note
    for raw in (None, "", "nonsense", TODAY.isoformat(), "2999-01-01", "2026-02-30"):
        assert as_of.day(db, then_and_now, raw).on is None, raw  # today: the latest data
    assert as_of.day(db, then_and_now, day(30)).query == {"on": day(30)}


def test_a_period_is_put_in_order_and_moved_into_the_data_held(db, then_and_now):
    first = TODAY - timedelta(days=60)
    assert as_of.period(db, then_and_now, None, day(5)) is None  # no start: the last N days
    swapped = as_of.period(db, then_and_now, day(5), day(20))
    assert (swapped.start, swapped.end) == (TODAY - timedelta(days=20), TODAY - timedelta(days=5))
    assert swapped.query == {"from": day(20), "to": day(5)}
    assert as_of.period(db, then_and_now, day(20), TODAY.isoformat()).end is None  # up to the latest data
    early = as_of.period(db, then_and_now, day(400), day(10))
    assert early.start == first and "starts on" in early.note
    # From the first day: as the data stood at its end, since nothing came before it.
    assert early.since == as_of.end_of(first)
    assert as_of.period(db, then_and_now, day(20), None).since == datetime.combine(
        TODAY - timedelta(days=20), datetime.min.time(), UTC
    )


def test_a_tenant_with_no_data_has_no_past_days(db, two_tenants):
    assert as_of.earliest(db, two_tenants[1]) is None
    assert as_of.day(db, two_tenants[1], day(10)).on is None


# -- Reading a past day ----------------------------------------------------------------------------


def test_a_past_day_is_rebuilt_from_the_history_once_the_snapshot_has_gone(db, then_and_now):
    tenant = store.get_tenant(db, then_and_now)
    then = TenantSnapshot(db, tenant, as_of=as_of.end_of(TODAY - timedelta(days=30)))
    assert {u.get("mail") or u["displayName"] for u in then.objects("users", "user")} == {
        "Ann",
        "Bob",
        "gwen@then.example",
    }
    assert {u.get("mail") or u["displayName"] for u in TenantSnapshot(db, tenant).objects("users", "user")} == {
        "Ann",
        "gwen@now.example",
    }
    # Before the first collection there is nothing to read.
    assert TenantSnapshot(db, tenant, as_of=NOW - timedelta(days=90)).objects("users", "user") is None


def test_change_feeds_compare_any_two_days(db, then_and_now):
    before, until = day(61) + "T00:00:00Z", day(1) + "T00:00:00Z"
    removed = [c for c in changes.user_changes(db, then_and_now, day(59) + "T00:00:00Z") if c.subject == "bob@x.com"]
    assert removed and removed[0].kind == "removed"
    # Up to a day before today's collection, no account had changed yet.
    assert not [c for c in changes.all_changes(db, then_and_now, before, until) if c.area == "Accounts"]


def test_a_change_report_between_two_past_days_reads_both_from_the_history(db, then_and_now):
    tenant = store.get_tenant(db, then_and_now)
    first = TODAY - timedelta(days=60)
    span = month_changes.build(db, tenant, since=as_of.end_of(first), until=as_of.end_of(TODAY - timedelta(days=1)))
    assert span.has_baseline and not [c for c in span.key_changes if c.area == "Accounts"]
    assert not sum(r.count for r in span.reports)  # every detail report read the same on both days
    since_then = month_changes.build(db, tenant, since=as_of.end_of(first))
    assert since_then.has_baseline and [c for c in since_then.key_changes if c.subject == "bob@x.com"]
    assert next(r for r in since_then.reports if r.key == "guests").count


# -- The pages -------------------------------------------------------------------------------------


def test_a_report_shows_the_tenant_as_it_was_on_a_chosen_day(admin, then_and_now):
    url = f"/tenants/{then_and_now}/data/guests"
    today = admin.get(url).get_data(as_text=True)
    assert "gwen@now.example" in today and "gwen@then.example" not in today
    assert 'name="on"' in today and f'min="{day(60)}"' in today and "As it was on" not in today

    page = admin.get(f"{url}?on={day(30)}").get_data(as_text=True)
    assert "gwen@then.example" in page and "gwen@now.example" not in page
    assert "As it was on" in page and "Back to today" in page
    assert f"/tenants/{then_and_now}/changes?from={day(29)}" in page  # from the end of the day shown
    # The downloads and every table link keep the day.
    assert f"guests.pdf?on={day(30)}" in page and f"guests.xlsx?on={day(30)}" in page
    assert f"on={day(30)}&amp;section=1" in page or f"section=1&amp;on={day(30)}" in page

    # So does the report picker, on laptops and phones.
    assert f'href="/tenants/{then_and_now}/data/mfa?on={day(30)}"' in page
    assert f'<input type="hidden" name="on" value="{day(30)}">' in page
    picked = admin.get(f"/tenants/{then_and_now}/data?key=mfa&on={day(30)}")
    assert picked.headers["Location"].endswith(f"/tenants/{then_and_now}/data/mfa?on={day(30)}")

    csv = admin.get(f"/tenants/{then_and_now}/data/guests.csv?section=2&on={day(30)}")
    assert "gwen@then.example" in csv.get_data(as_text=True) and "gwen@now.example" not in csv.get_data(as_text=True)
    assert f"-{day(30)}.csv" in csv.headers["Content-Disposition"]
    xlsx = admin.get(f"{url}.xlsx?on={day(30)}")
    assert xlsx.status_code == 200 and f"-{day(30)}.xlsx" in xlsx.headers["Content-Disposition"]


def test_what_changed_since_a_past_day_starts_after_what_that_day_shows(admin, db, then_and_now):
    # The second collection two days ago: that day already shows it, so nothing has changed since.
    _collected_again(db, then_and_now, 2)
    page = admin.get(f"/tenants/{then_and_now}/data/guests?on={day(2)}").get_data(as_text=True)
    assert "gwen@now.example" in page
    assert f'href="/tenants/{then_and_now}/changes?from={day(1)}"' in page
    period = as_of.period(db, then_and_now, day(1), None)
    since = month_changes.build(db, store.get_tenant(db, then_and_now), since=period.since)
    assert not [c for c in since.key_changes if c.area == "Accounts"]
    assert not next(r for r in since.reports if r.key == "guests").count
    assert as_of.day(db, then_and_now, day(1)).changes_from == TODAY  # never after today


def test_a_past_day_keeps_the_filters_and_saved_views(admin, db, then_and_now):
    page = admin.get(f"/tenants/{then_and_now}/data/guests?on={day(30)}&hide=guest-accounts/invited").get_data(
        as_text=True
    )
    assert '<input type="hidden" name="hide" value="guest-accounts/invited">' in page
    assert f'href="/tenants/{then_and_now}/data/guests?on={day(30)}"' in page  # the views menu's Reset


def test_a_day_before_the_data_shows_the_first_day_and_says_so(admin, then_and_now):
    page = admin.get(f"/tenants/{then_and_now}/data/guests?on={day(400)}").get_data(as_text=True)
    assert "gwen@then.example" in page and "starts on" in page


def test_the_changes_page_compares_any_two_days(admin, then_and_now):
    url = f"/tenants/{then_and_now}/changes"
    since = admin.get(f"{url}?from={day(60)}").get_data(as_text=True)
    assert "Chosen dates" in since and "gwen@now.example" in since and "bob@x.com" in since
    assert f'value="{day(60)}"' in since and f'value="{TODAY.isoformat()}"' in since
    assert f"changes.pdf?from={day(60)}" in since and f"changes.xlsx?from={day(60)}" in since

    quiet = admin.get(f"{url}?from={day(60)}&to={day(1)}").get_data(as_text=True)
    assert "gwen@now.example" not in quiet and "No findings appeared or were fixed" in quiet
    assert f"changes.csv?from={day(60)}&amp;to={day(1)}" in quiet

    pdf = admin.get(f"{url}.pdf?from={day(60)}&to={day(1)}")
    assert pdf.status_code == 200 and f"changes-{day(60)}-to-{day(1)}" in pdf.headers["Content-Disposition"]
    csv = admin.get(f"{url}.csv?from={day(60)}").get_data(as_text=True)
    assert "gwen@now.example" in csv
    assert "Last 30 days" in admin.get(f"{url}?from=nonsense").get_data(as_text=True)


def test_the_reports_a_period_ending_before_today_changed_open_at_its_end(admin, db, then_and_now):
    _collected_again(db, then_and_now, 2)
    url = f"/tenants/{then_and_now}/changes"
    ended = admin.get(f"{url}?from={day(60)}&to={day(1)}").get_data(as_text=True)
    assert "gwen@now.example" in ended
    assert f'href="/tenants/{then_and_now}/data/guests?on={day(1)}"' in ended
    latest = admin.get(f"{url}?from={day(60)}").get_data(as_text=True)
    assert f'href="/tenants/{then_and_now}/data/guests"' in latest  # up to today: today's data


def test_a_period_ending_before_today_names_its_end_day_in_the_spreadsheet(admin, db, then_and_now):
    tenant = store.get_tenant(db, then_and_now)
    end = TODAY - timedelta(days=20)
    ended = month_changes.build(db, tenant, since=as_of.end_of(TODAY - timedelta(days=60)), until=as_of.end_of(end))
    assert month_changes.to_tables(ended)[0].columns[2] == f"On {end.day} {end:%b %Y}"
    assert month_changes.to_tables(month_changes.build(db, tenant, 30))[0].columns[2] == "Now"
    xlsx = admin.get(f"/tenants/{then_and_now}/changes.xlsx?from={day(60)}&to={day(20)}")
    sheet = load_workbook(io.BytesIO(xlsx.data))["Key figures"]
    cells = {c.value for row in sheet.iter_rows() for c in row}
    assert f"On {end.day} {end:%b %Y}" in cells and "Now" not in cells


def test_a_past_day_is_only_shown_for_the_tenants_a_user_can_see(client, db, two_tenants, then_and_now):
    alpha, bravo = two_tenants
    make_user(db, "bravo-only", role="analyst", tenants=(bravo,))
    login(client, "bravo-only")
    for path in (
        f"/tenants/{alpha}/data/guests?on={day(30)}",
        f"/tenants/{alpha}/data/guests.csv?on={day(30)}",
        f"/tenants/{alpha}/data/guests.pdf?on={day(30)}",
        f"/tenants/{alpha}/changes?from={day(60)}",
        f"/tenants/{alpha}/changes.xlsx?from={day(60)}&to={day(1)}",
    ):
        assert client.get(path).status_code == 404, path


def test_every_report_builds_on_a_past_day(db, then_and_now):
    tenant = store.get_tenant(db, then_and_now)
    then = as_of.end_of(TODAY - timedelta(days=30))
    for key in data_reports.REPORTS:
        report = data_reports.build(db, tenant, key, as_of=then)
        if key == "sign-ins":  # sign-in detail isn't in the daily history, and its snapshot has gone
            assert "no longer held" in report.unavailable
        else:  # as available as today's
            assert report.unavailable == data_reports.build(db, tenant, key).unavailable, key


def test_legacy_sign_ins_retention_removed_are_said_to_be_gone_not_compared(db, then_and_now):
    # Kept for less time than the daily history: the day says so, and What changed doesn't list today's as new.
    tenant = store.get_tenant(db, then_and_now)
    then = data_reports.build(db, tenant, "legacy-auth", as_of=as_of.end_of(TODAY - timedelta(days=30)))
    assert "Legacy sign-ins from that day is no longer held" in then.notice
    legacy = next(r for r in month_changes.build(db, tenant, 30).reports if r.key == "legacy-auth")
    assert legacy.count == 0 and "Legacy sign-ins from the start of the period is no longer held" in legacy.note


def test_the_first_day_with_data_moves_on_when_the_daily_history_is_pruned(db, then_and_now):
    # A data source with nothing in it (no inbox rules, say) has nothing for pruning to remove.
    tenant = store.get_tenant(db, then_and_now)
    old = [
        r[0] for r in db.execute("SELECT id FROM runs WHERE collector = 'inbox_rules' AND finished_at < ?", (day(30),))
    ]
    marks = ",".join("?" * len(old))
    db.execute(f"UPDATE runs SET object_count = 0 WHERE id IN ({marks})", old)
    db.execute(f"UPDATE state_runs SET objects = 0 WHERE run_id IN ({marks})", old)
    assert as_of.earliest(db, then_and_now) == TODAY - timedelta(days=60)
    assert retention.prune_state(db, 1)  # Daily history: 1 month
    assert as_of.earliest(db, then_and_now) == TODAY
    assert "starts on" in as_of.day(db, then_and_now, day(45)).note
    # A day the daily history no longer holds says so, rather than blaming the snapshots' retention.
    gone = data_reports.build(db, tenant, "all-users", as_of=as_of.end_of(TODAY - timedelta(days=45))).unavailable
    assert gone.startswith("Users and ") and gone.endswith(
        " from that day are no longer held: the daily history doesn't go back that far."
    )
    assert "retention policy" not in gone


def test_admin_activity_on_a_day_before_the_activity_history_starts(db, then_and_now):
    # That day's snapshot has the changes; the activity history only has what came after it started.
    tenant = store.get_tenant(db, then_and_now)
    then = as_of.end_of(TODAY - timedelta(days=45))
    held = TenantSnapshot(db, tenant, as_of=then).objects("admin_activity", "admin_event")
    assert held
    db.execute(
        "UPDATE event_cursors SET since_at = ? WHERE tenant_id = ? AND source = 'audits'",
        (NOW.strftime("%Y-%m-%dT%H:%M:%SZ"), then_and_now),
    )
    report = data_reports.build(db, tenant, "admin-activity", as_of=then)
    every = next(s for s in report.sections if s.title == "All admin changes")
    assert len(every.rows) == len(held)


def test_end_of_a_day_is_its_last_second():
    assert as_of.end_of(date(2026, 3, 1)) == datetime(2026, 3, 1, 23, 59, 59, tzinfo=UTC)
    assert state_history.READABLE.count("EXISTS") == 2
