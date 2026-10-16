"""Exchange, inbox rules, apps, sharing: client, collectors, checks and coverage."""

import json
from datetime import UTC, datetime

import pytest
import responses

from officesentry import coverage, store
from officesentry.exchange import EXCHANGE_ROOT, ExchangeAccessDenied, ExchangeClient, ExchangeError
from officesentry.mailflow import inbox_rule_risk, is_external, smtp_address
from officesentry.reports.data_reports import build
from officesentry.rules import evaluate
from officesentry.snapshot import TenantSnapshot
from tests import graphdata as gd
from tests.seed import seed_tenant
from tests.stored import unpack
from tests.test_collectors import USERS

URL = f"{EXCHANGE_ROOT}/adminapi/beta/tid/InvokeCommand"


def client():
    return ExchangeClient("tid", lambda: "tok", sleep=lambda s: None)


# -- Exchange client ------------------------------------------------------------


def test_exchange_refuses_cmdlets_outside_the_read_only_list():
    for cmdlet in ("Set-Mailbox", "New-InboxRule", "Remove-Mailbox", "Get-MessageTrace"):
        with pytest.raises(ValueError):
            client().invoke(cmdlet)


@responses.activate
def test_exchange_pages_and_sends_cmdlet():
    responses.post(URL, json={"value": [{"n": 1}], "@odata.nextLink": URL + "?$skiptoken=2"})
    responses.post(URL + "?$skiptoken=2", json={"value": [{"n": 2}]})
    assert client().invoke("Get-Mailbox", {"ResultSize": "Unlimited"}) == [{"n": 1}, {"n": 2}]
    body = json.loads(responses.calls[0].request.body)
    assert body == {"CmdletInput": {"CmdletName": "Get-Mailbox", "Parameters": {"ResultSize": "Unlimited"}}}
    assert responses.calls[0].request.headers["Authorization"] == "Bearer tok"


@responses.activate
def test_exchange_cuts_an_endless_answer_at_the_page_cap_instead_of_failing():
    for n in (1, 2, 3):
        responses.post(
            URL + (f"?$skiptoken={n}" if n > 1 else ""),
            json={"value": [{"n": n}], "@odata.nextLink": URL + f"?$skiptoken={n + 1}"},
        )
    exo = client()
    assert exo.invoke("Get-Mailbox", max_pages=2) == [{"n": 1}, {"n": 2}]
    assert exo.truncated == {"Get-Mailbox": 2} and len(responses.calls) == 2


def test_a_cut_exchange_answer_makes_the_run_partial_and_says_so(db, two_tenants, monkeypatch):
    alpha, _ = two_tenants
    import re
    from unittest.mock import patch

    from officesentry import exchange
    from officesentry.graph import GraphClient
    from officesentry.worker import execute_run
    from tests.seed import exchange_factory, mock_tenant

    monkeypatch.setattr(exchange, "MAX_PAGES", 2)
    with (
        responses.RequestsMock(assert_all_requests_are_fired=False) as rsps,
        patch("officesentry.dnscheck.lookup", gd.fake_dns),
    ):
        mock_tenant(rsps)

        def endless(request):  # every cmdlet's answer goes on for ever
            page = int(request.url.rsplit("=", 1)[-1] if "skiptoken" in request.url else "0") + 1
            link = f"{EXCHANGE_ROOT}/adminapi/beta/tid/InvokeCommand?$skiptoken={page}"
            return 200, {}, json.dumps({"value": [], "@odata.nextLink": link})

        rsps.add_callback(responses.POST, re.compile(re.escape(EXCHANGE_ROOT) + r"/.*"), callback=endless)
        store.enqueue_run(db, alpha, "mailboxes")
        run = store.claim_next_run(db)
        status = execute_run(db, run, lambda t, c: GraphClient(lambda: "t", sleep=lambda s: None), exchange_factory)
    row = db.execute("SELECT error, truncated FROM runs WHERE id = ?", (run["id"],)).fetchone()
    assert status == "partial", row["error"]
    assert row["error"].startswith("Get-Mailbox returned more than 2 pages of results; the rest were left out; ")
    assert row["truncated"] == 1


@responses.activate
def test_exchange_refuses_next_link_to_another_host():
    responses.post(URL, json={"value": [], "@odata.nextLink": "https://evil.example/adminapi/x"})
    with pytest.raises(ValueError):
        client().invoke("Get-Mailbox")


@responses.activate
def test_exchange_retries_throttling_then_explains_access_denied():
    responses.post(URL, status=429, headers={"Retry-After": "1"})
    responses.post(URL, status=401, body="")
    with pytest.raises(ExchangeAccessDenied) as exc:
        client().invoke("Get-Mailbox")
    assert "Global Reader" in str(exc.value)


@responses.activate
def test_exchange_reports_cmdlet_errors():
    responses.post(URL, status=400, json={"error": {"code": "BadRequest", "details": [{"message": "not found"}]}})
    with pytest.raises(ExchangeError, match="not found"):
        client().invoke("Get-MailboxPermission", {"Identity": "x"})


# -- Mail-flow helpers ------------------------------------------------------------


def test_smtp_address_and_external():
    assert smtp_address("SMTP:Ann@Gmail.com") == "ann@gmail.com"
    assert smtp_address("not an address") is None
    assert is_external("a@gmail.com", {"x.com"})
    assert not is_external("a@sub.x.com", {"x.com"})


def test_inbox_rule_risk():
    hide = gd.INBOX_RULES["u1"][0] | {"moveToFolderName": "RSS Feeds"}
    risk = inbox_rule_risk(hide, {"x.com"})
    assert risk.severity == "high" and any("invoice" in r for r in risk.reasons)
    assert any("blank" in r for r in risk.reasons)
    internal = inbox_rule_risk(gd.INBOX_RULES["u7"][0], {"x.com"})
    assert internal.severity is None and internal.forwards_to == ["boss@x.com"]
    disabled = inbox_rule_risk(dict(gd.INBOX_RULES["u1"][1], isEnabled=False), {"x.com"})
    assert disabled.severity == "low"


def test_inbox_rule_moving_mail_to_an_unnamed_folder_is_raised():
    """Without Exchange the folder has no name, so a rule hiding invoices can't be shown to be harmless."""
    rule = {
        "id": "r1",
        "displayName": "Tidy",
        "isEnabled": True,
        "conditions": {"subjectContains": ["invoice"]},
        "actions": {"moveToFolder": "AAMkAD1", "markAsRead": True},
        "moveToFolderName": None,
    }
    risk = inbox_rule_risk(rule, {"x.com"})
    assert risk.severity == "medium"
    assert risk.reasons == [
        "Moves mail mentioning invoice to a folder Office Sentry couldn't name and marks it as read"
    ]
    named = inbox_rule_risk(rule | {"moveToFolderName": "Invoices"}, {"x.com"})
    assert named.severity is None  # a named, ordinary folder is a tidy-up rule
    plain = inbox_rule_risk(rule | {"conditions": {}, "actions": {"moveToFolder": "AAMkAD1"}}, {"x.com"})
    assert plain.severity is None


# -- Collectors and checks end to end ----------------------------------------------


@pytest.fixture
def seeded(db, two_tenants):
    alpha, _ = two_tenants
    seed_tenant(db, alpha, users=USERS + [gd.INFO_USER])
    tenant = store.get_tenant(db, alpha)
    return TenantSnapshot(db, tenant), {r.key: r for r in evaluate(TenantSnapshot(db, tenant))}


def test_collectors_store_phase4_data(seeded):
    snap, _ = seeded
    for collector in ("mailboxes", "inbox_rules", "app_consents", "tenant_policies"):
        assert snap.runs[collector]["status"] == "succeeded", collector
    perms = snap.objects("mailboxes", "mailbox_permission")
    assert [(p["mailbox"], p["User"]) for p in perms] == [("info@x.com", "ann@x.com")]
    assert [p["Trustee"] for p in snap.objects("mailboxes", "send_as")] == ["ann@x.com"]
    rules = {r["displayName"]: r for r in snap.objects("inbox_rules", "inbox_rule")}
    assert rules["."]["moveToFolderName"] == "RSS Feeds"
    perm = snap.objects("app_consents", "app_permission")
    assert {(p["clientId"], p["permission"]) for p in perm} == {("sp-evil", "Mail.Read"), ("sp-zoom", "User.Read.All")}
    groups = snap.objects("conditional_access", "ca_group")
    assert groups == []  # the default policies name no groups


def test_inbox_rule_folders_come_from_exchange(db, two_tenants):
    # Graph refuses to name a folder without mail access, so Get-InboxRule does. Graph's rule ID is Exchange's
    # RuleIdentity in base64 with - for / (as on the real test tenant), so it's found even when renamed.
    alpha, beta = two_tenants
    rule = {
        "id": "AQAAAIPT-P8=",
        "displayName": "Tidy",
        "isEnabled": True,
        "sequence": 1,
        "conditions": {"subjectContains": ["invoice"]},
        "actions": {"moveToFolder": "AAMkAD1", "markAsRead": True},
    }
    moved = {
        "Name": "Tidy (renamed)",
        "RuleIdentity": 18445850733396885505,
        "MoveToFolder": "ann@x.com:\\Deleted Items",
    }
    exchange = dict(gd.EXCHANGE, **{"Get-InboxRule": {"u1": [moved]}})
    seed_tenant(db, alpha, inbox_rules={"u1": [rule]}, exchange=exchange)
    snap = TenantSnapshot(db, store.get_tenant(db, alpha))
    assert [r["moveToFolderName"] for r in snap.objects("inbox_rules", "inbox_rule")] == ["Deleted Items"]
    assert {r.key: r for r in evaluate(snap)}["suspicious_inbox_rules"].status == "fail"

    # Without an Exchange role the rules are still read, just without folder names.
    refused = dict(gd.EXCHANGE, **{"Get-InboxRule": (403, {"error": {"code": "Forbidden", "message": "No"}})})
    seed_tenant(db, beta, inbox_rules={"u1": [rule]}, exchange=refused)
    snap = TenantSnapshot(db, store.get_tenant(db, beta))
    assert snap.runs["inbox_rules"]["status"] == "succeeded"
    assert [r["moveToFolderName"] for r in snap.objects("inbox_rules", "inbox_rule")] == [None]


def test_email_checks(seeded):
    _, r = seeded
    fwd = r["external_forwarding"]
    assert fwd.status == "fail"
    assert {f.object_id for f in fwd.findings} == {"mailbox:u1", "rule:u1:r2"}
    assert r["suspicious_inbox_rules"].status == "fail"
    assert [f.object_id for f in r["suspicious_inbox_rules"].findings] == ["rule:u1:r1"]
    assert r["auto_forward_policy"].status == "pass"
    assert r["shared_mailbox_signin"].status == "fail"
    assert r["shared_mailbox_signin"].findings[0].subject == "info@x.com"
    assert r["mailbox_auditing"].status == "pass"
    assert r["smtp_auth"].status == "fail"


def test_app_and_sharing_checks(seeded):
    _, r = seeded
    apps = r["risky_app_permissions"]
    assert apps.status == "fail"
    assert [(f.subject, f.severity) for f in apps.findings] == [("Mail Sync Pro", "high")]
    assert r["user_consent"].status == "fail"
    creds = r["app_credentials"]
    assert [(f.subject, f.severity) for f in creds.findings] == [("Old connector", "low")]
    assert r["sharepoint_sharing"].status == "fail"
    assert r["guest_invites"].status == "fail"


def test_admin_mfa_coverage_flags_excluded_breakglass(seeded):
    _, r = seeded
    result = r["admins_ca_mfa"]
    assert result.status == "fail"
    assert [f.object_id for f in result.findings] == ["u4"]
    assert "excluded from Require MFA for all users" in result.findings[0].detail


def test_no_exchange_licence_makes_mail_checks_unavailable(db, two_tenants):
    alpha, _ = two_tenants
    skus = [
        dict(gd_sku, servicePlans=[p for p in gd_sku["servicePlans"] if "EXCHANGE" not in p["servicePlanName"]])
        for gd_sku in __import__("tests.test_collectors", fromlist=["SKUS"]).SKUS
    ]
    seed_tenant(db, alpha, skus=skus)
    r = {x.key: x for x in evaluate(TenantSnapshot(db, store.get_tenant(db, alpha)))}
    for key in ("external_forwarding", "shared_mailbox_signin", "mailbox_auditing", "smtp_auth"):
        assert r[key].status == "unavailable", key


def test_exchange_access_denied_fails_only_that_run(db, two_tenants):
    alpha, _ = two_tenants
    import re
    from unittest.mock import patch

    from officesentry.collectors import REGISTRY
    from officesentry.graph import GraphClient
    from officesentry.worker import execute_run
    from tests.seed import exchange_factory, mock_tenant

    with (
        responses.RequestsMock(assert_all_requests_are_fired=False) as rsps,
        patch("officesentry.dnscheck.lookup", gd.fake_dns),
    ):
        mock_tenant(rsps, exchange=False)
        rsps.post(re.compile(re.escape(EXCHANGE_ROOT) + r"/.*"), status=401)
        statuses = {}
        for key in REGISTRY:
            store.enqueue_run(db, alpha, key)
        while (run := store.claim_next_run(db)) is not None:
            statuses[run["collector"]] = execute_run(
                db, run, lambda t, c: GraphClient(lambda: "t", sleep=lambda s: None), exchange_factory
            )
    assert statuses["mailboxes"] == "failed"
    assert statuses["inbox_rules"] == "succeeded"  # falls back to licensed users via Graph
    run = (
        store.latest_run(db, alpha, "mailboxes")
        or db.execute("SELECT * FROM runs WHERE collector = 'mailboxes' ORDER BY id DESC").fetchone()
    )
    assert "Global Reader" in run["error"]


def test_inbox_rules_fallback_skips_accounts_without_a_mailbox(db, two_tenants):
    """Without Exchange access every licensed account is tried; one with no Exchange Online mailbox has no
    rules to read, and mustn't make the run partial every night."""
    alpha, _ = two_tenants
    import re
    from unittest.mock import patch

    from officesentry.collectors import REGISTRY
    from officesentry.graph import GraphClient
    from officesentry.worker import execute_run
    from tests.seed import V1, exchange_factory, mock_tenant

    with (
        responses.RequestsMock(assert_all_requests_are_fired=False) as rsps,
        patch("officesentry.dnscheck.lookup", gd.fake_dns),
    ):
        rsps.get(
            f"{V1}/users/u1/mailFolders/inbox/messageRules",
            status=404,
            json={
                "error": {
                    "code": "MailboxNotEnabledForRESTAPI",
                    "message": "The mailbox is either inactive, soft-deleted, or is hosted on-premise.",
                }
            },
        )
        mock_tenant(rsps, exchange=False)
        rsps.post(re.compile(re.escape(EXCHANGE_ROOT) + r"/.*"), status=401)
        for key in REGISTRY:
            store.enqueue_run(db, alpha, key)
        while (run := store.claim_next_run(db)) is not None:
            execute_run(db, run, lambda t, c: GraphClient(lambda: "t", sleep=lambda s: None), exchange_factory)
    run = store.latest_run(db, alpha, "inbox_rules")
    assert run["status"] == "succeeded" and not run["error"]
    logs = [r["message"] for r in db.execute("SELECT message FROM run_events WHERE run_id = ?", (run["id"],))]
    assert any("1 account had no Exchange Online mailbox" in m for m in logs)


# -- Conditional Access coverage ----------------------------------------------------


def _directory(**groups):
    return coverage.Directory(
        groups={k: {"displayName": k, "members": v} for k, v in groups.items()},
        user_roles={"admin": {"role-ga"}},
        role_names={"role-ga": "Global Administrator"},
        names={},
    )


def _policy(**users):
    return {
        "id": "p",
        "displayName": "P",
        "state": "enabled",
        "conditions": {"users": users, "applications": {"includeApplications": ["All"]}, "clientAppTypes": ["all"]},
        "grantControls": {"builtInControls": ["mfa"]},
    }


def test_coverage_groups_roles_and_unknown_membership():
    d = _directory(staff=["u1"], unknown=None)
    assert coverage.applies(_policy(includeGroups=["staff"]), "u1", "Member", d)
    assert not coverage.applies(_policy(includeGroups=["staff"]), "u2", "Member", d)
    assert coverage.applies(_policy(includeRoles=["role-ga"]), "admin", "Member", d)
    assert not coverage.applies(_policy(includeUsers=["All"], excludeGroups=["staff"]), "u1", "Member", d)
    # Unreadable group: never counts as included, always counts as excluded.
    assert not coverage.applies(_policy(includeGroups=["unknown"]), "u1", "Member", d)
    assert not coverage.applies(_policy(includeUsers=["All"], excludeGroups=["unknown"]), "u1", "Member", d)
    assert coverage.applies(_policy(includeUsers=["GuestsOrExternalUsers"]), "g", "Guest", d)


def test_describe_policy_in_plain_words():
    d = _directory()
    row = coverage.describe(gd.CA_BLOCK_LEGACY, d)
    assert (row.state, row.users, row.apps, row.grant) == ("On", "All users", "All apps", "Block")
    assert row.conditions == "Clients: ActiveSync, Other legacy"


def test_user_coverage_lists_gaps_first(seeded):
    snap, _ = seeded
    cov = coverage.user_coverage(snap)
    assert [c.mfa for c in cov] == sorted(c.mfa for c in cov)
    by_id = {c.id: c for c in cov}
    assert by_id["u1"].mfa and not by_id["u1"].legacy_blocked
    assert "u2" not in by_id  # disabled accounts are left out


def test_mfa_exclusion_names_the_group(db, two_tenants):
    alpha, _ = two_tenants
    policy = dict(
        gd.CA_MFA_ALL,
        conditions=dict(
            gd.CA_MFA_ALL["conditions"],
            users={"includeUsers": ["All"], "excludeUsers": [], "excludeGroups": ["grp-bg"]},
        ),
    )
    seed_tenant(db, alpha, ca=[policy])
    r = {x.key: x for x in evaluate(TenantSnapshot(db, store.get_tenant(db, alpha)))}
    assert [f.subject for f in r["mfa_enforced"].findings] == ["Break glass"]
    assert [f.object_id for f in r["admins_ca_mfa"].findings] == ["u4"]  # u4 is in the excluded group


def test_expiring_credentials(db, two_tenants):
    from officesentry.rules import app_credentials

    class Snap:
        def objects(self, collector, kind):
            return [{"_id": "a", "displayName": "Soon", "credentials": [{"endDateTime": "2026-10-20T00:00:00Z"}]}]

    status, _, found = app_credentials(Snap(), datetime(2026, 10, 2, tzinfo=UTC))
    assert status == "fail" and found[0].severity == "medium"


# -- Detail reports ------------------------------------------------------------------


def test_detail_reports_build(db, two_tenants):
    from officesentry.reports.data_reports import REPORTS, build

    alpha, _ = two_tenants
    seed_tenant(db, alpha, users=USERS + [gd.INFO_USER])
    tenant = store.get_tenant(db, alpha)
    mail = {s.title: s for s in build(db, tenant, "mailboxes").sections}
    assert [(r[0], r[1], r[3]) for r in mail["Forwarding"].rows] == [
        ("ann@x.com", "ann@gmail.com", "Yes"),
        ("ann@x.com", "ann.home@gmail.com", "Yes"),
        ("info@x.com", "boss@x.com", "No"),
    ]
    assert mail["Forwarding"].flagged == 2
    assert [r[2] for r in mail["What stops forwarding"].rows] == ["Blocked (Microsoft's default)", "Allowed"]
    apps = {s.title: s for s in build(db, tenant, "apps").sections}
    third = apps["Third-party apps with access"]
    assert third.rows[0][0] == "Mail Sync Pro" and third.flags[0] == "high"
    # Every Risk value is a level, with what earns it after the colon.
    assert third.rows[0][5] == "High: Mail.Read" and {r[5] for r in third.rows[1:]} <= {"Low"}
    ca = {s.title: s for s in build(db, tenant, "conditional-access").sections}
    assert ca["Policies"].rows[0][:2] == ["Require MFA for all users", "On"]
    assert ca["Coverage by account"].intro.startswith("3 of 3")
    guests = {s.title: s for s in build(db, tenant, "guests").sections}
    assert ["Who can invite guests", "Everyone, including guests"] in guests["External access settings"].rows
    for key in REPORTS:
        assert build(db, store.get_tenant(db, two_tenants[1]), key).unavailable  # never collected


def test_detail_report_pages_and_access(client, db, two_tenants):
    from openpyxl import load_workbook

    from tests.conftest import login, make_user

    alpha, bravo = two_tenants
    seed_tenant(db, alpha)
    make_user(db, "client", role="client", tenants=[bravo])
    login(client, "client")
    assert client.get(f"/tenants/{alpha}/data/mailboxes").status_code == 404
    client.post("/logout")
    make_user(db, "admin")
    login(client, "admin")
    for key in ("mailboxes", "apps", "conditional-access", "guests"):
        page = client.get(f"/tenants/{alpha}/data/{key}")
        assert page.status_code == 200, key
    assert "Mail Sync Pro" in client.get(f"/tenants/{alpha}/data/apps").get_data(as_text=True)
    assert client.get(f"/tenants/{alpha}/data/nope").status_code == 404
    resp = client.get(f"/tenants/{alpha}/data/mailboxes.xlsx")
    assert resp.status_code == 200
    from io import BytesIO

    wb = load_workbook(BytesIO(resp.data))
    assert "Forwarding" in wb.sheetnames


def test_changes_include_new_forwarding_and_app_permissions(db, two_tenants):
    from datetime import timedelta

    from officesentry.changes import all_changes

    alpha, _ = two_tenants
    quiet = dict(
        gd.EXCHANGE, **{"Get-Mailbox": [dict(m, ForwardingSmtpAddress=None) for m in gd.EXCHANGE["Get-Mailbox"]]}
    )
    seed_tenant(db, alpha, exchange=quiet)
    old = (datetime.now(UTC) - timedelta(days=30)).strftime("%Y-%m-%dT%H:%M:%SZ")
    unpack(db)
    db.execute("DELETE FROM objects WHERE kind = 'app_permission' AND object_id = 'ara1'")
    db.execute("UPDATE runs SET finished_at = ?", (old,))
    seed_tenant(db, alpha)
    since = (datetime.now(UTC) - timedelta(days=28)).strftime("%Y-%m-%dT%H:%M:%SZ")
    changes = {(c.area, c.subject, c.detail) for c in all_changes(db, alpha, since)}
    assert ("Mail", "ann@x.com", "Now forwards to ann@gmail.com") in changes
    assert ("Apps", "Mail Sync Pro", "Granted Mail.Read") in changes


def test_exchange_allow_list_is_get_only():
    from officesentry.entra import EXCHANGE_PERMISSIONS, required_resource_access
    from officesentry.exchange import READ_ONLY_CMDLETS

    assert all(c.startswith("Get-") for c in READ_ONLY_CMDLETS)
    assert set(EXCHANGE_PERMISSIONS) == {"Exchange.ManageAsApp"}
    graph = {
        "appRoles": [
            {"value": p, "id": p, "allowedMemberTypes": ["Application"]}
            for p in __import__("officesentry.entra", fromlist=["x"]).GRAPH_PERMISSIONS
        ]
    }
    exo = {"appRoles": [{"value": "Exchange.ManageAsApp", "id": "exo-id", "allowedMemberTypes": ["Application"]}]}
    access = required_resource_access(graph, exo)
    assert access[1] == {
        "resourceAppId": "00000002-0000-0ff1-ce00-000000000000",
        "resourceAccess": [{"id": "exo-id", "type": "Role"}],
    }


def test_live_test_command_runs_every_collector(cfg, db, monkeypatch, capsys):
    from unittest.mock import patch

    from officesentry import livetest
    from tests.seed import mock_tenant

    for k in livetest.ENV:  # a developer machine may have real test-tenant variables set
        monkeypatch.delenv(k, raising=False)
    assert livetest.run(cfg) == 2  # variables missing
    for k in livetest.ENV:
        monkeypatch.setenv(k, "tid" if k.endswith("TENANT_ID") else "x")
    monkeypatch.setattr(livetest, "secret_token_provider", lambda *a: lambda: "tok")
    with (
        responses.RequestsMock(assert_all_requests_are_fired=False) as rsps,
        patch("officesentry.dnscheck.lookup", gd.fake_dns),
    ):
        mock_tenant(rsps)
        assert livetest.run(cfg) == 0
    out = capsys.readouterr().out
    assert "Mailboxes" in out and "0 collectors failed" in out


def test_live_test_runs_only_the_collectors_asked_for_on_a_new_tenant(cfg, db, monkeypatch, capsys):
    # A first success on an unconsented tenant queues every other collector; the live test mustn't run those.
    from unittest.mock import patch

    from officesentry import livetest
    from tests.seed import mock_tenant

    for k in livetest.ENV:
        monkeypatch.setenv(k, "tid" if k.endswith("TENANT_ID") else "x")
    monkeypatch.setattr(livetest, "secret_token_provider", lambda *a: lambda: "tok")
    with (
        responses.RequestsMock(assert_all_requests_are_fired=False) as rsps,
        patch("officesentry.dnscheck.lookup", gd.fake_dns),
    ):
        mock_tenant(rsps)
        assert livetest.run(cfg, ["users", "groups"]) == 0
    ran = [r["collector"] for r in db.execute("SELECT collector FROM runs ORDER BY id")]
    assert ran == ["users", "groups"]
    assert db.execute("SELECT consented_at FROM tenants").fetchone()[0]
    assert "Groups and Teams" in capsys.readouterr().out


@responses.activate
def test_live_test_can_look_dns_up_over_https():
    # For machines whose DNS is blocked: the same answers dnscheck.lookup gives, from DNS-over-HTTPS JSON.
    import pytest

    from officesentry import dnscheck
    from officesentry.livetest import DOH_URL, doh_lookup

    def answer(rdtype, *data, status=0):
        responses.get(
            DOH_URL,
            match=[responses.matchers.query_param_matcher({"name": "x.example", "type": rdtype})],
            json={
                "Status": status,
                "Answer": [
                    {"name": "x.example.", "type": {"TXT": 16, "MX": 15, "CNAME": 5}[rdtype], "data": d} for d in data
                ],
            },
        )

    answer("MX", "20 backup.example.", "10 X-Example.mail.protection.outlook.com.")
    answer("TXT", '"v=spf1 include:spf.protection.outlook.com" " -all"', '"say \\"hi\\""')
    answer("CNAME", "selector1-x._domainkey.contoso.onmicrosoft.com.")
    assert doh_lookup("x.example", "MX") == ["x-example.mail.protection.outlook.com", "backup.example"]
    assert doh_lookup("x.example", "TXT") == ["v=spf1 include:spf.protection.outlook.com -all", 'say "hi"']
    assert doh_lookup("x.example", "CNAME") == ["selector1-x._domainkey.contoso.onmicrosoft.com"]
    responses.reset()
    answer("TXT", status=3)  # NXDOMAIN
    assert doh_lookup("x.example", "TXT") == []
    responses.reset()
    answer("TXT", status=2)  # SERVFAIL is an error, not "no records"
    with pytest.raises(dnscheck.DnsError):
        doh_lookup("x.example", "TXT")


def test_shared_mailbox_that_needs_its_licence_is_not_waste(db, two_tenants):
    """A shared mailbox should be blocked from sign-in, and needs a licence once it has an archive, a hold or
    50 GB; flagging it as a disabled account wasting a licence would tell the MSP to break it."""
    alpha, beta = two_tenants
    shared = gd.INFO_USER | {"accountEnabled": False, "assignedLicenses": USERS[0]["assignedLicenses"]}
    boxes = [dict(m) for m in gd.EXCHANGE["Get-Mailbox"]]
    boxes[1]["LitigationHoldEnabled"] = True
    seed_tenant(db, alpha, users=USERS + [shared], exchange=dict(gd.EXCHANGE, **{"Get-Mailbox": boxes}))
    seed_tenant(db, beta, users=USERS + [shared])
    held = {r.key: r for r in evaluate(TenantSnapshot(db, store.get_tenant(db, alpha)))}["licensed_disabled"]
    assert "info@x.com" not in [f.subject for f in held.findings]
    assert "(not counting 1 shared mailbox that needs one)" in held.summary
    plain = {r.key: r for r in evaluate(TenantSnapshot(db, store.get_tenant(db, beta)))}["licensed_disabled"]
    assert "info@x.com" in [f.subject for f in plain.findings]  # no hold, archive or size: licence not needed

    # The reports agree with the check: Licence waste and Inactive accounts leave it out, and say so.
    def sections(tenant, key):
        report = build(db, store.get_tenant(db, tenant), key, datetime.now(UTC))
        return report, {s.title: s for s in report.sections}

    waste, kept = sections(alpha, "licences")
    assert "info@x.com" not in [r[0] for r in kept["Accounts to review"].rows]
    assert any("1 shared mailbox holds a licence it needs" in n for n in waste.notes)
    _, counted = sections(beta, "licences")
    assert "info@x.com" in [r[0] for r in counted["Accounts to review"].rows]
    on_disabled = [sum(r[4] for r in x["Waste by subscription"].rows) for x in (kept, counted)]
    assert on_disabled[1] == on_disabled[0] + 1
    _, inactive = sections(alpha, "inactive")
    disabled = inactive["Disabled accounts still holding licences or roles"]
    assert "info@x.com" not in [r[0] for r in disabled.rows]
    assert "so it is not listed" in disabled.intro
    _, inactive = sections(beta, "inactive")
    assert "info@x.com" in [r[0] for r in inactive["Disabled accounts still holding licences or roles"].rows]


def test_shared_mailbox_over_50_gb_is_not_waste(db, two_tenants):
    alpha, _ = two_tenants
    shared = gd.INFO_USER | {"accountEnabled": False, "assignedLicenses": USERS[0]["assignedLicenses"]}
    size = [{"TotalItemSize": "51 GB (54,760,833,024 bytes)", "TotalDeletedItemSize": "0 B (0 bytes)", "ItemCount": 9}]
    exchange = dict(gd.EXCHANGE, **{"Get-MailboxStatistics": {"g-u7": size}})
    seed_tenant(db, alpha, users=USERS + [shared], exchange=exchange)
    result = {r.key: r for r in evaluate(TenantSnapshot(db, store.get_tenant(db, alpha)))}["licensed_disabled"]
    assert "info@x.com" not in [f.subject for f in result.findings]
    assert "(not counting 1 shared mailbox that needs one)" in result.summary
