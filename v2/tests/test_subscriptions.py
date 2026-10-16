"""Licences and subscriptions: renewals, trials, how licences are assigned, group licensing errors."""

from datetime import UTC, datetime, timedelta
from io import BytesIO

from openpyxl import load_workbook

from officesentry import store
from officesentry.reports import data_reports, library, user_profile, views
from officesentry.rules import evaluate
from officesentry.snapshot import TenantSnapshot
from tests import groupdata as grp
from tests import licencedata as ld
from tests.conftest import login, make_user
from tests.seed import seed_tenant
from tests.stored import unpack


def _licensing(db, tenant_id, **kwargs):
    seed_tenant(
        db,
        tenant_id,
        **{
            "users": ld.USERS,
            "skus": ld.SKUS,
            "subscriptions": ld.TENANT_SUBSCRIPTIONS,
            "groups": grp.GROUPS + ld.LICENSING_GROUPS,
            **kwargs,
        },
    )
    return store.get_tenant(db, tenant_id)


def _report(db, tenant):
    return data_reports.build(db, tenant, "subscriptions")


def _table(report, title):
    section = next(s for s in report.sections if s.title == title)
    return [[str(c) for c in row] for row in section.rows], section.flags


def _results(db, tenant, now=None):
    return {r.key: r for r in evaluate(TenantSnapshot(db, tenant), now)}


def _price(db, part, price):
    db.execute(
        "INSERT INTO sku_prices (sku_part_number, monthly_price, updated_at) VALUES (?, ?, ?)",
        (part, price, "2026-10-01T00:00:00Z"),
    )


def test_products_with_cost_and_what_needs_action_first(db, two_tenants):
    alpha, _ = two_tenants
    tenant = _licensing(db, alpha)
    _price(db, "SPB", 18.10)
    _price(db, "Microsoft_365_Copilot", 25.0)
    report = _report(db, tenant)
    assert report.unavailable is None
    rows, flags = _table(report, "Subscriptions")
    by_name = {r[0]: (r, f) for r, f in zip(rows, flags, strict=True)}
    # Lapsing and trials first, free products last.
    assert [r[0] for r in rows][-2:] == ["Microsoft 365 Business Premium", "Power Automate Free"]
    visio, flag = by_name["Visio Plan 2"]
    assert (visio[1], visio[5], visio[4], flag) == ("Grace period", "3", "0", "high")
    exo, flag = by_name["Exchange Online (Plan 1)"]
    assert (exo[1], exo[3], exo[4], flag) == ("Suspended", "1", "-1", "high")  # one person, no working licence
    copilot, flag = by_name["Microsoft 365 Copilot"]
    assert (copilot[1], copilot[10], flag) == ("Trial", "Yes", "medium")
    bp, _ = by_name["Microsoft 365 Business Premium"]
    assert bp[2:5] == ["10", "10", "0"] and "181" in bp[8]  # 10 x 18.10 a month
    assert by_name["Power Automate Free"][0][7] == "Free"
    figures = {h.label: h for h in report.highlights}
    assert float(figures["Monthly cost"].value) == round(10 * 18.10, 2)  # the Copilot trial costs nothing yet
    assert figures["Monthly cost"].foot == "Licences bought, at your prices"  # lapsed Visio isn't bought any more
    # Suspended Exchange licences aren't in use, and the Copilot trial's aren't paid licences.
    assert figures["Licences in use"].value == "13 of 13"


def test_a_trial_beside_paid_licences_isnt_the_renewal_date():
    from officesentry.reports.subscriptions import _next_dates

    now = datetime(2026, 10, 15, tzinfo=UTC)
    subs = [
        {"skuId": "sku-bp", "status": "Enabled", "isTrial": False, "nextLifecycleDateTime": "2027-03-01T00:00:00Z"},
        {"skuId": "sku-bp", "status": "Enabled", "isTrial": True, "nextLifecycleDateTime": "2026-11-01T00:00:00Z"},
        {"skuId": "sku-cp", "status": "Enabled", "isTrial": True, "nextLifecycleDateTime": "2026-11-05T00:00:00Z"},
    ]
    dates = {sid: f"{when:%Y-%m-%d}" for sid, when in _next_dates(subs, now).items()}
    assert dates == {"sku-bp": "2027-03-01", "sku-cp": "2026-11-05"}


def test_renewals_list_trials_grace_and_suspended_first(db, two_tenants):
    alpha, _ = two_tenants
    report = _report(db, _licensing(db, alpha))
    rows, flags = _table(report, "Renewals and expiry")
    assert [(r[0], r[2], r[6]) for r in rows] == [
        ("Microsoft 365 Copilot", "Trial", "30 days"),
        ("Visio Plan 2", "Grace period", "30 days"),
        ("Exchange Online (Plan 1)", "Suspended", "90 days"),
        ("Microsoft 365 Business Premium", "Active", "Later"),
        ("Power Automate Free", "Active", ""),
    ]
    assert flags[:4] == ["medium", "high", "high", None]
    assert rows[0][1] == "c-s-copilot" and "trial ends" in rows[0][7]
    assert rows[4][7] == "Free: nothing to renew."
    figures = {h.label: h for h in report.highlights}
    assert figures["Lapsing within 30 days"].value == 2 and figures["Lapsing within 30 days"].tone == "bad"
    assert figures["Trials"].value == 1


def test_who_has_which_licence_and_how(db, two_tenants):
    alpha, _ = two_tenants
    report = _report(db, _licensing(db, alpha))
    rows, flags = _table(report, "Who has which licence")
    who = {(r[0], r[1]): (r, f) for r, f in zip(rows, flags, strict=True)}
    ann, flag = who[("ann@x.com", "Microsoft 365 Business Premium")]
    assert (ann[2], ann[3], ann[4], flag) == ("Directly and by group", "Licence - Business Premium", "Active", "low")
    dan, flag = who[("dan@x.com", "Microsoft 365 Business Premium")]
    assert (dan[2], dan[4], flag) == ("By group", "Not assigned (error)", "medium")
    assert who[("dan@x.com", "Exchange Online (Plan 1)")][0][4] == "Subscription suspended"
    assert who[("eve@x.com", "Visio Plan 2")][0][5] == ""  # no usage location
    figures = {h.label: h for h in report.highlights}
    assert figures["Assigned by group"].value == "4 of 9"

    problems, flags = _table(report, "Assignment problems")
    kinds = [(r[0], r[2]) for r in problems]
    assert kinds == [
        ("dan@x.com", "Licence error"),
        ("fay@x.com", "Licence error"),
        ("eve@x.com", "No usage location"),
        ("ann@x.com", "Directly and by group"),
    ]
    assert "Not enough licences" in problems[0][4] and problems[0][3] == "Licence - Business Premium"
    assert "base licence" in problems[1][4]
    assert flags == ["medium", "medium", "low", "low"]
    assert figures["Licence errors"].value == 2 and figures["Licence errors"].tone == "bad"

    groups, flags = _table(report, "Groups that assign licences")
    by_group = {r[0]: r for r in groups}
    assert by_group["Licence - Business Premium"][1:4] == ["Microsoft 365 Business Premium", "2", "1"]
    assert by_group["Licence - Copilot"][2:4] == ["0", "1"] and by_group["Licence - Copilot"][5] == "Dynamic"


def test_headline_figures_open_their_rows(client, db, two_tenants):
    alpha, _ = two_tenants
    tenant = _licensing(db, alpha)
    make_user(db, "admin")
    login(client, "admin")
    path = f"/tenants/{alpha}/data/subscriptions"
    report = _report(db, tenant)
    links = views.headline_links(report, path)["highlights"]
    assert all(links)
    errors = client.get(links[4].split("#")[0]).get_data(as_text=True)
    problems = errors.split('id="t-assignment-problems"')[1].split("<tbody")[1].split("</tbody>")[0]
    assert "dan@x.com" in problems and "eve@x.com" not in problems  # Eve's missing usage location is filtered out
    page = client.get(path)
    assert page.status_code == 200
    html = page.get_data(as_text=True)
    assert "Licences and subscriptions" in html and "Not enough licences" in html
    assert f'href="/tenants/{alpha}/users/u8"' in html  # account names open their page
    pdf = client.get(f"{path}.pdf")
    assert pdf.status_code == 200 and pdf.data[:4] in (b"%PDF", b"<!do", b"<!DO")
    sheets = load_workbook(BytesIO(client.get(f"{path}.xlsx").data)).sheetnames
    assert {"Subscriptions", "Renewals and expiry", "Who has which licence", "Assignment problems"} <= set(sheets)
    # Filtering by licence, as a technician would.
    view = client.get(f"{path}?f=who-has-which-licence/licence:Microsoft+365+Copilot").get_data(as_text=True)
    rows = view.split('id="t-who-has-which-licence"')[1].split("<tbody")[1].split("</tbody>")[0]
    assert "fay@x.com" in rows and "ann@x.com" in rows and "eve@x.com" not in rows


def test_default_seed_passes_and_shows_group_licensing(db, two_tenants):
    alpha, _ = two_tenants
    seed_tenant(db, alpha)
    tenant = store.get_tenant(db, alpha)
    results = _results(db, tenant)
    assert results["licence_assignment_errors"].status == "pass"
    assert results["subscriptions_lapsing"].status == "pass"  # the suspended product has nobody on it
    report = _report(db, tenant)
    rows, _ = _table(report, "Groups that assign licences")
    assert rows == [["All Staff", "Microsoft 365 Business Premium", "2", "0", "", "Dynamic"]]
    figures = {h.label: h.value for h in report.highlights}
    assert figures["Lapsing within 30 days"] == 0 and figures["Licence errors"] == 0


def test_checks_find_errors_and_lapsing_subscriptions(db, two_tenants):
    alpha, _ = two_tenants
    results = _results(db, _licensing(db, alpha))
    errors = results["licence_assignment_errors"]
    assert errors.status == "fail" and errors.summary == "2 licence assignments failed for 2 accounts"
    dan = next(f for f in errors.findings if f.subject == "dan@x.com")
    assert dan.severity == "medium" and dan.object_id == "u8"
    assert dan.detail == (
        "Microsoft 365 Business Premium from Licence - Business Premium wasn't assigned. Not enough "
        "licences: every one bought is already assigned."
    )
    lapsing = results["subscriptions_lapsing"]
    assert lapsing.status == "fail" and lapsing.summary == "3 subscriptions about to lapse or already lapsed"
    found = {f.subject: f for f in lapsing.findings}
    assert set(found) == {"Visio Plan 2", "Exchange Online (Plan 1)", "Microsoft 365 Copilot"}
    assert found["Visio Plan 2"].severity == "high" and "grace period" in found["Visio Plan 2"].detail
    assert found["Microsoft 365 Copilot"].detail.startswith("Trial ends on")
    assert lapsing.action  # the first instruction, for lists of failing checks


def test_paid_subscriptions_and_unused_trials_are_not_flagged(db, two_tenants):
    alpha, _ = two_tenants
    subs = [
        ld.subscription("s-bp", "sku-bp", "SPB", total=10, ends_in=5),  # renews in 5 days: normal
        ld.subscription("s-copilot", "sku-copilot", "Microsoft_365_Copilot", trial=True, ends_in=10),
    ]
    skus = [ld.SKUS[0], dict(ld.SKUS[1], consumedUnits=0)]
    users = [ld.user("u1", "Ann", [ld.state("sku-bp")])]
    tenant = _licensing(db, alpha, skus=skus, subscriptions=subs, users=users)
    result = _results(db, tenant)["subscriptions_lapsing"]
    assert result.status == "pass" and result.summary == "2 subscriptions checked; none is lapsing"
    rows, flags = _table(_report(db, tenant), "Renewals and expiry")
    assert (rows[0][2], flags[0]) == ("Trial", "low")  # listed first, but nobody relies on it


def _old_snapshot(db):
    """Objects as collected before licence assignment and renewal dates were read."""
    unpack(db)
    db.execute("UPDATE objects SET data = json_remove(data, '$.licenceStates', '$.usageLocation') WHERE kind = 'user'")
    db.execute("DELETE FROM objects WHERE kind IN ('subscription_status', 'company_subscription')")


def test_unreadable_or_old_data_is_unavailable_not_passed(db, two_tenants):
    alpha, bravo = two_tenants
    seed_tenant(db, alpha)
    tenant = store.get_tenant(db, alpha)
    _old_snapshot(db)
    results = _results(db, tenant)
    assert results["licence_assignment_errors"].status == "unavailable"
    assert results["subscriptions_lapsing"].status == "unavailable"
    report = _report(db, tenant)
    assert report.unavailable is None
    assert any("Renewal dates show after" in n for n in report.notes)
    assert any("directly or by a group" in n for n in report.notes)
    rows, _ = _table(report, "Who has which licence")
    assert {r[2] for r in rows} == {"Not known"}
    figures = {h.label: h.value for h in report.highlights}
    assert figures["Licence errors"] == "Not checked" and figures["Lapsing within 30 days"] == "Not checked"

    assert data_reports.build(db, store.get_tenant(db, bravo), "subscriptions").unavailable.startswith(
        "Licence data hasn't been collected"
    )


def test_account_page_shows_how_each_licence_is_assigned(db, two_tenants):
    alpha, _ = two_tenants
    tenant = _licensing(db, alpha)
    ann = user_profile.build(db, tenant, "u1")
    licences = next(p for p in ann.parts if p.key == "licences")
    assert licences.tables[0].columns == ["Licence", "Assigned", "Cost / month"]
    rows = {r[0]: r for r in licences.tables[0].rows}
    assert rows["Microsoft 365 Business Premium"][1] == "Directly and by group: Licence - Business Premium"
    assert ("Usage location", "GB") in licences.facts
    assert any("assigned directly and by Licence - Business Premium" in a.text for a in ann.attention)

    dan = user_profile.build(db, tenant, "u8")
    licences = next(p for p in dan.parts if p.key == "licences")
    rows = {r[0]: (r, f) for r, f in zip(licences.tables[0].rows, licences.tables[0].flags, strict=True)}
    bp, flag = rows["Microsoft 365 Business Premium"]
    assert (bp[1], bp[2], flag) == ("By group: Licence - Business Premium", "Not assigned", "medium")
    assert bp[3].startswith("Not enough licences")
    attention = [(a.severity, a.text) for a in dan.attention if a.part == "licences"]
    assert (
        "medium",
        "Microsoft 365 Business Premium from Licence - Business Premium was not assigned. Not enough "
        "licences: every one bought is already assigned. Buy more, or remove the licence from someone "
        "who doesn't need it.",
    ) in attention
    assert any(
        s == "medium" and t.startswith("Exchange Online (Plan 1) comes from a suspended subscription")
        for s, t in attention
    )

    eve = user_profile.build(db, tenant, "u5")
    assert ("Usage location", "Not set") in next(p for p in eve.parts if p.key == "licences").facts
    assert any("No usage location" in a.text for a in eve.attention)


def test_account_page_before_assignment_paths_were_collected(db, two_tenants):
    alpha, _ = two_tenants
    seed_tenant(db, alpha)
    _old_snapshot(db)
    profile = user_profile.build(db, store.get_tenant(db, alpha), "u1")
    licences = next(p for p in profile.parts if p.key == "licences")
    assert licences.tables[0].columns == ["Licence", "Cost / month"] and not licences.facts


def test_library_lists_it_for_renewals_and_billing(db):
    uses = {u.key: u for u in library.USES}
    assert "subscriptions" in uses["renewal"].reports and "subscriptions" in uses["billing"].reports
    entry = next(e for a in library.catalogue() for e in a.entries if e.key == "subscriptions")
    assert entry.area == "Tenant and licences" and "trial" in entry.search and "renewal" in entry.uses


def test_subscription_dates_count_from_now(db, two_tenants):
    alpha, _ = two_tenants
    tenant = _licensing(db, alpha)
    later = datetime.now(UTC) + timedelta(days=20)  # the Copilot trial has ended by then
    rows, _ = _table(data_reports.build(db, tenant, "subscriptions", now=later), "Renewals and expiry")
    assert (rows[0][0], rows[0][2], rows[0][6]) == ("Microsoft 365 Copilot", "Trial", "Passed")


def test_an_active_state_for_a_licence_no_longer_assigned_is_left_out():
    # Recorded from the test tenant: assignedLicenses has one SKU, licenseAssignmentStates an extra "Active" one
    # for a SKU the tenant no longer has. It isn't held, so it isn't listed as an "Unknown licence".
    from officesentry import licences

    def state(sku, state="Active", group=None):
        return {"skuId": sku, "state": state, "group": group}

    user = {"skuIds": ["flow"], "licenceStates": [state("flow"), state("gone")]}
    assert set(licences.assignments(user)) == {"flow"}
    # A group licence that failed to assign, or a direct one with an error, still shows.
    user = {"skuIds": [], "licenceStates": [state("e3", "Error", "g1"), {**state("e5", "Error"), "error": "x"}]}
    assert set(licences.assignments(user)) == {"e3", "e5"}
    no_teams = "Microsoft_365_ Business_ Premium_(no Teams)"  # Microsoft's own part number, spaces included
    assert licences.sku_name(no_teams) == "Microsoft 365 Business Premium (no Teams)"
    assert licences.sku_name("FLOW_FREE") == "Power Automate Free"
