"""Microsoft 365 usage and licence right-sizing: collector, rules and reports."""

import io
from datetime import UTC, datetime

from openpyxl import load_workbook

from officesentry import store, usage
from officesentry.reports import data_reports, rollups
from officesentry.snapshot import TenantSnapshot
from tests import usagedata as ud
from tests.conftest import login, make_user
from tests.seed import seed_tenant

PRICES = {"SPB": 18.10, "O365_BUSINESS_PREMIUM": 10.30, "O365_BUSINESS_ESSENTIALS": 4.90,
          "EXCHANGESTANDARD": 3.30, "ENTERPRISEPACK": 20.00}
HASH = "FCE5081112C8B5FD27675CC9AAC4AB09"


def _snap(db, tenant_id):
    return TenantSnapshot(db, store.get_tenant(db, tenant_id))


def _report(db, tenant_id, key):
    return data_reports.build(db, store.get_tenant(db, tenant_id), key)


def _section(report, title):
    return next(s for s in report.sections if s.title == title)


def _price(db):
    db.executemany("INSERT INTO sku_prices (sku_part_number, monthly_price, updated_at) VALUES (?, ?, '2026-10-01')",
                   list(PRICES.items()))


# -- Collector -------------------------------------------------------------------------------


def test_collector_merges_reports_per_person(db, two_tenants):
    alpha, _ = two_tenants
    seed_tenant(db, alpha)
    snap = _snap(db, alpha)
    people = usage.people(snap)
    assert "left@x.com" not in people  # deleted accounts are skipped
    carol = usage.person(snap, "Carol@X.com")
    assert carol["active"] == {"exchange": True, "teams": True, "onedrive": True, "sharepoint": False, "office": True}
    assert carol["platforms"] == {"windows": False, "mac": False, "mobile": True, "web": True}
    assert carol["apps"] == {"outlook": ["web"], "word": ["web"], "teams": ["mobile", "web"]}
    assert carol["email"] == {"sent": 12, "received": 80, "read": 80}
    assert carol["teamsCounts"] == {"messages": 15, "calls": 1, "meetings": 2}
    assert people["dan@x.com"]["active"]["exchange"] is False  # 120 days ago
    assert people["dan@x.com"]["lastActive"]["exchange"] == ud.days_ago(120)
    erin = people["erin@x.com"]
    assert erin["activations"][0] == {"product": "MICROSOFT 365 APPS FOR ENTERPRISE", "windows": 1, "mac": 1,
                                      "mobile": 1, "shared": False, "lastActivated": ud.days_ago(200)}
    assert usage.desktop_use(erin) and not usage.desktop_use(carol)
    status = usage.status(snap)
    assert status["available"] and not status["concealed"] and all(status["read"].values())
    assert snap.metric("usage.people") == 5  # Power Automate Free alone isn't a licence
    assert snap.metric("usage.active.exchange") == 3
    assert snap.metric("usage.desktop_apps") == 1
    assert snap.runs["m365_usage"]["status"] == "succeeded"


def test_concealed_names_still_line_up(db, two_tenants):
    alpha, _ = two_tenants
    hidden = dict(ud.active_user(HASH, ["MICROSOFT 365 BUSINESS STANDARD"], exchange=1), displayName="A" * 32)
    seed_tenant(db, alpha, usage_kw={"active": [hidden], "apps": [ud.app_user(HASH, 1, outlook=("Web",))]})
    snap = _snap(db, alpha)
    row = usage.person(snap, HASH)
    assert row["concealed"] and row["name"] is None and row["platforms"]["web"]
    assert usage.label(row) == "Hidden name (FCE508)"
    assert usage.status(snap)["concealed"]
    report = _report(db, alpha, "usage")
    assert "Untick" in report.notice and "Display concealed user" in report.notice
    # Licences come from the report's product names when people can't be matched to accounts.
    _price(db)
    moves = _section(_report(db, alpha, "licence-rightsizing"), "Could move to a cheaper licence")
    assert moves.rows[0][:4] == ["Hidden name (FCE508)", "Microsoft 365 Business Standard", "Email",
                                 "Exchange Online (Plan 1)"]


def test_missing_permission_and_optional_reports(db, two_tenants):
    alpha, bravo = two_tenants
    seed_tenant(db, alpha, usage_kw={"status": 403})
    status = usage.status(_snap(db, alpha))
    assert not status["available"] and status["reason"] == "permission"
    assert "Reports.Read.All" in _report(db, alpha, "usage").unavailable
    assert "hasn't been collected" in _report(db, bravo, "licence-rightsizing").unavailable
    # Copilot's report failing is only logged when nobody holds Copilot.
    seed_tenant(db, bravo, usage_kw={"copilot_status": 400})
    assert _snap(db, bravo).runs["m365_usage"]["status"] == "succeeded"
    assert usage.status(_snap(db, bravo))["read"]["copilot"] is False


# -- Right-sizing rules ----------------------------------------------------------------------


def _row(active=(), platforms=None, apps=None, copilot=None):
    return {"active": {k: k in active for k, _ in usage.SERVICES}, "platforms": platforms, "apps": apps or {},
            "copilotActive": copilot}


def test_rules():
    nothing = usage.rightsize(_row(), ["SPB", "Microsoft_365_Copilot"])
    assert [(s.current, s.target) for s in nothing] == [("SPB", ""), ("Microsoft_365_Copilot", "")]
    assert nothing[0].action == "Remove the licence"
    receiving = dict(_row(), email={"received": 36}, licensed={"exchange": True})
    [s] = usage.rightsize(receiving, ["O365_BUSINESS_PREMIUM"])
    assert s.action == "Convert to a shared mailbox" and "received 36" in s.why
    mail = _row({"exchange", "office"}, {"windows": False, "mac": False, "mobile": False, "web": True},
                {"outlook": ["web"]})
    [s] = usage.rightsize(mail, ["SPB"])
    assert (s.target, s.action) == ("EXCHANGESTANDARD", "Exchange Online (Plan 1)") and "Intune" in s.caution
    [s] = usage.rightsize(mail, ["O365_BUSINESS_PREMIUM"], {"totalBytes": 48 * 1024 ** 3})
    assert s.target == "EXCHANGEENTERPRISE" and not s.caution  # too big for Plan 1
    web = _row({"exchange", "teams", "office"}, {"windows": False, "mac": False, "mobile": False, "web": True},
               {"word": ["web"]})
    assert [s.target for s in usage.rightsize(web, ["O365_BUSINESS_PREMIUM", "ENTERPRISEPACK", "SPB",
                                                    "SMB_BUSINESS"])] == ["O365_BUSINESS_ESSENTIALS", "STANDARDPACK",
                                                                          None, ""]
    desktop = _row({"exchange", "office"}, {"windows": True, "mac": False, "mobile": False, "web": False})
    assert usage.rightsize(desktop, ["SPB"]) == []
    unknown = _row({"exchange", "teams"}, None)  # apps report not read: no desktop suggestion
    assert usage.rightsize(unknown, ["O365_BUSINESS_PREMIUM"]) == []
    assert usage.rightsize(desktop, ["Microsoft_365_Copilot"]) == []  # Copilot use not known
    desktop["copilotActive"] = False
    copilot = usage.rightsize(desktop, ["SPB", "Microsoft_365_Copilot"])
    assert [(s.current, s.target) for s in copilot] == [("Microsoft_365_Copilot", "")]
    p2 = usage.rightsize(desktop, ["EXCHANGEENTERPRISE"], {"totalBytes": 2 * 1024 ** 3, "archive": False})
    assert [(s.current, s.target) for s in p2] == [("EXCHANGEENTERPRISE", "EXCHANGESTANDARD")]
    assert usage.rightsize(desktop, ["EXCHANGEENTERPRISE"], {"totalBytes": 1, "archive": True}) == []
    assert usage.Suggestion("SPB", "EXCHANGESTANDARD", "").saving(PRICES) == 14.80
    assert usage.Suggestion("SPB", "", "").saving(PRICES) == 18.10
    assert usage.Suggestion("SPB", None, "").saving(PRICES) is None
    assert usage.Suggestion("SPB", "STANDARDPACK", "").saving(PRICES) is None  # target not priced


# -- Reports ---------------------------------------------------------------------------------


def test_usage_report(db, two_tenants):
    alpha, _ = two_tenants
    seed_tenant(db, alpha)
    report = _report(db, alpha, "usage")
    assert not report.unavailable and not report.notice
    tiles = {t.label: t for t in report.tiles}
    assert tiles["Active people"].value == "60%"  # Ann, Carol and Erin of five licensed
    assert tiles["Email"].value == "60%" and tiles["Email"].foot == "3 of 5 licensed"
    assert tiles["Desktop Office apps"].foot == "1 of 4 with Office used Windows or Mac apps"
    service = {r[0]: r for r in _section(report, "Usage by service").rows}
    assert service["Teams"] == ["Teams", 5, 2, 3, "40%"]
    last = {r[0]: r for r in _section(report, "Last active in each service").rows}
    assert last["dan@x.com"][1] != "Never" and last["dan@x.com"][2] == "Never"
    assert last["carol@x.com"][6] == "Mobile, Web" and last["erin@x.com"][6] == "Windows, Mac"
    assert last["carol@x.com"][7] == "Microsoft 365 Business Standard"
    flags = dict(zip([r[0] for r in _section(report, "Last active in each service").rows],
                     _section(report, "Last active in each service").flags, strict=True))
    assert flags["dan@x.com"] == "low" and flags["carol@x.com"] is None
    apps = {r[0]: r for r in _section(report, "Office apps by platform").rows}
    assert apps["Excel"] == ["Excel", 1, 1, 1, 0, 0] and apps["Outlook"][1] == 3
    assert _section(report, "Office installs").rows == [["erin@x.com", "Microsoft 365 Apps for enterprise", 1, 1, 1,
                                                         "No", report.sections[3].rows[0][6]]]
    assert not _section(report, "Activity counts").compare


def test_rightsizing_report(db, two_tenants):
    alpha, _ = two_tenants
    seed_tenant(db, alpha)
    _price(db)
    report = _report(db, alpha, "licence-rightsizing")
    moves = _section(report, "Could move to a cheaper licence")
    rows = {r[0]: r for r in moves.rows}
    assert "bob@x.com" not in rows  # disabled: in Licence waste
    assert rows["dan@x.com"][1:5] == ["Microsoft 365 Business Standard", "Nothing", "Remove the licence", 10.30]
    assert rows["ann@x.com"][1:5] == ["Microsoft 365 Business Premium", "Email",
                                      "Exchange Online (Plan 1)", 14.80]
    assert "Intune" in rows["ann@x.com"][5]
    assert rows["carol@x.com"][3:5] == ["Microsoft 365 Business Basic", 5.40]
    assert "erin@x.com" not in rows
    assert [r[0] for r in moves.rows] == ["ann@x.com", "dan@x.com", "carol@x.com"]  # biggest saving first
    unused = {r[0]: r for r in _section(report, "Licensed but not used").rows}
    assert unused["erin@x.com"][2] == "SharePoint"
    assert unused["ann@x.com"][2] == "Teams, OneDrive, SharePoint"
    assert "dan@x.com" not in unused  # nothing used at all is a move, not an adoption gap
    highlights = {h.label: h for h in report.highlights}
    assert highlights["Saving / month"].value == 30.50
    assert highlights["Saving / year"].value == 366.00
    assert highlights["People who could move"].value == 3
    assert not any("no price" in n for n in report.notes)


def test_rightsizing_skips_new_and_inactive_accounts_and_notes_missing_prices(db, two_tenants):
    alpha, _ = two_tenants
    from tests.test_collectors import USERS
    users = USERS + [
        {"id": "u8", "displayName": "Dan", "userPrincipalName": "dan@x.com", "userType": "Member",
         "accountEnabled": True, "assignedLicenses": [{"skuId": "sku-bp"}],
         "createdDateTime": ud.days_ago(5) + "T00:00:00Z", "signInActivity": {}},
    ]
    seed_tenant(db, alpha, users=users)
    report = _report(db, alpha, "licence-rightsizing")
    rows = {r[0]: r for r in _section(report, "Could move to a cheaper licence").rows}
    assert "dan@x.com" not in rows  # created five days ago
    assert rows["ann@x.com"][4] == "Not priced"
    assert any("Microsoft 365 Business Premium" in n and "Licence prices" in n for n in report.notes)
    assert report.highlights[0].value == "Not priced"


def test_rollup_and_pages(client, db, two_tenants):
    alpha, bravo = two_tenants
    seed_tenant(db, alpha)
    _price(db)
    rollup = rollups.build(db, [store.get_tenant(db, t) for t in (alpha, bravo)], "licence-rightsizing",
                           datetime.now(UTC))
    row = next(r for r in rollup.rows if r.tenant["id"] == alpha)
    assert row.values["c:Saving / month"] == 30.50
    assert next(r for r in rollup.rows if r.tenant["id"] == bravo).note
    make_user(db, "admin")
    login(client, "admin")
    for key in ("usage", "licence-rightsizing"):
        page = client.get(f"/tenants/{alpha}/data/{key}")
        assert page.status_code == 200
        wb = load_workbook(io.BytesIO(client.get(f"/tenants/{alpha}/data/{key}.xlsx").data))
        assert "Summary" in wb.sheetnames
        assert client.get(f"/tenants/{bravo}/data/{key}").status_code == 200  # empty state renders
        assert client.get(f"/across/{key}").status_code == 200
    page = client.get(f"/tenants/{alpha}/data/licence-rightsizing").get_data(as_text=True)
    assert "Exchange Online (Plan 1)" in page and "Licensed but not used" in page
    prices = client.get("/admin/prices").get_data(as_text=True)
    assert "EXCHANGESTANDARD" in prices  # cheaper plans can be priced before any tenant holds them
