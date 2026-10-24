"""Mobile devices, automatic replies, distribution lists and the mailbox size trend.

Exchange Online and Graph are simulated at the HTTP level (``tests/exchangedata.py``) and every collector
runs through the worker, as in production.
"""

import copy
import json
import re
from datetime import UTC, datetime, timedelta
from io import BytesIO

import pytest
import responses
from openpyxl import load_workbook

from officesentry import state_history, store
from officesentry.collectors import distribution_lists as dl_collector
from officesentry.collectors.auto_replies import plain_text
from officesentry.collectors.mobile_devices import owner_name
from officesentry.exchange import EXCHANGE_ROOT
from officesentry.formatting import day
from officesentry.graph import GraphClient
from officesentry.reports import library, views
from officesentry.reports.auto_replies import ENDED, ON, SCHEDULED_ON, UPCOMING, state
from officesentry.reports.data_reports import GROUPS, REPORTS, build
from officesentry.reports.mobile_devices import platform
from officesentry.worker import execute_run
from tests import exchangedata as xd
from tests import graphdata as gd
from tests.conftest import login, make_user
from tests.seed import exchange_callback, exchange_factory, mock_tenant, seed_tenant
from tests.stored import unpack
from tests.test_collectors import SKUS
from tests.test_collectors import USERS as COLLECTOR_USERS

KEYS = ("mobile-devices", "auto-replies", "distribution-lists", "mailbox-trend")
FORBIDDEN = ("356938035643809", "+44 7700 900123", "RecoveryPassword", "Private subject", "outside.example")


@pytest.fixture
def tenant(db, two_tenants):
    alpha, _ = two_tenants
    seed_tenant(db, alpha, users=USERS)
    return store.get_tenant(db, alpha)


USERS = [*COLLECTOR_USERS, gd.INFO_USER]


def GRAPH(tenant, connection):
    return GraphClient(lambda: "t", sleep=lambda s: None)


def _report(db, tenant, key):
    report = build(db, tenant, key)
    assert not report.unavailable, report.unavailable
    return report, {s.title: s for s in report.sections}


def _tiles(report):
    return {t.label: t for t in report.tiles}


def _opens(report, tile):
    """The first column of the rows a headline figure's link shows."""
    view, _ = views.rows_view(report, tile.rows)
    shown = views.apply(report, view)
    return [r[0] for s in shown.sections if s.title == tile.rows.table for r in s.rows]


def _collect(db, tenant_id, key, *, exchange=None, **kwargs):
    """Run one collector through the worker against the simulated tenant."""
    factory = GRAPH
    with responses.RequestsMock(assert_all_requests_are_fired=False) as rsps:
        mock_tenant(rsps, users=USERS, exchange=exchange, **kwargs)
        store.enqueue_run(db, tenant_id, key)
        run = store.claim_next_run(db)
        status = execute_run(db, run, factory, exchange_factory)
    return run["id"], status


def _backdate(db, run_id, days):
    when = (datetime.now(UTC) - timedelta(days=days)).strftime("%Y-%m-%dT%H:%M:%SZ")
    db.execute("UPDATE runs SET finished_at = ?, started_at = ? WHERE id = ?", (when, when, run_id))
    db.execute("UPDATE metrics SET captured_at = ? WHERE run_id = ?", (when, run_id))


# -- Collectors --------------------------------------------------------------------------------------------


def test_owner_name_from_identity_or_distinguished_name():
    assert owner_name(xd.ANN_IPHONE) == "u1"
    assert owner_name({"DistinguishedName": xd.BOB_IPHONE["DistinguishedName"]}) == "u2"
    assert owner_name({"Identity": "odd"}) == ""


def test_plain_text_strips_html_and_cuts():
    assert plain_text("<p>Away &amp; back <b>Monday</b></p><style>p{}</style>") == "Away & back Monday"
    assert plain_text(None) == ""
    assert len(plain_text("x" * 1000)) == 300 and plain_text("x" * 1000).endswith("…")


def test_nothing_private_is_stored(db, tenant):
    unpack(db)
    rows = db.execute(
        "SELECT o.data FROM objects o JOIN runs r ON r.id = o.run_id WHERE r.collector IN "
        "('mobile_devices', 'auto_replies', 'distribution_lists')"
    ).fetchall()
    assert len(rows) > 5
    for (data,) in rows:
        assert not any(word in data for word in FORBIDDEN), data


def test_only_get_cmdlets_and_documented_parameters(db, two_tenants):
    alpha, _ = two_tenants
    sent = []

    def recorder(request):
        sent.append(json.loads(request.body)["CmdletInput"])
        return exchange_callback(gd.EXCHANGE)(request)

    factory = GRAPH
    with responses.RequestsMock(assert_all_requests_are_fired=False) as rsps:
        mock_tenant(rsps, users=USERS, exchange=False)
        rsps.add_callback(
            responses.POST,
            re.compile(re.escape(EXCHANGE_ROOT) + r"/adminapi/beta/[^/]+/InvokeCommand"),
            callback=recorder,
            content_type="application/json",
        )
        for key in ("mailboxes", "mobile_devices", "distribution_lists"):
            store.enqueue_run(db, alpha, key)
            assert execute_run(db, store.claim_next_run(db), factory, exchange_factory) == "succeeded"
    names = {c["CmdletName"] for c in sent}
    assert {
        "Get-MobileDevice",
        "Get-MobileDeviceStatistics",
        "Get-DistributionGroup",
        "Get-DistributionGroupMember",
        "Get-MessageTraceV2",
    } <= names
    assert all(n.startswith("Get-") for n in names)
    trace = next(c["Parameters"] for c in sent if c["CmdletName"] == "Get-MessageTraceV2")
    assert trace["Status"] == ["Expanded"] and trace["ResultSize"] == 5000
    assert sorted(trace["RecipientAddress"]) == ["everyone@x.com", "oldproject@x.com", "sales@x.com"]  # no room list
    start = datetime.fromisoformat(trace["StartDate"].replace("Z", "+00:00"))
    end = datetime.fromisoformat(trace["EndDate"].replace("Z", "+00:00"))
    assert timedelta(days=9, hours=23) < end - start <= timedelta(days=10)  # never more than Microsoft answers
    stats = [c["Parameters"]["Mailbox"] for c in sent if c["CmdletName"] == "Get-MobileDeviceStatistics"]
    assert sorted(stats) == ["g-u1", "g-u7", "u2"]  # once per mailbox with a device


# -- Mobile devices ----------------------------------------------------------------------------------------


def test_mobile_devices_lead_with_what_needs_attention(db, tenant):
    report, s = _report(db, tenant, "mobile-devices")
    assert [x.title for x in report.sections] == ["Devices that need attention", "Every device", "Devices by platform"]
    tiles = _tiles(report)
    assert tiles["Mobile devices"].value == 4
    assert tiles["Mobile devices"].foot == "2 synced in the last 30 days"
    assert (tiles["Leavers' devices"].value, tiles["Leavers' devices"].tone) == (1, "bad")
    assert tiles["Waiting for approval"].value == 1
    assert tiles["Not synced in 90 days"].value == 1 and _opens(report, tiles["Not synced in 90 days"]) == ["ann@x.com"]

    attention = s["Devices that need attention"]
    assert attention.flags == ["high", "medium", "low"]
    bob, info, android = attention.rows
    assert bob[:2] == ["u2", "iPhone 13"]
    assert str(bob[2]) == "Still syncing mail although the person's sign-in is blocked"
    assert info[0] == "info@x.com" and "Waiting for an admin" in str(info[2])
    assert android[1] == "SM-G991B" and str(android[2]).startswith("Not synced for 200 days")


def test_not_synced_figure_counts_calendar_days_like_its_link(db, two_tenants):
    """Ann's old phone last synced at midnight 90 days ago: over 90 days by the clock, but not "Older than 90
    days" on the table, which counts calendar days. The figure counts what its link opens."""
    alpha, _ = two_tenants
    midnight = (datetime.now(UTC) - timedelta(days=90)).strftime("%Y-%m-%dT00:00:00.0000000Z")
    stats = {**xd.DEVICE_STATS, "g-u1": [xd.DEVICE_STATS["g-u1"][0], xd._stats(xd.ANN_ANDROID, midnight)]}
    seed_tenant(db, alpha, users=USERS, exchange={**gd.EXCHANGE, "Get-MobileDeviceStatistics": stats})
    report, _ = _report(db, store.get_tenant(db, alpha), "mobile-devices")
    stale = _tiles(report)["Not synced in 90 days"]
    assert stale.value == 0 and _opens(report, stale) == []


def test_every_device_row(db, tenant):
    _, s = _report(db, tenant, "mobile-devices")
    rows = {r[1]: r for r in s["Every device"].rows}
    assert rows["Ann's iPhone"][0] == "ann@x.com"
    assert rows["Ann's iPhone"][2:7] == ["iOS", "iOS 17.5.1 21F90", "Outlook app", "Allowed", "2025-11-03"]
    assert rows["Ann's iPhone"][7] == xd.ago(hours=3)[:10] and rows["Ann's iPhone"][8] == "Active"
    assert rows["SM-G991B"][2:6] == ["Android", "Android 12", "Mail app (ActiveSync)", "Allowed"]
    assert rows["iPad Air"][5] == "Waiting for approval" and rows["iPad Air"][7] == ""
    assert rows["iPhone 13"][8] == "Sign-in blocked"
    platforms = {r[0]: r[1:] for r in s["Devices by platform"].rows}
    assert platforms == {"iOS": [3, 1, 2, 2], "Android": [1, 0, 1, 0]}


def test_device_left_from_a_removed_mailbox(db, two_tenants):
    """Exchange keeps a partnership after its mailbox is deleted or recreated under a new name, then answers
    404 for that mailbox's statistics (seen on the live test tenant). That's a leftover device, not a failed read."""
    alpha, _ = two_tenants
    old = xd._device("OldName", "d-5", "MNO345", "iPhone 11", "iOS 15.1 19B74")
    not_found = (
        404,
        {
            "error": {
                "code": "NotFound",
                "message": "|Microsoft.Exchange.Configuration.Tasks."
                "ManagementObjectNotFoundException|The operation couldn't be performed because "
                'Identity:"CN=OldName,OU=x.onmicrosoft.com" couldn\'t be found.',
            }
        },
    )
    data = {
        **gd.EXCHANGE,
        "Get-MobileDevice": [*xd.DEVICES, old],
        "Get-MobileDeviceStatistics": {**xd.DEVICE_STATS, "OldName": not_found},
    }
    seed_tenant(db, alpha, users=USERS, exchange=data)
    run = db.execute(
        "SELECT status FROM runs WHERE tenant_id = ? AND collector = 'mobile_devices'", (alpha,)
    ).fetchone()
    assert run["status"] == "succeeded"
    _, s = _report(db, store.get_tenant(db, alpha), "mobile-devices")
    row = next(r for r in s["Every device"].rows if r[1] == "iPhone 11")
    assert row[0] == "OldName" and row[7] == "" and row[8] == "Mailbox no longer exists"
    why = next(r for r in s["Devices that need attention"].rows if r[1] == "iPhone 11")[2]
    assert "Remove-MobileDevice" in str(why)


def test_a_failed_statistics_read_still_warns(db, two_tenants):
    alpha, _ = two_tenants
    busy = (500, {"error": {"code": "InternalServerError", "message": "Try again later."}})
    seed_tenant(
        db, alpha, users=USERS, exchange={**gd.EXCHANGE, "Get-MobileDeviceStatistics": {**xd.DEVICE_STATS, "u2": busy}}
    )
    run = db.execute(
        "SELECT id, status FROM runs WHERE tenant_id = ? AND collector = 'mobile_devices'", (alpha,)
    ).fetchone()
    assert run["status"] == "partial"


def test_platform_names():
    assert [platform(v) for v in ("iOS 17", "Android 14", "Windows 10.0", "", None, "BlackBerry")] == [
        "iOS",
        "Android",
        "Windows",
        "Other",
        "Other",
        "Other",
    ]


# -- Automatic replies -------------------------------------------------------------------------------------


def test_reply_state():
    now = datetime(2026, 10, 10, 12, tzinfo=UTC)
    sched = {
        "status": "scheduled",
        "start": {"dateTime": "2026-10-09T08:00:00.0000000", "timeZone": "UTC"},
        "end": {"dateTime": "2026-10-12T08:00:00.0000000", "timeZone": "UTC"},
    }
    assert state({"status": "alwaysEnabled"}, now) == ON
    assert state(sched, now) == SCHEDULED_ON
    assert state(sched, now - timedelta(days=3)) == UPCOMING
    assert state(sched, now + timedelta(days=3)) == ENDED
    assert state({"status": None}, now) == "Unknown"


def test_auto_replies(db, tenant):
    report, s = _report(db, tenant, "auto-replies")
    every = s["Mailboxes with an automatic reply"]
    assert len(every.rows) == 1  # the disabled one isn't kept
    row = every.rows[0]
    assert row[:7] == ["ann@x.com", "User", "On", "", "", "Everyone", datetime.now(UTC).date().isoformat()]
    assert row[7:] == [
        "Active",
        "I'm out of the office until Monday. For urgent things call Bob.",
        "Thanks for your email & sorry, I'm away.",
    ]
    assert every.flags == [None]  # first seen today, so not forgotten
    tiles = _tiles(report)
    assert (tiles["Replying now"].value, tiles["Replying outside the organisation"].value) == (1, 1)
    assert tiles["On for over 30 days"].value == 0 and tiles["Scheduled to start"].value == 0
    assert "Checked 2 mailboxes" in every.intro


def test_a_reply_left_on_for_weeks_is_flagged(db, tenant):
    first = db.execute(
        "SELECT id FROM runs WHERE tenant_id = ? AND collector = 'auto_replies'", (tenant["id"],)
    ).fetchone()["id"]
    _backdate(db, first, 40)
    _collect(db, tenant["id"], "auto_replies")
    report, s = _report(db, tenant, "auto-replies")
    row = s["Mailboxes with an automatic reply"].rows[0]
    assert row[6] == (datetime.now(UTC) - timedelta(days=40)).date().isoformat()
    assert s["Mailboxes with an automatic reply"].flags == ["medium"]
    assert _tiles(report)["On for over 30 days"].value == 1

    # Turned off and on again: the unbroken run starts over.
    _collect(db, tenant["id"], "auto_replies", auto_replies={"u1": xd.auto_reply("disabled"), "u7": xd.auto_reply()})
    newest = _collect(db, tenant["id"], "auto_replies")[0]
    _, s = _report(db, tenant, "auto-replies")
    found = db.execute("SELECT finished_at FROM runs WHERE id = ?", (newest,)).fetchone()["finished_at"]
    assert s["Mailboxes with an automatic reply"].rows[0][6] == found[:10]
    assert s["Mailboxes with an automatic reply"].flags == [None]


def test_a_reply_keeps_its_first_day_once_retention_removes_that_snapshot(db, tenant):
    first = db.execute(
        "SELECT id FROM runs WHERE tenant_id = ? AND collector = 'auto_replies'", (tenant["id"],)
    ).fetchone()["id"]
    _backdate(db, first, 40)
    _collect(db, tenant["id"], "auto_replies")
    db.execute("DELETE FROM objects WHERE run_id = ?", (first,))  # as the retention policy does
    then = (datetime.now(UTC) - timedelta(days=40)).date().isoformat()
    _, s = _report(db, tenant, "auto-replies")
    assert s["Mailboxes with an automatic reply"].rows[0][6] == then  # read from the daily history
    past = build(db, tenant, "auto-replies", as_of=datetime.now(UTC) - timedelta(days=20))
    assert past.sections[0].rows[0][6] == then


def test_scheduled_reply_to_contacts(db, two_tenants):
    alpha, _ = two_tenants
    soon = (datetime.now(UTC) + timedelta(days=5)).strftime("%Y-%m-%dT08:00:00.0000000")
    later = (datetime.now(UTC) + timedelta(days=12)).strftime("%Y-%m-%dT08:00:00.0000000")
    seed_tenant(
        db,
        alpha,
        users=USERS,
        auto_replies={
            "u1": xd.auto_reply("disabled"),
            "u7": xd.auto_reply("scheduled", "contactsOnly", soon, later, "Closed for the holiday", "We're closed"),
        },
    )
    report, s = _report(db, store.get_tenant(db, alpha), "auto-replies")
    row = s["Mailboxes with an automatic reply"].rows[0]
    assert row[:7] == ["info@x.com", "Shared", "Scheduled", soon[:10], later[:10], "Contacts only", ""]
    assert _tiles(report)["Scheduled to start"].value == 1 and _tiles(report)["Replying now"].value == 0
    # Not replying yet, so not counted as replying outside, and its link doesn't show it either.
    outside = _tiles(report)["Replying outside the organisation"]
    assert outside.value == 0 and _opens(report, outside) == []


# -- Distribution lists ------------------------------------------------------------------------------------


def test_distribution_lists(db, tenant):
    report, s = _report(db, tenant, "distribution-lists")
    assert [x.title for x in report.sections] == ["Lists that need attention", "Every list", "Members"]
    every = {r[0]: r for r in s["Every list"].rows}
    assert set(every) == {"Everyone", "Sales", "Old project"}  # the room list is left out
    assert every["Everyone"][1:9] == [
        "everyone@x.com",
        "Distribution list",
        3,
        "ann@x.com",
        0,
        "Yes",
        3,
        xd.ago(days=1)[:10],
    ]
    assert every["Sales"][2:8] == ["Mail-enabled security", 2, "bob@x.com", 1, "No", 1]
    assert every["Old project"][3:9] == [0, "None", 0, "No", 0, ""]
    assert s["Every list"].columns[9] == "Hidden from address list" and every["Everyone"][9] == "No"
    assert str(every["Everyone"][4]) == "ann@x.com"

    attention = {r[0]: r for r in s["Lists that need attention"].rows}
    assert set(attention) == {"Sales", "Old project"}  # watched for 10 days only, so nothing is "unused" yet
    assert list(attention["Old project"][1].parts) == ["No members", "No owner"]
    assert list(attention["Sales"][1].parts) == ["Only leavers own it"]
    assert any("less than 30 days" in n for n in report.notes)

    tiles = _tiles(report)
    assert [
        tiles[k].value
        for k in (
            "Distribution lists",
            "Empty lists",
            "Without an owner",
            "No mail in 90 days",
            "Open to outside senders",
        )
    ] == [3, 1, 2, 0, 1]
    assert _opens(report, tiles["No mail in 90 days"]) == []  # Old project gets no mail, but isn't watched long enough
    members = s["Members"].rows
    assert ["Sales", "Partner", "pat@partner.example", "External contact"] in members
    assert ["Everyone", "Sales", "sales@x.com", "Security group"] in members


def _trace_callback(data, rows, calls):
    """Exchange, with a message trace that honours the dates and addresses asked for."""
    plain = exchange_callback(data)

    def callback(request):
        body = json.loads(request.body)["CmdletInput"]
        if body["CmdletName"] != "Get-MessageTraceV2":
            return plain(request)
        p = body["Parameters"]
        calls.append(p)
        found = [
            r
            for r in rows
            if r["RecipientAddress"] in p["RecipientAddress"]
            and p["StartDate"] <= r["Received"][:19] + "Z" < p["EndDate"]
        ]
        return 200, {}, json.dumps({"value": found})

    return callback


def test_list_activity_adds_up_across_runs_without_counting_twice(db, two_tenants):
    alpha, _ = two_tenants
    # The first run reads up to two days ago (moved back from "an hour ago", so the second has something new).
    older = [m for m in xd.MESSAGE_TRACE if m["Received"] < xd.ago(days=2)]
    seed_tenant(db, alpha, users=USERS, exchange={**gd.EXCHANGE, "Get-MessageTraceV2": older})
    first = db.execute("SELECT id FROM runs WHERE collector = 'distribution_lists'").fetchone()["id"]
    _backdate(db, first, 2)
    unpack(db)
    for rowid, data in db.execute(
        "SELECT rowid, data FROM objects WHERE run_id = ? AND kind = 'dl_activity'", (first,)
    ).fetchall():
        entry = json.loads(data)
        entry["readTo"] = xd.ago(days=2)[:19] + "Z"
        db.execute("UPDATE objects SET data = ? WHERE rowid = ?", (json.dumps(entry), rowid))

    calls = []
    every_message = [*xd.MESSAGE_TRACE, xd._trace("oldproject@x.com", xd.ago(days=1, hours=2), 9)]
    with responses.RequestsMock(assert_all_requests_are_fired=False) as rsps:
        mock_tenant(rsps, users=USERS, exchange=False)
        rsps.add_callback(
            responses.POST,
            re.compile(re.escape(EXCHANGE_ROOT) + r"/adminapi/beta/[^/]+/InvokeCommand"),
            callback=_trace_callback(gd.EXCHANGE, every_message, calls),
            content_type="application/json",
        )
        store.enqueue_run(db, alpha, "distribution_lists")
        assert execute_run(db, store.claim_next_run(db), GRAPH, exchange_factory) == "succeeded"
    assert calls and all(c["StartDate"] == xd.ago(days=2)[:19] + "Z" for c in calls)
    _, s = _report(db, store.get_tenant(db, alpha), "distribution-lists")
    every = {r[0]: r for r in s["Every list"].rows}
    # Everyone: one message from the first run and two since; the older ones aren't counted again.
    assert every["Everyone"][7] == 3 and every["Everyone"][8] == xd.ago(days=1)[:10]
    assert every["Sales"][7] == 1
    assert every["Old project"][7] == 1 and every["Old project"][8] == xd.ago(days=1, hours=2)[:10]


def test_unused_list_is_flagged_once_watched_long_enough(db, tenant):
    unpack(db)
    run = db.execute("SELECT id FROM runs WHERE collector = 'distribution_lists'").fetchone()["id"]
    for rowid, data in db.execute(
        "SELECT rowid, data FROM objects WHERE run_id = ? AND kind = 'dl_activity'", (run,)
    ).fetchall():
        entry = json.loads(data)
        entry["watchedFrom"] = xd.ago(days=120)[:19] + "Z"
        db.execute("UPDATE objects SET data = ? WHERE rowid = ?", (json.dumps(entry), rowid))
    report, s = _report(db, tenant, "distribution-lists")
    attention = {r[0]: r for r in s["Lists that need attention"].rows}
    assert "No mail in 90 days" in attention["Old project"][1].parts
    assert "Everyone" not in attention and _tiles(report)["No mail in 90 days"].value == 1
    assert _opens(report, _tiles(report)["No mail in 90 days"]) == ["Old project"]
    assert not report.notes


def test_lists_without_the_message_trace(db, two_tenants):
    alpha, _ = two_tenants
    data = copy.deepcopy(gd.EXCHANGE)
    data["Get-MessageTraceV2"] = (403, {"error": {"code": "Forbidden", "message": "Access denied."}})
    seed_tenant(db, alpha, users=USERS, exchange=data)
    status = db.execute("SELECT status FROM runs WHERE collector = 'distribution_lists'").fetchone()["status"]
    assert status == "succeeded"  # lists are still collected
    report, s = _report(db, store.get_tenant(db, alpha), "distribution-lists")
    assert len(s["Every list"].rows) == 3
    assert {r[7] for r in s["Every list"].rows} == {"Not read"}
    assert _tiles(report)["No mail in 90 days"].value == "Not read"
    assert report.notes and "message trace" in report.notes[0] and "Access denied" in report.notes[0]


def test_trace_query_budget_leaves_the_rest_for_next_run(db, two_tenants, monkeypatch):
    alpha, _ = two_tenants
    monkeypatch.setattr(dl_collector, "TRACE_ADDRESSES", 1)
    monkeypatch.setattr(dl_collector, "TRACE_QUERIES", 2)
    seed_tenant(db, alpha, users=USERS)
    run = db.execute("SELECT id, status FROM runs WHERE collector = 'distribution_lists'").fetchone()
    assert run["status"] == "partial"
    entries = {e["_id"]: e for e in store.run_objects(db, run["id"], "dl_activity")}
    read = [a for a, e in entries.items() if e["readTo"] != e["watchedFrom"]]
    assert sorted(read) == ["everyone@x.com", "oldproject@x.com"]  # alphabetical, two queries
    assert entries["sales@x.com"]["days"] == {}


# -- Mailbox size trend ------------------------------------------------------------------------------------


GB = 1024**3


def _sizes(ann_gb: float, info_gb: float):
    data = copy.deepcopy(gd.EXCHANGE)
    data["Get-Mailbox"][0]["ProhibitSendQuota"] = "49.5 GB (53,150,220,288 bytes)"
    data["Get-MailboxStatistics"] = {
        "g-u1": [{"TotalItemSize": f"{ann_gb} GB ({int(ann_gb * GB):,} bytes)", "ItemCount": 9000}],
        "g-u7": [{"TotalItemSize": f"{info_gb} GB ({int(info_gb * GB):,} bytes)", "ItemCount": 13}],
    }
    return data


def test_mailbox_trend_needs_history_first(db, two_tenants):
    alpha, _ = two_tenants
    seed_tenant(db, alpha, users=USERS, exchange=_sizes(40, 1))
    report, s = _report(db, store.get_tenant(db, alpha), "mailbox-trend")
    assert "at least a week old" in report.notice
    assert _tiles(report)["Change in the last month"].value == "Not yet"
    rows = {r[0]: r for r in s["Growth by mailbox"].rows}
    assert rows["ann@x.com"][2:5] == ["40.0 GB", "", ""]
    assert not report.chart  # one point isn't a line


def test_mailbox_trend(db, two_tenants):
    alpha, _ = two_tenants
    seed_tenant(db, alpha, users=USERS, exchange=_sizes(37, 1))
    old = db.execute("SELECT id FROM runs WHERE collector = 'mailbox_usage'").fetchone()["id"]
    _backdate(db, old, 95)
    _backdate(db, _collect(db, alpha, "mailbox_usage", exchange=_sizes(40, 1))[0], 30)
    _collect(db, alpha, "mailbox_usage", exchange=_sizes(45, 1))
    report, s = _report(db, store.get_tenant(db, alpha), "mailbox-trend")
    assert not report.notice
    growth = s["Growth by mailbox"]
    assert growth.compare is False
    ann, info = growth.rows
    assert ann[:8] == ["ann@x.com", "User", "45.0 GB", "40.0 GB", "+5.0 GB", "37.0 GB", "+5.0 GB", "49.5 GB"]
    # 4.5 GB left at 5 GB a month: full in about 27 days, 0.9 months.
    assert growth.columns[8:] == ["Months until full", "Forecast full"]
    full = datetime.fromisoformat(ann[9]).date()
    assert abs((full - datetime.now(UTC).date()).days - 27) <= 1 and ann[8] == 0.9
    assert growth.flags == ["high", None]
    assert info[4] == "No change" and info[7] == "No limit" and info[8:] == ["", ""]
    tiles = _tiles(report)
    assert tiles["Total mailbox size"].value == "46.0 GB"
    assert tiles["Change in the last month"].value == "+5.0 GB"
    assert (tiles["Fastest growing"].value, tiles["Fastest growing"].foot) == ("+5.0 GB a month", "ann@x.com")
    assert (tiles["Full within 90 days"].value, tiles["Full within 90 days"].tone) == (1, "bad")
    assert report.chart.startswith("<svg") and report.chart_title == "Mailbox storage over time"
    months = s["Total size by month"]
    assert months.rows[0][1] == "46.0 GB" and len(months.rows) >= 2


def test_mailbox_trend_caps_far_forecasts(db, two_tenants):
    """A mailbox growing by a few bytes a month is centuries from full: no date, and no date arithmetic."""
    alpha, _ = two_tenants
    seed_tenant(db, alpha, users=USERS, exchange=_sizes(37, 1))
    _backdate(db, db.execute("SELECT id FROM runs WHERE collector = 'mailbox_usage'").fetchone()["id"], 30)
    _collect(db, alpha, "mailbox_usage", exchange=_sizes(37 + 4 / GB, 1))
    report, s = _report(db, store.get_tenant(db, alpha), "mailbox-trend")
    growth = s["Growth by mailbox"]
    ann = growth.rows[0]
    assert ann[0] == "ann@x.com" and ann[4] == "+4 bytes" and ann[8:] == ["", "More than 2 years away"]
    assert growth.flags == [None, None]
    assert _tiles(report)["Full within 90 days"].value == 0
    assert f"30 days earlier ({day(datetime.now(UTC) - timedelta(days=30))})" in growth.intro


def test_mailbox_trend_reads_the_daily_history_once_snapshots_are_pruned(db, two_tenants):
    alpha, _ = two_tenants
    seed_tenant(db, alpha, users=USERS, exchange=_sizes(37, 1))
    old = db.execute("SELECT id FROM runs WHERE collector = 'mailbox_usage'").fetchone()["id"]
    _backdate(db, old, 95)
    month = _collect(db, alpha, "mailbox_usage", exchange=_sizes(40, 1))[0]
    _backdate(db, month, 30)
    _collect(db, alpha, "mailbox_usage", exchange=_sizes(45, 1))
    assert state_history.recorded(db, old) and state_history.recorded(db, month)
    db.execute("DELETE FROM objects WHERE run_id IN (?, ?)", (old, month))  # as the retention policy does
    report, s = _report(db, store.get_tenant(db, alpha), "mailbox-trend")
    assert not report.notice
    ann = next(r for r in s["Growth by mailbox"].rows if r[0] == "ann@x.com")
    assert ann[3:7] == ["40.0 GB", "+5.0 GB", "37.0 GB", "+5.0 GB"]


def test_trend_as_of_ignores_later_collections(db, two_tenants):
    alpha, _ = two_tenants
    seed_tenant(db, alpha, users=USERS, exchange=_sizes(37, 1))
    old = db.execute("SELECT id FROM runs WHERE collector = 'mailbox_usage'").fetchone()["id"]
    _backdate(db, old, 40)
    _collect(db, alpha, "mailbox_usage", exchange=_sizes(45, 1))
    as_of = datetime.now(UTC) - timedelta(days=20)
    report = build(db, store.get_tenant(db, alpha), "mailbox-trend", as_of=as_of)
    rows = {r[0]: r for r in report.sections[0].rows}
    assert rows["ann@x.com"][2] == "37.0 GB" and rows["ann@x.com"][4] == ""
    # Two collections 40 days apart: the month compares with the older one, three months with nothing.
    rows = {r[0]: r for r in _report(db, store.get_tenant(db, alpha), "mailbox-trend")[1]["Growth by mailbox"].rows}
    assert rows["ann@x.com"][3:6] == ["37.0 GB", "+8.0 GB", ""]


# -- Everywhere --------------------------------------------------------------------------------------------


def test_unavailable_before_collection_and_without_exchange(db, two_tenants):
    alpha, beta = two_tenants
    for key in KEYS:
        assert "haven't been collected" in build(db, store.get_tenant(db, beta), key).unavailable
    no_exchange = [{**SKUS[0], "servicePlans": [{"servicePlanName": "AAD_PREMIUM", "provisioningStatus": "Success"}]}]
    seed_tenant(db, alpha, skus=no_exchange)
    for key in KEYS:
        assert build(db, store.get_tenant(db, alpha), key).unavailable == "This tenant has no Exchange Online licence."


def test_registered_grouped_and_searchable():
    email = dict((k, keys) for k, _, keys in GROUPS)["email"]
    for key in KEYS:
        assert key in REPORTS and key in email and key in library.KEYWORDS
        assert any(key in use.reports for use in library.USES)


@pytest.mark.parametrize("key", KEYS)
def test_page_and_downloads(client, db, tenant, key):
    make_user(db, "admin")
    login(client, "admin")
    base = f"/tenants/{tenant['id']}/data/{key}"
    page = client.get(base)
    assert page.status_code == 200
    assert REPORTS[key][0] in page.get_data(as_text=True)
    assert load_workbook(BytesIO(client.get(base + ".xlsx").data)).sheetnames
    csv = client.get(base + ".csv")
    assert csv.status_code == 200 and csv.data
