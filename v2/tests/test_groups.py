"""Groups and Teams: collector, checks, the Groups and Teams report and changes."""

import json
import re
from datetime import UTC, datetime, timedelta
from io import BytesIO
from urllib.parse import parse_qs, urlparse

import pytest
import responses
from openpyxl import load_workbook

from officesentry import store
from officesentry.changes import all_changes
from officesentry.collectors import groups as collector
from officesentry.graph import GraphClient
from officesentry.reports import data_reports
from officesentry.rules import evaluate
from officesentry.snapshot import TenantSnapshot
from tests import groupdata as grp
from tests.conftest import login, make_user
from tests.seed import seed_tenant
from tests.test_collectors import SKUS, V1, _metrics, _run


@pytest.fixture
def factory():
    return lambda tenant, connection: GraphClient(lambda: "t", sleep=lambda s: None)


def _results(db, tenant_id):
    return {r.key: r for r in evaluate(TenantSnapshot(db, store.get_tenant(db, tenant_id)))}


def _groups(db, run_id):
    return {g["_id"]: g for g in store.run_objects(db, run_id, "group")}


def _one(db, run_id, kind):
    return store.run_objects(db, run_id, kind)[0]


def _mock(rsps, *, settings=None, lifecycle=None, page_size=999):
    rsps.get(f"{V1}/groups", json={"value": grp.GROUPS})

    def people(request):  # owners and members, paged like Graph
        gid, kind = request.path_url.split("?")[0].split("/")[3:5]
        page = int(parse_qs(urlparse(request.url).query).get("$skiptoken", ["0"])[0])
        rows = getattr(grp, kind.upper()).get(gid, [])
        body = {"value": rows[page * page_size:(page + 1) * page_size]}
        if (page + 1) * page_size < len(rows):
            body["@odata.nextLink"] = f"{V1}/groups/{gid}/{kind}?$skiptoken={page + 1}"
        return 200, {}, json.dumps(body)
    rsps.add_callback(responses.GET, re.compile(re.escape(V1) + r"/groups/[^/]+/(owners|members)"),
                      callback=people, content_type="application/json")
    for path, value, default in (("groupSettings", settings, grp.GROUP_SETTINGS),
                                 ("groupLifecyclePolicies", lifecycle, grp.LIFECYCLE_POLICIES)):
        if isinstance(value, dict):  # an error response
            rsps.get(f"{V1}/{path}", **value)
        else:
            rsps.get(f"{V1}/{path}", json={"value": default if value is None else value})


FORBIDDEN = {"status": 403, "json": {"error": {"code": "Authorization_RequestDenied", "message": "Insufficient"}}}


# -- Collector ---------------------------------------------------------------------------


@responses.activate
def test_collects_groups_owners_members_and_settings(db, two_tenants, factory):
    alpha, _ = two_tenants
    _mock(responses)
    run_id, status = _run(db, alpha, "groups", factory)
    assert status == "succeeded"
    m = _metrics(db, run_id)
    assert (m["groups.total"], m["groups.m365"], m["groups.teams"]) == (10, 5, 3)
    assert (m["groups.ownerless"], m["groups.with_guests"], m["groups.guest_memberships"]) == (3, 2, 3)

    urls = [c.request.url for c in responses.calls]
    assert "select=id%2CdisplayName%2Cdescription%2CgroupTypes" in urls[0] and "top=999" in urls[0]
    assert not any(f"/groups/{gid}/" in u for u in urls for gid in ("s1", "s2", "s3", "m1", "d1"))

    groups = _groups(db, run_id)
    phoenix = groups["g2"]
    assert (phoenix["isTeam"], phoenix["visibility"], phoenix["kind"]) == (True, "Public", "microsoft365")
    assert (phoenix["ownerCount"], phoenix["enabledOwnerCount"], phoenix["memberCount"]) == (1, 0, 3)
    assert phoenix["guests"] == ["lee_partner.org#EXT#@x.onmicrosoft.com", "pat@supplier.com"]
    assert phoenix["owners"] == [{"name": "Bob", "upn": "bob@x.com", "type": "user", "enabled": False}]
    assert groups["g5"]["owners"][0]["type"] == "servicePrincipal" and groups["g5"]["enabledOwnerCount"] == 0
    assert (groups["s1"]["kind"], groups["s1"]["dynamic"], groups["s1"]["examined"]) == ("security", True, False)
    assert groups["s2"]["synced"] and groups["m1"]["kind"] == "mail_security" and groups["d1"]["kind"] == "distribution"

    settings = {k: v for k, v in _one(db, run_id, "group_settings").items() if not k.startswith("_")}
    assert settings == {"available": True, "source": "tenant", "enableGroupCreation": False,
                        "allowedGroupId": "s3", "allowedGroupName": "Group creators",
                        "allowGuestsToAccessGroups": True, "allowToAddGuests": False,
                        "usageGuidelinesUrl": "https://intranet.x.com/teams"}
    policy = _one(db, run_id, "expiration_policy")
    assert (policy["managedGroupTypes"], policy["groupLifetimeInDays"]) == ("All", 365)


@responses.activate
def test_members_are_paged(db, two_tenants, factory):
    alpha, _ = two_tenants
    _mock(responses, page_size=2)
    run_id, status = _run(db, alpha, "groups", factory)
    assert status == "succeeded"
    sales = _groups(db, run_id)["g1"]
    assert (sales["memberCount"], sales["guestCount"], sales["guests"]) == (3, 1, ["pat@supplier.com"])
    pages = [c.request.url for c in responses.calls if "/groups/g1/members" in c.request.url]
    assert len(pages) == 2 and "skiptoken=1" in pages[1]


@responses.activate
def test_group_cap_warns_and_reads_teams_first(db, two_tenants, factory, monkeypatch):
    alpha, _ = two_tenants
    monkeypatch.setattr(collector, "GROUP_LIMIT", 2)
    _mock(responses)
    run_id, status = _run(db, alpha, "groups", factory)
    assert status == "partial"
    error = db.execute("SELECT error FROM runs WHERE id = ?", (run_id,)).fetchone()["error"]
    assert "first 2 of 5 teams and Microsoft 365 groups" in error
    examined = sorted(g["displayName"] for g in _groups(db, run_id).values() if g["examined"])
    assert examined == ["Board", "Project Phoenix"]  # teams, by name
    r = _results(db, alpha)["ownerless_groups"]
    assert r.summary == "1 team or group without an owner (2 of 5 checked)"


@responses.activate
def test_guest_list_is_bounded(db, two_tenants, factory, monkeypatch):
    alpha, _ = two_tenants
    monkeypatch.setattr(collector, "GUESTS_KEPT", 1)
    _mock(responses)
    run_id, _ = _run(db, alpha, "groups", factory)
    phoenix = _groups(db, run_id)["g2"]
    assert (phoenix["guestCount"], len(phoenix["guests"])) == (2, 1)
    report = data_reports.build(db, store.get_tenant(db, alpha), "groups")
    guests = {s.title: s for s in report.sections}["Guests in teams and groups"]
    assert ["1 more guest", "Project Phoenix", "Team"] in guests.rows


@responses.activate
def test_missing_settings_mean_microsoft_defaults(db, two_tenants, factory):
    alpha, _ = two_tenants
    _mock(responses, settings=[], lifecycle=[])
    run_id, status = _run(db, alpha, "groups", factory)
    assert status == "succeeded"
    settings = _one(db, run_id, "group_settings")
    assert settings["source"] == "default" and settings["enableGroupCreation"] is True
    assert settings["allowGuestsToAccessGroups"] is True and settings["allowToAddGuests"] is True
    assert settings["allowedGroupId"] is None and settings["usageGuidelinesUrl"] is None
    assert _one(db, run_id, "expiration_policy")["managedGroupTypes"] == "None"


@responses.activate
def test_unreadable_settings_are_partial(db, two_tenants, factory):
    alpha, _ = two_tenants
    _mock(responses, settings=FORBIDDEN, lifecycle={"status": 403, "json": {"error": {
        "code": "AadPremiumLicenseRequired", "message": "Needs a premium license"}}})
    run_id, status = _run(db, alpha, "groups", factory)
    assert status == "partial"
    assert len(_groups(db, run_id)) == 10
    r = _results(db, alpha)
    assert r["group_creation"].status == r["group_expiration"].status == "unavailable"
    assert r["group_creation"].summary == "Group settings couldn't be read"
    assert r["ownerless_groups"].status == "fail"


@responses.activate
def test_unreadable_owners_stop_and_warn(db, two_tenants, factory):
    alpha, _ = two_tenants
    responses.get(re.compile(re.escape(V1) + r"/groups/[^/]+/owners"), **FORBIDDEN)
    _mock(responses)
    run_id, status = _run(db, alpha, "groups", factory)
    assert status == "partial"
    assert sum("/owners" in c.request.url for c in responses.calls) == 1  # stopped after the first refusal
    assert not any(g["examined"] for g in _groups(db, run_id).values())
    assert _results(db, alpha)["ownerless_groups"].summary == "Owners of teams and groups couldn't be read"


# -- Checks ------------------------------------------------------------------------------


def test_group_checks(db, two_tenants):
    alpha, _ = two_tenants
    seed_tenant(db, alpha)
    r = _results(db, alpha)
    ownerless = r["ownerless_groups"]
    assert ownerless.status == "fail" and ownerless.category == "Groups"
    assert [(f.severity, f.subject, f.detail) for f in ownerless.findings] == [
        ("medium", "Project Phoenix", "Team, only owner is disabled (Bob)"),
        ("low", "Marketing", "Microsoft 365 group, no owner"),
        ("low", "Provisioned", "Microsoft 365 group, owned only by an app (Provisioning app)"),
    ]
    assert ownerless.summary == "3 teams and groups without an owner"
    assert r["group_creation"].status == "pass"
    assert r["group_creation"].summary == "Only admins and members of Group creators can create teams and groups"
    assert r["group_expiration"].status == "pass"
    assert r["group_expiration"].summary == "Unused teams and groups expire after 365 days"


def test_every_group_with_an_owner_passes(db, two_tenants, monkeypatch):
    alpha, _ = two_tenants
    owners = {gid: [grp.ANN] for gid in grp.OWNERS}
    monkeypatch.setattr(grp, "OWNERS", owners)
    seed_tenant(db, alpha)
    r = _results(db, alpha)["ownerless_groups"]
    assert (r.status, r.summary) == ("pass", "All 5 teams and groups have an owner")


def test_microsoft_defaults_fail(db, two_tenants, monkeypatch):
    alpha, _ = two_tenants
    monkeypatch.setattr(grp, "GROUP_SETTINGS", [])
    monkeypatch.setattr(grp, "LIFECYCLE_POLICIES", [])
    seed_tenant(db, alpha)
    r = _results(db, alpha)
    assert r["group_creation"].status == "fail"
    assert [(f.severity, f.detail) for f in r["group_creation"].findings] == [
        ("low", "Every user can create teams and Microsoft 365 groups (Microsoft's default)")]
    assert r["group_expiration"].status == "fail"
    assert [(f.severity, f.detail) for f in r["group_expiration"].findings] == [
        ("low", "No expiration policy: 5 teams and groups are kept until someone deletes them")]


def test_creation_limited_to_admins(db, two_tenants, monkeypatch):
    alpha, _ = two_tenants
    monkeypatch.setattr(grp, "GROUP_SETTINGS", [{"displayName": "Group.Unified", "values": [
        {"name": "EnableGroupCreation", "value": "False"}]}])
    seed_tenant(db, alpha)
    assert _results(db, alpha)["group_creation"].summary == "Only admins can create teams and groups"


def test_expiration_for_selected_groups(db, two_tenants, monkeypatch):
    alpha, _ = two_tenants
    monkeypatch.setattr(grp, "LIFECYCLE_POLICIES", [{"id": "lp1", "groupLifetimeInDays": 180,
                                                     "managedGroupTypes": "Selected"}])
    covered = [{**g, "expirationDateTime": "2027-03-01T00:00:00Z"} if g["id"] == "g1" else g for g in grp.GROUPS]
    monkeypatch.setattr(grp, "GROUPS", covered)
    seed_tenant(db, alpha)
    r = _results(db, alpha)["group_expiration"]
    assert (r.status, r.summary) == ("fail", "Expiration covers selected groups only")
    assert r.findings[0].detail == "4 of 5 teams and groups never expire"


def test_without_entra_p1_settings_checks_are_unavailable(db, two_tenants, monkeypatch):
    alpha, _ = two_tenants
    monkeypatch.setattr(grp, "GROUP_SETTINGS", [])
    monkeypatch.setattr(grp, "LIFECYCLE_POLICIES", [])
    no_p1 = [{**s, "servicePlans": [p for p in s["servicePlans"] if not p["servicePlanName"].startswith("AAD_")]}
             for s in SKUS]
    seed_tenant(db, alpha, skus=no_p1)
    r = _results(db, alpha)
    assert r["group_creation"].status == r["group_expiration"].status == "unavailable"
    assert "Entra ID P1" in r["group_creation"].summary and "Entra ID P1" in r["group_expiration"].summary
    assert r["ownerless_groups"].status == "fail"


def test_tenant_without_microsoft_365_groups(db, two_tenants, monkeypatch):
    alpha, _ = two_tenants
    monkeypatch.setattr(grp, "GROUPS", [g for g in grp.GROUPS if "Unified" not in g["groupTypes"]])
    monkeypatch.setattr(grp, "LIFECYCLE_POLICIES", [])
    seed_tenant(db, alpha)
    r = _results(db, alpha)
    assert r["ownerless_groups"].status == "unavailable"
    assert [f.severity for f in r["group_expiration"].findings] == ["info"]


def test_never_collected_is_unavailable(db, two_tenants):
    _, bravo = two_tenants
    r = _results(db, bravo)
    for key in ("ownerless_groups", "group_creation", "group_expiration"):
        assert r[key].status == "unavailable" and r[key].summary == "Group data hasn't been collected yet"


# -- Report ------------------------------------------------------------------------------


def test_groups_report(db, two_tenants):
    alpha, _ = two_tenants
    seed_tenant(db, alpha)
    report = data_reports.build(db, store.get_tenant(db, alpha), "groups", datetime.now(UTC))
    sections = {s.title: s for s in report.sections}
    assert list(sections) == ["Settings", "Needs attention", "Teams and groups", "Guests in teams and groups",
                              "Other groups"]
    settings = dict(sections["Settings"].rows)
    assert settings["Who can create teams and groups"] == "Admins and members of Group creators"
    assert (settings["Guests can access group content"], settings["Owners can add guests"]) == ("Yes", "No")
    assert settings["Expiration"] == "After 365 days unless in use or renewed"
    assert settings["Expiry notices for groups without an owner go to"] == "it@x.com"
    assert not any(sections["Settings"].flags)

    attention = sections["Needs attention"]
    assert [(row[0], row[2], flag) for row, flag in zip(attention.rows, attention.flags, strict=True)] == [
        ("Project Phoenix", "Only owner is disabled (Bob)", "medium"),
        ("Marketing", "No owner", "low"),
        ("Provisioned", "Owned only by an app (Provisioning app)", "low"),
        ("Board", "Single owner", "info"),
    ]
    assert attention.rows[0][3] == "Bob (disabled)" and attention.rows[2][3] == "Provisioning app (app)"

    every = sections["Teams and groups"]
    assert [row[0] for row in every.rows] == ["Board", "Marketing", "Project Phoenix", "Provisioned", "Sales"]
    sales = dict(zip(every.columns, every.rows[-1], strict=True))
    assert sales == {"Name": "Sales", "Team": "Yes", "Visibility": "Private", "Owners": "Ann, Eve", "Members": 3,
                     "Guests": 1, "Created": "2025-01-10"}

    assert sections["Guests in teams and groups"].rows == [
        ["lee_partner.org#EXT#@x.onmicrosoft.com", "Project Phoenix", "Team"],
        ["pat@supplier.com", "Project Phoenix", "Team"],
        ["pat@supplier.com", "Sales", "Team"],
    ]
    assert sections["Other groups"].rows == [["Security groups", 3, 1, 1], ["Mail-enabled security groups", 1, 0, 0],
                                             ["Distribution lists", 1, 0, 0]]


def test_groups_report_with_microsoft_defaults(db, two_tenants, monkeypatch):
    alpha, _ = two_tenants
    monkeypatch.setattr(grp, "GROUP_SETTINGS", [])
    monkeypatch.setattr(grp, "LIFECYCLE_POLICIES", [])
    seed_tenant(db, alpha)
    settings = data_reports.build(db, store.get_tenant(db, alpha), "groups").sections[0]
    assert "Microsoft's defaults apply" in settings.intro
    rows = {row[0]: (row[1], flag) for row, flag in zip(settings.rows, settings.flags, strict=True)}
    assert rows["Who can create teams and groups"] == ("Everyone", "low")
    assert rows["Expiration"] == ("Off: teams and groups are kept until someone deletes them", "low")
    assert rows["Usage guidelines"] == ("None", None)


def test_groups_report_without_data(db, two_tenants):
    _, bravo = two_tenants
    report = data_reports.build(db, store.get_tenant(db, bravo), "groups")
    assert report.unavailable and not report.sections


def test_groups_page_and_xlsx(client, db, two_tenants):
    alpha, _ = two_tenants
    seed_tenant(db, alpha)
    make_user(db, "admin")
    login(client, "admin")
    page = client.get(f"/tenants/{alpha}/data/groups")
    assert page.status_code == 200
    html = page.get_data(as_text=True)
    assert "Project Phoenix" in html and "Needs attention" in html and "Groups and Teams" in html
    resp = client.get(f"/tenants/{alpha}/data/groups.xlsx")
    assert resp.status_code == 200
    wb = load_workbook(BytesIO(resp.data))
    assert wb.sheetnames == ["Settings", "Needs attention", "Teams and groups", "Guests in teams and groups",
                             "Other groups"]


# -- Since last month --------------------------------------------------------------------


def test_group_changes_since_last_month(db, two_tenants, monkeypatch):
    alpha, _ = two_tenants
    seed_tenant(db, alpha)
    old = (datetime.now(UTC) - timedelta(days=30)).strftime("%Y-%m-%dT%H:%M:%SZ")
    db.execute("UPDATE runs SET finished_at = ?", (old,))
    launch = {**grp.GROUPS[0], "id": "g6", "displayName": "Launch", "visibility": "Public"}
    lab = {**grp.GROUPS[5], "id": "s9", "displayName": "New security group"}
    monkeypatch.setattr(grp, "GROUPS", [g for g in grp.GROUPS if g["id"] != "g5"] + [launch, lab])
    monkeypatch.setattr(grp, "OWNERS", {**grp.OWNERS, "g1": [], "g4": [{**grp.ANN, "accountEnabled": False}],
                                        "g6": [grp.ANN]})
    seed_tenant(db, alpha)
    since = (datetime.now(UTC) - timedelta(days=28)).strftime("%Y-%m-%dT%H:%M:%SZ")
    got = {(c.kind, c.subject, c.detail) for c in all_changes(db, alpha, since) if c.area == "Teams and groups"}
    assert got == {
        ("added", "Launch", "New public team"),
        ("changed", "Sales", "Last owner removed"),
        ("changed", "Board", "Last enabled owner disabled"),
        ("removed", "Provisioned", "Group deleted"),
    }
