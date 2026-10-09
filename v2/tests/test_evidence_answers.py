"""Suggested questionnaire answers in the evidence pack, and the technical lead's overrides."""

import copy
import csv
import io
from io import BytesIO

from openpyxl import load_workbook

from officesentry import store
from officesentry.db import MIGRATIONS_DIR
from officesentry.reports import evidence, evidence_answers, export
from officesentry.reports.evidence_answers import NO, NOT_KNOWN, PARTLY, YES
from tests import graphdata as gd
from tests.conftest import login, make_user
from tests.seed import seed_tenant
from tests.test_collectors import USERS


def _answers(db, tenant_id):
    return {a.key: a for a in evidence.build(db, store.get_tenant(db, tenant_id)).answers}


def test_migration_0019_creates_the_overrides_table(db):
    assert (MIGRATIONS_DIR / "0019_evidence_answers.sql").exists()
    columns = {r["name"] for r in db.execute("PRAGMA table_info(evidence_answer_overrides)")}
    assert {"tenant_id", "question_key", "answer", "note", "suggested", "set_by", "set_at"} <= columns
    assert db.execute("PRAGMA user_version").fetchone()[0] >= 19


def test_answers_from_seeded_data(db, two_tenants):
    alpha, bravo = two_tenants
    seed_tenant(db, alpha, users=USERS + [gd.INFO_USER])
    a = _answers(db, alpha)
    assert [q.key for q in evidence_answers.QUESTIONS] == list(a)
    assert a["legacy_auth"].suggested == NO  # no policy blocks legacy sign-in
    assert a["audit_logs"].suggested == YES
    assert a["forwarding"].suggested == PARTLY and "2 external forwards" in a["forwarding"].why
    assert a["separate_admins"].suggested == PARTLY and "ann@x.com" in a["separate_admins"].why
    assert a["devices_encrypted"].suggested == NO and "1 of 2 computers" in a["devices_encrypted"].why
    assert a["backups"].suggested == a["training"].suggested == NOT_KNOWN
    assert a["mfa_remote"].source.report == "mfa" and a["mfa_remote"].source.anchor == "t-accounts-with-a-gap"
    assert all(x.why and x.override is None for x in a.values())
    # Nothing collected: nothing is guessed.
    assert all(x.suggested == NOT_KNOWN for x in _answers(db, bravo).values())


def test_mfa_with_break_glass_exclusion_is_yes(db, two_tenants):
    """The seeded policy requires MFA for everyone except the break-glass Global Administrator."""
    alpha, _ = two_tenants
    seed_tenant(db, alpha)
    a = _answers(db, alpha)
    mfa = a["mfa_remote"]
    assert mfa.suggested == YES
    assert '"Require MFA for all users"' in mfa.why and "bg@x.com (break-glass)" in mfa.why
    assert "except 1 named exclusion" in mfa.why
    admins = a["mfa_privileged"]
    assert admins.suggested == YES and "break-glass" in admins.why and "bg@x.com" in admins.why


def test_mfa_not_enforced_and_group_exclusions(db, two_tenants):
    alpha, bravo = two_tenants
    seed_tenant(db, alpha, ca=[gd.CA_REPORT_ONLY])
    a = _answers(db, alpha)
    assert a["mfa_remote"].suggested == NO and "report-only" in a["mfa_remote"].why
    assert a["mfa_privileged"].suggested == NO

    by_group = copy.deepcopy(gd.CA_MFA_ALL)
    by_group["conditions"]["users"].update(excludeUsers=[], excludeGroups=["g-all-staff"])
    seed_tenant(db, bravo, ca=[by_group, gd.CA_BLOCK_LEGACY])
    b = _answers(db, bravo)
    assert b["mfa_remote"].suggested == PARTLY and "(group)" in b["mfa_remote"].why
    assert b["legacy_auth"].suggested == PARTLY and "SMTP AUTH" in b["legacy_auth"].why


def test_mfa_is_partly_when_anyone_but_break_glass_is_excluded(db, two_tenants):
    """A licensed Global Administrator or a plain user left out of MFA is a gap, not break-glass."""
    alpha, bravo = two_tenants
    people = copy.deepcopy(gd.CA_MFA_ALL)
    people["conditions"]["users"]["excludeUsers"] = ["u1", "u2", "u3"]
    seed_tenant(db, alpha, ca=[people])
    a = _answers(db, alpha)
    assert a["mfa_remote"].suggested == PARTLY and "ann@x.com (break-glass)" not in a["mfa_remote"].why
    assert a["mfa_privileged"].suggested in (PARTLY, NO)

    roles = copy.deepcopy(gd.CA_MFA_ALL)
    roles["conditions"]["users"].update(excludeUsers=[], excludeRoles=["role-1"], excludeGuestsOrExternalUsers={
        "guestOrExternalUserTypes": "b2bCollaborationGuest"})
    seed_tenant(db, bravo, ca=[roles])
    b = _answers(db, bravo)
    assert b["mfa_remote"].suggested == PARTLY
    assert "(role)" in b["mfa_remote"].why and "Guests and external users" in b["mfa_remote"].why


def test_security_defaults_answer_yes(db, two_tenants):
    alpha, _ = two_tenants
    seed_tenant(db, alpha, defaults=True)
    a = _answers(db, alpha)
    assert a["mfa_remote"].suggested == a["mfa_privileged"].suggested == YES
    assert "Security Defaults" in a["mfa_remote"].why


def test_override_set_and_clear(client, db, two_tenants):
    alpha, bravo = two_tenants
    seed_tenant(db, alpha)
    make_user(db, "lead", role="analyst", tenants=[alpha])
    login(client, "lead")
    url = f"/tenants/{alpha}/evidence/answers/legacy_auth"
    resp = client.post(url, data={"answer": "Yes", "note": "Basic auth is off tenant-wide since March."})
    assert resp.status_code == 302 and resp.headers["Location"].endswith("#answer-legacy_auth")
    a = _answers(db, alpha)["legacy_auth"]
    assert a.answer == YES and a.suggested == NO and a.override.set_by == "lead"
    assert a.override.note.startswith("Basic auth") and not a.data_changed
    logged = db.execute("SELECT detail FROM audit_log WHERE action = 'evidence_answer_set'").fetchone()
    assert logged["detail"] == "legacy_auth: Yes (suggested No)" and "March" not in logged["detail"]

    assert client.post(url, data={"answer": "Maybe"}).status_code == 400
    assert client.post(f"/tenants/{alpha}/evidence/answers/nope", data={"answer": "Yes"}).status_code == 404
    assert client.post(f"/tenants/{bravo}/evidence/answers/legacy_auth", data={"answer": "Yes"}).status_code == 404

    client.post(url, data={"action": "clear"})
    assert _answers(db, alpha)["legacy_auth"].override is None
    assert db.execute("SELECT COUNT(*) FROM audit_log WHERE action = 'evidence_answer_cleared'").fetchone()[0] == 1


def test_read_only_users_cannot_override(client, db, two_tenants):
    alpha, _ = two_tenants
    seed_tenant(db, alpha)
    make_user(db, "client", role="client", tenants=[alpha])
    login(client, "client")
    assert client.post(f"/tenants/{alpha}/evidence/answers/legacy_auth", data={"answer": "Yes"}).status_code == 403
    assert db.execute("SELECT COUNT(*) FROM evidence_answer_overrides").fetchone()[0] == 0
    page = client.get(f"/tenants/{alpha}/evidence").get_data(as_text=True)
    assert "Copy all answers" in page and "Change this answer" not in page and "Write this answer" not in page


def test_overrides_reach_page_pdf_and_exports(client, db, two_tenants, monkeypatch):
    alpha, _ = two_tenants
    seed_tenant(db, alpha)
    evidence_answers.set_override(db, alpha, "legacy_auth", YES, "Blocked at the firewall.", NO, "jack")
    make_user(db, "admin")
    make_user(db, "client", role="client", tenants=[alpha])

    login(client, "admin")
    page = client.get(f"/tenants/{alpha}/evidence").get_data(as_text=True)
    assert "Blocked at the firewall." in page and "Answer set by jack on" in page
    assert "Suggested from the data: <strong>No</strong>" in page  # staff still see the suggestion
    assert "Data changed since this was set" not in page
    assert f'href="/tenants/{alpha}/data/mfa#t-accounts-with-a-gap"' in page

    rows = list(csv.reader(io.StringIO(client.get(f"/tenants/{alpha}/evidence.csv").get_data(as_text=True))))
    legacy = next(r for r in rows if r[1] == "Is legacy authentication blocked?")
    assert legacy[2] == "Yes" and legacy[3] == "Blocked at the firewall." and legacy[4].startswith("jack, ")

    sheet = load_workbook(BytesIO(client.get(f"/tenants/{alpha}/evidence.xlsx").data))["Cyber insurance"]
    values = [[c.value for c in row] for row in sheet.iter_rows()]
    legacy = next(r for r in values if "Are legacy authentication protocols that can't use MFA blocked?" in r)
    assert "Yes" in legacy and "Blocked at the firewall." in legacy and any("jack" in str(v) for v in legacy)

    if export.pdf_available():
        assert client.get(f"/tenants/{alpha}/evidence.pdf").data[:4] == b"%PDF"
    monkeypatch.setattr(export, "pdf_available", lambda: False)  # the same template as a print page
    printed = client.get(f"/tenants/{alpha}/evidence.pdf").get_data(as_text=True)
    assert "Blocked at the firewall." in printed and "Set by jack on" in printed
    assert "Suggested from the data" not in printed  # documents never show the replaced suggestion

    client.post("/logout")
    login(client, "client")
    page = client.get(f"/tenants/{alpha}/evidence").get_data(as_text=True)
    assert "Blocked at the firewall." in page and "Suggested from the data" not in page


def test_data_changed_hint(client, db, two_tenants):
    alpha, _ = two_tenants
    seed_tenant(db, alpha)
    # Set against a suggestion of Yes; the data now suggests No, so the override stays but staff are told.
    evidence_answers.set_override(db, alpha, "legacy_auth", YES, "", YES, "jack")
    a = _answers(db, alpha)["legacy_auth"]
    assert a.answer == YES and a.data_changed
    make_user(db, "admin")
    login(client, "admin")
    page = client.get(f"/tenants/{alpha}/evidence").get_data(as_text=True)
    assert "Data changed since this was set" in page


def test_overrides_go_with_the_tenant(db, two_tenants):
    from officesentry import lifecycle
    alpha, _ = two_tenants
    evidence_answers.set_override(db, alpha, "backups", YES, "Daily to a separate provider.", NOT_KNOWN, "jack")
    lifecycle.delete_tenant(db, store.get_tenant(db, alpha), "jack")
    assert db.execute("SELECT COUNT(*) FROM evidence_answer_overrides").fetchone()[0] == 0
