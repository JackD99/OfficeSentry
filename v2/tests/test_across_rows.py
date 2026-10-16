"""All tenants: every row of a detail report table, totals on the rollup, and the checks matrix."""

from __future__ import annotations

import csv
import io
from datetime import UTC, datetime

import pytest
from openpyxl import load_workbook

from officesentry import findings
from officesentry.reports import data_reports, rollups
from officesentry.reports.data_reports import Section
from officesentry.reports.identity_reports import Money
from tests.conftest import login, make_user
from tests.seed import seed_tenant


def _csv(body: bytes) -> list[list[str]]:
    return list(csv.reader(io.StringIO(body.decode("utf-8-sig"))))


@pytest.fixture
def widgets(monkeypatch):
    """A detail report whose rows name their tenant: two widgets each, one flagged."""

    def builder(snap, now, report):
        name = snap.tenant["name"].split()[0]
        s = Section("Widgets", ["Widget", "State", "Count"])
        s.add([f"{name} w1", "ok", 1])
        s.add([f"{name} w2", "bad", 2], "high")
        report.sections = [s, Section("Settings", ["Setting", "Value"], rows=[["x", name]], flags=[None])]
        report.highlights = [
            data_reports.Highlight("Widgets", len(s.rows)),
            data_reports.Highlight("Saving / month", Money(10.5, "GBP")),
            data_reports.Highlight("Widget score", 7),
        ]

    monkeypatch.setitem(data_reports.REPORTS, "widgets", ("Widgets", "Test widgets", (), builder))
    return "widgets"


def _tenants(db):
    return db.execute("SELECT * FROM tenants ORDER BY name").fetchall()


# -- Every row -------------------------------------------------------------------------


def test_all_rows_stacks_each_tenants_rows_with_a_tenant_column(db, two_tenants, widgets):
    data = rollups.all_rows(db, _tenants(db), "widgets", href=lambda t: f"/tenants/{t['id']}/data/widgets")
    widgets_table = data.report.sections[0]
    assert widgets_table.columns == ["Client", "Widget", "State", "Count"]
    assert [r[:2] for r in widgets_table.rows] == [
        ["Alpha Ltd", "Alpha w1"],
        ["Alpha Ltd", "Alpha w2"],
        ["Bravo Inc", "Bravo w1"],
        ["Bravo Inc", "Bravo w2"],
    ]
    assert widgets_table.flags == [None, "high", None, "high"]
    assert widgets_table.rows[0][0].page_href.endswith("/data/widgets#t-widgets")  # opens the tenant's own table
    assert data.report.sections[1].rows == [["Alpha Ltd", "x", "Alpha"], ["Bravo Inc", "x", "Bravo"]]
    assert not data.truncated("Widgets")


def test_every_row_leaves_out_what_the_client_column_has_no_room_for(client, db, two_tenants, widgets, monkeypatch):
    """A column a table names in optional_across starts under Columns on Every row only: the tenant's own page
    has the room the Client column takes, and downloads keep it."""
    builder = data_reports.REPORTS["widgets"][3]

    def narrower(snap, now, report):
        builder(snap, now, report)
        report.sections[0].optional_across = ("Count",)

    monkeypatch.setitem(data_reports.REPORTS, "widgets", ("Widgets", "Test widgets", (), narrower))
    data = rollups.all_rows(db, _tenants(db), "widgets")
    assert data.report.sections[0].optional == ("Count",)
    make_user(db, "admin")
    login(client, "admin")
    alpha = _tenants(db)[0]["id"]
    rows = client.get("/across/widgets/rows/widgets").get_data(as_text=True)
    own = client.get(f"/tenants/{alpha}/data/widgets").get_data(as_text=True)
    assert ">Count<" not in rows.split("<thead>", 1)[1].split("</thead>", 1)[0]
    assert ">Count<" in own.split("<thead>", 1)[1].split("</thead>", 1)[0]
    assert "Count" in _csv(client.get("/across/widgets/rows/widgets.csv").data)[0]


def test_all_rows_caps_the_rows_and_says_so(client, db, two_tenants, widgets, monkeypatch):
    monkeypatch.setattr(rollups, "MAX_ROWS", 3)
    data = rollups.all_rows(db, _tenants(db), "widgets")
    assert len(data.report.sections[0].rows) == 3 and data.truncated("Widgets") == 4
    assert "first 3 of 4 rows" in rollups.one_table(data, "Widgets").notice
    make_user(db, "admin")
    login(client, "admin")
    page = client.get("/across/widgets/rows/widgets").get_data(as_text=True)
    assert "Only the first 3 of 4 rows are here." in page


def test_rows_page_shows_every_tenant_and_applies_filters(client, db, two_tenants, widgets):
    make_user(db, "admin")
    login(client, "admin")
    resp = client.get("/across/widgets/rows?f=widgets/state:bad")
    assert resp.status_code == 302
    assert resp.headers["Location"].endswith("/across/widgets/rows/widgets?f=widgets%2Fstate%3Abad")
    assert client.get("/across/widgets/rows?table=settings").headers["Location"].endswith("/rows/settings")
    page = client.get("/across/widgets/rows/widgets").get_data(as_text=True)
    assert ">Client</a>" in page and "Every row" in page
    assert all(w in page for w in ("Alpha w1", "Alpha w2", "Bravo w1", "Bravo w2"))
    page = client.get("/across/widgets/rows/widgets?f=widgets/state:bad").get_data(as_text=True)
    assert "Alpha w2" in page and "Bravo w2" in page and "Alpha w1" not in page and "Bravo w1" not in page
    page = client.get("/across/widgets/rows/widgets?f=widgets/client:Bravo Inc").get_data(as_text=True)
    assert "Bravo w1" in page and "Alpha w1" not in page
    # Links saved before the column was renamed still filter on it.
    page = client.get("/across/widgets/rows/widgets?f=widgets/tenant:Bravo Inc").get_data(as_text=True)
    assert "Bravo w1" in page and "Alpha w1" not in page
    assert client.get("/across/widgets/rows/nope").status_code == 404
    assert client.get("/across/nope/rows").status_code == 404
    # The one row per tenant page offers the every-row view.
    assert "/across/widgets/rows" in client.get("/across/widgets").get_data(as_text=True)


def test_rows_downloads_keep_the_view(client, db, two_tenants, widgets):
    make_user(db, "admin")
    login(client, "admin")
    resp = client.get("/across/widgets/rows/widgets.csv?f=widgets/state:bad&hide=widgets/count")
    assert resp.mimetype == "text/csv" and "attachment" in resp.headers["Content-Disposition"]
    rows = _csv(resp.data)
    assert rows[0] == ["Client", "Widget", "State", "Flag"]
    assert rows[1:] == [["Alpha Ltd", "Alpha w2", "bad", "high"], ["Bravo Inc", "Bravo w2", "bad", "high"]]
    resp = client.get("/across/widgets/rows/widgets.xlsx?sort=widgets/widget:desc")
    assert resp.status_code == 200
    wb = load_workbook(io.BytesIO(resp.data))
    assert wb["Summary"]["A2"].value == "All clients"
    sheet = wb.worksheets[1]
    values = [[c.value for c in row] for row in sheet.iter_rows()]
    header = next(i for i, row in enumerate(values) if row[0] == "Client")
    assert values[header][:3] == ["Client", "Widget", "State"]
    assert [row[1] for row in values[header + 1 : header + 5]] == ["Bravo w2", "Bravo w1", "Alpha w2", "Alpha w1"]
    assert len(wb.sheetnames) == 2  # only the table asked for
    assert client.get("/across/widgets/rows/widgets.pdf").status_code == 200
    assert db.execute("SELECT COUNT(*) FROM audit_log WHERE action = 'rollup_downloaded'").fetchone()[0] == 3


def test_rows_views_can_be_saved_and_open_again(client, db, two_tenants, widgets):
    make_user(db, "admin")
    login(client, "admin")
    q = "f=widgets/state:bad"
    resp = client.post("/across/widgets/rows/widgets/views", data={"name": "Bad ones", "query": q})
    assert resp.status_code == 302 and resp.headers["Location"].endswith(
        "/across/widgets/rows/widgets?" + q.replace("/", "%2F").replace(":", "%3A")
    )
    row = db.execute("SELECT * FROM saved_views").fetchone()
    assert row["report_key"] == "widgets" and row["name"] == "Bad ones"
    page = client.get("/across/widgets/rows/widgets").get_data(as_text=True)
    assert "Bad ones" in page
    assert client.post("/across/nope/rows/x/views", data={"name": "x", "query": q}).status_code == 404


def test_rows_only_include_visible_tenants(client, db, two_tenants, widgets):
    alpha, _ = two_tenants
    make_user(db, "analyst", role="analyst", tenants=[alpha])
    login(client, "analyst")
    page = client.get("/across/widgets/rows/widgets").get_data(as_text=True)
    assert "Alpha w1" in page and "Bravo" not in page
    assert b"Bravo" not in client.get("/across/widgets/rows/widgets.csv").data


def test_rows_mode_on_a_real_report(client, db, two_tenants):
    alpha, bravo = two_tenants
    seed_tenant(db, alpha)
    seed_tenant(db, bravo)
    make_user(db, "admin")
    login(client, "admin")
    report = data_reports.build(db, db.execute("SELECT * FROM tenants WHERE id = ?", (alpha,)).fetchone(), "guests")
    data = rollups.all_rows(db, _tenants(db), "guests")
    first = data.report.sections[0]
    assert len(first.rows) == 2 * len(report.sections[0].rows)
    resp = client.get("/across/guests/rows", follow_redirects=True)
    page = resp.get_data(as_text=True)
    assert resp.status_code == 200 and "Alpha Ltd" in page and "Bravo Inc" in page


# -- Totals -----------------------------------------------------------------------------


def test_rollup_totals_add_up_counts_and_money_only(db, two_tenants, widgets):
    rollup = rollups.build(db, _tenants(db), "widgets")
    total = rollup.total_row()
    assert total.values["_flagged"] == 2
    assert total.values["c:Widgets"] == 4
    assert total.values["c:Saving / month"] == 21.0 and isinstance(total.values["c:Saving / month"], Money)
    assert "c:Widget score" not in total.values  # a score doesn't add up

    cols = [
        rollups.Column("open", "Open", "sev"),
        rollups.Column("pct", "MFA", "pct"),
        rollups.Column("money", "Waste", "money"),
        rollups.Column("ctl", "Control", "status"),
        rollups.Column("txt", "Ratio", "text"),
    ]
    a = {
        "open": {"critical": 1, "high": 2, "medium": 0, "low": 0},
        "pct": 50,
        "money": 10.25,
        "ctl": "In place",
        "txt": "1 of 3",
    }
    b = {
        "open": {"critical": 0, "high": 1, "medium": 3, "low": 1},
        "pct": 70,
        "money": 5,
        "ctl": "Not in place",
        "txt": "2 of 3",
    }
    rows = [
        rollups.RollupRow({"name": "A"}, a, tones={"ctl": "ok"}),
        rollups.RollupRow({"name": "B"}, b, tones={"ctl": "bad"}),
        rollups.RollupRow({"name": "C"}, note="No data collected yet."),
    ]
    total = rollups.Rollup("x", "X", "", cols, rows, datetime.now(UTC)).total_row()
    assert total.values == {"open": {"critical": 1, "high": 3, "medium": 3, "low": 1}, "money": 15.25, "ctl": "1 of 2"}
    one = rollups.Rollup("x", "X", "", cols, rows[:1], datetime.now(UTC))
    assert one.total_row() is None


def test_rollup_page_shows_the_total_row(client, db, two_tenants, widgets):
    make_user(db, "admin")
    login(client, "admin")
    page = client.get("/across/widgets").get_data(as_text=True)
    assert 'class="total-row"' in page and "2 clients" in page
    assert client.get("/across/widgets.pdf").status_code == 200


# -- Checks matrix ------------------------------------------------------------------------


def test_check_matrix_reads_results_and_accepted_risks(db, two_tenants):
    alpha, bravo = two_tenants
    seed_tenant(db, alpha)
    tenant = db.execute("SELECT * FROM tenants WHERE id = ?", (alpha,)).fetchone()
    findings.sync(db, tenant)
    matrix = rollups.check_matrix(db, _tenants(db))
    by_id = {r.tenant["id"]: r for r in matrix.rows}
    assert by_id[bravo].note == "No data collected yet." and matrix.rows[-1].tenant["id"] == bravo
    alpha_row = by_id[alpha]
    assert set(alpha_row.cells) == {k for k, _t, _c in matrix.checks}
    failing = [k for k, c in alpha_row.cells.items() if c.status == "fail" and c.count]
    assert failing
    key = failing[0]
    for row in db.execute("SELECT id FROM findings WHERE tenant_id = ? AND rule_key = ?", (alpha, key)):
        findings.accept(db, row["id"], "admin", "Intended", None)
    again = {r.tenant["id"]: r for r in rollups.check_matrix(db, _tenants(db)).rows}[alpha]
    assert again.cells[key].status == "accepted" and again.cells[key].label == "Accepted"
    table = matrix.to_table()
    assert table.columns[:3] == ["Client", "Tenant ID", "Failing"] and len(table.rows) == 2
    assert all(len(r) == len(table.columns) for r in table.rows)


def test_check_matrix_page_lists_only_visible_tenants(client, db, two_tenants):
    alpha, bravo = two_tenants
    seed_tenant(db, alpha)
    seed_tenant(db, bravo)
    make_user(db, "analyst", role="analyst", tenants=[alpha])
    login(client, "analyst")
    page = client.get("/across/checks").get_data(as_text=True)
    assert "Alpha Ltd" in page and "Bravo" not in page
    assert f"/tenants/{alpha}/findings#check-" in page
    page = client.get("/across/checks?category=email").get_data(as_text=True)
    assert "Alpha Ltd" in page and "Every category" in page
    assert 'href="/across/checks.csv?category=email"' in page and "Email checks for every client" in page
    rows = _csv(client.get("/across/checks.csv").data)
    assert [r[0] for r in rows[1:]] == ["Alpha Ltd"]
    assert client.get("/across/checks.xlsx").status_code == 200
    # The downloads cover what the page shows: one category's checks, with the counts over those checks.
    email = _csv(client.get("/across/checks.csv?category=email").data)
    assert email[0][:6] == rows[0][:6] and email[0][6:] and all(c.startswith("Email: ") for c in email[0][6:])
    assert len(email[0]) < len(rows[0]) and int(email[1][2]) <= int(rows[1][2])
    resp = client.get("/across/checks.xlsx?category=email")
    assert resp.status_code == 200 and 'filename="all-clients-checks-email.xlsx"' in resp.headers["Content-Disposition"]
    assert _csv(client.get("/across/checks.csv?category=nonsense").data)[0] == rows[0]  # unknown: every check
    # Downloads without words of their own get the usual ones for their format.
    assert 'title="This page as a PDF"' in client.get("/changes").get_data(as_text=True)


def test_clients_cannot_open_cross_tenant_rows_or_checks(client, db, two_tenants, widgets):
    alpha, _ = two_tenants
    make_user(db, "client", role="client", tenants=[alpha])
    login(client, "client")
    for url in (
        "/across/checks",
        "/across/checks.csv",
        "/across/widgets/rows",
        "/across/widgets/rows/widgets",
        "/across/widgets/rows/widgets.csv",
        "/across/widgets/rows/widgets.xlsx",
        "/across/widgets/rows/widgets.pdf",
    ):
        assert client.get(url).status_code == 403, url
    assert client.post("/across/widgets/rows/widgets/views", data={"name": "x", "query": "f=a/b:c"}).status_code == 403


def test_library_links_findings_across_tenants_to_the_checks_matrix(client, db, two_tenants):
    make_user(db, "admin")
    login(client, "admin")
    page = client.get("/reports").get_data(as_text=True)
    assert "/across/checks" in page
