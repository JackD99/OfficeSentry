"""Hybrid identity: directory sync collector, sync and admin checks, and the Hybrid identity report."""

from datetime import UTC, datetime
from io import BytesIO

import pytest
import responses
from openpyxl import load_workbook

from officesentry import store
from officesentry.graph import GRAPH_ROOT, GraphClient
from officesentry.reports import data_reports, evidence
from officesentry.rules import evaluate
from officesentry.snapshot import TenantSnapshot
from tests import graphdata as gd
from tests import hybriddata as hd
from tests.conftest import login, make_user
from tests.seed import seed_tenant
from tests.test_collectors import V1, _metrics, _run

BETA = f"{GRAPH_ROOT}/beta"


@pytest.fixture
def factory():
    return lambda tenant, connection: GraphClient(lambda: "t", sleep=lambda s: None)


@pytest.fixture
def hybrid(monkeypatch):
    """Make the seed tenant a synced one: Entra Connect org, a federated domain, users from AD."""
    monkeypatch.setattr(gd, "ORGANIZATION", hd.ORGANIZATION)
    monkeypatch.setattr(hd, "DOMAINS", hd.HYBRID_DOMAINS)


def _results(db, tenant_id):
    return {r.key: r for r in evaluate(TenantSnapshot(db, store.get_tenant(db, tenant_id)))}


def _org(last_sync=0.3, password_sync=0.3):
    return {
        **hd.ORGANIZATION,
        "onPremisesLastSyncDateTime": None if last_sync is None else hd.hours_ago(last_sync),
        "onPremisesLastPasswordSyncDateTime": None if password_sync is None else hd.hours_ago(password_sync),
    }


def _collect(db, tenant_id, factory, org, beta=None, domains=None):
    with responses.RequestsMock(assert_all_requests_are_fired=False) as rsps:
        rsps.get(f"{V1}/organization", json={"value": [org]})
        rsps.get(f"{BETA}/organization", **(beta or {"json": {"value": [org]}}))
        rsps.get(f"{V1}/domains", json={"value": hd.HYBRID_DOMAINS if domains is None else domains})
        run_id, status = _run(db, tenant_id, "directory_sync", factory)
        return run_id, status, [c.request.url for c in rsps.calls]


# -- Collector ---------------------------------------------------------------------------


def test_collects_sync_status_and_domains(db, two_tenants, factory):
    alpha, _ = two_tenants
    run_id, status, urls = _collect(db, alpha, factory, hd.ORGANIZATION)
    assert status == "succeeded"
    sync = store.run_objects(db, run_id, "sync_status")[0]
    assert sync["enabled"] is True and sync["passwordSyncChecked"] is True
    assert sync["lastSyncDateTime"] == hd.ORGANIZATION["onPremisesLastSyncDateTime"]
    assert sync["lastPasswordSyncDateTime"] == hd.ORGANIZATION["onPremisesLastPasswordSyncDateTime"]
    domains = {d["_id"]: d for d in store.run_objects(db, run_id, "domain")}
    assert domains["x.co.uk"]["authenticationType"] == "Federated" and domains["x.com"]["isDefault"] is True
    m = _metrics(db, run_id)
    assert (m["sync.enabled"], m["sync.federated_domains"]) == (1, 1)
    assert m["sync.hours_since_last_sync"] < 1 and m["sync.hours_since_password_sync"] < 1
    # Timestamps from v1.0; only the password sync time comes from beta. Nothing else is read.
    assert urls[0].startswith(f"{V1}/organization") and "onPremisesLastSyncDateTime" in urls[0]
    assert urls[1].startswith(f"{BETA}/organization") and "onPremisesLastPasswordSyncDateTime" in urls[1]
    assert urls[2].startswith(f"{V1}/domains") and "authenticationType" in urls[2]
    assert not any("onPremisesSynchronization" in u for u in urls)


def test_cloud_only_tenant_skips_password_sync(db, two_tenants, factory):
    alpha, _ = two_tenants
    run_id, status, urls = _collect(db, alpha, factory, gd.ORGANIZATION, domains=hd.DOMAINS)
    assert status == "succeeded"
    assert not any(u.startswith(BETA) for u in urls)
    sync = store.run_objects(db, run_id, "sync_status")[0]
    assert sync["enabled"] is False and sync["onPremisesSyncEnabled"] is None
    assert _metrics(db, run_id) == {"sync.enabled": 0, "sync.federated_domains": 0}


def test_unreadable_password_sync_is_partial(db, two_tenants, factory):
    alpha, _ = two_tenants
    beta = {"status": 403, "json": {"error": {"code": "Forbidden", "message": "Insufficient privileges"}}}
    assert _collect(db, alpha, factory, hd.ORGANIZATION, beta=beta)[1] == "partial"
    r = _results(db, alpha)
    assert r["directory_sync_running"].status == "pass"
    assert r["password_hash_sync"].status == "unavailable"
    assert "couldn't be read" in r["password_hash_sync"].summary


# -- Checks ------------------------------------------------------------------------------


@pytest.mark.parametrize(
    "hours, status, severity, summary",
    [
        (0.3, "pass", None, "Last sync less than an hour ago"),
        (5, "fail", "high", "Not syncing: last sync 5 hours ago"),
        (30, "fail", "critical", "Not syncing: last sync 30 hours ago"),
        (24 * 9, "fail", "critical", "Not syncing: last sync 9 days ago"),
        (None, "fail", "high", "Turned on but has never synced"),
    ],
)
def test_directory_sync_running(db, two_tenants, factory, hours, status, severity, summary):
    alpha, _ = two_tenants
    _collect(db, alpha, factory, _org(last_sync=hours, password_sync=hours))
    result = _results(db, alpha)["directory_sync_running"]
    assert (result.status, result.severity, result.summary) == (status, severity, summary)
    assert result.category == "Hybrid identity"


def test_sync_checks_unavailable_for_cloud_only_and_uncollected(db, two_tenants, factory):
    alpha, bravo = two_tenants
    _collect(db, alpha, factory, gd.ORGANIZATION)
    r = _results(db, alpha)
    for key in ("directory_sync_running", "password_hash_sync"):
        assert r[key].status == "unavailable" and r[key].summary.startswith("Cloud-only tenant")
    r = _results(db, bravo)
    assert all(
        r[k].status == "unavailable" for k in ("directory_sync_running", "password_hash_sync", "cloud_only_admins")
    )


def test_password_hash_sync_on(db, two_tenants, factory):
    alpha, _ = two_tenants
    _collect(db, alpha, factory, hd.ORGANIZATION)
    result = _results(db, alpha)["password_hash_sync"]
    assert (result.status, result.summary) == ("pass", "Last password sync less than an hour ago")


def test_password_hash_sync_never_recorded(db, two_tenants, factory):
    alpha, _ = two_tenants
    _collect(db, alpha, factory, _org(password_sync=None))
    result = _results(db, alpha)["password_hash_sync"]
    assert (result.status, result.severity, result.summary) == ("fail", "medium", "No password hash sync recorded")
    assert "no record of a password hash sync" in result.findings[0].detail
    assert "x.co.uk is federated" in result.findings[0].detail  # no cloud fallback for that domain


def test_password_hash_sync_stopped(db, two_tenants, factory):
    alpha, _ = two_tenants
    _collect(db, alpha, factory, _org(password_sync=72), domains=hd.DOMAINS)
    result = _results(db, alpha)["password_hash_sync"]
    assert (result.status, result.summary) == ("fail", "Password hash sync has stopped")
    assert result.findings[0].detail == (
        "Last password sync 3 days ago, though directory sync ran less than an "
        "hour ago: password hash sync has stopped or been turned off"
    )


def test_whole_sync_stopped_is_flagged_once(db, two_tenants, factory):
    alpha, _ = two_tenants
    _collect(db, alpha, factory, _org(last_sync=30, password_sync=30))
    r = _results(db, alpha)
    assert r["directory_sync_running"].severity == "critical"
    assert r["password_hash_sync"].status == "pass"  # in step with directory sync; the outage is one finding


def test_synced_admins_fail(db, two_tenants, hybrid):
    alpha, _ = two_tenants
    seed_tenant(db, alpha, users=hd.USERS)
    result = _results(db, alpha)["cloud_only_admins"]
    assert result.status == "fail" and result.summary == "2 admin accounts synced from on-premises AD"
    assert [(f.severity, f.subject, f.detail) for f in result.findings] == [
        ("high", "ann@x.com", "Global Administrator; synced from on-premises AD"),
        ("high", "eve@x.com", "Global Administrator (eligible); synced from on-premises AD"),
    ]  # the break-glass account is cloud-only; the guest admin and the app aren't synced users


def test_disabled_synced_admin_is_medium(db, two_tenants, hybrid):
    alpha, _ = two_tenants
    users = [u if u["id"] != "u5" else {**u, "accountEnabled": False} for u in hd.USERS]
    seed_tenant(db, alpha, users=users)
    findings = {f.subject: f for f in _results(db, alpha)["cloud_only_admins"].findings}
    assert findings["eve@x.com"].severity == "medium" and "AD can re-enable it" in findings["eve@x.com"].detail


def test_cloud_only_admins_pass(db, two_tenants, hybrid):
    alpha, _ = two_tenants
    users = [{**u, "onPremisesSyncEnabled": None} if u["id"] in ("u1", "u5") else u for u in hd.USERS]
    seed_tenant(db, alpha, users=users)
    result = _results(db, alpha)["cloud_only_admins"]
    assert (result.status, result.summary) == ("pass", "All 4 privileged accounts are cloud-only")


def test_cloud_only_tenant_passes_admin_check(db, two_tenants):
    alpha, _ = two_tenants
    seed_tenant(db, alpha)
    r = _results(db, alpha)
    assert (r["cloud_only_admins"].status, r["cloud_only_admins"].summary) == (
        "pass",
        "Cloud-only tenant: no accounts are synced from on-premises AD",
    )
    assert r["directory_sync_running"].status == "unavailable"


@responses.activate
def test_cloud_only_admins_needs_user_data(db, two_tenants, factory):
    alpha, _ = two_tenants
    responses.get(f"{V1}/roleManagement/directory/roleDefinitions", json={"value": gd.ROLE_DEFINITIONS})
    responses.get(f"{V1}/roleManagement/directory/roleAssignments", json={"value": gd.ROLE_ASSIGNMENTS})
    responses.get(f"{V1}/roleManagement/directory/roleEligibilitySchedules", json={"value": gd.ROLE_ELIGIBILITY})
    assert _run(db, alpha, "admin_roles", factory)[1] == "succeeded"
    result = _results(db, alpha)["cloud_only_admins"]
    assert result.status == "unavailable" and "user data" in result.summary


def test_insurance_privileged_control_includes_synced_admins(db, two_tenants, hybrid):
    alpha, _ = two_tenants
    assert "cloud_only_admins" in {c.key: c for c in evidence.CONTROLS}["privileged"].rules
    seed_tenant(db, alpha, users=hd.USERS)
    control = {c.control.key: c for c in evidence.build(db, store.get_tenant(db, alpha)).controls}["privileged"]
    assert any("synced from on-premises AD" in g for g in control.gaps)


# -- Report ------------------------------------------------------------------------------


def test_hybrid_report(db, two_tenants, hybrid):
    alpha, _ = two_tenants
    seed_tenant(db, alpha, users=hd.USERS)
    report = data_reports.build(db, store.get_tenant(db, alpha), "hybrid", datetime.now(UTC))
    assert report.unavailable is None
    sections = {s.title: s for s in report.sections}
    assert list(sections) == ["Sync status", "Domains", "Synced admin accounts", "Accounts by source"]
    status = {r[0]: r[1:] for r in sections["Sync status"].rows}
    assert status["Directory sync"][0] == "On" and status["Federated domains"][0] == "x.co.uk"
    # The time has its own "Last ..." column, which What changed ignores, so a running sync isn't a change.
    assert status["Directory sync runs"][0] == "Running"
    assert status["Directory sync runs"][1].endswith("(less than an hour ago)")
    assert not sections["Sync status"].flagged
    figures = {t.label: (t.value, t.tone) for t in report.tiles}
    assert figures == {
        "Directory sync": ("Running", "ok"),
        "Synced accounts": (4, None),
        "Synced admins": (2, "bad"),
        "Federated domains": (1, None),
    }
    domains = sections["Domains"]
    assert [r[0] for r in domains.rows] == ["x.com", "x.co.uk", "x.onmicrosoft.com"]  # default first
    assert dict(zip(domains.columns, domains.rows[1], strict=True))["Sign-in"] == "Federated"
    admins = sections["Synced admin accounts"]
    assert [r[:2] for r in admins.rows] == [
        ["ann@x.com", "Global Administrator"],
        ["eve@x.com", "Global Administrator (eligible)"],
    ]
    assert admins.flags == ["high", "high"]
    assert sections["Accounts by source"].rows == [["Members", 4, 1, 5], ["Guests", 0, 1, 1], ["All accounts", 4, 2, 6]]


def test_hybrid_report_flags_stale_sync(db, two_tenants, factory):
    alpha, _ = two_tenants
    _collect(db, alpha, factory, _org(last_sync=30, password_sync=None))
    sync = data_reports.build(db, store.get_tenant(db, alpha), "hybrid").sections[0]
    flags = {row[0]: flag for row, flag in zip(sync.rows, sync.flags, strict=True)}
    assert flags["Directory sync runs"] == "critical" and flags["Password hash sync"] == "medium"
    status = {r[0]: r[1:] for r in sync.rows}
    assert status["Directory sync runs"][0] == "Stopped" and status["Password hash sync"] == ["No record", "Never"]


def test_hybrid_report_cloud_only(db, two_tenants):
    alpha, _ = two_tenants
    seed_tenant(db, alpha)
    report = data_reports.build(db, store.get_tenant(db, alpha), "hybrid")
    sections = {s.title: s for s in report.sections}
    assert list(sections) == ["Sync status", "Domains", "Accounts by source"]
    assert sections["Sync status"].intro.startswith("Cloud-only tenant")
    assert sections["Sync status"].rows[0][:2] == ["Directory sync", "Never turned on"]
    assert sections["Accounts by source"].rows[-1] == ["All accounts", 0, 3, 3]


def test_sync_turned_off(db, two_tenants, factory):
    alpha, _ = two_tenants
    _collect(db, alpha, factory, {**hd.ORGANIZATION, "onPremisesSyncEnabled": False})
    assert _results(db, alpha)["directory_sync_running"].summary.startswith("Cloud-only tenant")
    sync = data_reports.build(db, store.get_tenant(db, alpha), "hybrid").sections[0]
    assert sync.rows[0][:2] == ["Directory sync", "Turned off"] and "was turned off" in sync.intro


def test_hybrid_report_without_data(db, two_tenants):
    _, bravo = two_tenants
    report = data_reports.build(db, store.get_tenant(db, bravo), "hybrid")
    assert report.unavailable and not report.sections


def test_hybrid_page_and_xlsx(client, db, two_tenants, hybrid):
    alpha, _ = two_tenants
    seed_tenant(db, alpha, users=hd.USERS)
    make_user(db, "admin")
    login(client, "admin")
    page = client.get(f"/tenants/{alpha}/data/hybrid")
    assert page.status_code == 200
    html = page.get_data(as_text=True)
    assert "Hybrid identity" in html and "Synced admin accounts" in html and "eve@x.com" in html
    resp = client.get(f"/tenants/{alpha}/data/hybrid.xlsx")
    assert resp.status_code == 200
    wb = load_workbook(BytesIO(resp.data))
    assert wb.sheetnames == ["Summary", "Sync status", "Domains", "Synced admin accounts", "Accounts by source"]
