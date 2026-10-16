"""The Findings page on the shared report table: filters, groups, accepted risks, saved views and downloads."""

import csv
import io
from datetime import UTC, datetime, timedelta
from urllib.parse import quote

from openpyxl import load_workbook

from officesentry import findings, store
from officesentry.reports import findings_report, views
from officesentry.rules import RULES
from tests.conftest import login, make_user
from tests.seed import seed_tenant


def _finding(db, tenant_id, rule_key):
    return db.execute(
        "SELECT * FROM findings WHERE tenant_id = ? AND rule_key = ? ORDER BY id", (tenant_id, rule_key)
    ).fetchone()


def _sections(report):
    return {s.title: s for s in report.sections}


def test_builder_splits_open_and_accepted_risks(db, two_tenants):
    alpha, _ = two_tenants
    seed_tenant(db, alpha)
    tenant = store.get_tenant(db, alpha)
    guest = _finding(db, alpha, "guest_admins")
    findings.accept(db, guest["id"], "analyst", "Known partner", None)
    lapsed = _finding(db, alpha, "admin_mfa")
    findings.accept(db, lapsed["id"], "analyst", "Was fine", "2020-01-01T00:00:00Z")  # past its review date

    report = findings_report.build(db, tenant)
    sections = _sections(report)
    open_, accepted = sections["Open findings"], sections["Accepted risks"]
    assert open_.flag_name == "Severity" and open_.group == "Check" and open_.filter_columns == ("Check",)
    assert open_.columns == ["Account or item", "Check", "Detail", "Category", "First seen"]
    assert guest["id"] in accepted.keys and guest["id"] not in open_.keys
    assert lapsed["id"] in open_.keys  # a lapsed acceptance is open again
    assert accepted.rows[accepted.keys.index(guest["id"])][3:5] == ["analyst", "Known partner"]
    unresolved = db.execute(
        "SELECT COUNT(*) FROM findings WHERE tenant_id = ? AND resolved_at IS NULL", (alpha,)
    ).fetchone()[0]
    assert len(open_.rows) + len(accepted.rows) == unresolved == len(open_.keys) + len(accepted.keys)
    assert set(open_.flags) <= {"critical", "high", "medium", "low", "info"}
    # One check's findings sit together, worst check first.
    checks = [row[1] for row in open_.rows]
    assert checks == sorted(checks, key=checks.index)
    order = {s: i for i, s in enumerate(("critical", "high", "medium", "low", "info"))}
    firsts = [open_.flags[checks.index(c)] for c in dict.fromkeys(checks)]
    assert firsts == sorted(firsts, key=order.get)
    tiles = {t.label: t.value for t in report.tiles}
    assert tiles["Open findings"] == len(open_.rows) and tiles["Accepted risks"] == 1
    assert tiles["Critical or high"] == sum(f in ("critical", "high") for f in open_.flags)
    # Each tile leads to its rows; "Critical or high" filters the Severity column.
    links = views.headline_links(report, "/p")
    assert "f=open-findings%2Fflag%3Acritical" in links["tiles"][1] and "high" in links["tiles"][1]
    assert links["tiles"][4].endswith("#t-accepted-risks")


def test_group_rows_by_check_until_sorted(db, two_tenants):
    alpha, _ = two_tenants
    seed_tenant(db, alpha)
    report = findings_report.build(db, store.get_tenant(db, alpha))
    tv = views.table_view(report, 0, views.View(), "/p")
    assert tv.groups and sum(len(g.rows) for g in tv.groups) == len(tv.rows)
    assert all(len(cells) == len(tv.shown) for g in tv.groups for cells, _, _ in g.rows)
    assert "Check" not in [c.name for c in tv.shown]  # the heading says it
    assert {g.anchor for g in tv.groups} >= {"check-guest_admins"}
    assert [k for g in tv.groups for _, _, k in g.rows] == tv.keys
    sorted_tv = views.table_view(report, 0, views.View.parse("sort=open-findings/first-seen:asc"), "/p")
    assert not sorted_tv.groups and "Check" in [c.name for c in sorted_tv.shown]
    # The check filter is offered however many checks there are.
    check = next(c for c in tv.filters if c.name == "Check")
    assert len(check.options) == len(tv.groups)
    assert {c.name for c in tv.filters} >= {"Severity", "Check", "First seen"}


def test_page_filters_and_actions(client, db, two_tenants):
    alpha, bravo = two_tenants
    seed_tenant(db, alpha)
    make_user(db, "analyst", role="analyst", tenants=[alpha])
    login(client, "analyst")
    page = client.get(f"/tenants/{alpha}/findings").get_data(as_text=True)
    assert 'id="t-open-findings"' in page and 'id="t-accepted-risks"' in page
    assert 'id="check-guest_admins"' in page and 'data-group="check-guest_admins"' in page
    assert 'class="sheet-link"' in page and 'data-groups="tbl-open-findings"' in page  # each row opens its sheet
    assert f'href="/tenants/{alpha}/findings.pdf"' in page and f'href="/tenants/{alpha}/findings.xlsx"' in page
    assert f"/tenants/{alpha}/findings.csv?section=1" in page
    high = client.get(f"/tenants/{alpha}/findings?f=open-findings/flag:high").get_data(as_text=True)
    assert "Severity: High" in high
    body = high.split('id="tbl-open-findings"')[1].split("</table>")[0]
    assert 'class="flag-high"' in body and 'class="flag-medium"' not in body and 'class="flag-low"' not in body
    guest = _finding(db, alpha, "guest_admins")
    check = client.get(f"/tenants/{alpha}/findings?f=open-findings/check:{quote(_title(guest))}").get_data(as_text=True)
    assert check.split('id="tbl-open-findings"')[1].count('class="group"') == 1
    resp = client.post(f"/findings/{guest['id']}/accept", data={"note": "Known partner"})
    assert resp.headers["Location"].endswith(f"/tenants/{alpha}/findings#t-accepted-risks")
    page = client.get(f"/tenants/{alpha}/findings").get_data(as_text=True)
    accepted = page.split('id="t-accepted-risks"')[1]
    assert "Known partner" in accepted and f"/findings/{guest['id']}/unaccept" in accepted
    assert client.get(f"/tenants/{bravo}/findings").status_code == 404
    assert client.get(f"/tenants/{bravo}/findings.xlsx").status_code == 404


def _title(row):
    return {r.key: r.title for r in RULES}[row["rule_key"]]


def test_clients_see_findings_without_actions(client, db, two_tenants):
    alpha, _ = two_tenants
    seed_tenant(db, alpha)
    make_user(db, "client", role="client", tenants=[alpha])
    login(client, "client")
    page = client.get(f"/tenants/{alpha}/findings").get_data(as_text=True)
    assert 'id="t-open-findings"' in page and "Accept risk" not in page and "<th>Action</th>" not in page
    assert client.get(f"/tenants/{alpha}/findings.pdf").status_code == 200


def test_lapsed_acceptance_shows_for_review(client, db, two_tenants):
    alpha, _ = two_tenants
    seed_tenant(db, alpha)
    row = _finding(db, alpha, "guest_admins")
    findings.accept(db, row["id"], "analyst", "Was fine", "2020-01-01T00:00:00Z")
    make_user(db, "admin")
    login(client, "admin")
    page = client.get(f"/tenants/{alpha}/findings").get_data(as_text=True)
    assert "now due for review" in page and "Accept again" not in page  # the row says so; its sheet offers it
    assert page.split('id="t-accepted-risks"')[1].count('class="group"') == 0
    sheet = client.get(f"/tenants/{alpha}/findings/{row['id']}").get_data(as_text=True)
    assert "Accept again" in sheet and "now due for review" in sheet


def test_downloads_apply_the_view(client, db, two_tenants):
    alpha, _ = two_tenants
    seed_tenant(db, alpha)
    make_user(db, "admin")
    login(client, "admin")
    view = "f=open-findings/flag:high&hide=open-findings/category"
    rows = list(
        csv.reader(
            io.StringIO(
                client.get(f"/tenants/{alpha}/findings.csv?section=1&{view}").get_data(as_text=True).lstrip("\ufeff")
            )
        )
    )
    assert rows[0] == ["Account or item", "Check", "Detail", "First seen", "Severity"]
    assert len(rows) > 1 and {r[-1] for r in rows[1:]} == {"high"}
    book = load_workbook(io.BytesIO(client.get(f"/tenants/{alpha}/findings.xlsx?{view}").data))
    assert "Open findings" in book.sheetnames and "Accepted risks" in book.sheetnames
    sheet = book["Open findings"]
    assert [c.value for c in sheet[1]] == rows[0] and sheet.max_row == len(rows)
    pdf = client.get(f"/tenants/{alpha}/findings.pdf?{view}")
    assert pdf.status_code == 200
    assert client.get(f"/tenants/{alpha}/findings.csv?section=9").status_code == 404
    assert client.get(f"/tenants/{alpha}/findings.doc").status_code == 404
    actions = [r[0] for r in db.execute("SELECT detail FROM audit_log WHERE action = 'report_downloaded'")]
    assert {"findings.csv", "findings.xlsx", "findings.pdf"} <= set(actions)


def test_resolved_page_and_old_link(client, db, two_tenants):
    alpha, _ = two_tenants
    seed_tenant(db, alpha)
    row = _finding(db, alpha, "guest_admins")
    db.execute(
        "UPDATE findings SET resolved_at = ? WHERE id = ?",
        ((datetime.now(UTC) - timedelta(days=2)).strftime("%Y-%m-%dT%H:%M:%SZ"), row["id"]),
    )
    make_user(db, "admin")
    login(client, "admin")
    resp = client.get(f"/tenants/{alpha}/findings?view=resolved&f=resolved-findings/flag:high")
    assert resp.status_code == 302
    assert resp.headers["Location"].endswith("/findings/resolved?f=resolved-findings%2Fflag%3Ahigh")
    page = client.get(f"/tenants/{alpha}/findings/resolved").get_data(as_text=True)
    assert 'id="t-resolved-findings"' in page and "Accept risk" not in page and "to review" not in page
    assert client.get(f"/tenants/{alpha}/findings/resolved.xlsx").status_code == 200
    # The old Accepted risks tab is now a section of the open page.
    assert 'id="t-accepted-risks"' in client.get(f"/tenants/{alpha}/findings?view=accepted").get_data(as_text=True)


def test_saved_views_on_findings(client, db, two_tenants):
    alpha, bravo = two_tenants
    seed_tenant(db, alpha)
    make_user(db, "admin", all_tenants=True)
    login(client, "admin")
    query = "f=open-findings/flag:high"
    resp = client.post("/reports/findings/views", data={"name": "High only", "query": query, "tenant_id": alpha})
    assert resp.headers["Location"].endswith(f"/tenants/{alpha}/findings?f=open-findings%2Fflag%3Ahigh")
    saved = db.execute("SELECT report_key, query FROM saved_views").fetchone()
    assert tuple(saved) == ("findings", "f=open-findings%2Fflag%3Ahigh")
    page = client.get(f"/tenants/{bravo}/findings").get_data(as_text=True)
    assert "High only" in page and f"/tenants/{bravo}/findings?f=open-findings%2Fflag%3Ahigh" in page
    library = client.get(f"/reports?tenant={alpha}").get_data(as_text=True)
    assert f"/tenants/{alpha}/findings?f=open-findings%2Fflag%3Ahigh" in library
    resp = client.post(
        "/reports/findings-resolved/views",
        data={"name": "Old", "query": "sort=resolved-findings/resolved:desc", "tenant_id": alpha},
    )
    assert "/findings/resolved?" in resp.headers["Location"]
    assert client.post("/reports/nope/views", data={"name": "x", "query": query, "tenant_id": alpha}).status_code == 404


def test_an_impossible_review_date_is_refused(client, db, two_tenants):
    """A review date like 31 February never comes round, so accepting with it would hide the finding for good."""
    alpha, _ = two_tenants
    seed_tenant(db, alpha)
    make_user(db, "analyst1", role="analyst", tenants=[alpha])
    login(client, "analyst1")
    guest = _finding(db, alpha, "guest_admins")
    for until in ("2026-02-31", "2026-13-01"):
        resp = client.post(f"/findings/{guest['id']}/accept", data={"note": "Partner", "until": until})
        assert resp.status_code == 302 and "#t-accepted-risks" not in resp.headers["Location"]
        assert _finding(db, alpha, "guest_admins")["accepted_by"] is None
    page = client.get(f"/tenants/{alpha}/findings").get_data(as_text=True)
    assert "isn&#39;t a real date" in page
    next_year = datetime.now(UTC).year + 1  # a review date in the past is refused too (below)
    client.post(f"/findings/{guest['id']}/accept", data={"note": "Partner", "until": f"{next_year}-02-28"})
    assert _finding(db, alpha, "guest_admins")["accepted_by"] == "analyst1"


def test_a_review_date_of_today_or_earlier_is_refused(client, db, two_tenants):
    """An acceptance runs out at the start of its review date, so one of today or earlier would leave the finding
    open (still alerting and counted) while the page said the risk was accepted."""
    alpha, _ = two_tenants
    seed_tenant(db, alpha)
    make_user(db, "analyst1", role="analyst", tenants=[alpha])
    login(client, "analyst1")
    guest = _finding(db, alpha, "guest_admins")
    today = datetime.now(UTC).date()
    tomorrow = (today + timedelta(days=1)).isoformat()
    sheet = client.get(f"/tenants/{alpha}/findings/{guest['id']}").get_data(as_text=True)
    assert f'name="until" min="{tomorrow}"' in sheet
    for until in (today.isoformat(), (today - timedelta(days=30)).isoformat()):
        resp = client.post(f"/findings/{guest['id']}/accept", data={"note": "Partner", "until": until})
        assert resp.status_code == 302 and "#t-accepted-risks" not in resp.headers["Location"]
        assert _finding(db, alpha, "guest_admins")["accepted_by"] is None
        page = client.get(resp.headers["Location"]).get_data(as_text=True)
        assert "the review date has to be after today" in page and "Risk accepted." not in page
    resp = client.post(f"/findings/{guest['id']}/accept", data={"note": "Partner", "until": tomorrow})
    assert resp.headers["Location"].endswith("#t-accepted-risks")
    assert findings.lifecycle_map(db, alpha)[("guest_admins", guest["object_id"])].accepted
