"""A client's Timeline (``reports/timeline.py``, ``web/timeline.py``): the year month by month from what Office
Sentry keeps, with each person seeing only the parts they open on that client."""

from __future__ import annotations

from datetime import UTC, date, datetime, timedelta
from io import BytesIO

import pytest
from openpyxl import load_workbook

from officesentry import events, issued
from officesentry.reports import charts, timeline
from tests import adminactivitydata as ad
from tests.conftest import login, make_user
from tests.seed import seed_tenant
from tests.test_pdf_export import needs_pdf
from tests.test_roles import role_user

ISO = "%Y-%m-%dT%H:%M:%SZ"
NOW = datetime.now(UTC).replace(microsecond=0)


def ago(days: float) -> str:
    return (NOW - timedelta(days=days)).strftime(ISO)


@pytest.fixture
def alpha(db, two_tenants):
    """A seeded tenant whose first collection (and so its first look at findings) was 60 days ago."""
    alpha, _ = two_tenants
    seed_tenant(db, alpha)
    db.execute("UPDATE findings SET first_seen = ?, last_seen = ? WHERE tenant_id = ?", (ago(60), ago(0), alpha))
    db.execute("UPDATE runs SET created_at = ?, finished_at = ? WHERE tenant_id = ?", (ago(60), ago(60), alpha))
    db.execute("UPDATE metrics SET captured_at = ? WHERE tenant_id = ?", (ago(60), alpha))
    db.execute("UPDATE tenants SET created_at = ?, consented_at = ? WHERE id = ?", (ago(61), ago(61), alpha))
    db.commit()
    return alpha


def tenant(db, tenant_id: int):
    return db.execute("SELECT * FROM tenants WHERE id = ?", (tenant_id,)).fetchone()


def finding(db, tid, rule, subject, severity, first, resolved=None, reopened=None, accepted=None, until=None):
    db.execute(
        "INSERT INTO findings (tenant_id, rule_key, object_id, severity, subject, detail, first_seen, last_seen,"
        " resolved_at, reopened_at, accepted_by, accepted_at, accepted_until, accepted_note)"
        " VALUES (?, ?, ?, ?, ?, '', ?, ?, ?, ?, ?, ?, ?, ?)",
        (
            tid,
            rule,
            f"{rule}:{subject}",
            severity,
            subject,
            first,
            resolved or ago(0),
            resolved,
            reopened,
            "jane" if accepted else None,
            accepted,
            until,
            "Vendor fix due" if accepted else None,
        ),
    )
    db.commit()


def entries(tl: timeline.Timeline, kind: str | None = None) -> list[timeline.Entry]:
    return [e for m in tl.months for e in m.entries if kind is None or e.kind == kind]


def history(db, tid, *items):
    """Audit events in the activity history, with the coverage that says it has been read."""
    key = events.tenant_key(db, tid, create=True)
    events.add_audits(events.root(db), key, list(items))
    start = int((NOW - timedelta(days=400)).timestamp())
    events.save_coverage(db, tid, events.Coverage("audits", since=start, read_to=int(NOW.timestamp())))
    db.commit()


def at(event: dict, days: float) -> dict:
    return {**event, "activityDateTime": ago(days)}


# -- The builder ---------------------------------------------------------------------------------------


def test_findings_become_dated_lines_grouped_by_day_and_check(db, alpha):
    seeded = db.execute("SELECT COUNT(*) FROM findings WHERE tenant_id = ?", (alpha,)).fetchone()[0]
    finding(db, alpha, "admin_mfa", "a@alpha.example", "high", ago(40))
    finding(db, alpha, "admin_mfa", "b@alpha.example", "critical", ago(40))
    finding(db, alpha, "legacy_auth_blocked", "c@alpha.example", "medium", ago(45), resolved=ago(15))
    finding(db, alpha, "guest_access_level", "d@alpha.example", "low", ago(50), reopened=ago(9))
    finding(db, alpha, "external_forwarding", "e@alpha.example", "medium", ago(30), accepted=ago(3), until=ago(-90))

    tl = timeline.build(db, tenant(db, alpha), now=NOW)
    assert tl.total == sum(tl.counts.values()) and tl.counts["findings"] >= 5

    first = entries(tl, "first")
    assert len(first) == 1 and first[0].count == seeded and first[0].title.startswith(f"First look: {seeded} findings")
    assert first[0].day == (NOW - timedelta(days=60)).date() and " high" in first[0].detail

    new = [e for e in entries(tl, "finding") if e.title == timeline.RULE_TITLES["admin_mfa"]]
    assert len(new) == 1 and new[0].count == 2 and new[0].level == "critical" and new[0].word == "critical"
    assert new[0].detail == "a@alpha.example and b@alpha.example" and new[0].target == "findings"

    (fixed,) = entries(tl, "fixed")
    assert fixed.level == "ok" and fixed.word == "fixed" and fixed.target == "resolved"
    assert fixed.detail == "c@alpha.example · open 30 days" and fixed.day == (NOW - timedelta(days=15)).date()

    (back,) = entries(tl, "reopened")
    assert back.level == "low" and back.label == "Came back" and back.day == (NOW - timedelta(days=9)).date()

    (accepted,) = entries(tl, "accepted")
    assert accepted.detail.startswith("e@alpha.example · by jane · until ") and accepted.detail.endswith(
        "Vendor fix due"
    )

    # Months come newest first, each with its summary from every entry of the month.
    assert [m.start for m in tl.months] == sorted((m.start for m in tl.months), reverse=True)
    assert sum(e.count for e in entries(tl, "finding")) == 5  # a to e, whatever month each fell in
    month = next(m for m in tl.months if m.start == date(new[0].day.year, new[0].day.month, 1))
    assert any(p.endswith(" new findings") for p in month.summary)
    assert tl.months[0].current and tl.months[0].label.endswith(" so far")
    assert not tl.left_out and tl.kept_since == date(first[0].day.year, first[0].day.month, 1)


def test_admin_changes_come_from_the_activity_history(db, alpha, two_tenants):
    t = tenant(db, alpha)
    _, bravo = two_tenants
    assert not entries(timeline.build(db, tenant(db, bravo), now=NOW), "admin")  # no history read yet
    seeded = {(e.at, e.title) for e in entries(timeline.build(db, t, now=NOW), "admin")}  # the seed's own events

    history(
        db,
        alpha,
        at(ad.GLOBAL_ADMIN_GRANT, 20),
        at(ad.CA_DISABLED, 50),
        at(ad.PASSWORD_RESET, 5),
        at(ad.ROLE_GRANT_FAILED, 4),
    )
    tl = timeline.build(db, t, now=NOW)
    admin = [e for e in entries(tl, "admin") if (e.at, e.title) not in seeded]
    assert [e.day for e in admin] == [(NOW - timedelta(days=20)).date(), (NOW - timedelta(days=50)).date()]
    grant, ca = admin  # the password reset isn't a change worth a year's story; the failed grant never happened
    assert grant.title.startswith("Add member to role: ") and grant.level == "medium" and grant.word == "worth a look"
    assert "Global Administrator" in grant.detail and "Admin role granted" in grant.detail
    assert ca.title.startswith("Update conditional access policy") and "On → Off" in ca.detail
    assert grant.target == "admin-activity" and tl.counts["admin"] == len(seeded) + 2
    month = next(m for m in tl.months if m.start == date(ca.day.year, ca.day.month, 1))
    assert any(p.endswith("admin change, 1 worth a look") or ", 1 worth a look" in p for p in month.summary)

    narrowed = timeline.build(db, t, now=NOW, opens=lambda key: key != timeline.ADMIN_REPORT)
    assert not entries(narrowed, "admin") and "Admin changes" in narrowed.left_out


def test_reports_sent_and_approved_follow_each_files_area(db, alpha):
    t = tenant(db, alpha)
    files = [
        issued.IssuedFile("review", "Monthly review", "review.pdf", "application/pdf", b"%PDF"),
        issued.IssuedFile("data:mfa", "MFA coverage", "mfa.xlsx", "application/x", b"PK"),
    ]
    issued.archive(
        db, alpha, files, period="2026-09", trigger="schedule", user=None, to=("it@alpha.example",), now=ago(12)
    )
    db.execute(
        "INSERT INTO report_reviews (tenant_id, period, state, approved_by, approved_at, commentary, updated_at)"
        " VALUES (?, '2026-09', 'sent', 'sam', ?, 'A steady month.', ?)",
        (alpha, ago(13), ago(13)),
    )
    db.commit()

    tl = timeline.build(db, t, now=NOW)
    (sent,) = entries(tl, "report")
    assert sent.title == "Monthly review and MFA coverage sent" and sent.count == 2
    assert sent.detail == "September 2026 · sent to 1 recipient · by the schedule" and sent.target == "report-history"
    (approved,) = entries(tl, "approved")
    assert approved.title == "September 2026 report approved" and approved.detail == "by sam · A steady month."
    assert approved.target == "review:2026-09" and tl.counts["reports"] == 2

    no_mfa = timeline.build(db, t, now=NOW, opens=lambda key: key != "mfa")
    (sent,) = entries(no_mfa, "report")
    assert sent.title == "Monthly review sent" and sent.count == 1

    no_findings = timeline.build(db, t, now=NOW, opens=lambda key: key != timeline.FINDINGS)
    assert [e.title for e in entries(no_findings, "report")] == ["MFA coverage sent"]
    assert not entries(no_findings, "approved") and "Findings" in no_findings.left_out


def test_notes_and_next_steps_follow_who_may_see_them(db, alpha):
    t = tenant(db, alpha)
    for body, visibility in (("Client knows.", "client"), ("Vendor is slow.", "staff")):
        db.execute(
            "INSERT INTO finding_notes (tenant_id, rule_key, subject, body, visibility, created_by, created_at)"
            " VALUES (?, 'admin_mfa', 'a@alpha.example', ?, ?, 'sam', ?)",
            (alpha, body, visibility, ago(8)),
        )
    db.execute(
        "INSERT INTO report_actions (tenant_id, title, detail, status, target_month, done_at, created_by, created_at,"
        " updated_at) VALUES (?, 'Turn on MFA for admins', 'All four of them.', 'done', '2026-10', ?, 'sam', ?, ?)",
        (alpha, ago(2), ago(20), ago(2)),
    )
    db.commit()

    tl = timeline.build(db, t, now=NOW)
    notes = entries(tl, "note")
    assert [n.detail for n in notes] == ["Vendor is slow. · by sam · staff only", "Client knows. · by sam"] or [
        n.detail for n in notes
    ] == ["Client knows. · by sam", "Vendor is slow. · by sam · staff only"]
    assert notes[0].title == f"Note on {timeline.RULE_TITLES['admin_mfa']}: a@alpha.example"
    actions = entries(tl, "action")
    assert [a.title for a in actions] == ["Done: Turn on MFA for admins", "Recommended: Turn on MFA for admins"]
    assert actions[0].level == "ok" and actions[0].word == "done" and actions[1].target == "review:2026-10"
    assert tl.counts["notes"] == 4

    staff_hidden = timeline.build(db, t, now=NOW, reviews=False)
    assert [n.detail for n in entries(staff_hidden, "note")] == ["Client knows. · by sam"]
    no_findings = timeline.build(db, t, now=NOW, opens=lambda key: key != timeline.FINDINGS)
    assert [n.detail for n in entries(no_findings, "note")] == ["Vendor is slow. · by sam · staff only"]
    assert not entries(no_findings, "action")


def test_failed_collections_collapse_into_runs_of_nights(db, alpha):
    t = tenant(db, alpha)
    for days in (3, 4, 5, 15):
        for collector in ("users", "mailboxes") if days != 15 else ("users",):
            db.execute(
                "INSERT INTO runs (tenant_id, collector, status, trigger, created_at, started_at, finished_at, error)"
                " VALUES (?, ?, 'failed', 'schedule', ?, ?, ?, 'Sign-in failed: AADSTS700016')",
                (alpha, collector, ago(days), ago(days), ago(days)),
            )
    db.commit()

    tl = timeline.build(db, t, now=NOW)
    runs = entries(tl, "collection")
    assert len(runs) == 2 and tl.counts["collection"] >= 2
    three, one = runs
    assert three.count == 3 and three.until is not None and " to " in three.when and " to " in three.dates
    assert three.title.startswith("Collection failed: ") and "Users" in three.title
    assert (
        three.detail.startswith("3 nights running · Sign-in failed") and three.level == "bad" and three.word == "failed"
    )
    assert three.target.startswith("run:") and one.count == 1 and one.until is None
    month = next(m for m in tl.months if m.start == date(three.day.year, three.day.month, 1))
    assert any(p.startswith("collection failed ") and "night" in p for p in month.summary)

    assert not entries(timeline.build(db, t, now=NOW, collects=False), "collection")
    assert "Failed collections" in timeline.build(db, t, now=NOW, collects=False).left_out


def test_the_clients_own_milestones_are_lines_too(db, alpha):
    db.execute("UPDATE tenants SET paused_at = ? WHERE id = ?", (ago(7), alpha))
    db.execute(
        "INSERT INTO permission_requests (tenant_id, release, sent_at, approved_at) VALUES (?, '2.2', ?, ?)",
        (alpha, ago(30), ago(25)),
    )
    db.execute(
        "INSERT INTO audit_feeds (tenant_id, started_at, created_at) VALUES (?, ?, ?)", (alpha, ago(10), ago(11))
    )
    db.commit()
    tl = timeline.build(db, tenant(db, alpha), now=NOW)
    titles = {e.title for e in entries(tl, "client")}
    assert titles >= {"Added to Office Sentry", "Access granted", "Collection paused", "New permissions approved"}
    assert "Audit log feed started" in titles
    paused = next(e for e in entries(tl, "client") if e.title == "Collection paused")
    assert paused.level == "warn" and paused.word == "paused" and paused.target == "manage"
    assert "collection paused" in " · ".join(tl.months[0].summary) or tl.months[0].summary


def test_measures_read_each_months_last_value(db, alpha):
    t = tenant(db, alpha)
    db.execute("DELETE FROM metrics WHERE tenant_id = ? AND key = 'secure_score.pct'", (alpha,))
    for days, value in ((70, 31), (65, 32), (40, 35), (5, 40)):
        db.execute(
            "INSERT INTO metrics (tenant_id, key, value, captured_at) VALUES (?, 'secure_score.pct', ?, ?)",
            (alpha, value, ago(days)),
        )
    finding(db, alpha, "admin_mfa", "a@alpha.example", "high", ago(40))
    finding(db, alpha, "admin_mfa", "b@alpha.example", "critical", ago(40), resolved=ago(2))
    db.commit()

    tl = timeline.build(db, t, now=NOW)
    by_key = {m.key: m for m in tl.measures}
    assert list(by_key) == ["secure_score.pct", "mfa.registered_pct", "open_severe", "licences.unassigned_paid"]
    score = by_key["secure_score.pct"]
    assert len(score.months) == 12 and score.months[-1] == date(NOW.year, NOW.month, 1)
    assert score.latest == 40 and score.first in (31, 32) and score.trend == "better" and score.svg
    assert score.values[-1] == 40 and score.values[0] is None  # nothing a year ago
    # The last value captured in a month is that month's point, so the 70- and 65-day values give one point.
    assert len(score.points) == len({(NOW - timedelta(days=d)).strftime("%Y-%m") for d in (70, 65, 40, 5)})

    severe = by_key["open_severe"]
    live = db.execute(
        "SELECT COUNT(*) FROM findings WHERE tenant_id = ? AND severity IN ('critical', 'high') AND resolved_at IS NULL"
        " AND accepted_at IS NULL",
        (alpha,),
    ).fetchone()[0]
    assert severe.latest == live and severe.values[0] is None and severe.unit == "" and severe.text(live) == str(live)
    assert severe.points[0][0] <= date(*map(int, ago(60)[:7].split("-")), 1)  # open since the first look

    licences = by_key["licences.unassigned_paid"]
    assert licences.svg == "" or len(licences.points) >= 2  # one capture draws nothing yet

    narrowed = timeline.build(db, t, now=NOW, opens=lambda key: key not in ("secure-score", timeline.FINDINGS))
    assert [m.key for m in narrowed.measures] == ["mfa.registered_pct", "licences.unassigned_paid"]
    assert narrowed.left_out[:2] == ["Findings", "Secure Score"]
    assert "Secure Score" not in " ".join(" ".join(m.summary) for m in narrowed.months)


def test_the_period_and_the_filter(db, alpha, monkeypatch):
    t = tenant(db, alpha)
    history(db, alpha, at(ad.GLOBAL_ADMIN_GRANT, 20))
    everything = timeline.build(db, t, now=NOW)
    assert everything.months_back == 12 and everything.since == timeline._add_months(date(NOW.year, NOW.month, 1), -11)

    kept = timeline.build(db, t, now=NOW, months=0)
    assert kept.months_back == 0 and kept.since == kept.kept_since == date(*map(int, ago(61)[:7].split("-")), 1)
    assert len(kept.measures[0].months) <= 4

    admin = timeline.build(db, t, now=NOW, show="admin")
    assert admin.show == "admin" and {e.group for e in entries(admin)} == {"admin"}
    assert (
        admin.total == everything.total and admin.shown == everything.counts["admin"]
    )  # the counts still say what there is
    assert (
        timeline.build(db, t, now=NOW, show="bogus").show == ""
        and timeline.build(db, t, now=NOW, months=7).months_back == 12
    )

    monkeypatch.setattr(timeline, "MAX_ENTRIES", 2)
    cut = timeline.build(db, t, now=NOW)
    assert cut.shown == 2 and cut.truncated == everything.total - 2
    assert [e.at for e in entries(cut)] == sorted((e.at for e in entries(everything)), reverse=True)[:2]


def test_quiet_months_are_named_not_listed(db, alpha):
    tl = timeline.build(db, tenant(db, alpha), now=NOW)
    listed = {m.start for m in tl.months}
    assert not (listed & set(tl.quiet)) and all(m >= tl.kept_since for m in tl.quiet)
    first_look = next(m for m in tl.months if any(e.kind == "first" for e in m.entries))
    gap = timeline._add_months(first_look.start, 1)
    if gap < date(NOW.year, NOW.month, 1):  # a month between the first look and now with nothing in it
        assert gap in tl.quiet and f"{gap:%b %Y}" in tl.quiet_text


def test_month_line_breaks_at_gaps_and_ticks_every_month():
    months = [date(2026, m, 1) for m in range(1, 13)]
    values = [None, None, 38.0, 40.0, None, 42.0, 44.0, 44.5, 45.0, 46.0, 47.0, 47.5]
    svg = charts.month_line(months, values, label="Secure Score by month", unit="%")
    assert svg.count('class="spark-tick"') == 12 and svg.count('class="spark-line"') == 2  # two runs of months
    assert svg.count("<title>") == 13 and "<title>May 2026: nothing captured</title>" in svg
    assert "<title>December 2026: 47.5%</title>" in svg and 'class="spark-end"' in svg
    assert 'aria-label="Secure Score by month"' in svg and 'role="img"' in svg
    assert charts.month_line(months[:3], [None, 5.0, None], label="x") == ""  # one value is no trend
    lone = charts.month_line(months[:4], [3.0, None, 4.0, None], label="x")
    assert lone.count('class="spark-dot"') == 2 and 'class="spark-line"' not in lone


def test_tables_carry_every_entry_and_the_measures(db, alpha):
    history(db, alpha, at(ad.GLOBAL_ADMIN_GRANT, 20))
    tl = timeline.build(db, tenant(db, alpha), now=NOW)
    sheets = timeline.to_tables(tl)
    assert [s.title for s in sheets] == ["Timeline", "Months", "By month"]
    assert sheets[0].columns == ["Date", "Month", "Group", "Kind", "Level", "What", "Detail"]
    assert len(sheets[0].rows) == tl.shown and any(
        r[2] == "Admin changes" and r[4] == "worth a look" for r in sheets[0].rows
    )
    assert len(sheets[1].rows) == len(tl.months) and sheets[1].rows[0][0].endswith(" so far")
    assert sheets[2].columns[1] == "Secure Score (%)" and all(any(v is not None for v in r[1:]) for r in sheets[2].rows)


# -- The page and its downloads --------------------------------------------------------------------------


def test_the_page_sits_on_the_changes_tab_with_its_downloads(client, db, alpha):
    history(db, alpha, at(ad.GLOBAL_ADMIN_GRANT, 20))
    make_user(db, "admin")
    login(client, "admin")
    page = client.get(f"/tenants/{alpha}/timeline").get_data(as_text=True)
    assert f'<a href="/tenants/{alpha}/timeline" class="active" aria-current="page">Timeline</a>' in page
    assert "What happened at Alpha Ltd" in page and "First look: " in page and "Add member to role" in page
    assert (
        f'href="/tenants/{alpha}/timeline?show=admin"' in page
        and f'href="/tenants/{alpha}/timeline?months=all"' in page
    )
    assert "Month by month" in page and 'class="spark trend"' in page and "worth a look" in page
    assert f"/tenants/{alpha}/timeline.pdf" in page and f"/tenants/{alpha}/timeline.xlsx" in page

    filtered = client.get(f"/tenants/{alpha}/timeline?show=admin&months=24").get_data(as_text=True)
    assert '<span class="kind">Finding</span>' not in filtered and '<span class="kind">Admin change</span>' in filtered
    assert "No admin changes this month." in filtered  # a month with none of the chosen kind still has its summary
    assert f'href="/tenants/{alpha}/timeline?months=24"' in filtered  # Everything keeps the period
    assert client.get(f"/tenants/{alpha}/timeline?show=bogus&months=99").status_code == 200

    csv = client.get(f"/tenants/{alpha}/timeline.csv")
    assert csv.status_code == 200 and csv.headers["Content-Disposition"].endswith('timeline-12-months-alpha-ltd.csv"')
    assert csv.get_data(as_text=True).startswith("﻿Date,Month,Group,Kind,Level,What,Detail\r\n")
    xlsx = client.get(f"/tenants/{alpha}/timeline.xlsx?show=admin")
    assert xlsx.status_code == 200 and 'timeline-12-months-admin-alpha-ltd.xlsx"' in xlsx.headers["Content-Disposition"]
    assert load_workbook(BytesIO(xlsx.data)).sheetnames[1:] == ["Timeline", "Months", "By month"]
    assert (
        db.execute(
            "SELECT COUNT(*) FROM audit_log WHERE action = 'report_downloaded' AND detail LIKE 'timeline-%'"
        ).fetchone()[0]
        == 2
    )


@needs_pdf
def test_the_pdf_tells_the_same_story(client, db, alpha):
    history(db, alpha, at(ad.GLOBAL_ADMIN_GRANT, 20))
    make_user(db, "admin")
    login(client, "admin")
    pdf = client.get(f"/tenants/{alpha}/timeline.pdf")
    assert pdf.status_code == 200 and pdf.mimetype == "application/pdf" and pdf.data.startswith(b"%PDF")
    assert 'timeline-12-months-alpha-ltd.pdf"' in pdf.headers["Content-Disposition"]


def test_a_new_client_starts_with_being_added(client, db, two_tenants):
    _, bravo = two_tenants
    make_user(db, "admin")
    login(client, "admin")
    page = client.get(f"/tenants/{bravo}/timeline").get_data(as_text=True)
    assert "Added to Office Sentry" in page and "Access granted" in page and "First look" not in page
    assert page.count("Not enough history yet") == 4 and 'Everything <span class="count-badge">2</span>' in page


def test_each_person_sees_only_the_areas_they_open(client, db, alpha):
    history(db, alpha, at(ad.GLOBAL_ADMIN_GRANT, 20))
    db.execute(
        "INSERT INTO runs (tenant_id, collector, status, trigger, created_at, finished_at, error)"
        " VALUES (?, 'users', 'failed', 'schedule', ?, ?, 'Sign-in failed')",
        (alpha, ago(3), ago(3)),
    )
    db.commit()
    make_user(db, "admin")
    login(client, "admin")
    page = client.get(f"/tenants/{alpha}/timeline").get_data(as_text=True)
    assert "First look: " in page and "Add member to role" in page and "Collection failed" in page
    assert "Users with MFA" in page and "Left out, since" not in page

    # An account manager holds findings and tenant reports (admin activity is one), but no identity reports
    # and no collection runs.
    role_user(db, "am", "account_manager", tenants=[alpha])
    client.post("/logout")
    login(client, "am")
    page = client.get(f"/tenants/{alpha}/timeline").get_data(as_text=True)
    assert "First look: " in page and "Add member to role" in page and "Collection failed" not in page
    assert '<span class="label">Users with MFA</span>' not in page
    assert "Left out, since you can't open them for this client: Users with MFA, Failed collections." in page
    csv = client.get(f"/tenants/{alpha}/timeline.csv").get_data(as_text=True)
    assert "Collection failed" not in csv and "Add member to role" in csv

    # First line opens identity reports, not findings or the tenant's admin activity.
    role_user(db, "fl", "first_line", tenants=[alpha])
    client.post("/logout")
    login(client, "fl")
    page = client.get(f"/tenants/{alpha}/timeline").get_data(as_text=True)
    assert "First look: " not in page and "Add member to role" not in page
    assert '<span class="label">Users with MFA</span>' in page
    assert "Left out, since you can't open them for this client: Findings, Admin changes, Secure Score" in page

    # A client account never opens it: the page is staff only.
    role_user(db, "cl", "client_viewer", tenants=[alpha])
    client.post("/logout")
    login(client, "cl")
    assert client.get(f"/tenants/{alpha}/timeline").status_code == 403
    assert client.get(f"/tenants/{alpha}/timeline.csv").status_code == 403
