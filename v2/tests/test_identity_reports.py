"""Identity and licence detail reports: MFA coverage, admins, licence waste, inactive accounts, Secure Score."""

from datetime import UTC, datetime
from io import BytesIO

import pytest
import responses
from openpyxl import load_workbook

from officesentry import store
from officesentry.db import utcnow
from officesentry.graph import GRAPH_ROOT
from officesentry.reports.data_reports import REPORTS, build
from officesentry.snapshot import TenantSnapshot
from tests import graphdata as gd
from tests.conftest import login, make_user
from tests.seed import seed_tenant
from tests.stored import unpack
from tests.test_collectors import SKUS, USERS, V1, _run, factory  # noqa: F401

NOW = datetime(2026, 10, 2, tzinfo=UTC)
NEW_KEYS = ("mfa", "admins", "licences", "inactive", "secure-score")

CARL = {
    "id": "u6",
    "displayName": "Carl",
    "userPrincipalName": "carl@x.com",
    "userType": "Member",
    "accountEnabled": True,
    "createdDateTime": "2025-01-01T00:00:00Z",
    "assignedLicenses": [{"skuId": "sku-bp"}],
    "signInActivity": {"lastSignInDateTime": "2026-03-01T00:00:00Z"},
}
BREAKGLASS = {
    "id": "u4",
    "displayName": "Breakglass",
    "userPrincipalName": "bg@x.com",
    "userType": "Member",
    "accountEnabled": True,
    "createdDateTime": "2025-01-01T00:00:00Z",
    "assignedLicenses": [],
    "signInActivity": {"lastSignInDateTime": "2026-09-01T00:00:00Z"},
}
NO_P1_SKUS = [dict(SKUS[0], servicePlans=[p for p in SKUS[0]["servicePlans"] if p["servicePlanName"] != "AAD_PREMIUM"])]


def sections(report):
    return {s.title: s for s in report.sections}


def figures(report):
    return {h.label: h.value for h in report.highlights}


def price(db, part, amount):
    db.execute(
        "INSERT INTO sku_prices (sku_part_number, monthly_price, updated_at) VALUES (?, ?, ?)", (part, amount, utcnow())
    )


@pytest.fixture
def p1_tenant(db, two_tenants):
    alpha, _ = two_tenants
    seed_tenant(db, alpha, users=USERS + [CARL, BREAKGLASS], secure_scores=gd.SECURE_SCORES_HISTORY)
    return store.get_tenant(db, alpha)


@pytest.fixture
def free_tenant(db, two_tenants):
    """No Entra ID P1: no Conditional Access, no sign-in dates, per-user MFA instead."""
    alpha, _ = two_tenants
    users = [dict(u, signInActivity=None) for u in USERS + [CARL]]
    seed_tenant(db, alpha, users=users, skus=NO_P1_SKUS, ca=[])
    return store.get_tenant(db, alpha)


# -- MFA coverage ---------------------------------------------------------------------


def test_mfa_coverage_cross_checks_registration_against_policies(db, p1_tenant):
    report = build(db, p1_tenant, "mfa", NOW)
    s = sections(report)
    gaps = {r[0]: r for r in s["Accounts with a gap"].rows}
    # Carl is required by the all-users policy but hasn't set MFA up.
    assert gaps["carl@x.com"][2] == "Required, not set up"
    # The break-glass account is excluded from the only MFA policy and has no method: no MFA at all, and an admin.
    assert gaps["bg@x.com"][2] == "No MFA" and gaps["bg@x.com"][5] == "Require MFA for all users"
    assert s["Accounts with a gap"].flags[0] == "critical"  # admins first
    assert [r[0] for r in s["Protected accounts"].rows] == ["ann@x.com"]
    assert ["Report-only MFA policies (not enforced)", "MFA for admins (report only)"] in s["How MFA is required"].rows
    f = figures(report)
    assert f["Fully protected"] == "1 of 3" and f["No MFA at all"] == 1 and f["Admins with a gap"] == 1
    assert not report.notes


def test_mfa_coverage_counts_only_privileged_roles_as_admins(db, two_tenants, monkeypatch):
    """A read-only role (Directory Readers) doesn't make Carl an admin: his gap is medium, as on the Findings tab."""
    alpha, _ = two_tenants
    carl = {"@odata.type": "#microsoft.graph.user", "id": "u6", "userPrincipalName": "carl@x.com"}
    reader = {**gd.ROLE_ASSIGNMENTS[3], "id": "a-carl", "principalId": "u6", "principal": carl}
    monkeypatch.setattr(gd, "ROLE_ASSIGNMENTS", [*gd.ROLE_ASSIGNMENTS, reader])
    seed_tenant(db, alpha, users=USERS + [CARL, BREAKGLASS])
    gaps = sections(build(db, store.get_tenant(db, alpha), "mfa", NOW))["Accounts with a gap"]
    carl = next(i for i, r in enumerate(gaps.rows) if r[0] == "carl@x.com")
    assert gaps.rows[carl][1] == "Directory Readers" and gaps.flags[carl] == "medium"


def test_admins_report_names_the_group_a_role_comes_through(db, two_tenants, monkeypatch):
    from tests import groupdata as grp

    alpha, _ = two_tenants
    group = {
        "id": "a-grp",
        "roleDefinitionId": gd.GA,
        "principalId": "rg1",
        "directoryScopeId": "/",
        "principal": {"@odata.type": "#microsoft.graph.group", "id": "rg1", "displayName": "Tenant admins"},
    }
    monkeypatch.setattr(gd, "ROLE_ASSIGNMENTS", [*gd.ROLE_ASSIGNMENTS, group])
    monkeypatch.setitem(
        grp.MEMBERS,
        "rg1",
        [
            {
                "@odata.type": "#microsoft.graph.user",
                "id": "u6",
                "userPrincipalName": "carl@x.com",
                "accountEnabled": True,
            }
        ],
    )
    seed_tenant(db, alpha, users=USERS + [CARL, BREAKGLASS])
    report = build(db, store.get_tenant(db, alpha), "admins", NOW)
    rows = {r[0]: r for r in sections(report)["Privileged role holders"].rows}
    assert "Global Administrator" in rows["carl@x.com"][2]
    assert "Through the group Tenant admins" in rows["carl@x.com"][8]
    assert rows["Tenant admins"][1] == "Group"
    assert figures(report)["Global Administrators"] == 3  # Ann, the break-glass account and Carl


def test_mfa_coverage_without_p1_uses_per_user_mfa(db, free_tenant):
    report = build(db, free_tenant, "mfa", NOW)
    s = sections(report)
    assert ["Conditional Access", "Not available: needs Entra ID P1"] in s["How MFA is required"].rows
    assert ["Per-user MFA", "1 of 2 members"] in s["How MFA is required"].rows
    gaps = {r[0]: r for r in s["Accounts with a gap"].rows}
    assert gaps["ann@x.com"][4] == "Per-user MFA (enforced)" and gaps["ann@x.com"][2] == "Required, not set up"
    assert gaps["carl@x.com"][2] == "No MFA"


def test_empty_registration_report_falls_back_to_per_user_methods(db, two_tenants):
    """A newly P1 tenant can get an empty registration report: members are then checked one by one."""
    alpha, _ = two_tenants
    seed_tenant(db, alpha, users=USERS + [CARL], registration=[])
    tenant = store.get_tenant(db, alpha)
    rows = store.run_objects(db, store.latest_run(db, alpha, "mfa_registration")["id"], "mfa_registration")
    assert {r["_id"]: r["source"] for r in rows} == {"u1": "per_user_methods", "u6": "per_user_methods"}
    report = build(db, tenant, "mfa", NOW)
    assert not report.notes
    assert {r[0]: r[2] for r in sections(report)["Accounts with a gap"].rows} == {
        "ann@x.com": "Required, not set up",
        "carl@x.com": "Required, not set up",
    }


def test_mfa_checks_unavailable_without_registration_data(db, two_tenants):
    """No registration rows for anyone: the checks say so instead of passing, or failing with "0 admins"."""
    from officesentry.rules import evaluate

    alpha, _ = two_tenants
    seed_tenant(db, alpha, users=[USERS[2]], registration=[])  # only a guest: nobody to check one by one
    results = {r.key: r for r in evaluate(TenantSnapshot(db, store.get_tenant(db, alpha)), NOW)}
    assert results["users_mfa_registered"].status == "unavailable"
    assert results["admin_mfa"].status == "unavailable"
    report = build(db, store.get_tenant(db, alpha), "mfa", NOW)
    assert any("no MFA registration data" in n for n in report.notes)


def test_admin_mfa_unknown_registration_is_not_a_failure(db, two_tenants):
    from officesentry.rules import evaluate

    alpha, _ = two_tenants
    registered = [dict(r, isMfaRegistered=True) for r in gd.REGISTRATION_DETAILS if r["id"] == "u1"]
    seed_tenant(db, alpha, registration=registered)  # bg@x.com and the guest admin aren't in the data
    r = {r.key: r for r in evaluate(TenantSnapshot(db, store.get_tenant(db, alpha)), NOW)}["admin_mfa"]
    assert r.status == "unavailable" and "isn't known for 2 admins" in r.summary


def test_mfa_coverage_with_security_defaults(db, two_tenants):
    alpha, _ = two_tenants
    seed_tenant(db, alpha, users=USERS + [CARL], defaults=True, ca=[])
    s = sections(build(db, store.get_tenant(db, alpha), "mfa", NOW))
    assert ["Security Defaults", "On"] in s["How MFA is required"].rows
    assert s["Protected accounts"].rows[0][3] == "Security Defaults"


# -- Admins -----------------------------------------------------------------------------------


def test_admins_report_flags_problems_and_finds_break_glass(db, p1_tenant):
    report = build(db, p1_tenant, "admins", NOW)
    s = sections(report)
    holders = {
        r[0]: (r, flag)
        for r, flag in zip(s["Privileged role holders"].rows, s["Privileged role holders"].flags, strict=True)
    }
    assert holders["bg@x.com"][1] == "critical" and "No MFA set up" in holders["bg@x.com"][0][8]
    assert holders["g#EXT#@x.com"][1] == "high" and "Guest account" in holders["g#EXT#@x.com"][0][8]
    assert holders["eve@x.com"][0][3] == "Eligible (PIM)"
    assert holders["Some app"][1] == "medium"
    assert holders["ann@x.com"][1] is None
    assert [r[0] for r in s["Other roles"].rows] == ["bob@x.com"]

    glass = s["Break-glass accounts"]
    assert [r[0] for r in glass.rows] == ["bg@x.com"]
    assert "Excluded from Conditional Access" in glass.rows[0][1] and glass.flags[0] == "high"
    f = figures(report)
    assert f["Global Administrators"] == 2 and f["Admins without MFA"] == 1 and f["Break-glass accounts"] == 1
    assert f["Using PIM"] == "Not checked" and any("Entra ID P2" in n for n in report.notes)


def test_admins_report_shows_office_sentrys_own_app_with_a_privileged_role(db, p1_tenant):
    """Office Sentry needs only Global Reader. Holding a privileged role (here Exchange Administrator) is too
    much access, and the check behind this report flags it, so the report lists it too."""
    sp = db.execute("SELECT client_id FROM app_connections").fetchone()["client_id"]
    run = store.latest_run(db, p1_tenant["id"], "admin_roles")
    unpack(db)
    db.execute(
        "UPDATE objects SET data = json_set(data, '$.appId', ?) WHERE run_id = ? AND object_id = 'a5'", (sp, run["id"])
    )
    rows = sections(build(db, p1_tenant, "admins", NOW))["Privileged role holders"].rows
    assert "Some app" in [r[0] for r in rows]


# -- Licence waste --------------------------------------------------------------------------------


def test_licence_waste_in_money(db, p1_tenant):
    price(db, "SPB", 18.10)
    report = build(db, p1_tenant, "licences", NOW)
    s = sections(report)
    sub = s["Waste by subscription"].rows[0]
    # 3 unassigned + Bob (disabled) + Carl (no sign-in since March) = 5 wasted.
    assert sub[0] == "Microsoft 365 Business Premium" and sub[3:6] == [3, 1, 1]
    assert sub[7] == pytest.approx(90.5) and str(sub[7]) == "£90.50"
    accounts = {r[0]: r for r in s["Accounts to review"].rows}
    assert accounts["bob@x.com"][1] == "Disabled" and accounts["carl@x.com"][1] == "Inactive"
    assert figures(report)["Waste per year"] == pytest.approx(1086)
    assert not report.notes


def test_licence_waste_without_prices_or_p1(db, free_tenant):
    report = build(db, free_tenant, "licences", NOW)
    assert figures(report)["Waste per month"] == "Not priced"
    assert sections(report)["Waste by subscription"].rows[0][5] == "Not checked"
    assert any("Entra ID P1" in n for n in report.notes) and any("Licence prices" in n for n in report.notes)


def test_overlapping_licences(db, two_tenants):
    alpha, _ = two_tenants
    p1 = {
        "skuId": "sku-p1",
        "skuPartNumber": "AAD_PREMIUM",
        "appliesTo": "User",
        "capabilityStatus": "Enabled",
        "prepaidUnits": {"enabled": 1},
        "consumedUnits": 1,
        "servicePlans": [
            {"servicePlanId": "plan-p1", "servicePlanName": "AAD_PREMIUM", "provisioningStatus": "Success"}
        ],
    }
    bp = dict(SKUS[0], servicePlans=[dict(p, servicePlanId=f"plan-{i}") for i, p in enumerate(SKUS[0]["servicePlans"])])
    bp["servicePlans"][0]["servicePlanId"] = "plan-p1"
    users = [dict(USERS[0], assignedLicenses=[{"skuId": "sku-bp"}, {"skuId": "sku-p1"}])]
    seed_tenant(db, alpha, users=users, skus=[bp, p1])
    price(db, "AAD_PREMIUM", 5)
    rows = sections(build(db, store.get_tenant(db, alpha), "licences", NOW))["Overlapping licences"].rows
    assert rows == [["ann@x.com", "Microsoft Entra ID P1", "Microsoft 365 Business Premium", 5.0]]


# -- Inactive accounts ------------------------------------------------------------------------------


def test_inactive_accounts(db, p1_tenant):
    report = build(db, p1_tenant, "inactive", NOW)
    s = sections(report)
    assert [r[:2] for r in s["Inactive members"].rows] == [["carl@x.com", "6 to 12 months"]]
    assert s["Inactive members"].flags == ["medium"]  # licensed
    assert [r[0] for r in s["Inactive guests"].rows] == ["g#EXT#@x.com"]
    assert [r[0] for r in s["Disabled accounts still holding licences or roles"].rows] == ["bob@x.com"]
    assert figures(report)["Inactive members"] == 1


def test_inactive_never_signed_in_wording(db, two_tenants):
    alpha, _ = two_tenants
    new = dict(CARL, id="u7", userPrincipalName="new@x.com", createdDateTime="2026-08-01T00:00:00Z", signInActivity={})
    seed_tenant(db, alpha, users=USERS + [new])
    rows = sections(build(db, store.get_tenant(db, alpha), "inactive", NOW))["Inactive members"].rows
    assert rows[0][1] == "Never signed in (created 62 days ago)" and rows[0][2] == "Never"


def test_inactive_without_p1_lists_only_what_it_can(db, free_tenant):
    report = build(db, free_tenant, "inactive", NOW)
    assert "Inactive members" not in sections(report)
    assert figures(report)["Inactive members"] == "Not checked"
    assert any("Entra ID P1" in n for n in report.notes)


# -- Secure Score -------------------------------------------------------------------------------------


def test_secure_score_trend_and_control_changes(db, p1_tenant):
    report = build(db, p1_tenant, "secure-score", NOW)
    s = sections(report)
    changes = s["Changes in the last 30 days"]
    assert changes.intro == "Compared with 31 Aug 2026. Drops first."
    assert changes.rows[0][0] == "Enable Conditional Access policies to block legacy authentication"
    assert changes.rows[0][4] == "-4" and changes.flags[0] == "medium"
    assert changes.rows[1][4] == "+5"
    assert s["Where the points are"].rows == [
        [
            "Enable Conditional Access policies to block legacy authentication",
            "Identity",
            4,
            8,
            4,
            "Moderate",
            "Moderate",
        ]
    ]
    assert [r[0] for r in s["Month by month"].rows] == ["October 2026 (so far)", "September 2026", "August 2026"]
    f = figures(report)
    assert f["Secure Score"] == "42%" and f["Change in 30 days"] == "+7%" and f["Points available"] == "4"
    assert "<svg" in report.chart


@responses.activate
def test_secure_score_without_control_profiles(db, two_tenants, factory):  # noqa: F811
    alpha, _ = two_tenants
    responses.get(f"{V1}/security/secureScores", json={"value": gd.SECURE_SCORES_HISTORY})
    responses.get(
        f"{V1}/security/secureScoreControlProfiles",
        status=403,
        json={"error": {"code": "Authorization_RequestDenied", "message": "Denied"}},
    )
    run_id, status = _run(db, alpha, "secure_score", factory)
    assert status == "succeeded"
    assert len(store.run_objects(db, run_id, "secure_score_control_past")) == 2
    report = build(db, store.get_tenant(db, alpha), "secure-score", NOW)
    assert any("control names" in n for n in report.notes)
    rows = {r[0]: r for r in sections(report)["Where the points are"].rows}
    assert rows["BlockLegacyAuthentication"][4] == "Unknown"  # no maximum without the profile


# -- Per-user MFA collector ------------------------------------------------------------------------------


@responses.activate
def test_per_user_mfa_skipped_with_p1(db, two_tenants, factory):  # noqa: F811
    alpha, _ = two_tenants
    store.set_capabilities(db, alpha, {"entra_p1": True})
    run_id, status = _run(db, alpha, "per_user_mfa", factory)
    assert status == "succeeded"
    assert store.run_objects(db, run_id, "per_user_mfa_status")[0]["checked"] is False


@responses.activate
def test_per_user_mfa_reads_beta_requirements(db, two_tenants, factory):  # noqa: F811
    alpha, _ = two_tenants
    store.set_capabilities(db, alpha, {"entra_p1": False})
    responses.get(f"{V1}/users", json={"value": USERS})
    responses.get(f"{GRAPH_ROOT}/beta/users/u1/authentication/requirements", json={"perUserMfaState": "enforced"})
    run_id, status = _run(db, alpha, "per_user_mfa", factory)
    assert status == "succeeded"
    rows = {r["_id"]: r["state"] for r in store.run_objects(db, run_id, "per_user_mfa")}
    assert rows == {"u1": "enforced"}  # only enabled members are read


# -- Pages and exports --------------------------------------------------------------------------------------


def test_identity_reports_registered():
    assert set(NEW_KEYS) <= set(REPORTS)


def test_identity_report_pages_and_xlsx(client, db, p1_tenant, two_tenants):
    _, bravo = two_tenants
    make_user(db, "client", role="client", tenants=[bravo])
    login(client, "client")
    assert client.get(f"/tenants/{p1_tenant['id']}/data/admins").status_code == 404
    client.post("/logout")
    make_user(db, "admin")
    login(client, "admin")
    for key in NEW_KEYS:
        page = client.get(f"/tenants/{p1_tenant['id']}/data/{key}")
        assert page.status_code == 200, key
        html = page.get_data(as_text=True)
        assert 'class="tiles"' in html, key
        resp = client.get(f"/tenants/{p1_tenant['id']}/data/{key}.xlsx")
        assert resp.status_code == 200, key
        wb = load_workbook(BytesIO(resp.data))
        assert wb.sheetnames[0] == "Summary", key
    page = client.get(f"/tenants/{p1_tenant['id']}/data/secure-score").get_data(as_text=True)
    assert '<svg class="chart"' in page and "Month by month" in page
    for key in NEW_KEYS:
        report = build(db, store.get_tenant(db, bravo), key, NOW)
        assert report.unavailable, key  # never collected


def test_money_cells_are_numbers_in_xlsx(client, db, p1_tenant):
    price(db, "SPB", 10)
    make_user(db, "admin")
    login(client, "admin")
    wb = load_workbook(BytesIO(client.get(f"/tenants/{p1_tenant['id']}/data/licences.xlsx").data))
    row = [c.value for c in wb["Waste by subscription"][2]]
    assert row[6] == 10 and row[7] == 50
    assert "£50.00" in client.get(f"/tenants/{p1_tenant['id']}/data/licences").get_data(as_text=True)


def test_break_glass_needs_more_than_one_exclusion(db, two_tenants):
    """An admin excluded from the all-users policy but covered by an admin policy isn't a break-glass account;
    an account named like one but without Global Administrator is called out."""
    alpha, _ = two_tenants
    everyone = dict(
        gd.CA_MFA_ALL,
        conditions={
            "users": {"includeUsers": ["All"], "excludeUsers": ["u1", "u4"]},
            "applications": {"includeApplications": ["All"]},
            "clientAppTypes": ["all"],
        },
    )
    admins_policy = dict(
        gd.CA_MFA_ALL,
        id="ca9",
        displayName="Admins MFA",
        conditions={
            "users": {"includeUsers": ["u1"]},
            "applications": {"includeApplications": ["All"]},
            "clientAppTypes": ["all"],
        },
    )
    emergency = dict(CARL, id="u8", displayName="Emergency access", userPrincipalName="emergency@x.com")
    seed_tenant(db, alpha, users=USERS + [BREAKGLASS, emergency], ca=[everyone, admins_policy])
    glass = sections(build(db, store.get_tenant(db, alpha), "admins", NOW))["Break-glass accounts"]
    assert [r[0] for r in glass.rows] == ["bg@x.com", "emergency@x.com"]  # not ann@x.com
    assert "Doesn't hold Global Administrator" in glass.rows[1][7]
