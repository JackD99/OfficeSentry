"""SharePoint storage: collector, tenant quota, check and the cleanup report."""

import io

from openpyxl import load_workbook

from officesentry import store
from officesentry.collectors.sharepoint_storage import is_library, quota_row, version_number
from officesentry.reports import data_reports
from officesentry.rules import evaluate
from officesentry.snapshot import TenantSnapshot
from officesentry.storage import GB, tenant_quota
from tests import sharepoint_data as sd
from tests.conftest import login, make_user
from tests.seed import seed_tenant
from tests.test_collectors import SKUS

MB = 1024**2


def _snap(db, tenant_id):
    return TenantSnapshot(db, store.get_tenant(db, tenant_id))


def _report(db, tenant_id):
    return data_reports.build(db, store.get_tenant(db, tenant_id), "sharepoint-storage")


def _section(report, title):
    return next(s for s in report.sections if s.title == title)


def _column(section, name):
    i = section.columns.index(name)
    return [str(r[i]) for r in section.rows]


def test_collector_totals(db, two_tenants):
    alpha, _ = two_tenants
    seed_tenant(db, alpha)
    snap = _snap(db, alpha)
    assert store.latest_run(db, alpha, "sharepoint_storage")["status"] == "succeeded"
    sites = {s["_id"]: s for s in snap.objects("sharepoint_storage", "site_storage")}
    assert set(sites) == {"site-fin", "site-old"}  # OneDrive sites aren't SharePoint storage
    site = sites["site-fin"]
    assert site["used"] == 5 * GB and site["recycleBin"] == 1 * GB and site["limit"] is None
    # Current files and versions leave out the Preservation Hold Library, which is counted on its own.
    assert site["current"] == 2 * GB + 100 * MB + 500 * MB
    assert site["versions"] == 1 * GB - 100 * MB
    assert site["preservationHold"] == 300 * MB and site["staleBytes"] == 500 * MB
    libs = {x["library"]: x for x in snap.objects("sharepoint_storage", "library_storage") if not x["personal"]}
    assert set(libs) == {"Documents", "Preservation Hold Library"}  # the Events list isn't a library
    assert libs["Documents"]["files"] == 3 and libs["Documents"]["folders"] == 2
    assert libs["Documents"]["maxVersion"] == 320 and libs["Preservation Hold Library"]["preservationHold"]
    folder = {f["path"]: f for f in snap.objects("sharepoint_storage", "large_folder")}["Clients"]
    assert folder["total"] == 3 * GB and folder["files"] == 2 and folder["library"] == "Documents"
    (old,) = snap.objects("sharepoint_storage", "stale_file")
    assert old["path"] == "Old/Archive.zip"
    assert [f["path"] for f in snap.objects("sharepoint_storage", "version_file")] == ["Clients/Plan.xlsx"]
    assert {f["path"] for f in snap.objects("sharepoint_storage", "large_file")} == {
        "Clients/Mail export.pst",
        "Clients/Plan.xlsx",
        "Old/Archive.zip",
        "Plan_v1.xlsx",
    }
    # Ann's OneDrive is 93% full, so it was read file by file, separately from SharePoint.
    (ann,) = snap.objects("sharepoint_storage", "onedrive_storage")
    assert ann["scanned"] and ann["versions"] == 10 * GB and ann["limit"] == 1024 * GB
    assert [f["path"] for f in snap.objects("sharepoint_storage", "onedrive_file")] == ["Videos/raw.mp4"]
    names = {r["name"] for r in snap.objects("sharepoint_storage", "recycle_item")}
    assert names == {"Q1 scans.zip", "notes.txt"}
    assert snap.metric("sharepoint.storage.recycle_bin_bytes") == 1 * GB
    assert snap.metric("sharepoint.storage.versions_bytes") == 1 * GB - 100 * MB


def test_without_permission(db, two_tenants):
    alpha, _ = two_tenants
    seed_tenant(db, alpha, sharepoint_kw={"files": False})
    run = store.latest_run(db, alpha, "sharepoint_storage")
    assert run["status"] == "partial" and "Sites.Read.All" in run["error"]
    assert "consent link" in _report(db, alpha).unavailable


def test_report(db, two_tenants):
    alpha, _ = two_tenants
    seed_tenant(db, alpha)
    report = _report(db, alpha)
    assert report.unavailable is None
    tiles = {t.label: t for t in report.tiles}
    assert tiles["Storage used"].value == "Under 1%" and "of 1.10 TB (estimated)" in tiles["Storage used"].foot
    assert tiles["Recycle bins"].value == "1.0 GB"
    assert tiles["Untouched for 2 years"].foot == "1 file"
    parts = _section(report, "Where the space goes")
    assert _column(parts, "Part") == [
        "Current files",
        "Recycle bins",
        "Older versions of files",
        "Preservation Hold libraries",
        "Lists, pages and other",
    ]
    assert "1 TB + 10 GB × 10 licences" in parts.intro
    assert _column(_section(report, "Largest files"), "File")[0] == "Clients/Mail export.pst"  # largest first
    assert _column(_section(report, "Largest files"), "Location")[0] == "Finance / Documents"
    assert _column(_section(report, "Files with the most version history"), "Version") == ["320"]
    libraries = _section(report, "Libraries")
    assert libraries.flags == [None, "medium"]  # the Preservation Hold Library is called out
    assert _column(_section(report, "OneDrives nearly full"), "Share used") == ["93%"]
    assert _column(_section(report, "Largest files in nearly full OneDrives"), "Person") == ["ann@x.com"]
    # OneDrive files don't appear among SharePoint's files.
    assert "Videos/raw.mp4" not in _column(_section(report, "Largest files"), "File")


def test_site_limit_and_nearly_full_tenant(db, two_tenants):
    alpha, _ = two_tenants
    quota = {"site-fin": {"used": 1100 * GB, "deleted": 40 * GB, "total": 1200 * GB, "state": "critical"}}
    seed_tenant(db, alpha, sharepoint_kw={"site_quota": quota})
    report = _report(db, alpha)
    sites = _section(report, "Sites by storage")
    assert sites.flags[0] == "high" and _column(sites, "Site limit")[0] == "1.17 TB (92% used)"
    assert {t.label: t for t in report.tiles}["Storage used"].tone == "bad"  # 1,100 GB of 1,124 GB
    result = {r.key: r for r in evaluate(_snap(db, alpha))}["sharepoint_storage_space"]
    assert result.status == "fail" and result.findings[0].severity == "high"
    assert result.findings[1].subject == sd.FINANCE and "40.0 GB in the recycle bin" in result.findings[1].detail


def test_tenant_over_its_storage(db, two_tenants):
    alpha, _ = two_tenants
    quota = {"site-fin": {"used": 1200 * GB, "deleted": 0, "total": 1300 * GB, "state": "critical"}}
    seed_tenant(db, alpha, sharepoint_kw={"site_quota": quota})
    free = {t.label: t for t in _report(db, alpha).tiles}["Free space"]  # 1,200 GB and more of 1,124 GB
    assert free.value == "None" and free.tone == "bad" and free.foot.startswith("About 7")
    assert free.foot.endswith(" GB over; Microsoft can make sites read-only until there's room")


def test_check_passes_and_waits_for_data(db, two_tenants):
    alpha, bravo = two_tenants
    seed_tenant(db, alpha)
    assert {r.key: r for r in evaluate(_snap(db, alpha))}["sharepoint_storage_space"].status == "pass"
    assert {r.key: r for r in evaluate(_snap(db, bravo))}["sharepoint_storage_space"].status == "unavailable"


def test_tenant_quota():
    assert tenant_quota(SKUS).total == 1024 * GB + 10 * 10 * GB
    extra = {
        "skuPartNumber": "SHAREPOINTSTORAGE",
        "appliesTo": "Company",
        "capabilityStatus": "Enabled",
        "prepaidUnits": {"enabled": 250},
        "servicePlans": [{"servicePlanName": "SHAREPOINTSTORAGE"}],
    }
    free = dict(SKUS[0], skuPartNumber="FREE", prepaidUnits={"enabled": 10000})
    quota = tenant_quota([*SKUS, extra, free])
    assert quota.total == 1024 * GB + 100 * GB + 250 * GB and quota.extra_gb == 250
    assert "250 GB of extra storage" in quota.explained
    assert tenant_quota([SKUS[1]]) is None and tenant_quota(None) is None
    # A subscription in its grace period still gives its storage until Microsoft suspends it.
    lapsing = dict(SKUS[0], capabilityStatus="Warning", prepaidUnits={"enabled": 0, "warning": 10})
    assert tenant_quota([lapsing]).total == tenant_quota([SKUS[0]]).total
    suspended = dict(SKUS[0], capabilityStatus="Suspended", prepaidUnits={"enabled": 0, "suspended": 10})
    assert tenant_quota([suspended]) is None


def test_helpers():
    assert version_number("512.0") == 512 and version_number("3.2") == 3 and version_number(None) == 0
    assert is_library({"name": "Shared Documents", "list": {"template": "documentLibrary"}})
    assert is_library({"name": "PreservationHoldLibrary", "list": {"template": "101"}})
    assert not is_library({"name": "Events", "list": {"template": "events"}})
    assert quota_row({"quota": {"used": 1, "deleted": 2, "total": 27487790694400}})["limit"] is None
    assert quota_row({"quota": {"used": "5", "total": 100, "state": "normal"}}) == {
        "used": 5,
        "recycleBin": 0,
        "limit": 100,
        "state": "normal",
    }
    assert quota_row(None)["used"] == 0


def test_page_export_and_isolation(client, db, two_tenants):
    alpha, bravo = two_tenants
    seed_tenant(db, alpha)
    make_user(db, "admin")
    login(client, "admin")
    page = client.get(f"/tenants/{alpha}/data/sharepoint-storage").get_data(as_text=True)
    assert "Where the space goes" in page and "Largest folders" in page and 'class="tiles"' in page
    assert f"/tenants/{alpha}/data/sharepoint-storage.pdf" in page
    assert client.get(f"/tenants/{bravo}/data/sharepoint-storage").status_code == 200  # empty state renders
    wb = load_workbook(io.BytesIO(client.get(f"/tenants/{alpha}/data/sharepoint-storage.xlsx").data))
    assert wb.sheetnames[:3] == ["Summary", "Where the space goes", "Sites by storage"]


def test_client_only_sees_own_tenant(client, db, two_tenants):
    alpha, bravo = two_tenants
    seed_tenant(db, alpha)
    make_user(db, "client", role="client", tenants=[bravo])
    login(client, "client")
    assert client.get(f"/tenants/{alpha}/data/sharepoint-storage").status_code == 404
    assert client.get(f"/tenants/{alpha}/data/sharepoint-storage.xlsx").status_code == 404
