"""Intune policies: the collector, the report, the compliance coverage check, and the waiting and unlicensed
states. Also BitLocker recovery keys (Devices report and check) and PIM for groups (Admins report), the
other permissions of the 2.2 batch that nothing read until now."""

import json
from datetime import UTC, datetime
from io import BytesIO

from openpyxl import load_workbook

from officesentry import endpoints, store, support
from officesentry.reports import data_reports
from officesentry.rules import evaluate
from officesentry.snapshot import TenantSnapshot
from tests import devicedata as dd
from tests import graphdata as gd
from tests import groupdata as grp
from tests import intunedata as idata
from tests.conftest import login, make_user
from tests.seed import seed_tenant
from tests.test_collectors import SKUS

CONFIG = "DeviceManagementConfiguration.Read.All"
APPS = "DeviceManagementApps.Read.All"
BITLOCKER = "BitLockerKey.ReadBasic.All"
PIM_GROUPS = "PrivilegedEligibilitySchedule.Read.AzureADGroup"


def _report(db, tenant_id, key="intune-policies"):
    return data_reports.build(db, store.get_tenant(db, tenant_id), key, datetime.now(UTC))


def _table(report, title) -> list[dict]:
    section = next(s for s in report.sections if s.title == title)
    return [dict(zip(section.columns, row, strict=True)) for row in section.rows]


def _flags(report, title) -> dict[str, str | None]:
    section = next(s for s in report.sections if s.title == title)
    return {str(row[0]): flag for row, flag in zip(section.rows, section.flags, strict=True)}


def _results(db, tenant_id):
    return {r.key: r for r in evaluate(TenantSnapshot(db, store.get_tenant(db, tenant_id)))}


def _run(db, tenant_id, collector):
    return db.execute(
        "SELECT * FROM runs WHERE tenant_id = ? AND collector = ? ORDER BY id DESC LIMIT 1", (tenant_id, collector)
    ).fetchone()


def _asked(db, run_id) -> list[str]:
    return [r["endpoint"] for r in db.execute("SELECT endpoint FROM run_requests WHERE run_id = ?", (run_id,))]


# -- Collector ------------------------------------------------------------------------------------------------


def test_collects_every_kind_of_policy_with_its_assignments(db, two_tenants):
    alpha, _ = two_tenants
    seed_tenant(db, alpha)
    run = _run(db, alpha, "intune_policies")
    assert run["status"] == "succeeded"
    snap = TenantSnapshot(db, store.get_tenant(db, alpha))
    status = snap.objects("intune_policies", "intune_policy_status")[0]
    assert status["available"] is True and status["waiting"] == []
    assert set(status["parts"].values()) == {"read"}
    compliance = {p["_id"]: p for p in snap.objects("intune_policies", "compliance_policy")}
    win = compliance["cp-win"]
    assert (win["name"], win["type"]) == ("Windows baseline", "windows10CompliancePolicy")
    assert win["assignments"] == [{"kind": "all_devices", "groupId": None, "filter": None}]
    assert win["settings"]["bitLockerEnabled"] is True
    assert "earlyLaunchAntiMalwareDriverEnabled" not in win["settings"]  # only what the report describes
    assert compliance["cp-ios"]["assignments"] == []
    profiles = {p["_id"]: p for p in snap.objects("intune_policies", "configuration_profile")}
    assert profiles["ring-broad"]["assignments"] == [
        {"kind": "group", "groupId": "g1", "filter": None},
        {"kind": "exclude", "groupId": "g2", "filter": None},
    ]
    assert profiles["ring-pilot"]["assignments"][0]["filter"] == "include"
    assert profiles["ep-win"]["settings"]["firewallProfileDomain"] == "allowed"  # just whether it's on
    settings = {p["_id"]: p for p in snap.objects("intune_policies", "settings_policy")}
    assert settings["sc-av"]["templateFamily"] == "endpointSecurityAntivirus"
    assert settings["sc-filevault"]["isAssigned"] is False
    protection = {p["_id"]: p for p in snap.objects("intune_policies", "app_protection")}
    assert protection["T_ios-office"]["platform"] == "iOS" and protection["T_android-outlook"]["platform"] == "Android"
    metrics = {
        r["key"]: r["value"] for r in db.execute("SELECT key, value FROM metrics WHERE run_id = ?", (run["id"],))
    }
    assert metrics["intune.compliance_policies"] == 2
    assert metrics["intune.unassigned_policies"] == 3  # iPhone draft, Mac restrictions, Mac FileVault
    assert metrics["intune.paused_update_rings"] == 1
    endpoints = _asked(db, run["id"])
    assert "/deviceManagement/configurationPolicies" in " ".join(endpoints)


def test_waits_for_the_client_to_approve_without_asking_microsoft(db, two_tenants):
    alpha, _ = two_tenants
    seed_tenant(db, alpha, our_app={"missing": (CONFIG, APPS)})
    run = _run(db, alpha, "intune_policies")
    assert run["status"] == "succeeded"  # waiting isn't partial
    assert not _asked(db, run["id"])
    report = _report(db, alpha)
    assert report.notice.startswith("Waiting for Alpha Ltd to approve the new read-only permissions this needs")
    assert CONFIG in report.notice and APPS in report.notice
    tiles = {t.label: (t.value, t.foot) for t in report.tiles}
    assert tiles["Compliance policies"] == ("Waiting", "For the client's approval")
    assert tiles["App protection"] == ("Waiting", "For the client's approval")
    compliance = next(s for s in report.sections if s.title == "Compliance policies")
    assert compliance.empty == "Compliance policies will show once the client approves the new permission."
    coverage = _table(report, "Coverage by platform")
    assert {r["Platform"]: r["Enrolled devices"] for r in coverage} == {"Windows": 2, "iOS/iPadOS": 1}
    assert all(r["Compliance policies"] == "" for r in coverage)  # not "0": not read
    assert _results(db, alpha)["intune_compliance_coverage"].status == "unavailable"


def test_a_refusal_before_the_connection_check_knows_counts_as_waiting(db, two_tenants):
    alpha, _ = two_tenants
    denied = (403, idata.DENIED)
    seed_tenant(db, alpha, intune_kw={"status": {"iosManagedAppProtections": denied}})
    assert _run(db, alpha, "intune_policies")["status"] == "succeeded"
    status = TenantSnapshot(db, store.get_tenant(db, alpha)).objects("intune_policies", "intune_policy_status")[0]
    assert status["parts"] == {
        "compliance": "read",
        "configuration": "read",
        "settings": "read",
        "app_protection": "waiting",
    }
    assert status["waiting"] == [APPS]
    report = _report(db, alpha)
    assert APPS in report.notice and CONFIG not in report.notice
    assert not _table(report, "App protection for phones")


def test_tenant_without_intune_is_not_asked(db, two_tenants):
    alpha, _ = two_tenants
    no_intune = [
        {**s, "servicePlans": [p for p in s["servicePlans"] if p["servicePlanName"] != "INTUNE_A"]} for s in SKUS
    ]
    seed_tenant(db, alpha, skus=no_intune)
    run = _run(db, alpha, "intune_policies")
    assert run["status"] == "succeeded" and not _asked(db, run["id"])
    report = _report(db, alpha)
    assert report.unavailable == "This client has no Intune licence, so there are no Intune policies to report."
    assert _results(db, alpha)["intune_compliance_coverage"].status == "unavailable"


def test_a_failed_list_is_partial_and_says_so(db, two_tenants):
    alpha, _ = two_tenants
    broken = (500, {"error": {"code": "InternalServerError", "message": "Something went wrong."}})
    seed_tenant(db, alpha, intune_kw={"status": {"deviceConfigurations": broken}})
    assert _run(db, alpha, "intune_policies")["status"] == "partial"
    report = _report(db, alpha)
    rings = next(s for s in report.sections if s.title == "Windows update rings")
    assert rings.empty.startswith("Windows update rings couldn't be read")
    assert {t.label: t.value for t in report.tiles}["Update rings paused"] == "Not read"


# -- Report ---------------------------------------------------------------------------------------------------


def test_intune_policies_report(db, two_tenants):
    alpha, _ = two_tenants
    seed_tenant(db, alpha)
    report = _report(db, alpha)
    assert report.unavailable is None and not report.notice
    assert [s.title for s in report.sections] == [
        "Coverage by platform",
        "Needs attention",
        "Compliance policies",
        "Windows update rings",
        "Configuration profiles",
        "Endpoint security and settings catalog",
        "App protection for phones",
    ]

    coverage = {r["Platform"]: r for r in _table(report, "Coverage by platform")}
    assert list(coverage) == ["Windows", "iOS/iPadOS", "Android"]  # macOS: no devices, nothing assigned
    assert coverage["Windows"]["Enrolled devices"] == 2 and coverage["Windows"]["Compliance policies"] == 1
    assert coverage["Windows"]["Configuration profiles"] == 3  # both rings and the firewall
    assert coverage["Windows"]["App protection"] == ""
    assert coverage["iOS/iPadOS"]["Compliance policies"] == 0 and coverage["iOS/iPadOS"]["App protection"] == 1
    assert _flags(report, "Coverage by platform")["iOS/iPadOS"] == "medium"

    attention = [(r["Policy"], r["What to check"]) for r in _table(report, "Needs attention")]
    assert attention == [
        ("Android Outlook", "No PIN to open company apps"),
        ("iOS/iPadOS", "Devices are enrolled but no compliance policy is assigned (1 device)"),
        ("Update ring Pilot", "Security updates paused"),
        ("Android Outlook", "Company data can be shared with any app"),
        ("iPhone draft", "Assigned to nobody, so it does nothing"),
        ("Mac restrictions", "Assigned to nobody, so it does nothing"),
        ("Mac FileVault", "Assigned to nobody, so it does nothing"),
    ]

    compliance = {r["Policy"]: r for r in _table(report, "Compliance policies")}
    win = compliance["Windows baseline"]
    assert win["Platform"] == "Windows" and win["Assigned to"].parts == ("All devices",)
    assert win["Requires"].parts == (
        "Password of 8+ characters",
        "BitLocker",
        "Secure Boot",
        "Microsoft Defender on",
        "Version 10.0.22631 or later",
    )
    assert compliance["iPhone draft"]["Assigned to"] == "Nobody"

    rings = {r["Ring"]: r for r in _table(report, "Windows update rings")}
    broad = rings["Update ring Broad"]
    assert broad["Security updates after (days)"] == 7 and broad["Restart deadline (days)"] == 3
    assert broad["Assigned to"].parts == ("Sales", "Except Project Phoenix")
    assert rings["Update ring Pilot"]["Paused"] == "Security"
    assert rings["Update ring Pilot"]["Assigned to"] == "Board (filtered)"
    assert rings["Update ring Pilot"]["Drivers"] == "Left out"
    assert _flags(report, "Windows update rings") == {"Update ring Broad": None, "Update ring Pilot": "medium"}

    profiles = {r["Profile"]: r for r in _table(report, "Configuration profiles")}
    assert profiles["Windows firewall"]["Settings"].parts == ("Firewall on", "Encrypts with BitLocker")
    assert profiles["Windows firewall"]["Type"] == "Endpoint protection"
    assert profiles["Update ring Broad"]["Settings"] == (
        "Security updates after 7 days, feature updates after 60 days, restart deadline 3 days"
    )
    assert profiles["Mac restrictions"]["Platform"] == "macOS"

    catalog = {r["Policy"]: r for r in _table(report, "Endpoint security and settings catalog")}
    assert (catalog["Defender antivirus"]["Kind"], catalog["Defender antivirus"]["Settings"]) == ("Antivirus", 14)
    assert catalog["Mac FileVault"]["Assigned"] == "No"

    protection = {r["Policy"]: r for r in _table(report, "App protection for phones")}
    ios = protection["iOS Outlook and Teams"]
    assert (ios["PIN"], ios["Sharing data"], ios["Backups"], ios["Encryption"]) == (
        "Required (6+ digits)",
        "Managed apps only",
        "Blocked",
        "When locked",
    )
    assert ios["Assigned to"] == "All users"
    assert protection["Android Outlook"]["PIN"] == "Not required"
    assert _flags(report, "App protection for phones")["Android Outlook"] == "medium"

    tiles = {t.label: t for t in report.tiles}
    assert (tiles["Compliance policies"].value, tiles["Compliance policies"].foot) == (2, "1 assigned to someone")
    assert (tiles["Platforms without compliance"].value, tiles["Platforms without compliance"].tone) == (1, "bad")
    assert tiles["Policies assigned to nobody"].value == 3
    assert (tiles["Update rings paused"].value, tiles["Update rings paused"].foot) == (1, "Of 2 Windows update rings")
    assert (tiles["App protection"].value, tiles["App protection"].tone) == ("iOS and Android", "ok")


def test_intune_policies_report_before_collection(db, two_tenants):
    _, bravo = two_tenants
    report = _report(db, bravo)
    assert report.unavailable.startswith("Intune policies haven't been collected yet")


def test_intune_policies_page_and_downloads(client, db, two_tenants):
    alpha, _ = two_tenants
    seed_tenant(db, alpha)
    make_user(db, "admin")
    login(client, "admin")
    page = client.get(f"/tenants/{alpha}/data/intune-policies")
    assert page.status_code == 200
    html = page.get_data(as_text=True)
    assert "Update ring Pilot" in html and "Coverage by platform" in html
    for ext in ("xlsx", "csv"):
        assert client.get(f"/tenants/{alpha}/data/intune-policies.{ext}").status_code == 200


# -- Checks ---------------------------------------------------------------------------------------------------


def test_compliance_coverage_check(db, two_tenants):
    alpha, bravo = two_tenants
    seed_tenant(db, alpha)
    result = _results(db, alpha)["intune_compliance_coverage"]
    assert result.status == "fail" and result.summary == "No compliance policy for iOS/iPadOS"
    assert [(f.severity, f.subject, f.object_id) for f in result.findings] == [
        ("medium", "iOS/iPadOS", "platform:iOS/iPadOS")
    ]
    ios = {**idata.COMPLIANCE_POLICIES[1], "assignments": [idata.ALL_USERS]}
    seed_tenant(db, bravo, intune_kw={"compliance": [idata.COMPLIANCE_POLICIES[0], ios]})
    result = _results(db, bravo)["intune_compliance_coverage"]
    assert result.status == "pass" and result.summary == "All 2 platforms in use have a compliance policy"


# -- BitLocker recovery keys ----------------------------------------------------------------------------------


def test_recovery_keys_are_read_without_the_key(db, two_tenants):
    alpha, _ = two_tenants
    seed_tenant(db, alpha)
    run = _run(db, alpha, "devices")
    assert run["status"] == "succeeded"
    snap = TenantSnapshot(db, store.get_tenant(db, alpha))
    keys = {r["_id"]: r for r in snap.objects("devices", "recovery_key")}
    assert keys["d-ann"]["keys"] == 3 and keys["d-ann"]["osVolume"] is True
    assert keys["d-ann"]["latest"] == dd.RECOVERY_KEYS[1]["createdDateTime"]
    assert "key" not in keys["d-ann"]
    # Asking for one key's own record with $select=key is what reveals it; the list is all Office Sentry asks.
    assert [e for e in _asked(db, run["id"]) if "bitlocker" in e] == [
        "v1.0 /informationProtection/bitlocker/recoveryKeys"
    ]


# A recovery key's shape: eight groups of six digits.
RECOVERY_KEY = "314159-265358-979323-846264-338327-950288-419716-939937"


def test_a_recovery_key_never_reaches_storage_or_output(client, cfg, db, two_tenants):
    """Keys are never downloadable. The list Office Sentry reads (BitLockerKey.ReadBasic.All) carries no key, and
    even if Microsoft sent one, nothing would keep it: not the stored data, the report or its downloads, the
    support package, or the log."""
    alpha, _ = two_tenants
    seed_tenant(db, alpha, recovery_keys=[{**k, "key": RECOVERY_KEY} for k in dd.RECOVERY_KEYS])
    assert _run(db, alpha, "devices")["status"] == "succeeded"

    # Microsoft returns a key only from one key's own record (``/recoveryKeys/{id}?$select=key``). That path
    # isn't catalogued, so a call to it fails test_ratelimit, and the list is the only BitLocker call.
    bitlocker = [e.path for e in endpoints.GRAPH if "bitlocker" in e.path]
    assert bitlocker == ["/informationProtection/bitlocker/recoveryKeys"]

    # Stored: every object of every run, unpacked, and every text column of every table.
    runs = [r["id"] for r in db.execute("SELECT id FROM runs WHERE tenant_id = ?", (alpha,))]
    for run_id in runs:
        for (kind,) in db.execute("SELECT DISTINCT kind FROM objects WHERE run_id = ?", (run_id,)).fetchall():
            assert RECOVERY_KEY not in json.dumps(store.run_objects(db, run_id, kind)), kind
    for (table,) in db.execute("SELECT name FROM sqlite_master WHERE type = 'table'").fetchall():
        for row in db.execute(f"SELECT * FROM {table}"):  # noqa: S608 (table names from sqlite_master)
            assert not any(isinstance(v, str) and RECOVERY_KEY in v for v in tuple(row)), table

    # Shown: the report, its page and downloads, and the support package.
    report = _report(db, alpha, "devices")
    assert RECOVERY_KEY not in json.dumps([s.rows for s in report.sections] + report.notes, default=str)
    make_user(db, "admin")
    login(client, "admin")
    for ext in ("", ".csv", ".xlsx"):
        resp = client.get(f"/tenants/{alpha}/data/devices{ext}")
        assert resp.status_code == 200
        body = resp.data
        if ext == ".xlsx":
            book = load_workbook(BytesIO(body))
            body = "\n".join(str(c) for ws in book for row in ws.iter_rows(values_only=True) for c in row).encode()
        assert RECOVERY_KEY.encode() not in body, ext
    package = support.build(db, cfg)
    assert not any(RECOVERY_KEY.encode() in data for data in package.files.values())


def test_devices_report_shows_which_computers_have_a_saved_key(db, two_tenants, monkeypatch):
    alpha, bravo = two_tenants
    seed_tenant(db, alpha)
    report = _report(db, alpha, "devices")
    assert [s.title for s in report.sections][-2:] == ["BitLocker recovery keys", "Entra ID device records"]
    keys = {r["Device"]: r for r in _table(report, "BitLocker recovery keys")}
    assert list(keys) == ["DESK-RECEPTION", "LAPTOP-ANN", "LAPTOP-BOB"]  # not OLD-PC (gone), not the iPhone
    assert (keys["LAPTOP-ANN"]["Recovery key saved"], keys["LAPTOP-ANN"]["Keys saved"]) == ("Yes", 3)
    assert keys["DESK-RECEPTION"]["Encrypted"] == "" and keys["DESK-RECEPTION"]["Recovery key saved"] == "No"
    assert not any(_flags(report, "BitLocker recovery keys").values())

    # Ann's laptop is encrypted, so without a saved key it can't be unlocked if it asks for one.
    monkeypatch.setattr(dd, "RECOVERY_KEYS", [k for k in dd.RECOVERY_KEYS if k["deviceId"] != "d-ann"])
    seed_tenant(db, bravo)
    report = _report(db, bravo, "devices")
    assert _flags(report, "BitLocker recovery keys")["LAPTOP-ANN"] == "medium"
    attention = {r["Device"]: r for r in _table(report, "Needs attention")}
    assert attention["LAPTOP-ANN"]["Problems"] == "Encrypted, but no BitLocker recovery key saved in Entra ID"
    result = _results(db, bravo)["bitlocker_keys"]
    assert [(f.severity, f.subject) for f in result.findings] == [("medium", "LAPTOP-ANN"), ("low", "DESK-RECEPTION")]


def test_bitlocker_check(db, two_tenants, monkeypatch):
    alpha, bravo = two_tenants
    seed_tenant(db, alpha)
    result = _results(db, alpha)["bitlocker_keys"]
    # LAPTOP-BOB isn't encrypted at all (device_encryption says so); DESK-RECEPTION isn't in Intune.
    assert result.status == "fail"
    assert result.summary == "1 of 3 joined Windows computers without a saved recovery key"
    assert [(f.severity, f.detail) for f in result.findings] == [
        ("low", "No recovery key saved in Entra ID (encryption not known)")
    ]
    monkeypatch.setattr(dd, "ENTRA_DEVICES", [dd.ENTRA_DEVICES[0], dd.ENTRA_DEVICES[4]])
    seed_tenant(db, bravo)
    result = _results(db, bravo)["bitlocker_keys"]
    assert (
        result.status == "pass" and result.summary == "The 1 joined Windows computer is covered by a saved recovery key"
    )


def test_recovery_keys_wait_for_approval(db, two_tenants):
    alpha, _ = two_tenants
    seed_tenant(db, alpha, our_app={"missing": (BITLOCKER,)})
    run = _run(db, alpha, "devices")
    assert run["status"] == "succeeded"
    assert not [e for e in _asked(db, run["id"]) if "bitlocker" in e]
    report = _report(db, alpha, "devices")
    assert "BitLocker recovery keys" not in [s.title for s in report.sections]
    assert any(BITLOCKER in n and n.startswith("Whether each computer's BitLocker") for n in report.notes)
    assert _results(db, alpha)["bitlocker_keys"].status == "unavailable"


def test_a_refused_recovery_key_list_counts_as_waiting(db, two_tenants, monkeypatch):
    alpha, _ = two_tenants
    import tests.seed as seed

    original = seed.paged

    def refusing(rsps, url, items, **extra):
        if url.endswith("/bitlocker/recoveryKeys"):
            rsps.get(url, status=403, json=idata.DENIED)
        else:
            original(rsps, url, items, **extra)

    monkeypatch.setattr(seed, "paged", refusing)
    seed_tenant(db, alpha)
    assert _run(db, alpha, "devices")["status"] == "succeeded"
    snap = TenantSnapshot(db, store.get_tenant(db, alpha))
    assert snap.objects("devices", "bitlocker_status")[0]["state"] == "waiting"
    assert snap.objects("devices", "recovery_key") == []


# -- PIM for groups -------------------------------------------------------------------------------------------


ROLE_GROUP = {
    "id": "a-grp",
    "roleDefinitionId": gd.GA,
    "principalId": "rg1",
    "directoryScopeId": "/",
    "principal": {"@odata.type": "#microsoft.graph.group", "id": "rg1", "displayName": "Tenant admins"},
}
STANDBY = {
    "id": "s1",
    "groupId": "rg1",
    "principalId": "u6",
    "accessId": "member",
    "memberType": "direct",
    "principal": {
        "@odata.type": "#microsoft.graph.user",
        "id": "u6",
        "displayName": "Carl",
        "userPrincipalName": "carl@x.com",
    },
}


# PIM for groups needs Entra ID P2.
P2_SKUS = [
    {
        **SKUS[0],
        "servicePlans": [
            *SKUS[0]["servicePlans"],
            {"servicePlanName": "AAD_PREMIUM_P2", "provisioningStatus": "Success"},
        ],
    },
    *SKUS[1:],
]


def test_admins_report_lists_eligible_members_of_role_groups(db, two_tenants, monkeypatch):
    alpha, bravo = two_tenants
    monkeypatch.setattr(gd, "ROLE_ASSIGNMENTS", [*gd.ROLE_ASSIGNMENTS, ROLE_GROUP])
    monkeypatch.setitem(grp.MEMBERS, "rg1", [])
    monkeypatch.setattr(gd, "GROUP_ELIGIBILITY", {"rg1": [STANDBY]})
    seed_tenant(db, alpha, skus=P2_SKUS)
    report = _report(db, alpha, "admins")
    holders = {r["Account"]: r for r in _table(report, "Privileged role holders")}
    assert "Tenant admins (eligible member)" in str(holders["carl@x.com"])
    assert not any("PIM for groups" in n for n in report.notes)
    seed_tenant(db, bravo)  # Business Premium: no P2, so nobody can be eligible and nothing is asked
    holders = {r["Account"] for r in _table(_report(db, bravo, "admins"), "Privileged role holders")}
    assert "carl@x.com" not in holders


def test_admins_report_says_when_pim_for_groups_waits_for_approval(db, two_tenants, monkeypatch):
    alpha, _ = two_tenants
    monkeypatch.setattr(gd, "ROLE_ASSIGNMENTS", [*gd.ROLE_ASSIGNMENTS, ROLE_GROUP])
    monkeypatch.setitem(grp.MEMBERS, "rg1", [])
    seed_tenant(db, alpha, skus=P2_SKUS, our_app={"missing": (PIM_GROUPS,)})
    report = _report(db, alpha, "admins")
    note = next(n for n in report.notes if "PIM for groups" in n)
    assert PIM_GROUPS in note and "Waiting for Alpha Ltd" in note
