"""Sign-in activity: collector, report, the risky sign-in check and the account page."""

import copy
import json
from datetime import UTC, datetime, timedelta
from io import BytesIO
from urllib.parse import parse_qs, urlsplit

import responses
from openpyxl import load_workbook

from officesentry import signins, store
from officesentry.collectors import sign_ins
from officesentry.graph import GRAPH_ROOT, GraphClient
from officesentry.reports import data_reports, export, library, user_profile, views
from officesentry.rules import evaluate
from officesentry.snapshot import TenantSnapshot
from officesentry.worker import execute_run
from tests import graphdata as gd
from tests.conftest import login, make_user
from tests.seed import seed_tenant
from tests.test_collectors import SKUS, USERS

BETA = f"{GRAPH_ROOT}/beta"
P2_SKUS = [dict(SKUS[0], servicePlans=SKUS[0]["servicePlans"] + [
    {"servicePlanName": "AAD_PREMIUM_P2", "provisioningStatus": "Success"}])]


def _seed(db, tenant_id, **kwargs):
    kwargs.setdefault("signins", list(gd.SIGNINS))  # the counts below leave out the seeded legacy sign-ins
    seed_tenant(db, tenant_id, users=[*USERS, gd.INFO_USER], **kwargs)
    return store.get_tenant(db, tenant_id)


def _sections(report):
    return {s.title: s for s in report.sections}


def _run(db, tenant_id, caps=None):
    if caps is not None:
        store.set_capabilities(db, tenant_id, caps)
    store.enqueue_run(db, tenant_id, "sign_ins")
    run = store.claim_next_run(db)
    graph = lambda t, c: GraphClient(lambda: "t", sleep=lambda s: None)  # noqa: E731
    return run["id"], execute_run(db, run, graph)


# -- What the codes mean ----------------------------------------------------------------------


def test_error_codes_in_plain_english():
    assert signins.meaning(50126) == "Wrong user name or password"
    assert signins.meaning("53003").startswith("Blocked by Conditional Access")
    assert signins.meaning(50057) == "Account is disabled"
    assert signins.meaning(99999, "Something Microsoft said.") == "Something Microsoft said."
    assert signins.meaning(99999, "Other.") == "Error 99999"
    # Seen on the test tenant: Microsoft's text with its placeholders unfilled.
    template = "The {redirectTerm} '{replyAddress}' specified in the request does not match. {detail}"
    assert signins.meaning(99999, template) == "Error 99999"
    assert "isn't set up with" in signins.meaning(50011, template)
    assert signins.meaning(None) == "Signed in"
    assert [signins.outcome(c) for c in (0, 50074, 50076, 50140, 50126, 53003)] == \
        ["success", "interrupted", "interrupted", "interrupted", "failure", "failure"]
    assert signins.risk_severity("high", "atRisk") == "high"
    assert signins.risk_severity("low", "confirmedCompromised") == "high"
    assert signins.risk_severity("medium", "atRisk") == "medium"
    assert signins.risk_severity("medium", "remediated") is None and signins.risk_severity("low", "atRisk") is None


# -- Collector --------------------------------------------------------------------------------


def test_collector_summarises_the_week(db, two_tenants):
    alpha, _ = two_tenants
    tenant = _seed(db, alpha)
    snap = TenantSnapshot(db, tenant)
    assert snap.runs["sign_ins"]["status"] == "succeeded"
    status = signins.status(snap)
    assert status["available"] and status["days"] == 7 and status["signIns"] == len(gd.SIGNINS)
    since = datetime.fromisoformat(status["since"].replace("Z", "+00:00"))
    assert timedelta(days=6, hours=23) < datetime.now(UTC) - since < timedelta(days=7, hours=1)
    assert status["riskVisible"] is False
    accounts = {a["user"]: a for a in snap.objects("sign_ins", "sign_in_account")}
    ann, bob, info = accounts["ann@x.com"], accounts["bob@x.com"], accounts["info@x.com"]
    assert (ann["succeeded"], ann["mfa"], ann["singleFactor"], ann["interrupted"], ann["failed"]) == (27, 26, 1, 2, 1)
    assert ann["countries"][0] == "GB" and "FR" in ann["countries"] and ann["risky"] == 0
    assert (bob["failed"], bob["credentialFailures"], bob["lockouts"], bob["succeeded"]) == (9, 8, 1, 0)
    assert bob["lastFailureCode"] == 50057 and bob["failureIps"][0] == "198.51.100.7" and bob["failureIpCount"] == 3
    assert bob["failureCodes"] == {"50126": 7, "50053": 1, "50057": 1}
    assert (info["singleFactor"], info["mfa"]) == (2, 0)
    assert accounts["nobody@x.com"]["failed"] == 1
    countries = {c["country"]: c for c in snap.objects("sign_ins", "sign_in_country")}
    assert set(countries) == {"GB", "FR", "RU", "BR", "US"}
    assert countries["BR"]["accounts"] == 5 and countries["BR"]["failed"] == 5
    spray = {i["ip"]: i for i in snap.objects("sign_ins", "sign_in_failure_ip")}["203.0.113.66"]
    assert spray["accounts"] == 5 and spray["credentialFailures"] == 4 and spray["succeeded"] == 0
    failures = {f["code"]: f for f in snap.objects("sign_ins", "sign_in_failure")}
    assert failures[50126]["signIns"] == 10 and failures[50074]["signIns"] == 1
    assert failures[50126]["reason"].startswith("Invalid username or password")
    apps = {a["app"]: a for a in snap.objects("sign_ins", "sign_in_app")}
    assert apps["Azure Portal"]["singleFactor"] == 1
    # Only summaries are kept: no raw sign-in, and no risk without Entra ID P2.
    assert snap.objects("sign_ins", "risky_sign_in") == []
    assert snap.metric("signins.total") == len(gd.SIGNINS) and snap.metric("signins.repeated_failures") == 1
    assert snap.metric("signins.risky") is None


@responses.activate
def test_collector_reads_the_log_page_by_page(db, two_tenants):
    alpha, _ = two_tenants
    page2 = f"{BETA}/auditLogs/signIns?$skiptoken=abc"
    responses.get(page2, json={"value": gd.SIGNINS[25:]})
    responses.get(f"{BETA}/auditLogs/signIns", json={"value": gd.SIGNINS[:25], "@odata.nextLink": page2})
    run_id, status = _run(db, alpha)
    assert status == "succeeded"
    first = parse_qs(urlsplit(responses.calls[0].request.url).query)
    assert first["$filter"][0].startswith("createdDateTime ge ") and first["$top"] == ["1000"]
    # The first run reads back as far as Microsoft keeps the log.
    since = datetime.fromisoformat(first["$filter"][0].split(" ")[2].replace("Z", "+00:00"))
    assert abs(datetime.now(UTC) - since - timedelta(days=sign_ins.LOG_DAYS)) < timedelta(minutes=5)
    assert all(c.request.method == "GET" for c in responses.calls) and len(responses.calls) == 2
    [st] = store.run_objects(db, run_id, "sign_in_status")
    assert st["signIns"] == len(gd.SIGNINS) and not st["truncated"]
    assert st["historyFrom"][:10] == since.strftime("%Y-%m-%d")


@responses.activate
def test_a_run_that_reaches_the_page_budget_carries_on_next_time(db, two_tenants, monkeypatch):
    alpha, bravo = two_tenants
    rows = _recent_signins()
    _signin_log(rows)
    expected = _stored(db, _run(db, alpha)[0])

    # Three pages a run: the newest sign-ins come first, the summary says what's missing, and each later run
    # carries on from where the last stopped until the week is complete, with nothing counted twice.
    monkeypatch.setattr(sign_ins, "SIGNIN_PAGES", 3)
    run_id, status = _run(db, bravo)
    assert status == "partial"
    [st] = store.run_objects(db, run_id, "sign_in_status")
    assert st["truncated"] and st["signIns"] == 3 * 7
    error = db.execute("SELECT error, truncated FROM runs WHERE id = ?", (run_id,)).fetchone()
    assert "the next run carries on from there" in error["error"] and error["truncated"] == 1
    for _ in range(5):
        run_id, status = _run(db, bravo)
        if status == "succeeded":
            break
    assert status == "succeeded"
    assert _stored(db, run_id) == expected
    assert store.run_objects(db, run_id, "sign_in_status")[0]["signIns"] == len(rows)


@responses.activate
def test_without_entra_p1_the_log_is_unavailable(db, two_tenants):
    alpha, bravo = two_tenants
    # Known from the licences: nothing is asked of Microsoft.
    run_id, status = _run(db, alpha, {"entra_p1": False})
    assert status == "succeeded" and not responses.calls
    assert store.run_objects(db, run_id, "sign_in_status")[0] == {
        "_id": "status", "_kind": "sign_in_status", "available": False, "reason": "licence"}
    report = data_reports.build(db, store.get_tenant(db, alpha), "sign-ins")
    assert "Entra ID P1" in report.unavailable and not report.sections
    results = {r.key: r for r in evaluate(TenantSnapshot(db, store.get_tenant(db, alpha)))}
    assert results["risky_sign_ins"].status == "unavailable"

    # Not known: Microsoft says so.
    responses.get(f"{BETA}/auditLogs/signIns", status=403, json={"error": {
        "code": "Authentication_RequestFromNonPremiumTenantOrB2CTenant", "message": "Needs a premium licence"}})
    run_id, status = _run(db, bravo)
    assert status == "partial"
    assert store.run_objects(db, run_id, "sign_in_status")[0]["reason"] == "licence"
    assert "Entra ID P1" in data_reports.build(db, store.get_tenant(db, bravo), "sign-ins").unavailable


def test_not_collected_yet(db, two_tenants):
    alpha, _ = two_tenants
    report = data_reports.build(db, store.get_tenant(db, alpha), "sign-ins")
    assert "hasn't been collected" in report.unavailable
    assert user_profile.sign_ins_part(user_profile.Person(
        TenantSnapshot(db, store.get_tenant(db, alpha)), {"_id": "u1"}, report.generated_at)) is None


# -- Report -----------------------------------------------------------------------------------


def test_report_leads_with_what_to_act_on(db, two_tenants):
    alpha, _ = two_tenants
    tenant = _seed(db, alpha)
    report = data_reports.build(db, tenant, "sign-ins")
    assert report.unavailable is None
    s = _sections(report)
    assert list(s) == ["Accounts with repeated failures", "Addresses with many failures",
                       "Successful sign-ins from unusual countries", "Successful sign-ins without MFA",
                       "Failure reasons", "Sign-ins by country", "Sign-ins by app", "Every account"]
    assert not any(sec.compare for sec in report.sections)
    repeated = s["Accounts with repeated failures"]
    assert [r[0] for r in repeated.rows] == ["bob@x.com"] and repeated.flags == ["high"]
    assert repeated.rows[0][1:5] == [9, 8, 1, 0] and repeated.rows[0][-1] == "Account is disabled"
    addresses = {r[0]: (r, f) for r, f in zip(s["Addresses with many failures"].rows,
                                              s["Addresses with many failures"].flags, strict=True)}
    assert addresses["203.0.113.66"][0][3] == 5 and addresses["203.0.113.66"][1] == "high"
    assert addresses["198.51.100.7"][1] == "medium"
    unusual = s["Successful sign-ins from unusual countries"]
    assert [(r[0], r[1], r[2]) for r in unusual.rows] == [("ann@x.com", "FR", 1)]
    single = {r[0]: f for r, f in zip(s["Successful sign-ins without MFA"].rows,
                                      s["Successful sign-ins without MFA"].flags, strict=True)}
    assert single == {"info@x.com": "medium", "ann@x.com": "low"}
    reasons = {r[0]: r for r in s["Failure reasons"].rows}
    assert reasons["50126"][1:4] == ["Wrong user name or password", "Failure", 10]
    assert reasons["50074"][2] == "Prompt"
    countries = {r[0]: r for r in s["Sign-ins by country"].rows}
    assert countries["FR"][-1] == "Yes" and countries["GB"][-1] == "No"
    tiles = {t.label: t for t in report.tiles}
    assert tiles["Sign-ins"].value == len(gd.SIGNINS)
    assert tiles["Repeated failures"].value == 1 and tiles["Repeated failures"].tone == "bad"
    assert tiles["Without MFA"].value == 3 and tiles["Unusual countries"].value == 1
    assert tiles["Failure rate"].value == "33%"  # 14 of 43, prompts left out
    assert "Risky sign-ins" not in tiles and any("Entra ID P2" in n for n in report.notes)
    # Headline figures open their rows.
    links = views.headline_links(report, "/r")
    assert "#t-failure-reasons" in links["tiles"][1] and "kind%3AFailure" in links["tiles"][1]
    # Names link to the account's page.
    assert repeated.rows[0][0].page_href == f"/tenants/{alpha}/users/u2"


def test_report_page_pdf_and_spreadsheet(client, db, two_tenants):
    alpha, _ = two_tenants
    _seed(db, alpha)
    make_user(db, "admin")
    login(client, "admin")
    page = client.get(f"/tenants/{alpha}/data/sign-ins")
    assert page.status_code == 200
    text = page.get_data(as_text=True)
    assert "Accounts with repeated failures" in text and "Wrong user name or password" in text
    wb = load_workbook(BytesIO(client.get(f"/tenants/{alpha}/data/sign-ins.xlsx").data))
    assert "Summary" in wb.sheetnames and "Accounts with repeated failures" in wb.sheetnames
    pdf = client.get(f"/tenants/{alpha}/data/sign-ins.pdf")
    assert pdf.status_code == 200
    if export.pdf_available():
        assert pdf.data[:4] == b"%PDF"
    assert client.get("/across/sign-ins").status_code == 200


def test_report_needs_a_tenant_the_user_can_see(client, db, two_tenants):
    alpha, bravo = two_tenants
    _seed(db, alpha)
    make_user(db, "analyst", role="analyst", tenants=[bravo])
    login(client, "analyst")
    assert client.get(f"/tenants/{alpha}/data/sign-ins").status_code == 404
    assert client.get(f"/tenants/{alpha}/data/sign-ins.pdf").status_code == 404


def test_library_lists_the_report_for_compromised_accounts():
    use = next(u for u in library.USES if u.key == "compromise")
    assert use.reports[0] == "sign-ins"
    entry = next(e for a in library.catalogue() for e in a.entries if e.key == "sign-ins")
    assert entry.area == "Identity and access" and "password spray" in entry.search


# -- Risky sign-ins (Entra ID P2) and the check ----------------------------------------------


def test_risky_sign_ins_with_entra_p2(db, two_tenants):
    alpha, _ = two_tenants
    tenant = _seed(db, alpha, skus=P2_SKUS)
    snap = TenantSnapshot(db, tenant)
    assert signins.status(snap)["riskVisible"] is True
    [risky] = snap.objects("sign_ins", "risky_sign_in")
    assert (risky["user"], risky["level"], risky["state"], risky["succeeded"]) == ("ann@x.com", "medium", "atRisk",
                                                                                   True)
    result = {r.key: r for r in evaluate(snap)}["risky_sign_ins"]
    assert result.status == "fail" and result.summary.startswith("1 account signed in despite")
    [finding] = result.findings
    assert (finding.severity, finding.subject, finding.object_id) == ("medium", "ann@x.com", "u1")
    assert f"{gd.SIGNIN_BASE - timedelta(days=2)} from FR (medium risk, at risk)" in finding.detail
    report = data_reports.build(db, tenant, "sign-ins")
    s = _sections(report)
    assert report.sections[0].title == "Risky sign-ins" and s["Risky sign-ins"].flags == ["medium"]
    tiles = {t.label: t for t in report.tiles}
    assert tiles["Risky sign-ins"].value == 1 and tiles["Risky sign-ins"].tone == "bad"
    assert not any("Entra ID P2" in n for n in report.notes)


def test_risky_check_passes_and_is_unavailable_without_p2(db, two_tenants):
    alpha, bravo = two_tenants
    quiet = copy.deepcopy(gd.SIGNINS)
    for s in quiet:
        s["riskLevelDuringSignIn"], s["riskState"] = "none", "none"
    snap = TenantSnapshot(db, _seed(db, alpha, skus=P2_SKUS, signins=quiet))
    result = {r.key: r for r in evaluate(snap)}["risky_sign_ins"]
    assert result.status == "pass" and result.summary == "No risky sign-ins succeeded in the last 7 days"
    result = {r.key: r for r in evaluate(TenantSnapshot(db, _seed(db, bravo)))}["risky_sign_ins"]
    assert result.status == "unavailable" and "Entra ID P2" in result.summary


# -- Account page ------------------------------------------------------------------------------


def test_account_page_shows_recent_sign_ins(client, db, two_tenants):
    alpha, _ = two_tenants
    tenant = _seed(db, alpha)
    bob = user_profile.build(db, tenant, "u2")
    part = {p.key: p for p in bob.parts}["signins"]
    assert part.summary == "9 sign-ins, 9 failed" and part.tone == "warn"
    facts = dict(part.facts)
    assert facts["Successful"] == "0" and facts["Failed"].startswith("9 (8 wrong password, 1 lockout")
    assert "Account is disabled" in facts["Last failure"]
    why = {r[0]: r for r in part.tables[1].rows}
    assert why["50126"][1:] == ["Wrong user name or password", 7]
    assert any(a.severity == "high" and "9 failed sign-ins" in a.text and a.part == "signins" for a in bob.attention)

    ann = user_profile.build(db, tenant, "u1")
    part = {p.key: p for p in ann.parts}["signins"]
    assert "FR (unusual for this tenant)" in dict(part.facts)["Countries"]
    places = part.tables[0]
    assert places.title == "Where this account signed in" and places.flags[places.rows.index(
        next(r for r in places.rows if r[0] == "FR"))] == "medium"
    texts = [a.text for a in ann.attention]
    assert any("Signed in from FR" in t for t in texts)
    info = user_profile.build(db, tenant, "u7")
    assert any("with only a password" in a.text and "never with MFA" in a.text for a in info.attention)

    make_user(db, "admin")
    login(client, "admin")
    page = client.get(f"/tenants/{alpha}/users/u2").get_data(as_text=True)
    assert "Recent sign-ins" in page and "Wrong user name or password" in page


def test_account_page_flags_risky_sign_ins(db, two_tenants):
    alpha, _ = two_tenants
    ann = user_profile.build(db, _seed(db, alpha, skus=P2_SKUS), "u1")
    part = {p.key: p for p in ann.parts}["signins"]
    assert part.tone == "bad" and part.tables[-1].title == "Risky sign-ins"
    assert any(a.severity == "medium" and "rated the sign-in medium risk" in a.text for a in ann.attention)


def test_account_page_without_entra_p1(db, two_tenants):
    alpha, _ = two_tenants
    _run(db, alpha, {"entra_p1": False})
    store.enqueue_run(db, alpha, "users")
    with responses.RequestsMock() as rsps:
        rsps.get(f"{GRAPH_ROOT}/v1.0/users", json={"value": USERS})
        execute_run(db, store.claim_next_run(db), lambda t, c: GraphClient(lambda: "t", sleep=lambda s: None))
    profile = user_profile.build(db, store.get_tenant(db, alpha), "u1")
    part = {p.key: p for p in profile.parts}["signins"]
    assert "Entra ID P1" in part.unavailable


# -- When Microsoft times out -----------------------------------------------------------------

SIGNIN_KINDS = ("sign_in_account", "sign_in_location", "sign_in_country", "sign_in_app", "sign_in_failure",
                "sign_in_failure_ip", "risky_sign_in")


def _recent_signins():
    """Sign-ins spread over the last six days, newest first, several sharing a timestamp."""
    now = datetime.now(UTC).replace(microsecond=0)
    out = []
    for i in range(40):
        when = (now - timedelta(hours=1 + 3.5 * i)).strftime("%Y-%m-%dT%H:%M:%SZ")
        for _ in range(1 + (i % 3 == 0)):  # every third time has two sign-ins
            n = len(out)
            out.append(gd._signin(n, when, ["ann@x.com", "bob@x.com", "cat@x.com"][n % 3], f"u{n % 3}",
                                  [0, 0, 50126, 50074][n % 4], ip=f"198.51.100.{n % 5}",
                                  country=["GB", "FR", "RU"][n % 3], app=["Teams", "Outlook"][n % 2],
                                  mfa=bool(n % 2)))
    return out


def _signin_log(rows, fail=lambda flt, skip: False, page=7):
    """A fake /beta/auditLogs/signIns honouring the createdDateTime filter, newest first, in pages."""
    def handler(request):
        query = parse_qs(urlsplit(request.url).query)
        flt, skip = query["$filter"][0], int(query.get("$skiptoken", ["0"])[0])
        if fail(flt, skip):
            return 504, {}, json.dumps({"error": {"code": "GatewayTimeout", "message": "Request timed out"}})
        clauses = [c.split(" ") for c in flt.split(" and ")]
        ops = {"ge": str.__ge__, "le": str.__le__, "lt": str.__lt__}
        match = [s for s in rows if all(ops[op](s["createdDateTime"], value) for _, op, value in clauses)]
        match.sort(key=lambda s: s["createdDateTime"], reverse=True)
        body = {"value": match[skip:skip + page]}
        if skip + page < len(match):
            body["@odata.nextLink"] = (f"{BETA}/auditLogs/signIns?$filter={flt}&$skiptoken={skip + page}")
        return 200, {}, json.dumps(body)
    responses.add_callback(responses.GET, f"{BETA}/auditLogs/signIns", callback=handler,
                           content_type="application/json")


def _stored(db, run_id):
    return {kind: store.run_objects(db, run_id, kind) for kind in SIGNIN_KINDS}


def _is_week_query(flt):
    since = flt.split(" ")[2]
    return " and " not in flt and since < (datetime.now(UTC) - timedelta(days=3)).strftime("%Y-%m-%dT")


@responses.activate
def test_a_timed_out_week_is_read_a_day_at_a_time_with_the_same_result(db, two_tenants):
    alpha, bravo = two_tenants
    rows = _recent_signins()
    _signin_log(rows)
    base_id, status = _run(db, alpha)
    assert status == "succeeded"
    expected = _stored(db, base_id)
    assert store.run_objects(db, base_id, "sign_in_status")[0]["signIns"] == len(rows)

    responses.reset()
    _signin_log(rows, fail=lambda flt, skip: _is_week_query(flt))
    run_id, status = _run(db, bravo)
    assert status == "succeeded"
    assert _stored(db, run_id) == expected
    assert store.run_objects(db, run_id, "sign_in_status")[0]["signIns"] == len(rows)
    filters = [parse_qs(urlsplit(c.request.url).query)["$filter"][0] for c in responses.calls]
    assert sum(" and " in f for f in filters) >= 6  # one-day windows after the first
    assert all(c.request.method == "GET" for c in responses.calls)


@responses.activate
def test_a_timeout_part_way_through_carries_on_from_the_oldest_sign_in(db, two_tenants):
    alpha, bravo = two_tenants
    rows = _recent_signins()
    _signin_log(rows)
    expected = _stored(db, _run(db, alpha)[0])

    # Page 1 of the week arrives, then Microsoft times out. Page 1 ends part-way through a
    # timestamp shared by two sign-ins, which must be neither lost nor counted twice.
    responses.reset()
    _signin_log(rows, fail=lambda flt, skip: _is_week_query(flt) and skip > 0, page=13)
    run_id, status = _run(db, bravo)
    assert status == "succeeded"
    assert _stored(db, run_id) == expected
    assert store.run_objects(db, run_id, "sign_in_status")[0]["signIns"] == len(rows)
    resumed = [parse_qs(urlsplit(c.request.url).query)["$filter"][0] for c in responses.calls]
    assert any(" and createdDateTime le " in f for f in resumed)


@responses.activate
def test_a_day_that_still_times_out_is_left_out_with_a_warning(db, two_tenants):
    alpha, bravo = two_tenants
    rows = _recent_signins()
    day = (datetime.now(UTC) - timedelta(days=2)).strftime("%Y-%m-%dT")

    def fail(flt, skip):
        return _is_week_query(flt) or (" and " in flt and flt.split(" ")[2].startswith(day))
    _signin_log(rows, fail=fail)
    run_id, status = _run(db, alpha)
    assert status == "partial"
    [st] = store.run_objects(db, run_id, "sign_in_status")
    assert st["available"] and 0 < st["signIns"] < len(rows)
    error = db.execute("SELECT error FROM runs WHERE id = ?", (run_id,)).fetchone()[0]
    assert "timed out for 1 one-day window" in error

    # Nothing at all comes back: the run fails as before rather than claiming an empty week.
    responses.reset()
    _signin_log(rows, fail=lambda flt, skip: True)
    run_id, status = _run(db, bravo)
    assert status == "failed"
    assert store.run_objects(db, run_id, "sign_in_status") == []


@responses.activate
def test_a_sign_in_log_that_is_down_gives_up_instead_of_holding_the_worker(db, two_tenants, monkeypatch):
    # Seen on the test tenant: every sign-in log query, however small, timed out after two minutes, and
    # the run retried each one six times and then tried every day: over an hour on one collector.
    alpha, bravo = two_tenants
    _signin_log(_recent_signins(), fail=lambda flt, skip: True)
    run_id, status = _run(db, alpha)
    assert status == "failed"
    # One retry for the month, then one for each one-day window (the last can be a sliver): not seven each.
    assert len(responses.calls) <= 2 * (2 + sign_ins.LOG_DAYS)

    # When even the fallback runs out of time, the windows left are not asked for.
    now = [0.0]

    def clock():
        now[0] += sign_ins.FALLBACK_SECONDS / 3
        return now[0]
    monkeypatch.setattr(sign_ins, "clock", clock)
    responses.reset()
    _signin_log(_recent_signins(), fail=lambda flt, skip: _is_week_query(flt))
    run_id, status = _run(db, bravo)
    assert status == "partial"
    [st] = store.run_objects(db, run_id, "sign_in_status")
    assert st["available"] and 0 < st["signIns"] < len(_recent_signins())
    error = db.execute("SELECT error FROM runs WHERE id = ?", (run_id,)).fetchone()[0]
    assert "too slow to answer" in error
    filters = {parse_qs(urlsplit(c.request.url).query)["$filter"][0] for c in responses.calls}
    assert len(filters) == 1 + 2  # the week, then two days
