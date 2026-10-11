"""One page per SharePoint site: storage and what fills it, sharing, the team behind it, activity, history,
the Sites page, links from reports and the group page, search and access."""

import copy
from datetime import UTC, datetime
from html import unescape
from io import BytesIO

from openpyxl import load_workbook

from officesentry import store
from officesentry.reports import data_reports, group_page, site_page
from tests import sharepoint_data as sd
from tests.conftest import login, make_user
from tests.test_group_page import SALES_SITE, _seed

NOW = datetime(2026, 10, 10, 12, tzinfo=UTC)
FINANCE = sd.SITE_USAGE[0]["siteId"]
SALES = SALES_SITE["siteId"]
OLD = sd.SITE_USAGE[1]["siteId"]


def _tenant(db, tenant_id):
    return store.get_tenant(db, tenant_id)


def _parts(page):
    return {p.key: p for p in page.parts}


def test_site_page_brings_storage_and_sharing_together(db, two_tenants):
    alpha, _ = two_tenants
    _seed(db, alpha)
    page = site_page.build(db, _tenant(db, alpha), FINANCE, NOW)
    assert page.name == "Finance" and page.subtitle == sd.FINANCE
    assert page.path == f"/tenants/{alpha}/sites/{FINANCE}" and page.list_path == f"/tenants/{alpha}/sites"
    assert page.external == ("Open in SharePoint", sd.FINANCE)
    parts = _parts(page)
    assert list(parts) == ["storage", "sharing", "owner", "activity"]

    storage = parts["storage"]
    facts = dict(storage.facts)
    assert facts["Storage used"] == "5.0 GB" and facts["Recycle bin"] == "1.0 GB (20% of the site)"
    assert facts["Site limit"].startswith("None of its own")
    assert facts["Older versions of files"].startswith("924.0 MB") and facts["Files"] == "10"
    libraries, largest = storage.tables
    assert [str(r[0]) for r in libraries.rows] == ["Documents", "Preservation Hold Library"]
    assert libraries.flags == [None, "medium"]
    assert str(largest.rows[0][0]) == "Clients/Mail export.pst" and largest.rows[0][2] == "2.0 GB"
    assert storage.note.startswith("Growth shows once")

    sharing = parts["sharing"]
    assert sharing.summary == "2 anyone links" and sharing.tone == "bad"
    assert dict(sharing.facts)["External sharing"] == "On"
    shared = sharing.tables[0]
    assert [(str(r[0]), r[2].parts) for r in shared.rows] == [
        ("Clients (folder)", ("Anyone with the link",)),
        ("Budget.xlsx", ("partner@other.com",)),
        ("Policy.docx", ("The whole organisation",)),
    ]
    assert shared.flags[0] == "high"
    assert "permission reports" in sharing.note

    owner = parts["owner"]
    assert owner.summary == "Group not found" and dict(owner.facts) == {"Primary owner": "owner@x.com"}
    assert parts["activity"].summary == "Active"

    texts = [a.text for a in page.attention]
    assert [a.severity for a in page.attention] == ["high", "high", "low"]
    assert texts[0].startswith("2 anyone links on the site") and "changed by anyone" in texts[1]
    labels = {t.label: t.value for t in page.tiles}
    assert labels["Needs attention"] == 3 and labels["Storage"] == "5.0 GB" and labels["Sharing"] == "2 anyone links"
    assert page.badges == [("Group or Team site", None), ("Shared outside", "warn")]
    # The finding about the folder shared with anyone is on this site's page.
    assert [f["rule_key"] for f in page.findings] == ["anyone_links"]
    reports = dict(page.reports)
    assert reports["SharePoint sites"].startswith(f"/tenants/{alpha}/data/sharepoint?f=all-sites%2Fsite%3A")
    assert (
        reports["SharePoint storage"] == f"/tenants/{alpha}/data/sharepoint-storage?f=sites-by-storage%2Fsite%3AFinance"
    )
    assert "Shared files" in reports
    assert site_page.build(db, _tenant(db, alpha), "nope", NOW) is None


def test_team_site_names_its_team_and_links_both_ways(db, two_tenants):
    alpha, _ = two_tenants
    _seed(db, alpha)
    tenant = _tenant(db, alpha)
    page = site_page.build(db, tenant, SALES, NOW)
    team = _parts(page)["team"]
    assert team.title == "Team" and team.summary == "Sales" and team.tone == "warn"  # a guest is a member
    facts = dict(team.facts)
    assert facts["Group"].page_href == f"/tenants/{alpha}/groups/g1"
    assert facts["Owners"] == "2" and facts["Members"] == "3" and facts["Guests"] == "1"
    assert {t.label: t.value for t in page.tiles}["Team"] == "Sales"
    assert dict(page.reports)["Groups and Teams"] == f"/tenants/{alpha}/data/groups?f=teams-and-groups%2Fname%3ASales"
    # Microsoft counts two anyone links, but the files scan found nothing on the site: the page says so.
    assert "found nothing on this site" in _parts(page)["sharing"].note
    # And the group's page opens the site's page.
    files = dict(_parts(group_page.build(db, tenant, "g1", NOW))["files"].facts)
    assert files["Site"].page_href == f"/tenants/{alpha}/sites/{SALES}"


def test_inactive_site_full_site_and_history(db, two_tenants):
    alpha, _ = two_tenants
    _seed(db, alpha)
    later = copy.deepcopy([*sd.SITE_USAGE, SALES_SITE])
    later[0].update({"anonymousLinkCount": 0, "secureLinkForGuestCount": 3, "siteSensitivityLabelId": "label-1"})
    later[1].update({"externalSharing": True, "storageUsedInBytes": 3 * 1024**3})
    quota = {"site-fin": {"used": 48 * sd.GB, "deleted": 1 * sd.GB, "total": 50 * sd.GB, "state": "nearing"}}
    _seed(db, alpha, sharepoint_kw={"sites": later, "site_quota": quota})
    tenant = _tenant(db, alpha)

    finance = site_page.build(db, tenant, FINANCE, NOW)
    storage = _parts(finance)["storage"]
    assert dict(storage.facts)["Storage used"] == "48.0 GB of 50.0 GB (96%)" and storage.tone == "warn"
    assert storage.chart and storage.tables[0].title == "Storage by month"
    assert any(a.severity == "medium" and "96% full" in a.text for a in finance.attention)
    events = {(e.area, e.text, e.tone) for e in finance.history}
    assert ("Sharing", "Anyone links went from 2 to 0", None) in events
    assert ("Sharing", "Links for people outside went from 1 to 3", "warn") in events
    assert ("Settings", "Sensitivity label set", None) in events
    assert ("Labelled", None) in finance.badges

    old = site_page.build(db, tenant, OLD, NOW)
    activity = _parts(old)["activity"]
    assert activity.summary == "Inactive" and activity.tone == "warn"
    assert any(a.severity == "low" and "Nobody has used the site" in a.text for a in old.attention)
    events = {(e.area, e.text) for e in old.history}
    assert ("Sharing", "External sharing turned on") in events
    assert ("Storage", "+2.0 GB to 3.0 GB") in events


def test_sites_page_lists_every_site(db, two_tenants):
    alpha, bravo = two_tenants
    _seed(db, alpha)
    rows, tiles, notice = site_page.directory(db, _tenant(db, alpha), NOW)
    assert [r.name for r in rows] == ["Finance", "Sales", "Old Projects", "https://x-admin.sharepoint.com"]
    assert rows[0].findings == 1 and rows[0].anyone == 2 and rows[2].inactive and rows[-1].type == "System"
    assert "gone" not in {r.url.rsplit("/", 1)[-1] for r in rows}  # deleted sites have no page
    assert {t.label: t.value for t in tiles} == {
        "Sites": 3,
        "Storage used": "11.0 GB",
        "With anyone links": 2,
        "Inactive": 1,
        "Near their limit": 0,
    }
    assert notice == ""
    assert site_page.directory(db, _tenant(db, bravo), NOW) is None


def test_concealed_sites_have_no_page_and_the_list_says_why(db, two_tenants):
    alpha, _ = two_tenants
    _seed(db, alpha, sharepoint_kw={"sites": [sd.CONCEALED_SITE, sd.SITE_USAGE[1]]})
    rows, _, notice = site_page.directory(db, _tenant(db, alpha), NOW)
    assert [r.name for r in rows] == ["Old Projects"] and "concealing" in notice.lower() or "hiding" in notice


def test_report_cells_link_to_the_site_page(db, two_tenants):
    alpha, _ = two_tenants
    _seed(db, alpha)
    tenant = _tenant(db, alpha)
    report = data_reports.build(db, tenant, "sharepoint")
    every = next(s for s in report.sections if s.title == "All sites")
    cells = {str(r[0]): r[0] for r in every.rows}
    finance = cells[sd.FINANCE]
    assert finance.page_href == f"/tenants/{alpha}/sites/{FINANCE}"
    assert finance.safe_href == sd.FINANCE  # PDFs and spreadsheets keep the address
    storage = data_reports.build(db, tenant, "sharepoint-storage")
    by_site = next(s for s in storage.sections if s.title == "Sites by storage")
    assert {str(r[0]): r[0].page_href for r in by_site.rows}["Finance"] == f"/tenants/{alpha}/sites/{FINANCE}"
    # Files and libraries keep opening SharePoint.
    files = next(s for s in storage.sections if s.title == "Largest files")
    assert not files.rows[0][0].page_href


def test_pages_render_and_are_tenant_scoped(client, db, two_tenants):
    alpha, bravo = two_tenants
    _seed(db, alpha)
    make_user(db, "admin")
    make_user(db, "client", role="client", tenants=[bravo])
    login(client, "admin")
    html = client.get(f"/tenants/{alpha}/sites").get_data(as_text=True)
    assert f'href="/tenants/{alpha}/sites/{FINANCE}"' in html and "Every site" in html
    html = unescape(client.get(f"/tenants/{alpha}/sites/{FINANCE}").get_data(as_text=True))
    for heading in ("Needs attention", "Storage and growth", "Sharing", "Owner", "Activity", "In the reports"):
        assert heading in html, heading
    assert f'href="{sd.FINANCE}"' in html  # Open in SharePoint
    assert client.get(f"/tenants/{alpha}/sites/nobody").status_code == 404
    xlsx = client.get(f"/tenants/{alpha}/sites/{FINANCE}.xlsx")
    assert xlsx.status_code == 200
    sheets = load_workbook(BytesIO(xlsx.data)).sheetnames
    assert "Needs attention" in sheets and "Site" in sheets and "Libraries" in sheets
    assert client.get(f"/tenants/{alpha}/sites/{FINANCE}.pdf").status_code == 200
    assert "No sites yet" in client.get(f"/tenants/{bravo}/sites").get_data(as_text=True)
    html = client.get(f"/tenants/{alpha}/data/sharepoint").get_data(as_text=True)
    assert f'href="/tenants/{alpha}/sites/{FINANCE}"' in html

    client.post("/logout")
    login(client, "client")
    for path in (
        f"/tenants/{alpha}/sites",
        f"/tenants/{alpha}/sites/{FINANCE}",
        f"/tenants/{alpha}/sites/{FINANCE}.pdf",
    ):
        assert client.get(path).status_code == 404, path


def test_search_finds_sites(db, two_tenants, app):
    from officesentry.web.search import search_tenant

    alpha, _ = two_tenants
    _seed(db, alpha)
    with app.test_request_context():
        hits = [h for h in search_tenant(db, _tenant(db, alpha), "finance").hits if h.label == "Site"]
    assert [h.href for h in hits] == [f"/tenants/{alpha}/sites/{FINANCE}"]
    assert hits[0].title == sd.FINANCE and hits[0].detail == "group or team site, 5.0 GB, 2 anyone links"


def test_unknown_values_do_not_break_the_page(db, two_tenants):
    alpha, _ = two_tenants
    odd = copy.deepcopy([*sd.SITE_USAGE, SALES_SITE])
    for s in odd:
        s.update(
            {
                "lastActivityDate": "nope",
                "rootWebTemplate": 5,
                "ownerPrincipalName": ["x"],
                "unmanagedDevicePolicy": 3,
                "siteSensitivityLabelId": 9,
            }
        )
    _seed(db, alpha, sharepoint_kw={"sites": odd})
    tenant = _tenant(db, alpha)
    for key in (FINANCE, SALES, OLD):
        page = site_page.build(db, tenant, key, NOW)
        assert page is not None and page.parts, key
        assert copy.deepcopy(page.as_report()).sections
    rows, _, _ = site_page.directory(db, tenant, NOW)
    assert len(rows) == 4
