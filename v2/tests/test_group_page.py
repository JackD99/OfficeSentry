"""One page per group: teams, security groups and distribution lists, the Groups and teams page, links to it
from reports and the person page, search, access and history."""

import copy
from datetime import UTC, datetime
from html import unescape
from io import BytesIO
from unittest.mock import patch

from openpyxl import load_workbook

from officesentry import store
from officesentry.reports import data_reports, group_page, user_profile
from officesentry.reports.data_reports import PageLink, UserLink
from tests import exchangedata as xd
from tests import graphdata as gd
from tests import groupdata as grp
from tests import sharepoint_data as sd
from tests.conftest import login, make_user
from tests.seed import seed_tenant
from tests.test_collectors import USERS
from tests.test_mail_reports import _exchange

NOW = datetime(2026, 10, 10, 12, tzinfo=UTC)
SALES_SITE = sd._site(
    "https://x.sharepoint.com/sites/sales", "Group", storage=5 * 1024**3, anyone=2, guest=1, owner="g1"
)
# A policy that covers the All Staff group and leaves VPN Users out, and Global Administrator held through a group.
CA_GROUPS = copy.deepcopy(gd.CA_MFA_ALL)
CA_GROUPS["conditions"]["users"].update({"includeGroups": ["s1"], "excludeGroups": ["s2"]})
ROLE_VIA_GROUP = {
    "id": "a-grp",
    "roleDefinitionId": gd.GA,
    "principalId": "s3",
    "directoryScopeId": "/",
    "principal": {"@odata.type": "#microsoft.graph.group", "id": "s3", "displayName": "Group creators"},
}


def _seed(db, tenant_id, users=None, **kwargs):
    kwargs.setdefault("exchange", _exchange())
    kwargs.setdefault("sharepoint_kw", {"sites": [*sd.SITE_USAGE, SALES_SITE]})
    seed_tenant(db, tenant_id, users=[*USERS, gd.INFO_USER] if users is None else users, **kwargs)


def _tenant(db, tenant_id):
    return store.get_tenant(db, tenant_id)


def _parts(page):
    return {p.key: p for p in page.parts}


def _links(sections):
    """Every linked cell in a report or page, by text: the last one wins."""
    return {str(v): v for s in sections for row in s.rows for v in row if isinstance(v, PageLink)}


def test_team_page_brings_the_team_together(db, two_tenants):
    alpha, _ = two_tenants
    _seed(db, alpha)
    page = group_page.build(db, _tenant(db, alpha), "g1", NOW)
    assert page.name == "Sales" and page.subtitle == "g1@x.com"
    assert page.path == f"/tenants/{alpha}/groups/g1" and page.list_path == f"/tenants/{alpha}/groups"
    parts = _parts(page)
    assert list(parts) == ["owners", "members", "access", "mail", "files", "structure", "admin"]
    owners = parts["owners"]
    assert owners.summary == "2 owners" and [r[:2] for r in owners.tables[0].rows] == [
        ["Ann", "ann@x.com"],
        ["Eve", "eve@x.com"],
    ]
    members = parts["members"]
    assert members.summary == "3 members, 1 guest"
    rows = {row[0]: row for row in members.tables[0].rows}
    assert rows["ann@x.com"][2:] == ["Member", "Active"] and rows["pat@supplier.com"][2] == "Guest"
    assert members.note == "1 other member (groups, devices or contacts) not shown."
    assert parts["access"].summary == "Nothing recorded"
    assert dict(parts["mail"].facts) == {"Address": "g1@x.com"} and parts["mail"].summary == "Mail-enabled"
    # The team site, matched by the group's address, with the headline sharing figures.
    files = dict(parts["files"].facts)
    assert files["Site"].safe_href == SALES_SITE["siteUrl"] and not files["Site"].page_href
    assert files["Storage"].startswith("5.0 GB") and files["Files"] == "10"
    assert files["Sharing links"] == "2 with anyone, 1 with guests" and files["External sharing"] == "On"
    assert parts["files"].tone == "warn"
    structure = dict(parts["structure"].facts)
    assert structure["Membership"] == "Assigned" and structure["Visibility"] == "Private" and structure["Team"] == "Yes"
    # Group membership changes aren't kept from the audit log (History covers them), so nothing names Sales.
    assert parts["admin"].summary == "None" and parts["admin"].empty.startswith("No admin role, Conditional Access")
    assert page.badges == [("Team", None), ("Private", None)]
    labels = {t.label: t.value for t in page.tiles}
    assert labels["Owners"] == "2 owners" and labels["Members"] == "3 members, 1 guest" and labels["Files"] == "5.0 GB"
    about = dict(page.about)
    assert about["Type"] == "Team" and about["Email"] == "g1@x.com" and about["Members"] == "3"
    reports = dict(page.reports)
    assert reports["Groups and Teams"] == f"/tenants/{alpha}/data/groups?f=teams-and-groups%2Fname%3ASales"
    assert reports["Group membership"] == f"/tenants/{alpha}/data/group-membership?f=by-team-or-group%2Fname%3ASales"
    assert reports["SharePoint sites"].startswith(f"/tenants/{alpha}/data/sharepoint?f=all-sites%2Fsite%3A")
    assert "Conditional Access" not in reports and "Admin roles" not in reports  # the group isn't used by either
    assert page.history == [] and page.history_since
    assert [title for title, _ in page.sources][:2] == ["Groups and Teams", "Group nesting"]
    assert group_page.build(db, _tenant(db, alpha), "nope", NOW) is None


def test_ownerless_team_with_guests_leads_with_the_problem(db, two_tenants):
    alpha, _ = two_tenants
    _seed(db, alpha)
    page = group_page.build(db, _tenant(db, alpha), "g2", NOW)
    parts = _parts(page)
    assert parts["owners"].summary == "1 owner, 0 active" and parts["owners"].tone == "bad"
    assert parts["owners"].tables[0].rows == [["Bob", "bob@x.com", "Sign-in blocked"]]
    assert parts["owners"].tables[0].flags == ["medium"]
    members = parts["members"]
    assert members.summary == "3 members, 2 guests" and members.note == ""
    assert sorted(row[2] for row in members.tables[0].rows) == ["Guest", "Guest", "Member"]
    texts = [a.text for a in page.attention]
    assert texts[0].startswith("Only owner is disabled (Bob): nobody active approves members and guests")
    assert page.attention[0].severity == "medium" and page.attention[0].part == "owners"
    assert ("No active owner", "bad") in page.badges and ("Public", None) in page.badges
    tiles = {t.label: t for t in page.tiles}
    assert tiles["Owners"].value == "1 owner, 0 active" and tiles["Owners"].tone == "bad"
    assert "Files" not in tiles  # no team site matched
    assert _parts(page)["files"].summary == "No site found"


def test_security_groups_show_what_membership_gives(db, two_tenants):
    alpha, _ = two_tenants
    with patch.object(gd, "ROLE_ASSIGNMENTS", [*gd.ROLE_ASSIGNMENTS, ROLE_VIA_GROUP]):
        _seed(db, alpha, ca=[CA_GROUPS])
    tenant = _tenant(db, alpha)
    # All Staff: dynamic, licensed and inside a policy.
    staff = group_page.build(db, tenant, "s1", NOW)
    assert staff.badges == [("Security group", None), ("Dynamic", None)]
    parts = _parts(staff)
    assert list(parts) == ["owners", "members", "access", "structure", "admin"]
    assert parts["owners"].summary == "Not read" and parts["members"].summary == "2 members"
    access = {t.title: t.rows for t in parts["access"].tables}
    assert access["Conditional Access policies"] == [["Require MFA for all users", "Included", "On"]]
    assert access["Licences given to members"] == [["Microsoft 365 Business Premium"]]
    assert parts["access"].summary == "1 policy, 1 licence" and parts["access"].tone is None
    structure = dict(parts["structure"].facts)
    assert structure["Rule"] == '(user.userType -eq "Member")' and structure["Rule processing"] == "Running"
    assert parts["structure"].summary == "Dynamic"
    assert dict(staff.reports)["Conditional Access"] == f"/tenants/{alpha}/data/conditional-access"
    # VPN Users: excluded from an enforced policy, holds Group creators, sits inside Finance alerts.
    vpn = group_page.build(db, tenant, "s2", NOW)
    parts = _parts(vpn)
    assert ("Synced from on-premises AD", None) in vpn.badges
    assert parts["access"].tables[0].rows == [["Require MFA for all users", "Excluded", "On"]]
    assert parts["access"].tone == "warn"
    assert [a.text for a in vpn.attention if a.part == "access"] == [
        "Members are excluded from 1 Conditional Access policy (Require MFA for all users). Check who is in the group."
    ]
    inside = {t.title: t for t in parts["members"].tables}["Groups inside it"]
    assert inside.rows[0][1] == "Security group" and inside.rows[0][0].page_href == f"/tenants/{alpha}/groups/s3"
    parents = parts["structure"].tables[0]
    assert parents.title == "Inside of" and str(parents.rows[0][0]) == "Finance alerts"
    assert parents.rows[0][0].page_href == f"/tenants/{alpha}/groups/m1"
    # Group creators: empty, and everyone in it would be a Global Administrator.
    creators = group_page.build(db, tenant, "s3", NOW)
    parts = _parts(creators)
    roles = parts["access"].tables[0]
    assert roles.title == "Admin roles held through it" and roles.rows == [["Global Administrator", "Active", "Yes"]]
    assert roles.flags == ["high"] and parts["access"].tone == "bad"
    texts = [(a.severity, a.text) for a in creators.attention]
    assert texts[0] == (
        "high",
        "Everyone in the group holds Global Administrator; adding a member makes them an admin.",
    )
    assert ("low", "Nobody is a member and it's older than a month: delete it if nobody needs it.") in texts
    assert dict(creators.reports)["Admin roles"] == f"/tenants/{alpha}/data/admins"
    assert {t.label: t.value for t in creators.tiles}["Gives members"] == "1 admin role"


def _list(address="d1@x.com", *, eid="d1", owners=("u1", "u2"), outside=True):
    data = _exchange()
    info = xd._group("Info", address, "dl-9", owners=list(owners), outside=outside)
    info["ExternalDirectoryObjectId"] = eid
    data["Get-DistributionGroup"] = [info]
    data["Get-DistributionGroupMember"] = {
        "dl-9": [xd._member("Ann", "ann@x.com", eid="u1"), xd._member("Partner", "pat@partner.example", "MailContact")]
    }
    data["Get-MessageTraceV2"] = [xd._trace(address, xd.ago(days=1), 1), xd._trace(address, xd.ago(days=3), 2)]
    return data


def test_distribution_list_page_joins_exchange_to_the_group(db, two_tenants):
    alpha, bravo = two_tenants
    _seed(db, alpha, exchange=_list())
    page = group_page.build(db, _tenant(db, alpha), "d1", NOW)
    assert page.badges == [("Distribution list", None)]
    parts = _parts(page)
    assert list(parts) == ["owners", "members", "access", "mail", "structure", "admin"]
    owners = parts["owners"]
    assert owners.summary == "2 owners, 1 active" and [r[2] for r in owners.tables[0].rows] == [
        "Active",
        "Sign-in blocked",
    ]
    members = parts["members"]
    assert members.summary == "2 members, 1 outside"
    assert [row[1:] for row in members.tables[0].rows] == [
        ["ann@x.com", "Mailbox"],
        ["pat@partner.example", "External contact"],
    ]
    # Finance alerts is inside it too (from Group nesting).
    assert [t.title for t in members.tables] == ["Members", "Groups inside it"]
    mail = dict(parts["mail"].facts)
    assert mail["Address"] == "d1@x.com" and mail["Accepts mail from outside"] == "Yes"
    assert mail["Joining"] == "Owners add members" and mail["Leaving"] == "Owners remove members"
    assert "Mail in the last 30 days" in mail and "Last mail received" in mail
    assert {t.label for t in page.tiles} >= {"Owners", "Members", "Email"}
    assert dict(page.reports)["Distribution lists"] == (
        f"/tenants/{alpha}/data/distribution-lists?f=every-list%2Flist%3AInfo"
    )
    # The list's row in the report opens the page.
    report = data_reports.build(db, _tenant(db, alpha), "distribution-lists")
    assert _links(report.sections)["Info"].page_href == f"/tenants/{alpha}/groups/d1"
    # When Exchange doesn't give the directory id, the address is enough.
    _seed(db, bravo, exchange=_list(eid="eid-dl-9", outside=False))
    other = group_page.build(db, _tenant(db, bravo), "d1", NOW)
    assert dict(_parts(other)["mail"].facts)["Accepts mail from outside"] == "No"
    # Closing the list to outside senders shows in its history.
    _seed(db, alpha, exchange=_list(outside=False))
    page = group_page.build(db, _tenant(db, alpha), "d1", NOW)
    assert ("Email", "No longer accepts mail from outside") in {(e.area, e.text) for e in page.history}


def test_history_spots_renames_owners_guests_and_nesting(db, two_tenants):
    alpha, _ = two_tenants
    _seed(db, alpha)
    later = copy.deepcopy(grp.GROUPS)
    later[0]["displayName"] = "Sales team"
    later[2].update({"resourceProvisioningOptions": ["Team"], "visibility": "Private"})
    owners = {**grp.OWNERS, "g1": [grp.ANN]}
    members = {
        **grp.MEMBERS,
        "g1": [*grp.MEMBERS["g1"], grp._guest("gx2", "lee@partner.org", mail=False)],
        "s2": [grp._member("u1")],
        "s3": [grp._member("u1")],
    }
    with patch.object(grp, "OWNERS", owners), patch.object(grp, "MEMBERS", members):
        _seed(db, alpha, groups=later)
    tenant = _tenant(db, alpha)
    sales = group_page.build(db, tenant, "g1", NOW)
    assert sales.name == "Sales team"
    events = {(e.area, e.text, e.tone) for e in sales.history}
    assert ("Name", "Renamed from Sales to Sales team", None) in events
    assert ("Owners", "eve@x.com no longer an owner", "warn") in events
    assert ("Guests", "lee_partner.org#ext#@x.onmicrosoft.com added as a guest", "warn") in events
    assert ("Members", "3 → 4 members", None) in events
    marketing = group_page.build(db, tenant, "g3", NOW)
    events = {(e.area, e.text) for e in marketing.history}
    assert ("Type", "Became a team") in events and ("Settings", "Visibility changed to Private") in events
    vpn = group_page.build(db, tenant, "s2", NOW)
    assert ("Nesting", "Group creators removed from inside it") in {(e.area, e.text) for e in vpn.history}
    creators = group_page.build(db, tenant, "s3", NOW)
    assert ("Members", "0 → 1 members") in {(e.area, e.text) for e in creators.history}
    assert not [a for a in creators.attention if "Nobody is a member" in a.text]


def test_groups_page_lists_every_group(db, two_tenants):
    alpha, bravo = two_tenants
    _seed(db, alpha)
    rows, tiles = group_page.directory(db, _tenant(db, alpha), NOW)
    by_id = {r.id: r for r in rows}
    assert len(rows) == 10 and rows[0].id == "g2"  # the ownerless team first
    sales = by_id["g1"]
    assert (sales.type, sales.mail, sales.owners, sales.members, sales.guests) == ("Team", "g1@x.com", "2", 3, 1)
    assert sales.membership == "Assigned" and sales.created == "2025-01-10" and sales.owners_ok
    phoenix = by_id["g2"]
    assert phoenix.owners == "0 of 1 active" and not phoenix.owners_ok and phoenix.guests == 2
    assert by_id["g5"].owners == "0 of 1 active"  # an app is the only owner
    assert by_id["s1"].membership == "Dynamic" and by_id["s2"].membership == "Synced"
    assert by_id["s3"].members == 0 and by_id["s3"].owners == ""
    assert by_id["d1"].type == "Distribution list" and by_id["m1"].type == "Mail-enabled security group"
    counts = {t.label: t.value for t in tiles}
    assert counts["Groups"] == 10 and counts["With guests"] == 2 and counts["Empty"] == 1 and counts["Dynamic"] == 1
    assert counts["Without an active owner"] == 3  # Project Phoenix, Marketing and Provisioned
    assert group_page.directory(db, _tenant(db, bravo), NOW) is None


def test_report_cells_link_to_the_group_page(db, two_tenants):
    alpha, _ = two_tenants
    _seed(db, alpha)
    tenant = _tenant(db, alpha)
    groups = data_reports.build(db, tenant, "groups")
    every = next(s for s in groups.sections if s.title == "Teams and groups")
    cells = _links([every])
    assert cells["Sales"].page_href == f"/tenants/{alpha}/groups/g1" and cells["Board"].page_href.endswith("/g4")
    assert not isinstance(cells["Sales"], UserLink) and cells["Sales"].text == "Sales"
    guests = next(s for s in groups.sections if s.title == "Guests in teams and groups")
    assert _links([guests])["Project Phoenix"].page_href == f"/tenants/{alpha}/groups/g2"
    membership = data_reports.build(db, tenant, "group-membership")
    assert _links(membership.sections)["Marketing"].page_href == f"/tenants/{alpha}/groups/g3"
    structure = data_reports.build(db, tenant, "group-structure")
    cells = _links(structure.sections)
    assert cells["VPN Users"].page_href == f"/tenants/{alpha}/groups/s2"
    assert cells["Group creators"].page_href == f"/tenants/{alpha}/groups/s3"
    # A department that happens to share a group's name is text, not a link...
    users = data_reports.build(db, tenant, "all-users")
    departments = [v for s in users.sections for row in s.rows for v in row if str(v) == "Sales"]
    assert departments and not any(isinstance(v, PageLink) for v in departments)
    # ...and a person's own page links the teams they're in.
    profile = user_profile.build(db, tenant, "u1")
    teams = _parts(profile)["groups"].tables[0]
    assert _links([teams])["Sales"].page_href == f"/tenants/{alpha}/groups/g1"


def test_pages_render_and_are_tenant_scoped(client, db, two_tenants):
    alpha, bravo = two_tenants
    _seed(db, alpha)
    make_user(db, "admin")
    make_user(db, "client", role="client", tenants=[bravo])
    login(client, "admin")
    html = client.get(f"/tenants/{alpha}/groups").get_data(as_text=True)
    assert f'href="/tenants/{alpha}/groups/g2"' in html and "Every group" in html
    html = unescape(client.get(f"/tenants/{alpha}/groups/g1").get_data(as_text=True))
    for heading in ("Needs attention", "Owners", "Members", "What it gives its members", "Files", "How it's set up"):
        assert heading in html, heading
    assert "In the reports" in html and "History" in html
    assert f'href="/tenants/{alpha}/users/u1"' in html  # Ann, the owner
    assert f'href="{SALES_SITE["siteUrl"]}"' in html  # the team site, opened in a new tab
    assert client.get(f"/tenants/{alpha}/groups/nobody").status_code == 404
    xlsx = client.get(f"/tenants/{alpha}/groups/g1.xlsx")
    assert xlsx.status_code == 200
    sheets = load_workbook(BytesIO(xlsx.data)).sheetnames
    assert "Needs attention" in sheets and "Group" in sheets and "Owners" in sheets
    assert client.get(f"/tenants/{alpha}/groups/g1.pdf").status_code == 200
    assert "No groups yet" in client.get(f"/tenants/{bravo}/groups").get_data(as_text=True)

    client.post("/logout")
    login(client, "client")
    for path in (f"/tenants/{alpha}/groups", f"/tenants/{alpha}/groups/g1", f"/tenants/{alpha}/groups/g1.pdf"):
        assert client.get(path).status_code == 404, path


def test_search_finds_groups(db, two_tenants, app):
    from officesentry.web.search import search_tenant

    alpha, _ = two_tenants
    _seed(db, alpha)
    with app.test_request_context():
        hits = [h for h in search_tenant(db, _tenant(db, alpha), "sales").hits if h.label == "Group"]
    assert [h.href for h in hits] == [f"/tenants/{alpha}/groups/g1"]
    assert hits[0].title == "Sales" and hits[0].detail == "g1@x.com, team, 3 members, 1 guest"


def test_unknown_values_do_not_break_the_page(db, two_tenants):
    alpha, _ = two_tenants
    odd = copy.deepcopy(grp.GROUPS)
    for g in odd:
        g.update(
            {
                "displayName": None,
                "description": ["x"],
                "visibility": 7,
                "createdDateTime": "nope",
                "expirationDateTime": 12,
                "membershipRule": 5,
            }
        )
    _seed(db, alpha, groups=odd)
    tenant = _tenant(db, alpha)
    for gid in ("g1", "s1", "d1"):
        page = group_page.build(db, tenant, gid, NOW)
        assert page is not None and page.parts, gid
        assert copy.deepcopy(page.as_report()).sections
    rows, _ = group_page.directory(db, tenant, NOW)
    assert len(rows) == 10
