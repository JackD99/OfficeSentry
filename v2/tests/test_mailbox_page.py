"""One page per mailbox: what it brings together, the Mailboxes page, links to it from reports, access and history."""

import copy
from datetime import UTC, datetime
from io import BytesIO

from openpyxl import load_workbook

from officesentry import store
from officesentry.reports import data_reports, mailbox_page, user_profile
from officesentry.reports.data_reports import PageLink, UserLink
from tests import graphdata as gd
from tests.conftest import login, make_user
from tests.seed import seed_tenant
from tests.test_collectors import USERS
from tests.test_mail_reports import _exchange

NOW = datetime(2026, 10, 10, 12, tzinfo=UTC)


def _seed(db, tenant_id, users=None, **kwargs):
    seed_tenant(db, tenant_id, users=[*USERS, gd.INFO_USER] if users is None else users, exchange=_exchange(), **kwargs)


def _tenant(db, tenant_id):
    return store.get_tenant(db, tenant_id)


def _parts(page):
    return {p.key: p for p in page.parts}


def test_user_mailbox_page_brings_the_mailbox_together(db, two_tenants):
    alpha, _ = two_tenants
    _seed(db, alpha)
    page = mailbox_page.build(db, _tenant(db, alpha), "u1", NOW)
    assert page.name == "Ann" and page.subtitle == "ann@x.com"
    assert page.path == f"/tenants/{alpha}/mailboxes/u1" and page.list_path == f"/tenants/{alpha}/mailboxes"
    parts = _parts(page)
    assert list(parts) == ["size", "forwarding", "access", "settings", "auto_reply", "devices", "activity"]
    # Size, with the sending limit and how full it is.
    facts = dict(parts["size"].facts)
    assert facts["Size"] == "48.0 GB of 49.5 GB (97%)" and facts["Stops sending at"] == "49.5 GB"
    assert not parts["size"].chart  # growth needs two collections
    assert ("Forwarding set by an admin", "ann@gmail.com (outside the organisation)") in parts["forwarding"].facts
    assert [t.title for t in parts["forwarding"].tables] == ["Inbox rules"]
    assert dict(parts["settings"].facts)["Older protocols allowed"] == "IMAP"
    assert parts["settings"].tables[0].title == "Legacy sign-ins in the last 30 days"
    assert dict(parts["auto_reply"].facts)["State"] == "On"
    assert [row[0] for row in parts["devices"].tables[0].rows] == ["Ann's iPhone", "SM-G991B"]
    assert dict(parts["activity"].facts)["Sent"] == "40"
    # The answer first: the worst thing on the page, and the tiles.
    texts = [a.text for a in page.attention]
    assert texts[0].startswith("All mail is forwarded to ann@gmail.com")
    assert any("97% full" in t for t in texts)
    assert page.tiles[0].label == "Needs attention" and page.tiles[0].value == len(page.attention)
    assert {t.label: t.value for t in page.tiles}["Size"] == "48.0 GB"
    assert ("Forwards outside", "bad") in page.badges and ("Archive on", None) in page.badges
    # The side column names the account, linked to its page.
    about = dict(page.about)
    assert isinstance(about["Account"], UserLink) and about["Account"].page_href == f"/tenants/{alpha}/users/u1"
    assert about["Sign-in"] == "Allowed"
    # Findings that name the mailbox, and the reports the facts came from, filtered to it.
    assert "No mail is forwarded outside the organisation" in {f["title"] for f in page.findings}
    reports = dict(page.reports)
    assert reports["Inbox rules"] == f"/tenants/{alpha}/data/inbox-rules?f=all-inbox-rules%2Fmailbox%3Aann%40x.com"
    assert reports["Mailbox size trend"] == f"/tenants/{alpha}/data/mailbox-trend"
    assert page.history == [] and page.history_since
    assert [title for title, _ in page.sources][:2] == ["Mailboxes", "Mailbox size"]


def test_shared_mailbox_page_shows_who_can_open_it(db, two_tenants):
    alpha, _ = two_tenants
    _seed(db, alpha)
    page = mailbox_page.build(db, _tenant(db, alpha), "u7", NOW)
    assert page.badges[0] == ("Shared mailbox", None) and ("On hold", None) in page.badges
    access = _parts(page)["access"]
    people = {(str(row[0]), row[1]): row[2] for row in access.tables[0].rows}
    assert people[("ann@x.com", "Full access")] == "Active"
    assert people[("bob@x.com", "Full access")] == "Sign-in blocked"
    assert people[("ann@x.com", "Send as")] == "Active" and people[("ann@x.com", "Send on behalf")] == "Active"
    ann = next(row[0] for row in access.tables[0].rows if str(row[0]) == "ann@x.com")
    assert isinstance(ann, UserLink) and ann.page_href == f"/tenants/{alpha}/users/u1"
    texts = [a.text for a in page.attention]
    assert any(t.startswith("The shared mailbox's account allows sign-in") for t in texts)
    assert any("bob@x.com (sign-in blocked) has full access" in t for t in texts)
    assert {t.label: t.value for t in page.tiles}["Others with access"] == "2 people"
    # Nothing is known about Microsoft's own mailboxes, and nothing is a page.
    assert mailbox_page.build(db, _tenant(db, alpha), "nope", NOW) is None
    assert mailbox_page.build(db, _tenant(db, alpha), "g-disc", NOW) is None


def test_size_growth_and_history_across_collections(db, two_tenants):
    alpha, _ = two_tenants
    _seed(db, alpha)
    later = _exchange()
    ann, info = later["Get-Mailbox"][:2]
    ann.pop("ForwardingSmtpAddress")  # the admin removed the forwarding...
    ann["LitigationHoldEnabled"] = True
    later["Get-MailboxStatistics"]["g-u1"][0]["TotalItemSize"] = "49 GB (52,613,349,376 bytes)"
    later["Get-MailboxPermission"]["g-u7"] = [  # ...and gave Carol Bob's access to info@
        {"User": "ann@x.com", "AccessRights": ["FullAccess"], "IsInherited": False},
        {"User": "carol@x.com", "AccessRights": ["FullAccess"], "IsInherited": False},
    ]
    seed_tenant(db, alpha, users=[*USERS, gd.INFO_USER], exchange=later)

    page = mailbox_page.build(db, _tenant(db, alpha), "u1", NOW)
    size = _parts(page)["size"]
    assert size.chart.startswith("<svg") and size.tables[0].title == "Size by month"
    # A month per month the two collections ran in, by the clock (NOW only says where the page looks from).
    ran = {r[0][:7] for r in db.execute("SELECT finished_at FROM runs WHERE collector = 'mailbox_usage'")}
    assert [row[0] for row in size.tables[0].rows] == sorted(ran)
    assert ("Forwarding set by an admin", "None") in _parts(page)["forwarding"].facts
    events = {(e.area, e.text) for e in page.history}
    assert ("Forwarding", "Removed") in events and ("Settings", "Litigation hold turned on") in events

    info_page = mailbox_page.build(db, _tenant(db, alpha), "u7", NOW)
    events = {(e.area, e.text) for e in info_page.history}
    assert ("Access", "Full access for carol@x.com added") in events
    assert ("Access", "Full access for bob@x.com removed") in events


def test_mailboxes_page_lists_every_mailbox(db, two_tenants):
    alpha, _ = two_tenants
    _seed(db, alpha)
    rows, tiles = mailbox_page.directory(db, _tenant(db, alpha), NOW)
    assert [(r.id, r.type) for r in rows] == [("u1", "User"), ("u7", "Shared")]
    ann, info = rows
    assert ann.size == "48.0 GB" and ann.share == 97 and ann.forwarding == "Outside" and ann.flagged_rules == 2
    assert info.access == 2 and info.account == "Active"
    counts = {t.label: t.value for t in tiles}
    assert counts["Mailboxes"] == 2 and counts["Over 90% full"] == 1 and counts["Shared mailboxes with sign-in"] == 1
    assert mailbox_page.directory(db, _tenant(db, two_tenants[1]), NOW) is None  # nothing collected yet


def test_report_cells_link_to_the_mailbox_page(db, two_tenants):
    alpha, _ = two_tenants
    # Without an account in the directory, info@ is reached through its mailbox page...
    _seed(db, alpha, users=USERS)
    report = data_reports.build(db, _tenant(db, alpha), "shared-mailboxes")
    cells = {str(v): v for s in report.sections for row in s.rows for v in row if isinstance(v, PageLink)}
    assert cells["info@x.com"].page_href == f"/tenants/{alpha}/mailboxes/u7"
    assert not isinstance(cells["info@x.com"], UserLink)
    assert cells["info@x.com"].text == "info@x.com" and cells["info@x.com"].safe_href is None
    # ...and with one, through the account's page, as before.
    _seed(db, alpha)
    report = data_reports.build(db, _tenant(db, alpha), "shared-mailboxes")
    cells = {str(v): v for s in report.sections for row in s.rows for v in row if isinstance(v, PageLink)}
    assert cells["info@x.com"].page_href == f"/tenants/{alpha}/users/u7"
    # The account's page points at the mailbox's.
    profile = user_profile.build(db, _tenant(db, alpha), "u7")
    more = dict(_parts(profile)["mailbox"].facts)["More"]
    assert isinstance(more, PageLink) and more.page_href == f"/tenants/{alpha}/mailboxes/u7"


def test_pages_render_and_are_tenant_scoped(client, db, two_tenants):
    alpha, bravo = two_tenants
    _seed(db, alpha)
    make_user(db, "admin")
    make_user(db, "client", role="client", tenants=[bravo])
    login(client, "admin")
    html = client.get(f"/tenants/{alpha}/mailboxes").get_data(as_text=True)
    assert f'href="/tenants/{alpha}/mailboxes/u7"' in html and "Every mailbox" in html
    html = client.get(f"/tenants/{alpha}/mailboxes/u1").get_data(as_text=True)
    for heading in ("Needs attention", "Size and growth", "Forwarding and inbox rules", "In the reports", "History"):
        assert heading in html, heading
    assert f'href="/tenants/{alpha}/users/u1"' in html  # the account, in the side column
    assert client.get(f"/tenants/{alpha}/mailboxes/nobody").status_code == 404
    xlsx = client.get(f"/tenants/{alpha}/mailboxes/u1.xlsx")
    assert xlsx.status_code == 200
    sheets = load_workbook(BytesIO(xlsx.data)).sheetnames
    assert "Needs attention" in sheets and "Mailbox" in sheets and "Inbox rules" in sheets
    assert client.get(f"/tenants/{alpha}/mailboxes/u1.pdf").status_code == 200
    empty = client.get(f"/tenants/{bravo}/mailboxes").get_data(as_text=True)
    assert "No mailboxes yet" in empty

    client.post("/logout")
    login(client, "client")
    for path in (f"/tenants/{alpha}/mailboxes", f"/tenants/{alpha}/mailboxes/u1", f"/tenants/{alpha}/mailboxes/u1.pdf"):
        assert client.get(path).status_code == 404, path


def test_search_finds_a_mailbox_without_an_account(db, two_tenants, app):
    from officesentry.web.search import search_tenant

    alpha, _ = two_tenants
    _seed(db, alpha, users=USERS)
    with app.test_request_context():
        hits = search_tenant(db, _tenant(db, alpha), "info@").hits
    assert [h.href for h in hits if h.label == "Mailbox"] == [f"/tenants/{alpha}/mailboxes/u7"]


def test_unknown_values_do_not_break_the_page(db, two_tenants):
    alpha, _ = two_tenants
    odd = _exchange()
    for box in odd["Get-Mailbox"]:
        box.update({"ForwardingSmtpAddress": None, "GrantSendOnBehalfTo": None, "RecipientTypeDetails": None})
    odd["Get-MailboxStatistics"] = {"g-u1": [{"TotalItemSize": None, "ItemCount": None}]}
    seed_tenant(db, alpha, users=[*USERS, gd.INFO_USER], exchange=odd, auto_replies={"u1": {"status": 7}})
    page = mailbox_page.build(db, _tenant(db, alpha), "u1", NOW)
    assert page is not None and page.parts
    assert copy.deepcopy(page.as_report()).sections
