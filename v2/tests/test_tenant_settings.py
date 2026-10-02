"""Tenant settings: default user permissions, guest directory access and the admin consent workflow."""

from datetime import UTC, datetime
from io import BytesIO

import pytest
import responses
from openpyxl import load_workbook

from officesentry import rules, store
from officesentry.graph import GraphClient
from officesentry.reports import data_reports
from officesentry.rules import evaluate
from officesentry.snapshot import TenantSnapshot
from tests import graphdata as gd
from tests import tenantsettingsdata as tsd
from tests.conftest import login, make_user
from tests.seed import seed_tenant
from tests.test_collectors import V1, _run

KEYS = ("tenant_creation", "app_registration", "guest_access_level", "email_verified_join", "admin_consent_workflow")
NOW = datetime(2026, 10, 2, tzinfo=UTC)


@pytest.fixture
def factory():
    return lambda tenant, connection: GraphClient(lambda: "t", sleep=lambda s: None)


def _results(db, tenant_id):
    return {r.key: r for r in evaluate(TenantSnapshot(db, store.get_tenant(db, tenant_id)))}


class Snap:
    """Stand-in snapshot holding just the tenant_policies objects."""

    def __init__(self, policy=None, consent=None):
        self.data = {"authorization_policy": policy, "admin_consent_policy": consent}

    def objects(self, collector, kind):
        return self.data.get(kind) if collector == "tenant_policies" else None


def _check(key, policy, consent=None):
    fn = next(r.check for r in rules.RULES if r.key == key)
    return fn(Snap([policy] if policy is not None else None, consent), NOW)


# -- Collector ---------------------------------------------------------------------------


@responses.activate
def test_collects_settings_and_admin_consent_policy(db, two_tenants, factory):
    alpha, _ = two_tenants
    responses.get(f"{V1}/policies/authorizationPolicy", json=tsd.OPEN_POLICY)
    responses.get(f"{V1}/policies/adminConsentRequestPolicy", json=tsd.ADMIN_CONSENT_ON)
    responses.get(f"{V1}/admin/sharepoint/settings", json=gd.SHAREPOINT_SETTINGS)
    assert _run(db, alpha, "tenant_policies", factory)[1] == "succeeded"
    snap = TenantSnapshot(db, store.get_tenant(db, alpha))
    policy = snap.objects("tenant_policies", "authorization_policy")[0]
    assert {k: policy[k] for k in (
        "allowedToCreateTenants", "allowedToCreateApps", "allowedToCreateSecurityGroups", "allowedToReadOtherUsers",
        "guestUserRoleId", "allowEmailVerifiedUsersToJoinOrganization", "allowedToSignUpEmailBasedSubscriptions",
        "allowedToUseSSPR", "blockMsolPowerShell")} == {
        "allowedToCreateTenants": True, "allowedToCreateApps": True, "allowedToCreateSecurityGroups": True,
        "allowedToReadOtherUsers": True, "guestUserRoleId": tsd.MEMBER_GUESTS,
        "allowEmailVerifiedUsersToJoinOrganization": True, "allowedToSignUpEmailBasedSubscriptions": True,
        "allowedToUseSSPR": True, "blockMsolPowerShell": False}
    # Fields the existing checks read are still stored as before.
    assert policy["allowInvitesFrom"] == "everyone"
    assert policy["permissionGrantPolicies"] == ["ManagePermissionGrantsForSelf.microsoft-user-default-low"]
    consent = snap.objects("tenant_policies", "admin_consent_policy")[0]
    assert consent["isEnabled"] and consent["requestDurationInDays"] == 30
    assert consent["reviewers"][0] == "/users/u1" and len(consent["reviewers"]) == 2
    assert all(c.request.method == "GET" for c in responses.calls)


@responses.activate
def test_admin_consent_policy_forbidden_is_partial(db, two_tenants, factory):
    alpha, _ = two_tenants
    responses.get(f"{V1}/policies/authorizationPolicy", json=tsd.OPEN_POLICY)
    responses.get(f"{V1}/policies/adminConsentRequestPolicy", status=403,
                  json={"error": {"code": "Authorization_RequestDenied", "message": "Insufficient privileges"}})
    responses.get(f"{V1}/admin/sharepoint/settings", json=gd.SHAREPOINT_SETTINGS)
    assert _run(db, alpha, "tenant_policies", factory)[1] == "partial"
    r = _results(db, alpha)
    assert r["admin_consent_workflow"].status == "unavailable"
    assert "didn't return" in r["admin_consent_workflow"].summary  # not "refresh": that wouldn't help
    assert r["tenant_creation"].status == "fail"  # the rest of the run still counts
    assert r["sharepoint_sharing"].status == "fail"
    report = data_reports.build(db, store.get_tenant(db, alpha), "tenant-settings")
    consent = {s.title: s for s in report.sections}["Admin consent requests"]
    assert consent.rows == [] and "didn't return" in consent.empty


# -- Checks ------------------------------------------------------------------------------


def test_open_tenant_fails_each_setting_check(db, two_tenants):
    alpha, _ = two_tenants
    seed_tenant(db, alpha, auth_policy=tsd.OPEN_POLICY)
    r = _results(db, alpha)
    assert {k: (r[k].status, r[k].severity) for k in KEYS} == {
        "tenant_creation": ("fail", "low"),
        "app_registration": ("fail", "low"),
        "guest_access_level": ("fail", "medium"),
        "email_verified_join": ("fail", "low"),
        # Microsoft's default (low-risk consent, requests off) is advice, not a standing failure.
        "admin_consent_workflow": ("pass", None),
    }
    assert all(r[k].category == "Tenant settings" for k in KEYS)
    assert all(r[k].recommendation for k in KEYS)


def test_locked_down_tenant_passes(db, two_tenants, monkeypatch):
    alpha, _ = two_tenants
    monkeypatch.setattr(tsd, "ADMIN_CONSENT_POLICY", tsd.ADMIN_CONSENT_ON)
    seed_tenant(db, alpha, auth_policy=tsd.LOCKED_POLICY)
    r = _results(db, alpha)
    for key in KEYS:
        assert r[key].status == "pass", (key, r[key].summary)
    assert "most restrictive" in r["guest_access_level"].summary
    assert r["admin_consent_workflow"].summary == "Users can ask an admin: requests are on, with 2 reviewer(s)"


def test_guest_access_levels():
    for role, status in ((tsd.MEMBER_GUESTS, "fail"), (tsd.LIMITED_GUESTS, "pass"), (tsd.RESTRICTED_GUESTS, "pass"),
                         ("not-a-role", "unavailable")):
        assert _check("guest_access_level", {"guestUserRoleId": role})[0] == status, role
    assert "the default" in _check("guest_access_level", {"guestUserRoleId": tsd.LIMITED_GUESTS})[1]


def test_admin_consent_workflow_depends_on_user_consent():
    off = [{"isEnabled": False, "reviewers": []}]
    legacy = {"permissionGrantPolicies": ["ManagePermissionGrantsForSelf.microsoft-user-default-legacy"]}
    status, summary, _ = _check("admin_consent_workflow", legacy, off)
    assert status == "pass" and "consent to most apps themselves" in summary
    low = {"permissionGrantPolicies": ["ManagePermissionGrantsForSelf.microsoft-user-default-low"]}
    status, summary, found = _check("admin_consent_workflow", low, off)
    assert (status, found) == ("pass", []) and "can't ask for anything else" in summary
    status, _, found = _check("admin_consent_workflow", {"permissionGrantPolicies": []}, off)
    assert status == "fail" and found[0].severity == "low"  # users blocked from every app
    status, _, found = _check("admin_consent_workflow", {"permissionGrantPolicies": []},
                              [{"isEnabled": True, "reviewers": []}])
    assert status == "fail" and "nobody can approve" in found[0].detail


def test_missing_settings_are_unavailable_not_passed(db, two_tenants):
    _, bravo = two_tenants
    r = _results(db, bravo)
    assert all(r[k].status == "unavailable" for k in KEYS)  # never collected
    # Stored before this setting was recorded, or not returned by Microsoft.
    status, summary, _ = _check("email_verified_join", {"guestUserRoleId": tsd.LIMITED_GUESTS})
    assert status == "unavailable" and "refresh" in summary
    assert _check("tenant_creation", {"allowedToCreateTenants": None})[0] == "unavailable"
    # Collected by an older version, with no admin consent policy stored: refreshing fixes it.
    status, summary, _ = _check("admin_consent_workflow", {"permissionGrantPolicies": []}, [])
    assert status == "unavailable" and "refresh Sharing and consent settings" in summary
    # Collected by this version, but Microsoft didn't return it: refreshing won't help.
    status, summary, _ = _check("admin_consent_workflow", {"permissionGrantPolicies": [],
                                                           "allowedToCreateSecurityGroups": True}, [])
    assert status == "unavailable" and "didn't return" in summary and "refresh" not in summary


def test_existing_fixture_without_new_fields(db, two_tenants):
    """The shared graphdata policy predates these fields; checks it can't answer stay unavailable."""
    alpha, _ = two_tenants
    seed_tenant(db, alpha)
    r = _results(db, alpha)
    assert r["app_registration"].status == "fail"
    assert r["guest_access_level"].status == "pass"
    assert r["tenant_creation"].status == r["email_verified_join"].status == "unavailable"
    assert r["admin_consent_workflow"].status == "pass"  # legacy consent: users can consent themselves


# -- Report ------------------------------------------------------------------------------


def test_tenant_settings_report(db, two_tenants, monkeypatch):
    alpha, _ = two_tenants
    monkeypatch.setattr(tsd, "ADMIN_CONSENT_POLICY", tsd.ADMIN_CONSENT_ON)
    seed_tenant(db, alpha, auth_policy=tsd.OPEN_POLICY)
    report = data_reports.build(db, store.get_tenant(db, alpha), "tenant-settings")
    assert report.title == "Tenant settings" and not report.unavailable
    sections = {s.title: s for s in report.sections}
    assert list(sections) == ["Settings", "Admin consent requests"]
    settings = sections["Settings"]
    rows = {row[0]: (row[1:], flag) for row, flag in zip(settings.rows, settings.flags, strict=True)}
    assert rows["Users can create tenants"] == (["Yes", "No"], "low")
    assert rows["Guest access to the directory"] == (["Same as members", "Limited or most restrictive"], "medium")
    assert rows["People can join by verifying an email address"][1] == "low"
    assert rows["Users can look up other users in the directory"] == (["Yes", "Yes"], None)
    assert rows["Legacy MSOnline PowerShell blocked"] == (["No", "Yes"], None)
    assert settings.flagged == 1
    consent = {row[0]: row[1] for row in sections["Admin consent requests"].rows}
    assert consent == {"Admin consent requests": "On", "Reviewers": 2,
                       "Reviewers are emailed about new requests": "Yes", "Reviewers get reminders": "Yes",
                       "Requests expire after": "30 days"}


def test_tenant_settings_report_workflow_off(db, two_tenants):
    alpha, _ = two_tenants
    seed_tenant(db, alpha, auth_policy=tsd.LOCKED_POLICY)
    report = data_reports.build(db, store.get_tenant(db, alpha), "tenant-settings")
    sections = {s.title: s for s in report.sections}
    assert all(f is None for f in sections["Settings"].flags)
    assert dict(zip(sections["Settings"].columns, sections["Settings"].rows[2], strict=True)) == {
        "Setting": "Guest access to the directory", "Value": "Only their own objects (most restrictive)",
        "Recommended": "Limited or most restrictive"}
    consent = sections["Admin consent requests"]
    assert consent.rows == [["Admin consent requests", "Off", "On"]]
    assert consent.flags == ["low"]  # user consent is off, so users have nowhere to ask
    assert _results(db, alpha)["admin_consent_workflow"].severity == "low"  # the check agrees


def test_report_and_check_agree_on_default_consent(db, two_tenants):
    """Low-risk user consent with requests off is Microsoft's default: neither the check nor the report flags it."""
    alpha, _ = two_tenants
    seed_tenant(db, alpha, auth_policy=tsd.OPEN_POLICY)
    report = data_reports.build(db, store.get_tenant(db, alpha), "tenant-settings")
    consent = {s.title: s for s in report.sections}["Admin consent requests"]
    assert consent.rows == [["Admin consent requests", "Off", "On"]] and consent.flags == [None]
    assert _results(db, alpha)["admin_consent_workflow"].status == "pass"


def test_stale_run_says_refresh(db, two_tenants):
    """A run from before this version stored the authorization policy without the admin consent settings."""
    alpha, _ = two_tenants
    seed_tenant(db, alpha)
    db.execute("DELETE FROM objects WHERE kind = 'admin_consent_policy'")
    new_fields = ("allowedToCreateSecurityGroups", "allowedToReadOtherUsers",
                  "allowEmailVerifiedUsersToJoinOrganization", "allowedToUseSSPR", "blockMsolPowerShell")
    for field in new_fields:
        db.execute("UPDATE objects SET data = json_remove(data, ?) WHERE kind = 'authorization_policy'",
                   (f"$.{field}",))
    r = _results(db, alpha)["admin_consent_workflow"]
    assert r.status == "unavailable" and "refresh Sharing and consent settings" in r.summary
    report = data_reports.build(db, store.get_tenant(db, alpha), "tenant-settings")
    consent = {s.title: s for s in report.sections}["Admin consent requests"]
    assert consent.rows == [] and "refresh Sharing and consent settings" in consent.empty


def test_guests_report_uses_the_same_consent_reading(db, two_tenants):
    """Only ManagePermissionGrantsForSelf policies are user consent; a Teams owner policy alone means it's off."""
    alpha, _ = two_tenants
    policy = {**tsd.LOCKED_POLICY, "defaultUserRolePermissions": {
        **tsd.LOCKED_POLICY["defaultUserRolePermissions"],
        "permissionGrantPoliciesAssigned": [
            "ManagePermissionGrantsForOwnedResource.microsoft-dynamically-managed-permissions-for-team"]}}
    seed_tenant(db, alpha, auth_policy=policy)
    report = data_reports.build(db, store.get_tenant(db, alpha), "guests")
    rows = {row[0]: row[1] for row in report.sections[0].rows}
    assert rows["User consent to apps"] == "Off"
    assert _results(db, alpha)["user_consent"].summary == "User consent is turned off"


def test_tenant_settings_report_without_data(db, two_tenants):
    _, bravo = two_tenants
    report = data_reports.build(db, store.get_tenant(db, bravo), "tenant-settings")
    assert report.unavailable and not report.sections


def test_tenant_settings_page_and_xlsx(client, db, two_tenants):
    alpha, bravo = two_tenants
    seed_tenant(db, alpha, auth_policy=tsd.OPEN_POLICY)
    make_user(db, "client", role="client", tenants=[bravo])
    login(client, "client")
    assert client.get(f"/tenants/{alpha}/data/tenant-settings").status_code == 404
    client.post("/logout")
    make_user(db, "admin")
    login(client, "admin")
    page = client.get(f"/tenants/{alpha}/data/tenant-settings")
    assert page.status_code == 200
    html = page.get_data(as_text=True)
    assert "Guest access to the directory" in html and "Admin consent requests" in html
    resp = client.get(f"/tenants/{alpha}/data/tenant-settings.xlsx")
    assert resp.status_code == 200
    assert load_workbook(BytesIO(resp.data)).sheetnames == ["Settings", "Admin consent requests"]
