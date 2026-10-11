"""Site permissions: SharePoint groups, permission levels and inheritance from SharePoint's own API
(client, collector, check and report). SharePoint is simulated at the HTTP level with ``responses``."""

import io
import json

import pytest
import responses
from openpyxl import load_workbook

from officesentry import clouds, endpoints, msauth, store
from officesentry.collectors import site_permissions as sp
from officesentry.graph import GraphClient
from officesentry.reports import data_reports
from officesentry.rules import evaluate
from officesentry.sharepoint_api import SharePointAccessDenied, SharePointClient, SharePointError, beside
from officesentry.snapshot import TenantSnapshot
from officesentry.worker import execute_run
from tests import sharepoint_data as sd
from tests.conftest import login, make_user
from tests.seed import graph_factory, mock_tenant, seed_tenant

HOST = "https://x.sharepoint.com"


def _snap(db, tenant_id):
    return TenantSnapshot(db, store.get_tenant(db, tenant_id))


def _report(db, tenant_id):
    return data_reports.build(db, store.get_tenant(db, tenant_id), "site-permissions")


def _section(report, title):
    return next(s for s in report.sections if s.title == title)


def _rows(section):
    return [[str(c) for c in r] for r in section.rows]


# -- The client ------------------------------------------------------------------------------------


def test_sharepoint_hosts_per_cloud():
    assert clouds.get("commercial").is_sharepoint("https://contoso.sharepoint.com")
    assert clouds.get("gcc_high").is_sharepoint("https://contoso.sharepoint.us")
    assert clouds.get("dod").is_sharepoint("https://contoso.sharepoint-mil.us")
    assert clouds.get("china").is_sharepoint("https://contoso.sharepoint.cn")
    for bad in (
        "https://contoso.sharepoint.us",  # another cloud's
        "https://evil.com/.sharepoint.com",
        "https://a.b.sharepoint.com",
        "http://contoso.sharepoint.com",
        "https://contoso.sharepoint.com/sites/x",
        "https://contoso.sharepoint.com.evil.com",
    ):
        assert not clouds.get("commercial").is_sharepoint(bad), bad
    msauth._check_resource("https://contoso.sharepoint.com", "commercial")  # a token may be asked for it
    with pytest.raises(ValueError):
        msauth._check_resource("https://contoso.sharepoint.us", "commercial")


def test_client_refuses_other_hosts_and_follows_paging():
    client = SharePointClient(HOST, lambda: "tok", sleep=lambda s: None)
    with pytest.raises(ValueError):
        client.url("https://y.sharepoint.com/sites/a", "/web")
    with pytest.raises(ValueError):
        SharePointClient("https://x.sharepoint.us", lambda: "t")
    with responses.RequestsMock() as rsps:
        rsps.get(
            f"{HOST}/sites/a/_api/web/sitegroups",
            json={"value": [{"Id": 1}], "odata.nextLink": f"{HOST}/sites/a/_api/web/sitegroups?$skiptoken=2"},
        )
        rsps.get(f"{HOST}/sites/a/_api/web/sitegroups?$skiptoken=2", json={"value": [{"Id": 2}]})
        result = client.get_all(f"{HOST}/sites/a", "/web/sitegroups")
        assert [g["Id"] for g in result.items] == [1, 2] and result.pages == 2
        sent = rsps.calls[0].request
        assert sent.method == "GET" and sent.headers["Authorization"] == "Bearer tok"
        assert sent.headers["Accept"] == "application/json;odata=nometadata"
    with responses.RequestsMock() as rsps:  # a next link to another host is never followed with the token
        rsps.get(f"{HOST}/sites/a/_api/web/webs", json={"value": [], "odata.nextLink": "https://evil.com/_api/web"})
        with pytest.raises(ValueError):
            client.get_all(f"{HOST}/sites/a", "/web/webs")


def test_client_retries_throttling_and_explains_refusals():
    sleeps = []
    client = SharePointClient(HOST, lambda: "t", sleep=sleeps.append)
    with responses.RequestsMock() as rsps:
        rsps.get(f"{HOST}/_api/web", status=429, headers={"Retry-After": "7"})
        rsps.get(f"{HOST}/_api/web", status=503, headers={"Retry-After": "2"})
        rsps.get(f"{HOST}/_api/web", json={"Title": "Home"})
        assert client.get(HOST, "/web")["Title"] == "Home" and sleeps == [7.0, 2.0]
        assert client.telemetry.endpoints["SharePoint /web"].calls == 3
    with responses.RequestsMock() as rsps:
        rsps.get(f"{HOST}/_api/web", status=401, json={"error_description": "Unsupported app only token."})
        with pytest.raises(SharePointAccessDenied):
            client.get(HOST, "/web")
        body = {"odata.error": {"code": "-1, Microsoft.SharePoint.SPException", "message": {"value": "Locked"}}}
        rsps.get(f"{HOST}/_api/web/lists", status=400, json=body)
        with pytest.raises(SharePointError, match="Locked"):
            client.get(HOST, "/web/lists")


def test_beside_signs_in_to_the_site_host_as_the_same_app():
    asked = []
    graph = GraphClient(lambda: "g", resource_tokens=lambda host: asked.append(host) or (lambda: "sp"))
    client = beside(graph, HOST)
    assert client.host == HOST and asked == [HOST] and client.telemetry is graph.telemetry
    assert beside(graph, "https://evil.com") is None  # never asks for a token for another host
    assert beside(GraphClient(lambda: "g"), HOST) is None  # a client from a secret can't sign in to SharePoint


def test_catalogue_names_each_rest_call():
    assert endpoints.sharepoint(f"{HOST}/sites/a/_api/web/lists(guid'l-1')/roleassignments").key() == (
        "SharePoint /web/lists(guid'{list}')/roleassignments"
    )
    assert endpoints.sharepoint(f"{HOST}/_api/web/sitegroups").limits == ("sharepoint",)
    assert endpoints.sharepoint(f"{HOST}/_api/web/folders") is None


def test_principals_and_levels():
    assert sp.principal({"Id": 3, "Title": "X Owners", "LoginName": "X Owners", "PrincipalType": 8}) == {
        "name": "X Owners",
        "kind": "sp_group",
        "id": 3,
        "outside": False,
    }
    guest = sp.principal({"Title": "Pat", "LoginName": sd.PAT, "PrincipalType": 1, "Email": "Pat@Supplier.com"})
    assert guest == {"name": "Pat", "kind": "guest", "id": "pat@supplier.com", "outside": True}
    assert sp.principal({"LoginName": "c:0(.s|true", "PrincipalType": 4})["kind"] == "everyone"
    assert sp.principal({"LoginName": "SHAREPOINT\\system", "PrincipalType": 1}) is None
    assert sp.principal("nonsense") is None
    assert sp.levels([{"Name": "Edit"}, {"Name": "Limited Access", "Hidden": True}, {"Name": "Edit"}]) == ["Edit"]
    assert sp.levels([{"Name": "Limited Access"}]) == [] and sp.levels(None) == []


# -- The collector ---------------------------------------------------------------------------------


def test_collector_reads_groups_levels_lists_and_subsites(db, two_tenants):
    alpha, _ = two_tenants
    seed_tenant(db, alpha)
    assert store.latest_run(db, alpha, "site_permissions")["status"] == "succeeded"
    snap = _snap(db, alpha)
    sites = {s["_id"]: s for s in snap.objects("site_permissions", "site")}
    assert set(sites) == {"site-fin", "site-old"}  # Ann's OneDrive isn't a team site

    fin = sites["site-fin"]
    assert [(g["principal"]["name"], g["principal"]["kind"], g["levels"]) for g in fin["grants"]] == [
        ("Sales Owners", "sp_group", ["Full Control"]),
        ("Sales Members", "sp_group", ["Edit"]),
        ("Sales Visitors", "sp_group", ["Read"]),
        ("Bob", "person", ["Contribute"]),
        ("Dan Contractor", "guest", ["Full Control"]),
    ]  # Eve's Limited Access and the system account left out
    assert [(g["name"], g["role"], g["levels"], g["memberCount"], g["used"]) for g in fin["groups"]] == [
        ("Sales Owners", "owners", ["Full Control"], 1, True),
        ("Sales Members", "members", ["Edit"], 2, True),
        ("Sales Visitors", "visitors", ["Read"], 1, True),
        ("Board pack readers", None, [], 1, True),  # granted on a library only
    ]
    assert [m["kind"] for m in fin["groups"][1]["members"]] == ["group_members", "guest"]
    assert [(x["title"], x["library"], x["items"]) for x in fin["lists"]] == [
        ("Board pack", True, 12),
        ("Expenses", False, 230),
    ]  # Documents inherits and Style Library is a system list
    assert fin["lists"][0]["url"] == sd.BOARD_URL
    assert fin["listCount"] == 3 and fin["uniqueListCount"] == 2
    assert [(w["title"], w["unique"], w["grants"] and len(w["grants"])) for w in fin["subsites"]] == [
        ("Archive", True, 2),
        ("Team wiki", False, None),
    ]
    status = snap.objects("site_permissions", "permissions_status")[0]
    assert status["available"] and status["sites"] == 2 and status["unreadable"] == 0 and status["complete"]
    assert snap.metric("site_permissions.direct_grants") == 2 and snap.metric("site_permissions.unique_lists") == 2


def test_collector_asks_sharepoint_only_with_get_and_reuses_the_site_list(db, two_tenants):
    alpha, _ = two_tenants
    with responses.RequestsMock(assert_all_requests_are_fired=False) as rsps:
        mock_tenant(rsps)
        store.enqueue_run(db, alpha, "site_permissions")
        assert execute_run(db, store.claim_next_run(db), graph_factory) == "succeeded"
        rest = [c.request for c in rsps.calls if "/_api/" in c.request.url]
    assert rest and {r.method for r in rest} == {"GET"}
    paths = [endpoints.sharepoint_key(r.url) for r in rest]
    assert paths.count("SharePoint /web/lists(guid'{list}')/roleassignments") == 2  # unique lists only
    assert not any("personal" in r.url for r in rest)
    run = store.latest_run(db, alpha, "site_permissions")
    made = {r["endpoint"] for r in db.execute("SELECT endpoint FROM run_requests WHERE run_id = ?", (run["id"],))}
    assert "SharePoint /web/sitegroups" in made


def test_waiting_for_the_sharepoint_permission(db, two_tenants):
    """Until the client approves the SharePoint permission, the run succeeds quietly and the report says why."""
    alpha, _ = two_tenants
    seed_tenant(db, alpha, sharepoint_kw={"rest": "consent"})
    run = store.latest_run(db, alpha, "site_permissions")
    assert run["status"] == "succeeded" and run["error"] is None
    assert _snap(db, alpha).objects("site_permissions", "permissions_status")[0]["reason"] == "consent"
    assert "Waiting for the client to approve the SharePoint permission" in _report(db, alpha).unavailable
    result = {r.key: r for r in evaluate(_snap(db, alpha))}["site_guest_full_control"]
    assert result.status == "unavailable" and "approve" in result.summary


def test_without_a_certificate_or_graph_permission(db, two_tenants):
    alpha, bravo = two_tenants
    with responses.RequestsMock(assert_all_requests_are_fired=False) as rsps:
        mock_tenant(rsps)
        store.enqueue_run(db, alpha, "site_permissions")
        plain = lambda t, c: GraphClient(lambda: "t", sleep=lambda s: None)  # noqa: E731
        assert execute_run(db, store.claim_next_run(db), plain) == "succeeded"
    assert "certificate" in _report(db, alpha).unavailable
    seed_tenant(db, bravo, sharepoint_kw={"files": False})
    run = store.latest_run(db, bravo, "site_permissions")
    assert run["status"] == "partial" and "Sites.Read.All" in run["error"]
    assert "consent link" in _report(db, bravo).unavailable


def test_a_site_that_refuses_is_counted_not_fatal(db, two_tenants):
    alpha, _ = two_tenants
    rest = json.loads(json.dumps(sd.SP_REST))
    del rest[sd.OLD]  # 404 everywhere on Old projects
    with responses.RequestsMock(assert_all_requests_are_fired=False) as rsps:
        mock_tenant(rsps, sharepoint_kw={"rest": False})
        sd.mock_sharepoint_rest(rsps, rest=rest)
        store.enqueue_run(db, alpha, "site_permissions")
        assert execute_run(db, store.claim_next_run(db), graph_factory) == "partial"
    status = _snap(db, alpha).objects("site_permissions", "permissions_status")[0]
    assert status["sites"] == 1 and status["unreadable"] == 1
    assert any("couldn't be read" in n for n in _report(db, alpha).notes)


def test_collector_stops_at_its_budget(db, two_tenants, monkeypatch):
    alpha, _ = two_tenants
    monkeypatch.setattr(sp, "MAX_REQUESTS", 3)
    seed_tenant(db, alpha)
    run = store.latest_run(db, alpha, "site_permissions")
    assert run["status"] == "partial" and "1 site of 2" in run["error"] and run["truncated"]
    assert any("stopped" in n for n in _report(db, alpha).notes)


# -- Check and report ------------------------------------------------------------------------------


def test_guest_full_control_check(db, two_tenants):
    alpha, bravo = two_tenants
    seed_tenant(db, alpha)
    result = {r.key: r for r in evaluate(_snap(db, alpha))}["site_guest_full_control"]
    assert result.status == "fail" and result.summary == "1 guest with full control of 1 site"
    assert [(f.severity, f.subject, f.detail, f.object_id) for f in result.findings] == [
        (
            "high",
            "dan@contractor.com on Finance",
            "Full Control, given Full Control directly",
            "site-fin|dan@contractor.com",
        )
    ]
    assert {r.key: r for r in evaluate(_snap(db, bravo))}["site_guest_full_control"].status == "unavailable"


def test_guest_in_an_owners_group_fails_the_check(db, two_tenants):
    alpha, _ = two_tenants
    rest = json.loads(json.dumps(sd.SP_REST))
    fin = rest[sd.FINANCE]
    fin["/web/roleassignments"] = [a for a in fin["/web/roleassignments"] if a["Member"]["Id"] != 14]
    fin["/web/sitegroups"][0]["Users"].append(sd.DAN_GUEST)
    with responses.RequestsMock(assert_all_requests_are_fired=False) as rsps:
        mock_tenant(rsps, sharepoint_kw={"rest": False})
        sd.mock_sharepoint_rest(rsps, rest=rest)
        store.enqueue_run(db, alpha, "site_permissions")
        execute_run(db, store.claim_next_run(db), graph_factory)
    result = {r.key: r for r in evaluate(_snap(db, alpha))}["site_guest_full_control"]
    assert [f.detail for f in result.findings] == ["Full Control, in Sales Owners"]


def test_report(db, two_tenants):
    alpha, _ = two_tenants
    seed_tenant(db, alpha)
    report = _report(db, alpha)
    assert report.unavailable is None
    tiles = {t.label: (t.value, t.tone) for t in report.tiles}
    assert tiles == {
        "Guests with full control": (1, "bad"),
        "People given access directly": (2, "warn"),
        "Open to everyone": (2, "warn"),
        "Lists and libraries with their own permissions": (2, None),
        "Sites": (2, None),
    }

    access = _section(report, "Who has access to each site")
    fin = [r[1:] for r in _rows(access) if r[0] == "Finance"]
    assert fin == [
        ["Owners of Sales", "Microsoft 365 group owners", "Through Sales Owners", "Full Control", "No"],
        ["Members of Sales", "Microsoft 365 group", "Through Sales Members", "Edit", "No"],
        ["pat@supplier.com", "Guest", "Through Sales Members", "Edit", "Yes"],
        ["Everyone except external users", "Everyone", "Through Sales Visitors", "Read", "No"],
        ["bob@x.com", "Person", "Given directly", "Contribute", "No"],
        ["dan@contractor.com", "Guest", "Given directly", "Full Control", "Yes"],
    ]
    assert access.flags[:6] == [None, None, "low", "low", None, "high"]
    assert access.rows[0][0].safe_href == sd.FINANCE

    groups = _section(report, "SharePoint groups")
    assert [r[1:6] for r in _rows(groups) if r[0] == "Finance"] == [
        ["Sales Owners", "Owners", "Full Control", "1", "None"],
        ["Sales Members", "Members", "Edit", "2", "pat@supplier.com"],
        ["Sales Visitors", "Visitors", "Read", "1", "None"],
        ["Board pack readers", "Other", "Lists or subsites only", "1", "None"],
    ]
    lists = _section(report, "Lists and libraries with their own permissions")
    assert _rows(lists) == [
        ["Finance", "Board pack", "Library", "12", "Sales Owners (Full Control), Board pack readers (Read)"],
        ["Finance", "Expenses", "List", "230", "Sales Owners (Full Control), Ann (Edit), Pat Supplier (Read)"],
    ]
    assert lists.flags == [None, "low"] and lists.drifts == ("Items",)
    subsites = _section(report, "Subsites with their own permissions")
    assert _rows(subsites) == [["Finance", "Archive", "Sales Owners (Full Control), Bob (Read)"]]

    people = _section(report, "Access by person")
    rows = _rows(people)
    assert rows[:2] == [
        ["dan@contractor.com", "Yes", "1", "Full Control", "Finance"],
        ["pat@supplier.com", "Yes", "1", "Edit", "Finance"],
    ]  # guests first
    assert ["bob@x.com", "No", "2", "Full Control", "Finance, Old projects"] in rows
    assert people.flags[:2] == ["high", "low"]


def test_report_before_collection(db, two_tenants):
    _, bravo = two_tenants
    assert "yet" in _report(db, bravo).unavailable


def test_page_and_exports(client, db, two_tenants):
    alpha, bravo = two_tenants
    seed_tenant(db, alpha)
    make_user(db, "admin")
    login(client, "admin")
    page = client.get(f"/tenants/{alpha}/data/site-permissions")
    assert page.status_code == 200
    text = page.get_data(as_text=True)
    assert 'class="tiles"' in text and "Guests with full control" in text and "Board pack readers" in text
    wb = load_workbook(io.BytesIO(client.get(f"/tenants/{alpha}/data/site-permissions.xlsx").data))
    assert "Access by person" in wb.sheetnames
    assert client.get(f"/tenants/{bravo}/data/site-permissions").status_code == 200  # the empty state renders
