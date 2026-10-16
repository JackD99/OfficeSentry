"""How client-facing exports read: PDF table cells, XLSX workbooks and branded emails."""

import base64
import io
import zipfile
from datetime import UTC, date, datetime
from xml.etree import ElementTree

from openpyxl import load_workbook

from officesentry import delivery, mailer
from officesentry.reports import branding, cells, data_reports, docx_writer, tables, workbook
from officesentry.reports.export import to_xlsx as review_xlsx
from officesentry.reports.monthly_review import build
from tests.conftest import login, make_user
from tests.seed import seed_tenant
from tests.test_email import _schedule

PNG = base64.b64decode(
    "iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAYAAAAfFcSJAAAADUlEQVR42mNkYPhfDwAChwGA60e6kgAAAABJRU5ErkJggg=="
)


def test_cells_read_as_words_in_pdfs_and_as_values_in_spreadsheets():
    assert cells.display("2026-08-20") == "20 Aug 2026"
    assert cells.display("2026-08-20 (45d)") == "20 Aug 2026 (45 days ago)"
    assert cells.display("2026-10-03 (0d)") == "3 Oct 2026"
    assert cells.display("2026-10-01 22:42") == "1 Oct 2026, 22:42"
    assert cells.display("2025-01-01 (expired)") == "1 Jan 2025 (expired)"
    # A device seen today (data_reports._seen) reads like one seen days ago, not as the raw date.
    assert cells.display("2026-10-15 (today)") == "15 Oct 2026 (today)"
    assert cells.xlsx_value("2026-10-15 (today)") == date(2026, 10, 15)
    assert cells.display("Windows 10") == "Windows 10" and cells.display(None) == ""
    assert cells.xlsx_value("2026-08-20 (45d)") == date(2026, 8, 20)
    assert cells.xlsx_value("2026-10-01T22:42:05Z") == datetime(2026, 10, 1, 22, 42, 5)
    assert cells.xlsx_value("60%") == 0.6 and cells.xlsx_value("+2") == 2 and cells.xlsx_value("320") == 320
    for text in ("007", "Under 1%", "=1+1", "2026-13-45", "1.0 GB"):
        assert cells.xlsx_value(text) == text
    # Long addresses get break points at their dots and @; everything is escaped.
    out = str(cells.breakable("someone.long@contoso-holdings.example <b>"))
    assert "someone.<wbr>long@<wbr>contoso-<wbr>holdings.<wbr>example" in out and "&lt;b&gt;" in out
    assert str(cells.breakable("bob@x.com")) == "bob@x.com"


def test_workbook_sheets_are_styled_filterable_and_typed():
    book = workbook.Book("#0d7a5f")
    book.table(
        "Devices",
        ["Device", "Last check-in", "Share", "Count"],
        [["LAPTOP-1", "2026-08-20 (45d)", "60%", 3], ["=HYPERLINK(1)", "2026-10-01 22:42", "Under 1%", 0]],
        flags=["high", None],
    )
    data = book.save(
        workbook.Cover(
            "Devices",
            "Alpha Ltd",
            datetime(2026, 10, 4, 9, 30, tzinfo=UTC),
            "Harbour IT",
            "alpha.onmicrosoft.com",
            "#0d7a5f",
            facts=[("Devices", 2, "")],
        )
    )
    wb = load_workbook(io.BytesIO(data))
    assert wb.sheetnames == ["Summary", "Devices"]
    summary = [c for row in wb["Summary"].iter_rows(values_only=True) for c in row if c]
    assert {"Devices", "Alpha Ltd"} <= set(summary)
    assert any("Harbour IT" in str(c) and "4 Oct 2026" in str(c) for c in summary)
    ws = wb["Devices"]
    assert ws.freeze_panes == "A2" and ws.auto_filter.ref == "A1:E3"
    assert ws["A1"].font.bold and ws["A1"].fill.fgColor.rgb.endswith("0D7A5F")
    assert ws["B2"].value == datetime(2026, 8, 20) and ws["B2"].number_format == workbook.DATE_FORMAT
    assert ws["C2"].value == 0.6 and ws["C2"].number_format == "0%" and ws["C3"].value == "Under 1%"
    assert ws["E1"].value == "Flag" and ws["E2"].value == "high" and ws["E3"].value is None
    assert ws["A3"].data_type == "s" and ws["A3"].value == "=HYPERLINK(1)"  # never a formula
    assert ws.print_title_rows == "$1:$1"


def test_changes_percentages_and_written_dates_are_numbers_and_dates_in_spreadsheets():
    assert cells.xlsx_value("+4.6%") == 0.046 and cells.xlsx_value("-4.6%") == -0.046
    assert cells.xlsx_value("+2.2") == 2.2 and cells.xlsx_value("-1.5") == -1.5 and cells.xlsx_value("+4") == 4
    assert cells.xlsx_value("6 Feb 2026") == date(2026, 2, 6) and cells.display("6 Feb 2026") == "6 Feb 2026"
    for text in ("31 Feb 2026", "6 Feb 2026, 14:00", "+1.2.3", "+1e5"):
        assert cells.xlsx_value(text) == text
    book = workbook.Book()
    book.table(
        "Month by month",
        ["Month", "Score", "Change", "Share", "Last sign-in"],
        [["Sep", "29%", "+2.2", "+4.6%", "6 Feb 2026"], ["Oct", "33.3%", "+4", "-4.6%", "12 Jan 2025"]],
    )
    data = book.save(workbook.Cover("Secure Score", facts=[("Score", "29%", ""), ("Change", "+2.2", "")]))
    wb = load_workbook(io.BytesIO(data))
    ws = wb["Month by month"]
    assert [ws["B2"].number_format, ws["B3"].number_format] == ["0%", "0.0%"]  # 29%, not 29.0%
    assert [ws["C2"].value, ws["C3"].value] == [2.2, 4] and ws["C2"].number_format == "+0.0;-0.0;0.0"
    assert ws["C3"].number_format == "+0;-0;0"
    assert [ws["D2"].value, ws["D3"].value] == [0.046, -0.046] and ws["D2"].number_format == "+0.0%;-0.0%;0.0%"
    assert ws["E2"].value == datetime(2026, 2, 6) and ws["E2"].number_format == workbook.DATE_FORMAT
    facts = {row[0].value: row[1] for row in wb["Summary"].iter_rows() if row[0].value in ("Score", "Change")}
    assert [facts["Score"].number_format, facts["Change"].number_format] == ["0%", "+0.0;-0.0;0.0"]


def test_tenant_text_with_control_characters_still_makes_a_workbook_and_a_document():
    rule = "Move\x0bto RSS"  # an inbox rule's name as Exchange returned it
    book = workbook.Book()
    book.table("Inbox\x07rules", ["Rule", "Mail\x01box"], [[rule, [rule, "Archive"]]], description=rule)
    data = book.save(workbook.Cover(rule, rule, intro=rule, facts=[(rule, rule, rule)], notes=[(rule, rule)]))
    wb = load_workbook(io.BytesIO(data))
    assert wb.sheetnames == ["Summary", "Inbox rules"]
    assert [c.value for c in wb["Inbox rules"][1]] == ["Rule", "Mailbox"]
    assert [c.value for c in wb["Inbox rules"][2]] == ["Moveto RSS", "Moveto RSS, Archive"]
    assert "Moveto RSS" in [c for row in wb["Summary"].iter_rows(values_only=True) for c in row]
    doc = docx_writer.Document(rule, rule)
    doc.heading(rule)
    doc.paragraph(rule, lead=rule)
    doc.table([(rule, rule)])
    with zipfile.ZipFile(io.BytesIO(doc.save())) as z:
        for part in ("word/document.xml", "docProps/core.xml"):  # well-formed XML, so Word opens it
            assert "Moveto RSS" in ElementTree.fromstring(z.read(part)).itertext()  # noqa: S314 (our own file)


def test_summary_links_open_sheets_with_quotes_and_no_sheet_takes_excels_history_name():
    book = workbook.Book()
    book.table("OneDrives of people who've left", ["Name"], [["Ann"]])
    book.table("History", ["When"], [["2026-10-01"]])
    wb = load_workbook(io.BytesIO(book.save(workbook.Cover("Ann's account"))))
    assert wb.sheetnames == ["Summary", "OneDrives of people who've left", "History (2)"]
    links = [c.hyperlink.target for c in wb["Summary"]["A"] if c.hyperlink]
    assert links == ["#'OneDrives of people who''ve left'!A1", "#'History (2)'!A1"]
    assert tables.sheet_name("HISTORY", set()) == "HISTORY (2)"
    assert tables.sheet_name("History of sign-ins", set()) == "History of sign-ins"


def test_plain_table_downloads_have_no_cover_unless_asked():
    body = tables.to_xlsx([tables.Table("Rows", ["A"], [["x"]])])
    assert load_workbook(io.BytesIO(body)).sheetnames == ["Rows"]


def test_report_workbooks_lead_with_a_summary(db, two_tenants):
    alpha, _ = two_tenants
    seed_tenant(db, alpha)
    tenant = db.execute("SELECT * FROM tenants WHERE id = ?", (alpha,)).fetchone()
    brand = branding.Brand(name="Harbour IT", accent="#0d7a5f")
    wb = load_workbook(io.BytesIO(review_xlsx(build(db, tenant), brand=brand)))
    summary = [c for row in wb["Summary"].iter_rows(values_only=True) for c in row if c is not None]
    assert "Alpha Ltd" in summary and "Secure Score" in summary
    assert "Needs attention" in [c.value for c in wb["Checks"]["C"]]  # results in words, not "fail"
    report = data_reports.build(db, tenant, "devices")
    wb = load_workbook(io.BytesIO(data_reports.to_xlsx(report, brand=brand)))
    listed = [r[0] for r in wb["Summary"].iter_rows(values_only=True)]
    assert all(name in listed for name in wb.sheetnames[1:])  # the summary lists every sheet
    dates = [c.value for c in wb["Entra ID device records"]["F"][1:]]
    assert dates and all(isinstance(v, datetime) for v in dates if v is not None)


def test_report_email_carries_the_brand_and_counts_every_failing_check(client, db, two_tenants):
    alpha, _ = two_tenants
    seed_tenant(db, alpha)
    logo = branding.logo_data_uri(PNG)
    for key, value in (
        ("brand_name", "Harbour IT"),
        ("brand_accent", "#0d7a5f"),
        ("brand_logo", logo),
        ("brand_contact", "help@harbourit.example"),
    ):
        db.execute("INSERT INTO settings (key, value) VALUES (?, ?)", (key, value))
    sid = _schedule(db, alpha, [delivery.add_contact(db, alpha, "IT", "it@client.com")])
    tenant = db.execute("SELECT * FROM tenants WHERE id = ?", (alpha,)).fetchone()
    schedule = db.execute("SELECT * FROM email_schedules WHERE id = ?", (sid,)).fetchone()
    built = delivery.build_report_email(
        client.application, db, schedule, tenant, delivery.recipients(db, "schedule", sid), attach=False
    )
    email = built.email
    failing = len(build(db, tenant).failing)
    assert failing > 6 and f"{failing} checks need attention" in email.html  # not just the 6 listed
    assert 'src="cid:brand-logo"' in email.html and "#0d7a5f" in email.html and "help@harbourit.example" in email.html
    assert email.inline == [("brand-logo", "image/png", PNG)]
    assert "data:image/png;base64," in email.html_for_browser()
    # SMTP: the logo travels as a related part; Graph: as an inline attachment.
    msg = mailer.to_mime(email, "r@msp.com")
    related = [p for p in msg.walk() if p.get("Content-ID") == "<brand-logo>"]
    assert related and related[0].get_content_type() == "image/png"
    inline = [a for a in mailer.graph_payload(email)["message"]["attachments"] if a.get("isInline")]
    assert inline == [
        {
            "@odata.type": "#microsoft.graph.fileAttachment",
            "name": "brand-logo.png",
            "contentType": "image/png",
            "contentBytes": base64.b64encode(PNG).decode(),
            "isInline": True,
            "contentId": "brand-logo",
        }
    ]
    make_user(db, "admin")
    login(client, "admin")
    body = client.get(f"/schedules/{sid}/preview/body").get_data(as_text=True)
    assert "cid:" not in body and "data:image/png;base64," in body


def test_demo_data_builds_three_tenants_with_history(db, vault):
    import importlib.util
    from pathlib import Path

    path = Path(__file__).resolve().parents[1] / "scripts" / "demo_data.py"
    spec = importlib.util.spec_from_file_location("demo_data", path)
    demo = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(demo)
    ids = demo.populate(db, vault)
    assert len(ids) == 3 and branding.load(db).name == "Harbour IT"
    tenant = db.execute("SELECT * FROM tenants WHERE id = ?", (ids[0],)).fetchone()
    report = build(db, tenant)
    assert not report.first_report and report.licences.monthly_waste  # history and priced waste
    users = data_reports.build(db, tenant, "mfa")
    text = str([s.rows for s in users.sections])
    assert "@northwindtraders.example" in text and "@x.com" not in text  # fixture names swapped out
    from tests import graphdata

    assert "ann@x.com" in str(graphdata.REGISTRATION_DETAILS)  # fixtures restored for other tests


def test_logo_sizes_for_email():
    assert branding.image_size(PNG) == (1, 1)
    assert branding.fit((720, 180), 32, 200) == (128, 32)
    assert branding.fit((2000, 100), 32, 200) == (200, 10)
    assert branding.fit(None, 32, 200) is None
