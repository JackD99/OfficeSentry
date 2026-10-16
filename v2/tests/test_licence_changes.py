"""Licence changes: licences assigned and removed between collections, who made them, and counts by month."""

import copy
from datetime import UTC, datetime, timedelta
from io import BytesIO

import responses
from openpyxl import load_workbook

from officesentry import changes, store
from officesentry.graph import GraphClient
from officesentry.reports import data_reports, library
from officesentry.worker import execute_run
from tests import adminactivitydata as ad
from tests import groupdata as grp
from tests import licencedata as ld
from tests.conftest import login, make_user
from tests.seed import exchange_factory, mock_tenant, seed_tenant

KEY = "licence-changes"
GROUPS = grp.GROUPS + ld.LICENSING_GROUPS


def GRAPH(tenant, connection):
    return GraphClient(lambda: "t", sleep=lambda s: None)


def _collect(db, tenant_id, key, **kwargs):
    """Run one collector through the worker against the simulated tenant."""
    with responses.RequestsMock(assert_all_requests_are_fired=False) as rsps:
        mock_tenant(rsps, **{"skus": ld.SKUS, "subscriptions": ld.TENANT_SUBSCRIPTIONS, "groups": GROUPS, **kwargs})
        store.enqueue_run(db, tenant_id, key)
        run = store.claim_next_run(db)
        assert execute_run(db, run, GRAPH, exchange_factory) in ("succeeded", "partial")
    return run["id"]


def _backdate_all(db, tenant_id, days):
    when = (datetime.now(UTC) - timedelta(days=days)).strftime("%Y-%m-%dT%H:%M:%SZ")
    db.execute("UPDATE runs SET finished_at = ?, started_at = ? WHERE tenant_id = ?", (when, when, tenant_id))
    db.execute("UPDATE metrics SET captured_at = ? WHERE tenant_id = ?", (when, tenant_id))


def _price(db, part, price):
    db.execute(
        "INSERT INTO sku_prices (sku_part_number, monthly_price, updated_at) VALUES (?, ?, ?)",
        (part, price, "2026-10-01T00:00:00Z"),
    )


def _report(db, tenant_id, **kwargs):
    report = data_reports.build(db, store.get_tenant(db, tenant_id), KEY, **kwargs)
    assert report.unavailable is None, report.unavailable
    return report, {s.title: s for s in report.sections}


def _rows(section):
    return [[str(c) for c in row] for row in section.rows]


# A month later: Ann's Copilot licence is taken away and Eve is given one (and Power Automate Free, which costs
# nothing), Dan's account is deleted, Gus joins and gets Business Premium from the licensing group, Fay is blocked
# from signing in but given Business Premium anyway, and two more Business Premium licences are bought.
LATER_USERS = [
    ld.user("u1", "Ann", [ld.state("sku-bp"), ld.state("sku-bp", "l1")]),
    ld.user(
        "u5", "Eve", [ld.state("sku-bp", "l1"), ld.state("sku-visio"), ld.state("sku-copilot"), ld.state("sku-flow")]
    ),
    ld.user("u9", "Fay", [ld.state("sku-visio"), ld.state("sku-flow"), ld.state("sku-bp")], enabled=False),
    ld.user("u10", "Gus", [ld.state("sku-bp", "l1")]),
]
LATER_SKUS = copy.deepcopy(ld.SKUS)
LATER_SKUS[0]["prepaidUnits"]["enabled"] = 12
LATER_SKUS[0]["consumedUnits"] = 12


def _licence_event(eid, activity, days, by, uid, upn, props=()):
    return ad.event(
        eid,
        "UserManagement",
        activity,
        ad.ago(days),
        ad.by_user(by),
        [ad.target("User", upn=upn, props=props, tid=uid)],
    )


def _audits() -> list[dict]:
    """Dated when the test reads them (a run may pass midnight UTC after this module loads)."""
    return [
        _licence_event(
            "e1",
            "Update user",
            2,
            "admin@x.com",
            "u5",
            "eve@x.com",
            [ad.prop("AssignedLicense", ["[SkuName=Microsoft_365_Copilot, SkuId=sku-copilot, DisabledPlans=[]]"], [])],
        ),
        _licence_event("e2", "Change user license", 3, "helpdesk@x.com", "u1", "ann@x.com"),
        _licence_event("e3", "Update user", 1, "admin@x.com", "u5", "eve@x.com", [ad.prop("DisplayName", "Eve", "E")]),
    ]


def _a_month_of_changes(db, tenant_id):
    seed_tenant(db, tenant_id, users=ld.USERS, skus=ld.SKUS, subscriptions=ld.TENANT_SUBSCRIPTIONS, groups=GROUPS)
    _backdate_all(db, tenant_id, 40)
    _collect(db, tenant_id, "users", users=LATER_USERS)
    _collect(db, tenant_id, "tenant_profile", users=LATER_USERS, skus=LATER_SKUS)
    _collect(db, tenant_id, "admin_activity", users=LATER_USERS, audits=_audits())


def test_needs_two_collections_to_compare(db, two_tenants):
    alpha, _ = two_tenants
    seed_tenant(db, alpha, users=ld.USERS, skus=ld.SKUS, subscriptions=ld.TENANT_SUBSCRIPTIONS, groups=GROUPS)
    report, s = _report(db, alpha)
    assert "two different days" in report.notice
    figures = {h.label: h.value for h in report.highlights}
    assert figures["Assigned in the last 30 days"] == "Not yet"
    # 10 Business Premium and 3 Visio in their grace period; the 5 Copilot licences are a trial's.
    assert figures["Licences bought"] == 13
    assert s["Assigned and removed"].rows == [] and "collected once" in s["Assigned and removed"].empty
    month = _rows(s["Licences by month"])
    assert {r[1] for r in month} == {
        "Microsoft 365 Business Premium",
        "Microsoft 365 Copilot",
        "Visio Plan 2",
        "Exchange Online (Plan 1)",
    }  # Power Automate Free is left out
    assert all(r[5] == "" for r in month)  # nothing to compare the first month with
    assert not report.chart


def test_licences_assigned_and_removed_with_who_and_how(db, two_tenants):
    alpha, _ = two_tenants
    _a_month_of_changes(db, alpha)
    report, s = _report(db, alpha)
    assert not report.notice and not report.notes
    section = s["Assigned and removed"]
    assert section.compare is False
    rows = {(r[1], r[2]): r for r in _rows(section)}
    assert set(rows) == {
        ("eve@x.com", "Microsoft 365 Copilot"),
        ("ann@x.com", "Microsoft 365 Copilot"),
        ("dan@x.com", "Exchange Online (Plan 1)"),
        ("gus@x.com", "Microsoft 365 Business Premium"),
        ("fay@x.com", "Microsoft 365 Business Premium"),
    }  # Eve's Power Automate Free is free, and Dan's failed group licence was never assigned
    today = datetime.now(UTC)
    day = lambda n: (today - timedelta(days=n)).strftime("%Y-%m-%d")  # noqa: E731
    eve = rows[("eve@x.com", "Microsoft 365 Copilot")]
    assert eve[0] == day(2) and eve[3:7] == ["Assigned", "Directly", "admin@x.com", ""]
    ann = rows[("ann@x.com", "Microsoft 365 Copilot")]
    assert ann[0] == day(3) and ann[3:6] == ["Removed", "Directly", "helpdesk@x.com"]  # the audit log's time
    dan = rows[("dan@x.com", "Exchange Online (Plan 1)")]
    assert dan[0] == day(0) and dan[3] == "Removed" and dan[5] == "" and dan[6] == "Account deleted"
    gus = rows[("gus@x.com", "Microsoft 365 Business Premium")]
    assert gus[3:7] == ["Assigned", "By group: Licence - Business Premium", "", "New account"]
    fay = rows[("fay@x.com", "Microsoft 365 Business Premium")]
    assert fay[6] == "Sign-in blocked"
    flags = dict(zip([(r[1], r[2]) for r in _rows(section)], section.flags, strict=True))
    assert flags[("fay@x.com", "Microsoft 365 Business Premium")] == "low"
    assert flags[("gus@x.com", "Microsoft 365 Business Premium")] is None
    assert [r[0] for r in _rows(section)] == sorted((r[0] for r in _rows(section)), reverse=True)  # newest first

    figures = {h.label: h for h in report.highlights}
    assert figures["Assigned in the last 30 days"].value == 3
    assert figures["Removed in the last 30 days"].value == 2
    assert figures["Licences bought"].value == "+2" and figures["Licences bought"].foot.startswith("15 now, 13 on ")
    assert figures["Licences bought"].foot.endswith(", plus 5 on trial")


def test_months_and_cost_at_your_prices(db, two_tenants):
    alpha, _ = two_tenants
    _price(db, "SPB", 18.10)
    _a_month_of_changes(db, alpha)
    report, s = _report(db, alpha)
    section = s["Licences by month"]
    assert section.group == "Month" and section.compare is False
    rows = _rows(section)
    now, then = datetime.now(UTC), datetime.now(UTC) - timedelta(days=40)
    current = {r[1]: r for r in rows if r[0] == f"{now:%B %Y}"}
    earlier = {r[1]: r for r in rows if r[0] == f"{then:%B %Y}"}
    assert rows[0][0] == f"{now:%B %Y}"  # newest month first
    bp = current["Microsoft 365 Business Premium"]
    assert bp[2:7] == ["12", "12", "0", "+2", "£217.20"]
    assert earlier["Microsoft 365 Business Premium"][2:6] == ["10", "10", "0", ""]
    assert current["Microsoft 365 Copilot"][5] == "No change" and current["Microsoft 365 Copilot"][6] == "Trial"
    assert current["Visio Plan 2"][6] == "Not priced"
    exo_flag = section.flags[rows.index(current["Exchange Online (Plan 1)"])]
    assert exo_flag == "medium"  # one assigned, none bought
    cost = {h.label: h for h in report.highlights}["Monthly cost"]
    assert cost.value == "+£36.20" and cost.foot.startswith("£217.20 now, £181.00 on ")
    assert report.chart.startswith("<svg") and report.chart_title == "Paid licences assigned over the last year"


def test_as_of_reads_only_what_was_collected_by_then(db, two_tenants):
    alpha, _ = two_tenants
    _a_month_of_changes(db, alpha)
    report, s = _report(db, alpha, as_of=datetime.now(UTC) - timedelta(days=20))
    assert "two different days" in report.notice
    assert s["Assigned and removed"].rows == []
    assert {r[2] for r in _rows(s["Licences by month"])} == {"10", "5", "3", "0"}


def test_without_the_audit_log_says_why_nobody_is_named(db, two_tenants):
    alpha, _ = two_tenants
    seed_tenant(db, alpha, users=ld.USERS, skus=ld.SKUS, subscriptions=ld.TENANT_SUBSCRIPTIONS, groups=GROUPS)
    db.execute("DELETE FROM event_cursors")  # the directory audit log has never been read
    _backdate_all(db, alpha, 40)
    _collect(db, alpha, "users", users=LATER_USERS)
    report, s = _report(db, alpha)
    assert any("directory audit log" in n for n in report.notes)
    assert "Changed by" not in s["Assigned and removed"].columns  # nobody to name, so no empty column


def test_unavailable_before_anything_is_collected(db, two_tenants):
    _, beta = two_tenants
    report = data_reports.build(db, store.get_tenant(db, beta), KEY)
    assert "haven't been collected" in report.unavailable


def test_since_last_month_lists_licence_changes(db, two_tenants):
    alpha, _ = two_tenants
    _a_month_of_changes(db, alpha)
    since = (datetime.now(UTC) - timedelta(days=35)).strftime("%Y-%m-%dT%H:%M:%SZ")
    found = {c.subject: c for c in changes.licence_changes(db, alpha, since)}
    assert set(found) == {"ann@x.com", "eve@x.com", "dan@x.com", "fay@x.com", "gus@x.com"}
    assert (found["eve@x.com"].kind, found["eve@x.com"].detail) == ("added", "Given Microsoft 365 Copilot")
    assert (found["ann@x.com"].kind, found["ann@x.com"].detail) == ("removed", "Lost Microsoft 365 Copilot")
    assert found["dan@x.com"].area == "Licences"
    assert any(c.area == "Licences" for c in changes.all_changes(db, alpha, since))
    assert changes.licence_changes(db, alpha, datetime.now(UTC).strftime("%Y-%m-%dT%H:%M:%SZ")) == []


def test_page_xlsx_and_billing_job(client, db, two_tenants):
    alpha, _ = two_tenants
    _a_month_of_changes(db, alpha)
    make_user(db, "admin")
    login(client, "admin")
    page = client.get(f"/tenants/{alpha}/data/{KEY}")
    assert page.status_code == 200
    html = page.get_data(as_text=True)
    assert "Licence changes" in html and "Assigned and removed" in html and "helpdesk@x.com" in html
    xlsx = client.get(f"/tenants/{alpha}/data/{KEY}.xlsx")
    assert xlsx.status_code == 200
    sheets = load_workbook(BytesIO(xlsx.data)).sheetnames
    assert any("Assigned and removed" in name for name in sheets)
    billing = next(u for u in library.USES if u.key == "billing")
    assert KEY in billing.reports
