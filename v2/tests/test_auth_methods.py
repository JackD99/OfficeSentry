"""Passwords and sign-in methods: the collector, the report, its two checks and the account page."""

from datetime import UTC, datetime
from io import BytesIO

import pytest
import responses
from openpyxl import load_workbook

from officesentry import authmethods, store
from officesentry.reports import data_reports, evidence, user_profile
from officesentry.rules import evaluate
from officesentry.snapshot import TenantSnapshot
from tests import graphdata as gd
from tests.conftest import login, make_user
from tests.seed import seed_tenant
from tests.test_collectors import V1, _run, factory  # noqa: F401

NOW = datetime(2026, 10, 2, tzinfo=UTC)


def _user(uid, name, upn, password, signed_in=None, created="2024-01-01T00:00:00Z", never_expires=False):
    return {
        "id": uid,
        "displayName": name,
        "userPrincipalName": upn,
        "userType": "Member",
        "accountEnabled": True,
        "createdDateTime": created,
        "assignedLicenses": [{"skuId": "sku-bp"}],
        "lastPasswordChangeDateTime": password,
        "passwordPolicies": "DisablePasswordExpiration" if never_expires else "None",
        "signInActivity": {"lastSignInDateTime": signed_in} if signed_in else {},
    }


USERS = [
    _user("u1", "Ann", "ann@x.com", "2026-01-01T00:00:00Z", "2026-09-30T10:00:00Z"),
    {
        "id": "u3",
        "displayName": "Guest",
        "userPrincipalName": "g#EXT#@x.com",
        "userType": "Guest",
        "accountEnabled": True,
        "assignedLicenses": [],
    },
    _user("u4", "Breakglass", "bg@x.com", "2025-01-01T00:00:00Z", "2026-09-01T00:00:00Z"),
    _user("u5", "Eve", "eve@x.com", "2026-05-01T00:00:00Z", "2026-09-20T00:00:00Z"),
    _user("u6", "Carl", "carl@x.com", "2024-06-01T00:00:00Z", "2026-09-01T00:00:00Z", never_expires=True),
    _user("u8", "Dana", "dana@x.com", "2026-02-01T00:00:00Z", "2026-09-29T00:00:00Z"),
    _user("u9", "Nick", "nick@x.com", "2026-06-01T00:00:00Z", None, created="2026-06-01T00:00:00Z"),
]


def _reg(uid, upn, methods, admin=False, **extra):
    return {
        "id": uid,
        "userPrincipalName": upn,
        "userDisplayName": upn,
        "userType": "member",
        "isAdmin": admin,
        "isMfaRegistered": bool(methods),
        "isMfaCapable": bool(methods),
        "methodsRegistered": methods,
        "isSsprRegistered": False,
        "isPasswordlessCapable": False,
        **extra,
    }


REGISTRATION = [
    _reg("u1", "ann@x.com", ["password", "microsoftAuthenticatorPush"], admin=True, isSsprRegistered=True),
    _reg("u4", "bg@x.com", ["fido2SecurityKey"], admin=True, isPasswordlessCapable=True),
    _reg("u5", "eve@x.com", ["mobilePhone", "email"], admin=True),
    _reg("u6", "carl@x.com", []),
    _reg(
        "u8",
        "dana@x.com",
        ["microsoftAuthenticatorPush", "mobilePhone"],
        isSsprRegistered=True,
        isPasswordlessCapable=True,
        isSystemPreferredAuthenticationMethodEnabled=False,
        userPreferredMethodForSecondaryAuthentication="sms",
    ),
]  # Nick isn't in the report yet, so his methods are read one by one (none)

PR_ADMINS = {
    "id": "ca4",
    "displayName": "Admins need phishing-resistant MFA",
    "state": "enabledForReportingButNotEnforced",
    "conditions": {
        "users": {"includeRoles": [gd.GA], "excludeUsers": ["u4"]},
        "applications": {"includeApplications": ["All"]},
        "clientAppTypes": ["all"],
    },
    "grantControls": {
        "operator": "OR",
        "builtInControls": [],
        "authenticationStrength": {
            "id": "00000000-0000-0000-0000-000000000004",
            "displayName": "Phishing-resistant MFA",
            "policyType": "builtIn",
            "allowedCombinations": ["windowsHelloForBusiness", "fido2", "x509CertificateMultiFactor"],
        },
    },
}


@pytest.fixture
def tenant(db, two_tenants):
    alpha, _ = two_tenants
    seed_tenant(db, alpha, users=USERS, registration=REGISTRATION, ca=[gd.CA_MFA_ALL, PR_ADMINS])
    return store.get_tenant(db, alpha)


def _sections(report):
    return {s.title: s for s in report.sections}


def _rows(section):
    return {row[0]: (row, flag) for row, flag in zip(section.rows, section.flags, strict=True)}


# -- Strength of one person's methods ------------------------------------------------------------


def test_strength_tiers():
    s = authmethods.strength({"methodsRegistered": ["password", "mobilePhone", "email"]})
    assert s.strongest == authmethods.PHONE and s.phone_only and not s.phishing_resistant
    assert s.labels == ["Mobile phone (SMS or call)", "Email"]
    s = authmethods.strength({"methodsRegistered": ["microsoftAuthenticatorPush", "passKeyDeviceBoundAuthenticator"]})
    assert s.strongest == authmethods.PHISHING_RESISTANT and s.labels[0] == "Passkey in Authenticator"
    # Method types read one user at a time name methods differently.
    assert authmethods.strength({"methodsRegistered": ["password", "phone"]}).phone_only
    assert authmethods.strength({"methodsRegistered": ["windowsHelloForBusiness"]}).phishing_resistant
    # A passkey kept in a password manager, as the test tenant's registration report lists it.
    synced = authmethods.strength({"methodsRegistered": ["passKeySynced"]})
    assert synced.phishing_resistant and synced.labels == ["Synced passkey"]
    tap = authmethods.strength({"methodsRegistered": ["temporaryAccessPass"]})
    assert tap.strongest == authmethods.NO_SECOND_FACTOR
    assert authmethods.strength(None) is None
    assert (
        authmethods.preferred(
            {
                "isSystemPreferredAuthenticationMethodEnabled": True,
                "systemPreferredAuthenticationMethods": ["Fido2"],
                "userPreferredMethodForSecondaryAuthentication": "sms",
            }
        )
        == "Passkey or security key"
    )
    assert authmethods.asks_for_phone({"userPreferredMethodForSecondaryAuthentication": "voiceMobile"})


# -- Collectors ----------------------------------------------------------------------------------


@responses.activate
def test_policy_collector_names_groups(db, two_tenants, factory):  # noqa: F811
    alpha, _ = two_tenants
    responses.get(f"{V1}/policies/authenticationMethodsPolicy", json=gd.AUTH_METHODS_POLICY)
    responses.get(f"{V1}/groups/grp-admins", json={"id": "grp-admins", "displayName": "IT admins"})
    run_id, status = _run(db, alpha, "auth_methods_policy", factory)
    assert status == "succeeded"
    assert all(c.request.method == "GET" for c in responses.calls)
    assert "%24select=id%2CdisplayName" in responses.calls[1].request.url
    snap = TenantSnapshot(db, store.get_tenant(db, alpha))
    [settings] = snap.objects("auth_methods_policy", "auth_methods_policy")
    assert settings["registrationCampaign"]["state"] == "default"
    assert settings["systemCredentialPreferences"]["state"] == "disabled"
    methods = {m["_id"]: m for m in snap.objects("auth_methods_policy", "auth_method_config")}
    assert methods["Sms"]["state"] == "enabled" and methods["Sms"]["excludeTargets"][0]["id"] == "grp-admins"
    assert methods["MicrosoftAuthenticator"]["features"]["displayLocationInformationRequiredState"] == "disabled"
    assert methods["Fido2"]["settings"] == {"isSelfServiceRegistrationAllowed": True, "isAttestationEnforced": False}
    assert [(g["_id"], g["displayName"]) for g in snap.objects("auth_methods_policy", "auth_policy_group")] == [
        ("grp-admins", "IT admins")
    ]
    assert snap.metric("auth_methods.sms_or_voice_enabled") == 1


@responses.activate
def test_policy_collector_without_permission_is_partial(db, two_tenants, factory):  # noqa: F811
    alpha, _ = two_tenants
    responses.get(
        f"{V1}/policies/authenticationMethodsPolicy",
        status=403,
        json={"error": {"code": "Authorization_RequestDenied", "message": "Insufficient privileges"}},
    )
    _, status = _run(db, alpha, "auth_methods_policy", factory)
    assert status == "partial"


def test_registration_and_domains_keep_what_the_report_needs(db, tenant):
    snap = TenantSnapshot(db, tenant)
    reg = {r["_id"]: r for r in snap.objects("mfa_registration", "mfa_registration")}
    assert reg["u8"]["userPreferredMethodForSecondaryAuthentication"] == "sms"
    assert reg["u8"]["isSystemPreferredAuthenticationMethodEnabled"] is False
    assert reg["u1"]["isSsprRegistered"] is True and reg["u9"]["source"] == "per_user_methods"
    domains = {d["_id"]: d for d in snap.objects("directory_sync", "domain")}
    assert domains["parked.com"]["passwordValidityPeriodInDays"] == 90
    assert domains["x.com"]["passwordNotificationWindowInDays"] == 14


# -- The report ----------------------------------------------------------------------------------


def test_report_ranks_methods_by_strength(db, tenant):
    report = data_reports.build(db, tenant, "auth-methods", NOW)
    assert report.unavailable is None
    figures = {h.label: h.value for h in report.highlights}
    assert figures["Phishing-resistant"] == "1 of 6"
    assert figures["Admins without phishing-resistant"] == 2  # Ann (app) and Eve (phone), eligible admins count
    assert figures["Phone is the only second factor"] == 1
    assert figures["Passwordless ready"] == "2 of 6"
    assert figures["No self-service reset"] == 3
    people = _sections(report)["Sign-in methods by person"]
    assert [r[0] for r in people.rows] == [
        "eve@x.com",
        "ann@x.com",
        "bg@x.com",
        "carl@x.com",
        "nick@x.com",
        "dana@x.com",
    ]
    rows = _rows(people)
    assert rows["eve@x.com"][1] == "high" and rows["eve@x.com"][0][3] == authmethods.PHONE
    assert rows["eve@x.com"][0][2] == "Global Administrator (eligible)"
    assert rows["ann@x.com"][1] == "medium" and "passkey" in rows["ann@x.com"][0][8]
    assert rows["bg@x.com"][1] is None and rows["bg@x.com"][0][6] == "Yes"
    assert rows["dana@x.com"][1] == "low" and rows["dana@x.com"][0][5] == "Text message (SMS)"
    assert rows["nick@x.com"][0][6] == "Unknown"  # read one by one: passwordless isn't known
    assert "MFA coverage" in rows["carl@x.com"][0][8]
    assert "g#EXT#@x.com" not in rows  # guests sign in at home
    spread = _rows(_sections(report)["Strongest method"])
    assert spread[authmethods.NO_SECOND_FACTOR][0][1:3] == [2, 0]
    methods = _rows(_sections(report)["Methods registered"])
    assert methods["Mobile phone (SMS or call)"][0][2:] == [2, 1] and "Email" in methods


def test_report_compares_the_methods_policy_with_advice(db, tenant):
    report = data_reports.build(db, tenant, "auth-methods", NOW)
    policy = _rows(_sections(report)["Sign-in methods policy"])
    assert policy["Text message (SMS)"][0][1] == "On for All users except Break glass"
    assert policy["Text message (SMS)"][1] == "low"
    assert policy["Passkeys and security keys (FIDO2)"][1] is None
    assert policy["Voice call"][0][1] == "Off" and policy["Voice call"][1] is None
    system = next(v for k, v in policy.items() if k.startswith("System-preferred"))
    assert system[0][1] == "Off" and system[1] == "low"
    assert policy["Authenticator shows where the sign-in is from"][1] == "low"
    strengths = _sections(report)["Authentication strengths in Conditional Access"]
    assert strengths.rows == [
        [
            "Admins need phishing-resistant MFA",
            "Report-only",
            "Phishing-resistant MFA",
            "Yes",
            "Global Administrator",
            "bg@x.com",
        ]
    ]


def test_report_password_hygiene(db, tenant):
    report = data_reports.build(db, tenant, "auth-methods", NOW)
    figures = {h.label: h.value for h in report.highlights}
    assert figures["Password never expires"] == 1
    assert figures["Password over a year old"] == 2  # Breakglass and Carl
    assert figures["Never signed in, old password"] == 1
    passwords = _rows(_sections(report)["Password age"])
    assert passwords["nick@x.com"][1] == "medium" and passwords["nick@x.com"][0][4] == "Never"
    assert passwords["carl@x.com"][1] == "medium" and "no MFA" in passwords["carl@x.com"][0][6]
    assert passwords["bg@x.com"][1] is None and "MFA protects it" in passwords["bg@x.com"][0][6]
    assert passwords["carl@x.com"][0][1:3] == ["2024-06-01", "Yes"]
    expiry = _rows(_sections(report)["Password expiry policy"])
    assert expiry["x.com"][0][2] == "Never" and expiry["x.com"][1] is None
    assert expiry["parked.com"][0][2:4] == ["90 days", "14 days"] and expiry["parked.com"][1] == "low"


def test_report_without_registration_or_policy(db, two_tenants):
    alpha, _ = two_tenants
    tenant = store.get_tenant(db, alpha)
    assert data_reports.build(db, tenant, "auth-methods", NOW).unavailable
    seed_tenant(db, alpha, users=USERS, registration=REGISTRATION, ca=[])
    db.execute("DELETE FROM runs WHERE collector IN ('auth_methods_policy', 'directory_sync')")
    report = data_reports.build(db, tenant, "auth-methods", NOW)
    assert report.unavailable is None
    assert any("Sign-in methods policy" in n for n in report.notes)
    assert "Sign-in methods policy" not in _sections(report)


def test_headline_figures_open_their_rows(client, db, tenant):
    make_user(db, "admin")
    login(client, "admin")
    path = f"/tenants/{tenant['id']}/data/auth-methods"
    from officesentry.reports import views

    report = data_reports.build(db, tenant, "auth-methods", NOW)
    links = [h for h in views.headline_links(report, path)["highlights"] if h]
    assert len(links) == len(report.highlights)
    assert all(client.get(link.split("#")[0]).status_code == 200 for link in links)
    query = links[1].split("?")[1].split("#")[0]  # admins without phishing-resistant
    shown = views.apply(report, views.View.parse(query))
    people = next(s for s in shown.sections if s.title == "Sign-in methods by person")
    assert [r[0] for r in people.rows] == ["eve@x.com", "ann@x.com"]


def test_page_pdf_and_spreadsheet(client, db, tenant):
    make_user(db, "admin")
    login(client, "admin")
    path = f"/tenants/{tenant['id']}/data/auth-methods"
    page = client.get(path)
    assert page.status_code == 200
    html = page.get_data(as_text=True)
    assert "Passwords and sign-in methods" in html and "Admins without phishing-resistant" in html
    assert f'href="/tenants/{tenant["id"]}/users/u5"' in html
    pdf = client.get(f"{path}.pdf")
    assert pdf.status_code == 200
    xlsx = client.get(f"{path}.xlsx")
    assert "Sign-in methods by person" in load_workbook(BytesIO(xlsx.data)).sheetnames
    assert client.get("/reports?q=passkey").status_code == 200


# -- Checks ----------------------------------------------------------------------------------


def test_checks(db, tenant):
    results = {r.key: r for r in evaluate(TenantSnapshot(db, tenant), NOW)}
    admins = results["admin_phishing_resistant"]
    assert admins.status == "fail" and admins.summary == "2 admins without a phishing-resistant method"
    assert {f.subject: f.detail for f in admins.findings} == {
        "ann@x.com": "Global Administrator; signs in with Authenticator app",
        "eve@x.com": "Global Administrator; signs in with Mobile phone (SMS or call), Email",
    }
    phone = results["phone_only_mfa"]
    assert phone.status == "fail" and phone.summary == "1 person can only use a text or call as a second factor"
    assert [(f.severity, f.subject, f.detail) for f in phone.findings] == [
        ("medium", "eve@x.com", "Second factor is only Mobile phone (SMS or call) (holds an admin role)")
    ]


def test_checks_pass_and_are_unavailable_without_data(db, two_tenants):
    alpha, bravo = two_tenants
    strong = [
        _reg(r["id"], r["userPrincipalName"], ["fido2SecurityKey", "microsoftAuthenticatorPush"]) for r in REGISTRATION
    ]
    seed_tenant(db, alpha, users=USERS, registration=strong)
    results = {r.key: r for r in evaluate(TenantSnapshot(db, store.get_tenant(db, alpha)), NOW)}
    assert results["admin_phishing_resistant"].status == "pass"
    assert results["admin_phishing_resistant"].summary == "All 3 admins have a phishing-resistant method"
    assert results["phone_only_mfa"].status == "pass"
    empty = {r.key: r for r in evaluate(TenantSnapshot(db, store.get_tenant(db, bravo)), NOW)}
    assert empty["admin_phishing_resistant"].status == "unavailable"
    assert empty["phone_only_mfa"].status == "unavailable"


def test_evidence_mentions_phishing_resistant_admins(db, tenant):
    pack = evidence.build(db, tenant, NOW)
    admins = next(c for c in pack.controls if c.control.key == "mfa_admins")
    assert any("Phishing-resistant MFA" in line and "2 admins" in line for line in admins.evidence)


# -- The account page --------------------------------------------------------------------------


def test_account_page_shows_method_strength(db, tenant):
    eve = user_profile.build(db, tenant, "u5", NOW)
    mfa = next(p for p in eve.parts if p.key == "mfa")
    facts = dict(mfa.facts)
    assert facts["Strongest method"] == authmethods.PHONE
    assert facts["Can reset their own password"] == "No"
    texts = [(a.severity, a.text) for a in eve.attention if a.part == "mfa"]
    assert any(s == "medium" and t.startswith("Admin without a phishing-resistant method") for s, t in texts)
    assert any(s == "medium" and t.startswith("A text or call is the only second factor") for s, t in texts)
    bg = user_profile.build(db, tenant, "u4", NOW)
    assert not any("phishing-resistant" in a.text for a in bg.attention)
    assert dict(next(p for p in bg.parts if p.key == "mfa").facts)["Can sign in without a password"] == "Yes"
    dana = user_profile.build(db, tenant, "u8", NOW)
    assert dict(next(p for p in dana.parts if p.key == "mfa").facts)["Asked for first"] == "Text message (SMS)"
    assert not any("phishing-resistant" in a.text for a in dana.attention)  # not an admin
