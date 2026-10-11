"""Question sets: insurance, supplier and Cyber Essentials answers with text to paste, and the gaps left to fill."""

import csv
import io
from io import BytesIO

from openpyxl import load_workbook

from officesentry import store
from officesentry.reports import evidence, evidence_answers, export, questionnaires
from officesentry.reports.evidence_base import BLANK, NOT_KNOWN, YES
from officesentry.reports.questionnaires import INPUT, SETS
from tests import graphdata as gd
from tests.conftest import login, make_user
from tests.seed import seed_tenant

# Every Cyber Essentials question number used, as printed in IASME's Danzell booklet (version 16, April 2026).
DANZELL = {
    "A2.4.2",
    "A2.6",
    "A2.8",
    "A2.9",
    "A4.1",
    "A4.1.1",
    "A4.2",
    "A4.6",
    "A4.7",
    "A4.9",
    "A5.1",
    "A5.2",
    "A5.3",
    "A5.8",
    "A5.9",
    "A6.1",
    "A6.2",
    "A6.4",
    "A6.5",
    "A7.1",
    "A7.2",
    "A7.3",
    "A7.4",
    "A7.5",
    "A7.6",
    "A7.7",
    "A7.8",
    "A7.9",
    "A7.10",
    "A7.13",
    "A7.14",
    "A7.16",
    "A7.17",
    "A8.1",
}


def _sheet(db, tenant_id, key):
    return questionnaires.sheet(evidence.build(db, store.get_tenant(db, tenant_id)), key)


def _rows(db, tenant_id, key):
    return {(r.item.ref or r.item.topic): r for r in _sheet(db, tenant_id, key).rows}


def test_every_question_names_a_topic_and_every_gap_has_a_template():
    topics = {q.key: q for q in evidence_answers.QUESTIONS}
    assert len(topics) == len(evidence_answers.QUESTIONS)  # keys are unique
    for s in SETS.values():
        assert s.items and s.title and s.short and s.intro
        for item in s.items:
            assert item.topic in topics, (s.key, item.topic)
            assert item.question.endswith(("?", ".")), item.question
    for q in topics.values():
        assert q.kind in ("yesno", "text")
        if q.outside:
            assert q.template and BLANK.search(q.template), q.key
    used = {i.topic for s in SETS.values() for i in s.items}
    assert used == set(topics), f"topics no question set asks: {set(topics) - used}"


def test_a_set_is_checked_when_registered():
    part = questionnaires.Part("Access", (questionnaires.Item("mfa_remote", "MFA?", ref="1"),))
    for bad, why in (
        (questionnaires.QuestionSet("insurance", "T", "T", "I", (part,)), "already exists"),
        (questionnaires.QuestionSet("controls", "T", "T", "I", (part,)), "already exists"),
        (
            questionnaires.QuestionSet(
                "x", "T", "T", "I", (questionnaires.Part("P", (questionnaires.Item("nope", "Q?"),)),)
            ),
            "unknown topics: nope",
        ),
        (questionnaires.QuestionSet("x", "T", "T", "I", (part, part)), "repeats"),
    ):
        try:
            questionnaires.register(bad)
        except ValueError as e:
            assert why in str(e)
        else:
            raise AssertionError(f"registered {bad.key}")
    assert "x" not in SETS


def test_cyber_essentials_numbers_follow_danzell():
    ce = SETS["cyber-essentials"]
    refs = [i.ref for i in ce.items]
    assert set(refs) == DANZELL and len(refs) == len(set(refs))
    assert "Danzell" in ce.title and "version 16" in ce.source and ce.source_url.startswith("https://iasme.co.uk/")
    assert {i.ref for i in ce.items if i.autofail} == {"A6.4", "A6.5", "A7.16", "A7.17"}


def test_answers_from_seeded_data(db, two_tenants):
    alpha, _ = two_tenants
    seed_tenant(db, alpha)
    ins = _rows(db, alpha, "insurance")
    mfa = ins["mfa_remote"]
    assert mfa.state == YES and mfa.answer.text.startswith("Multi-factor authentication (MFA) is enforced")
    assert "bg@x.com (break-glass)" in mfa.answer.text and "VPN" in mfa.beyond
    assert mfa.answer.as_of and mfa.based_on == ""  # the text already carries the evidence
    backups = ins["backups"]
    assert backups.state == INPUT and backups.tone == "input" and "[backup product or provider]" in backups.answer.text
    assert backups.answer.blanks and backups.based_on.startswith("Backups aren't visible")
    assert (
        any(blank for _, blank in backups.segments) and "".join(t for t, _ in backups.segments) == backups.answer.text
    )
    assert "Exchange Online Protection" in ins["email_filtering"].answer.text

    ce = _rows(db, alpha, "cyber-essentials")
    computers = ce["A2.6"]
    assert computers.state == questionnaires.ANSWERED
    assert "1 Dell computer running Windows 11 25H2" in computers.answer.text
    assert "1 HP computer running Windows 10" in computers.answer.text and "edition" in computers.beyond
    assert "1 Apple device running iOS 18.6" in ce["A2.8"].answer.text
    cloud = ce["A2.9"].answer.text
    assert "Zoom" in cloud and "Mail Sync Pro" in cloud and "Microsoft 365 (" in cloud
    assert ce["A7.17"].item.autofail and ce["A7.17"].state == YES
    assert ce["A6.4"].state == INPUT and "14 days" in ce["A6.4"].answer.text

    s = _sheet(db, alpha, "insurance")
    n = s.counts
    assert n["data"] + n["team"] + n["input"] == len(s.rows) and n["team"] == 0
    assert sum(c for _, c, _ in s.breakdown) == len(s.rows)


def test_nothing_collected_answers_nothing(db, two_tenants):
    _, bravo = two_tenants
    for key in SETS:
        s = _sheet(db, bravo, key)
        assert all(r.state == INPUT for r in s.rows), key
        assert s.counts["data"] == 0 and s.counts["input"] == len(s.rows)


def test_mail_gateway_is_named(db, two_tenants):
    alpha, _ = two_tenants
    dns = dict(gd.DNS)
    dns[("x.com", "MX")] = ["eu-smtp-inbound-1.mimecast.com"]
    seed_tenant(db, alpha, dns=lambda name, rdtype: list(dns.get((name, rdtype), [])))
    text = _rows(db, alpha, "insurance")["email_filtering"].answer.text
    assert "passes through Mimecast" in text and "Exchange Online Protection" not in text


def test_one_answer_fills_every_set(db, two_tenants):
    alpha, _ = two_tenants
    seed_tenant(db, alpha)
    note = "Datto SaaS Protection backs up Microsoft 365 three times a day, kept for a year. Last restore test 2 Oct."
    evidence_answers.set_override(db, alpha, "backups", YES, note, NOT_KNOWN, "jack")
    for key in ("insurance", "supplier"):
        row = next(r for r in _sheet(db, alpha, key).rows if r.item.topic == "backups")
        assert row.state == YES and row.answer.text == note and not row.answer.blanks
        assert row.based_on == "Answered by jack." and questionnaires.answered_by(row.answer).startswith("jack, ")
    assert _sheet(db, alpha, "insurance").counts["team"] == 1


def test_questions_asked_the_other_way_round_swap_the_answer(db, two_tenants):
    alpha, _ = two_tenants
    evidence_answers.set_override(db, alpha, "firewall_inbound", YES, "Inbound is blocked.", NOT_KNOWN, "jack")
    ce = _rows(db, alpha, "cyber-essentials")
    assert ce["A4.7"].state == "No" and "other way round" in ce["A4.7"].beyond


SCHEME_OF = {"insurance": "insurance", "supplier": "supplier", "cyber-essentials": "cyber-essentials"}


def test_pages_for_staff_and_clients(client, db, two_tenants):
    alpha, bravo = two_tenants
    seed_tenant(db, alpha)
    make_user(db, "lead", role="analyst", tenants=[alpha])
    make_user(db, "client", role="client", tenants=[alpha])
    login(client, "lead")
    for key, s in SETS.items():
        page = client.get(f"/tenants/{alpha}/evidence/{SCHEME_OF[key]}").get_data(as_text=True)
        assert s.title.split(" (")[0] in page or s.short in page
        assert "Copy all answers" in page and 'id="copy-all"' in page and 'data-copy="#q-' in page
        assert '<mark class="blank">[' in page and "Write this answer for Alpha Ltd" in page
        assert f'name="scheme" value="{SCHEME_OF[key]}"' in page
        assert client.get(f"/tenants/{bravo}/evidence/{SCHEME_OF[key]}").status_code == 404
    ce = client.get(f"/tenants/{alpha}/evidence/cyber-essentials").get_data(as_text=True)
    assert "A7.17" in ce and "Auto-fail" in ce and "Danzell" in ce
    only_input = client.get(f"/tenants/{alpha}/evidence/insurance?show=input").get_data(as_text=True)
    assert "Write this answer for Alpha Ltd" in only_input and 'name="show" value="input"' in only_input
    assert only_input.count("Write this answer for") < page.count("Write this answer for") or "Your input" in only_input

    resp = client.post(
        f"/tenants/{alpha}/evidence/answers/backups",
        data={"answer": "Yes", "note": "Backed up nightly by Acme Backup.", "set": "supplier"},
    )
    assert resp.headers["Location"].endswith(f"/tenants/{alpha}/evidence/supplier#answer-backups")
    page = client.get(f"/tenants/{alpha}/evidence/cyber-essentials").get_data(as_text=True)
    assert "Backed up nightly by Acme Backup." not in page  # Cyber Essentials doesn't ask about backups
    assert "Backed up nightly by Acme Backup." in client.get(f"/tenants/{alpha}/evidence/insurance").get_data(
        as_text=True
    )

    client.post("/logout")
    login(client, "client")
    page = client.get(f"/tenants/{alpha}/evidence/supplier").get_data(as_text=True)
    assert "Backed up nightly by Acme Backup." in page and "Copy all answers" in page
    assert "Write this answer" not in page and "Change this answer" not in page
    assert "Suggested from the data" not in page


def test_downloads_for_each_set(client, db, two_tenants, monkeypatch):
    from officesentry.reports import scheme_export, schemes

    alpha, bravo = two_tenants
    seed_tenant(db, alpha)
    evidence_answers.set_override(db, alpha, "training", YES, "Annual training with KnowBe4.", NOT_KNOWN, "jack")
    make_user(db, "client", role="client", tenants=[alpha])
    login(client, "client")
    for key, s in SETS.items():
        scheme = schemes.SCHEMES[SCHEME_OF[key]]
        text = client.get(f"/tenants/{alpha}/evidence/{scheme.key}.txt")
        assert text.mimetype == "text/plain" and f"-{scheme.key}.txt" in text.headers["Content-Disposition"]
        body = text.get_data(as_text=True)
        assert body.startswith(f"{scheme.title}: Alpha Ltd") and "Answer: Your input" in body
        assert f"Answers for the form: {s.title}".upper() in body

        rows = list(
            csv.reader(io.StringIO(client.get(f"/tenants/{alpha}/evidence/{scheme.key}.csv").get_data(as_text=True)))
        )
        assert [rows[0][0].lstrip("\ufeff")] + rows[0][1:] == scheme_export.COLUMNS
        assert len(rows) == len(scheme.controls) + 1

        wb = load_workbook(BytesIO(client.get(f"/tenants/{alpha}/evidence/{scheme.key}.xlsx").data))
        assert s.short in wb.sheetnames and "Controls" in wb.sheetnames
        assert wb[s.short].max_row == len(s.items) + 1

        for ext in ("txt", "csv", "xlsx", "pdf", "docx"):
            assert client.get(f"/tenants/{bravo}/evidence/{scheme.key}.{ext}").status_code == 404
    ce = client.get(f"/tenants/{alpha}/evidence/cyber-essentials.txt").get_data(as_text=True)
    assert "A7.17 Has MFA been applied to all users of your cloud services? (auto-fail)" in ce
    ins = client.get(f"/tenants/{alpha}/evidence/insurance.txt").get_data(as_text=True)
    assert "Annual training with KnowBe4.\nAnswered by jack on " in ins
    logged = {r["detail"] for r in db.execute("SELECT detail FROM audit_log WHERE action = 'report_downloaded'")}
    assert {"evidence-insurance.txt", "evidence-supplier.csv", "evidence-cyber-essentials.xlsx"} <= logged

    if export.pdf_available():
        assert client.get(f"/tenants/{alpha}/evidence/cyber-essentials.pdf").data[:4] == b"%PDF"
    monkeypatch.setattr(export, "pdf_available", lambda: False)  # the same template as a print page
    printed = client.get(f"/tenants/{alpha}/evidence/insurance.pdf").get_data(as_text=True)
    assert "Cyber insurance proposal" in printed and "Annual training with KnowBe4." in printed
    assert "Set by jack on" in printed and "Suggested from the data" not in printed
