"""Our work this quarter (reports/our_work.py, /across/our-work): what the team got done across its clients."""

from datetime import UTC, date, datetime, timedelta
from io import BytesIO

from openpyxl import load_workbook

from officesentry import store
from officesentry.reports import export, our_work
from tests.conftest import login, make_user
from tests.test_roles import role_user

NOW = datetime(2026, 10, 24, 9, 0, tzinfo=UTC)
ISO = "%Y-%m-%dT%H:%M:%SZ"


def finding(db, tenant_id, n, severity="high", rule="admin_mfa", resolved=None, accepted=None, handed_to=None):
    db.execute(
        "INSERT INTO findings (tenant_id, rule_key, object_id, severity, subject, detail, first_seen, last_seen,"
        " resolved_at, accepted_at, handed_to) VALUES (?, ?, ?, ?, ?, '', '2026-06-01T00:00:00Z',"
        " '2026-06-01T00:00:00Z', ?, ?, ?)",
        (tenant_id, rule, f"obj{n}", severity, f"user{n}@x.com", resolved, accepted, handed_to),
    )


def metric(db, tenant_id, key, value, at):
    db.execute(
        "INSERT INTO metrics (tenant_id, key, value, captured_at) VALUES (?, ?, ?, ?)", (tenant_id, key, value, at)
    )


def seed(db, alpha, bravo, sam):
    finding(db, alpha, 1, "critical", resolved="2026-08-03T02:00:00Z", handed_to=sam)
    finding(db, alpha, 2, "high", resolved="2026-10-02T02:00:00Z", handed_to=sam)
    finding(db, alpha, 3, "medium", rule="pending_guests", resolved="2026-10-20T02:00:00Z")
    finding(db, alpha, 4, "high", resolved="2026-07-30T02:00:00Z")  # the quarter before
    finding(db, bravo, 5, "low", rule="pending_guests", resolved="2026-09-09T02:00:00Z")
    finding(db, bravo, 6, "medium", accepted="2026-09-15T10:00:00Z")
    finding(db, bravo, 7, "high")  # still open
    # Alpha's Secure Score from 40 (read in July) to 55; Bravo's first reading is in the quarter, so no change.
    metric(db, alpha, "secure_score.pct", 40, "2026-07-28T02:00:00Z")
    metric(db, alpha, "secure_score.pct", 48, "2026-09-01T02:00:00Z")
    metric(db, alpha, "secure_score.pct", 55, "2026-10-23T02:00:00Z")
    metric(db, bravo, "secure_score.pct", 60, "2026-10-23T02:00:00Z")
    metric(db, alpha, "licences.unassigned_paid", 9, "2026-07-31T02:00:00Z")
    metric(db, alpha, "licences.unassigned_paid", 3, "2026-10-23T02:00:00Z")
    metric(db, bravo, "licences.unassigned_paid", 2, "2026-08-02T02:00:00Z")
    metric(db, bravo, "licences.unassigned_paid", 4, "2026-10-23T02:00:00Z")


def test_the_quarters_work_by_month_client_check_and_person(db, two_tenants):
    alpha, bravo = two_tenants
    sam = make_user(db, "sam")
    seed(db, alpha, bravo, sam)
    tenants = [store.get_tenant(db, alpha), store.get_tenant(db, bravo)]
    work = our_work.build(db, tenants, NOW)

    assert work.label == "August to October 2026" and work.so_far
    assert (work.fixed, work.severe, work.accepted) == (4, 2, 1)
    assert [(m.month, m.fixed, m.severe, m.accepted) for m in work.months] == [
        (date(2026, 8, 1), 1, 1, 0),
        (date(2026, 9, 1), 1, 0, 1),
        (date(2026, 10, 1), 2, 1, 0),
    ]
    a, b = work.clients  # most fixed first
    assert a.tenant["id"] == alpha and (a.fixed, a.severe, a.accepted) == (3, 2, 0)
    assert (b.fixed, b.accepted) == (1, 1) and len(work.worked) == 2
    assert [(k.key, k.fixed, len(k.clients), k.severity) for k in work.checks] == [
        ("admin_mfa", 2, 1, "critical"),
        ("pending_guests", 2, 2, "medium"),
    ]
    assert [(p.name, p.fixed, p.severe, len(p.clients)) for p in work.people] == [("sam", 2, 2, 1)]

    score = a.figures["secure_score.pct"]
    assert (score.start, score.end, score.change) == (40, 55, 15)
    only_now = b.figures["secure_score.pct"]
    assert (only_now.start, only_now.end, only_now.change) == (None, 60, None)
    assert [c.tenant["id"] for c in work.improved("secure_score.pct")] == [alpha]
    assert a.unused_removed == 6 and b.unused_removed == 0 and work.unused_removed == 6
    assert [f.key for f in work.figures] == ["secure_score.pct", "licences.unassigned_paid"]  # no MFA readings

    head = work.to_tables()[0].columns
    assert head[:5] == ["Client", "Tenant ID", "Findings fixed", "Critical and high fixed", "Risks accepted"]
    assert "Secure Score at the start (%)" in head and "Unused paid licences change" in head


def test_the_quarter_before_and_figures_the_person_cant_see(db, two_tenants):
    alpha, bravo = two_tenants
    seed(db, alpha, bravo, make_user(db, "sam"))
    tenants = [store.get_tenant(db, alpha), store.get_tenant(db, bravo)]
    before = our_work.build(db, tenants, NOW, back=1)
    assert before.label == "May to July 2026" and not before.so_far
    assert before.end == datetime(2026, 8, 1, tzinfo=UTC)
    assert before.fixed == 1 and before.clients[0].tenant["id"] == alpha

    hidden = our_work.build(db, tenants, NOW, sees=lambda tid, report: not (tid == alpha and report == "licences"))
    a = next(c for c in hidden.clients if c.tenant["id"] == alpha)
    assert "licences.unassigned_paid" not in a.figures and "secure_score.pct" in a.figures
    assert our_work.build(db, [], NOW).clients == []


def test_page_downloads_and_a_pdf_without_names(client, db, two_tenants, monkeypatch):
    monkeypatch.setattr(export, "pdf_available", lambda: False)  # the PDF's HTML, as the browser prints it
    alpha, bravo = two_tenants
    now = datetime.now(UTC)
    finding(db, alpha, 1, "critical", resolved=(now - timedelta(hours=2)).strftime(ISO))
    finding(db, bravo, 2, "medium", accepted=(now - timedelta(hours=3)).strftime(ISO))
    make_user(db, "admin")
    login(client, "admin")

    page = client.get("/across/our-work").get_data(as_text=True)
    assert (
        "Our work this quarter" in page
        and "1 finding fixed, 1 of them critical or high, and 1 risk accepted with the client, across both clients"
        in page
    )
    assert "Every admin has MFA registered" in page and 'href="/across/our-work.pdf?names=0"' in page
    assert client.get("/across?key=our-work").headers["Location"].endswith("/across/our-work")
    assert "The quarter before" in client.get("/across/our-work?back=1").get_data(as_text=True)

    book = load_workbook(BytesIO(client.get("/across/our-work.xlsx").data))
    assert {"By client", "By month", "What was fixed", "Who fixed it"} <= set(book.sheetnames)
    named = client.get("/across/our-work.pdf").get_data(as_text=True)
    assert "Alpha Ltd" in named and "What was fixed" in named
    unnamed = client.get("/across/our-work.pdf?names=0").get_data(as_text=True)
    assert "Alpha Ltd" not in unnamed and "Bravo Inc" not in unnamed and "Client 1" in unnamed
    rows = db.execute("SELECT detail FROM audit_log WHERE action = 'rollup_downloaded'").fetchall()
    assert any(r[0].endswith("-without-names.pdf") for r in rows)


def test_who_sees_it(client, db, two_tenants):
    alpha, bravo = two_tenants
    finding(db, bravo, 1, "high", resolved=(datetime.now(UTC) - timedelta(hours=1)).strftime(ISO))
    role_user(db, "lead", "second_line", tenants=[alpha])
    login(client, "lead")
    page = client.get("/across/our-work").get_data(as_text=True)
    assert "Alpha Ltd" in page and "Bravo Inc" not in page and "nothing fixed or accepted yet across 1 client" in page
    client.post("/logout")
    role_user(db, "first", "first_line", all_tenants=True)  # compares clients, but doesn't see findings
    login(client, "first")
    assert client.get("/across/our-work").status_code == 403
    assert "/across/our-work" not in client.get("/across/over-time").get_data(as_text=True)
