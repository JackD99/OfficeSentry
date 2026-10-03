"""Small page behaviours from the UI review: clickable headline numbers, compact empty tables, coloured flags."""

from __future__ import annotations

from io import BytesIO

from openpyxl import load_workbook

from officesentry import store
from officesentry.reports import data_reports
from officesentry.rules import RULES, RuleResult, evaluate
from officesentry.snapshot import TenantSnapshot
from officesentry.text import every, nouns, plural, verb
from tests.conftest import login, make_user
from tests.seed import seed_tenant


def test_overview_numbers_open_their_reports(client, db, two_tenants):
    alpha, _ = two_tenants
    seed_tenant(db, alpha)
    make_user(db, "admin")
    login(client, "admin")
    page = client.get(f"/tenants/{alpha}").get_data(as_text=True)
    for key in ("secure-score", "mfa", "admins", "licences", "inactive"):
        assert f'class="tile link" href="/tenants/{alpha}/data/{key}"' in page, key


def test_empty_tables_are_one_line(client, db, two_tenants):
    alpha, _ = two_tenants
    seed_tenant(db, alpha)
    make_user(db, "admin")
    login(client, "admin")
    tenant = db.execute("SELECT * FROM tenants WHERE id = ?", (alpha,)).fetchone()
    report = data_reports.build(db, tenant, "mfa")
    empty = [s for s in report.sections if not s.rows]
    assert empty, "the seeded MFA report has an empty table"
    page = client.get(f"/tenants/{alpha}/data/mfa").get_data(as_text=True)
    assert page.count('class="card flush empty-section"') == len(empty)
    assert empty[0].empty in page


def test_xlsx_flags_are_coloured(db, two_tenants):
    alpha, _ = two_tenants
    seed_tenant(db, alpha)
    tenant = db.execute("SELECT * FROM tenants WHERE id = ?", (alpha,)).fetchone()
    report = data_reports.build(db, tenant, "admins")
    wb = load_workbook(BytesIO(data_reports.to_xlsx(report)))
    flagged = [c for ws in wb.worksheets for row in ws.iter_rows(min_row=2) for c in row[-1:]
               if c.value in data_reports.FLAG_FILLS]
    assert flagged
    for cell in flagged:
        assert cell.fill.fgColor.rgb.endswith(data_reports.FLAG_FILLS[cell.value][0])


def test_reports_not_compared_are_grouped_by_reason(client, db, two_tenants):
    alpha, _ = two_tenants
    seed_tenant(db, alpha)  # collected today only: nothing from 30 days ago
    make_user(db, "admin")
    login(client, "admin")
    page = client.get(f"/tenants/{alpha}/changes").get_data(as_text=True)
    assert page.count("Not compared (no data from the start of the period to compare with)") == 1


# -- Wording ---------------------------------------------------------------------------


def test_counted_nouns_read_naturally():
    assert plural(1, "admin") == "1 admin" and plural(3, "admin") == "3 admins"
    assert plural(2, "MFA policy") == "2 MFA policies" and plural(2, "day") == "2 days"
    assert plural(2, "mailbox") == "2 mailboxes" and plural(0, "library") == "0 libraries"
    assert plural(2, "team or group") == "2 teams and groups" and plural(1, "team or group") == "1 team or group"
    assert nouns(5, "mail domain") == "mail domains" and verb(1, "is", "are") == "is"
    assert every(4, "domain") == "all 4 domains" and every(1, "domain", start=True) == "The 1 domain"


def test_check_summaries_never_say_s_in_brackets(db, two_tenants):
    alpha, _ = two_tenants
    seed_tenant(db, alpha)
    results = evaluate(TenantSnapshot(db, store.get_tenant(db, alpha)))
    text = [r.summary for r in results] + [f.detail for r in results for f in r.findings]
    assert not [t for t in text if "(s)" in t]


def test_every_check_has_a_short_action():
    for r in RULES:
        action = RuleResult(r.key, r.title, r.category, "fail", "", r.recommendation).action
        assert action and not action.endswith((".", ",")) and len(action) < 200, r.key
    rec = "Publish an MTA-STS policy so other servers only deliver over TLS: a TXT record (v=STSv1). Then wait."
    assert RuleResult("k", "t", "c", "fail", "", rec).action == "Publish an MTA-STS policy"


def test_failing_checks_lead_with_the_problem(client, db, two_tenants):
    """The check name is the goal ("No inactive accounts"); under Needs attention it read as the opposite."""
    alpha, _ = two_tenants
    seed_tenant(db, alpha)
    make_user(db, "admin")
    login(client, "admin")
    page = client.get(f"/tenants/{alpha}").get_data(as_text=True)
    attention = page.split("<h2>Needs attention</h2>")[1].split("</section>")[0]
    assert '<div class="title">1 admin without MFA' in attention
    assert "Have each listed admin register a strong method." in attention
    assert "Every admin has MFA registered" not in attention
    review = client.get(f"/tenants/{alpha}/report").get_data(as_text=True)
    steps = review.split('<ol class="actions">')[1].split("</ol>")[0]
    assert "Have each listed admin register a strong method" in steps and "Every admin has" not in steps


def test_tenants_page_explains_the_coloured_counts(client, db, two_tenants):
    make_user(db, "admin")
    login(client, "admin")
    page = client.get("/").get_data(as_text=True)
    assert 'class="sev-legend"' in page
    for level in ("Critical", "High", "Medium", "Low"):
        assert f"{level}</span>" in page
