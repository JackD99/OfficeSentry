"""Exchange detail reports: inbox rules, mailbox permissions, shared mailboxes, size and legacy sign-ins."""

import copy
import json
from datetime import UTC, datetime, timedelta
from io import BytesIO
from urllib.parse import parse_qs, urlsplit

import pytest
import responses
from openpyxl import load_workbook

from officesentry import store
from officesentry.collectors.mail_usage import LEGACY_CLIENTS, summarise_signins
from officesentry.exchange import ExchangeClient
from officesentry.graph import GRAPH_ROOT, GraphClient
from officesentry.mailflow import size_bytes
from officesentry.reports.data_reports import REPORTS, build
from officesentry.snapshot import TenantSnapshot
from officesentry.worker import execute_run
from tests import graphdata as gd
from tests.conftest import login, make_user
from tests.seed import seed_tenant
from tests.test_collectors import SKUS, USERS

GB = 1024 ** 3
MAIL_REPORTS = ("inbox-rules", "mailbox-permissions", "shared-mailboxes", "mailbox-size", "legacy-auth")


def _exchange():
    data = copy.deepcopy(gd.EXCHANGE)
    ann, info = data["Get-Mailbox"]
    ann.update({"Name": "u1", "Alias": "ann", "ProhibitSendQuota": "49.5 GB (53,150,220,288 bytes)",
                "ArchiveStatus": "Active"})
    info.update({"Name": "Info", "Alias": "info", "ProhibitSendQuota": "Unlimited", "LitigationHoldEnabled": True})
    data["Get-Mailbox"].append({
        "ExternalDirectoryObjectId": "", "Guid": "g-disc", "PrimarySmtpAddress": "DiscoverySearchMailbox@x.com",
        "Name": "DiscoverySearchMailbox", "RecipientTypeDetails": "DiscoveryMailbox", "GrantSendOnBehalfTo": []})
    data["Get-MailboxPermission"]["g-u7"].append(
        {"User": "bob@x.com", "AccessRights": ["FullAccess"], "IsInherited": False})
    data["Get-RecipientPermission"].append({"Identity": "ExchangeOnline-123", "Trustee": "NULL SID",
                                            "AccessRights": ["SendAs"]})
    data["Get-MailboxStatistics"] = {
        "g-u1": [{"TotalItemSize": "48 GB (51,539,607,552 bytes)", "TotalDeletedItemSize": "1 GB (1,073,741,824 bytes)",
                  "ItemCount": 90000}],
        "g-u7": [{"TotalItemSize": "586.6 KB (600,640 bytes)", "TotalDeletedItemSize": "0 B (0 bytes)",
                  "ItemCount": 13}],
    }
    data["Get-CASMailbox"] = [
        {"ExternalDirectoryObjectId": "u1", "Guid": "g-u1", "PrimarySmtpAddress": "ann@x.com",
         "SmtpClientAuthenticationDisabled": True, "PopEnabled": False, "ImapEnabled": True},
        {"ExternalDirectoryObjectId": "u7", "Guid": "g-u7", "PrimarySmtpAddress": "info@x.com",
         "SmtpClientAuthenticationDisabled": None, "PopEnabled": True, "ImapEnabled": True},
        {"ExternalDirectoryObjectId": "", "Guid": "g-disc", "PrimarySmtpAddress": "DiscoverySearchMailbox@x.com"},
    ]
    return data


@pytest.fixture
def tenant(db, two_tenants):
    alpha, _ = two_tenants
    seed_tenant(db, alpha, users=USERS + [gd.INFO_USER], exchange=_exchange())
    return store.get_tenant(db, alpha)


def _sections(db, tenant, key):
    report = build(db, tenant, key)
    assert not report.unavailable, report.unavailable
    return report, {s.title: s for s in report.sections}


def test_size_bytes():
    assert size_bytes("49.5 GB (53,150,220,288 bytes)") == 53150220288
    assert size_bytes("0 B (0 bytes)") == 0
    assert size_bytes("Unlimited") is None
    assert size_bytes(None) is None and size_bytes(["x"]) is None and size_bytes(True) is None
    assert size_bytes(1234) == 1234


def test_new_cmdlets_are_read_only():
    client = ExchangeClient("tid", lambda: "t", sleep=lambda s: None)
    for cmdlet in ("Set-CASMailbox", "Set-Mailbox", "New-MailboxExportRequest"):
        with pytest.raises(ValueError):
            client.invoke(cmdlet)


def test_collectors_store_sizes_protocols_and_signins(db, tenant):
    snap = TenantSnapshot(db, tenant)
    for key in ("mailboxes", "mailbox_usage", "legacy_signins"):
        assert snap.runs[key]["status"] == "succeeded", key
    sizes = {s["mailbox"]: s for s in snap.objects("mailbox_usage", "mailbox_size")}
    assert set(sizes) == {"ann@x.com", "info@x.com"}  # the discovery mailbox is skipped
    assert sizes["ann@x.com"]["totalBytes"] == 48 * GB
    assert sizes["ann@x.com"]["sendQuotaBytes"] == 53150220288 and sizes["ann@x.com"]["archive"] is True
    assert sizes["info@x.com"]["sendQuotaBytes"] is None
    assert {c["PrimarySmtpAddress"] for c in snap.objects("mailboxes", "cas_mailbox")} == {"ann@x.com", "info@x.com"}
    assert [p["Trustee"] for p in snap.objects("mailboxes", "send_as")] == ["ann@x.com"]  # no NULL SID
    signins = {(s["user"], s["client"]): s for s in snap.objects("legacy_signins", "legacy_signin")}
    assert signins[("info@x.com", "Authenticated SMTP")]["succeeded"] == 1
    assert signins[("ann@x.com", "IMAP4")]["failed"] == 1


def test_summarise_signins_groups_by_account_and_client():
    rows = dict(summarise_signins([
        {"createdDateTime": "2026-09-01T00:00:00Z", "userPrincipalName": "A@x.com", "clientAppUsed": "POP3",
         "ipAddress": "1.1.1.1", "status": {"errorCode": 0}, "location": {"countryOrRegion": "GB"}},
        {"createdDateTime": "2026-09-03T00:00:00Z", "userPrincipalName": "a@x.com", "clientAppUsed": "POP3",
         "ipAddress": "2.2.2.2", "status": {"errorCode": 50126, "failureReason": "Bad password"},
         "location": {"countryOrRegion": "US"}},
        {"createdDateTime": "2026-09-02T00:00:00Z", "clientAppUsed": None, "status": None, "location": None},
    ]))
    pop = rows["a@x.com|POP3"]
    assert (pop["succeeded"], pop["failed"]) == (1, 1)
    assert pop["lastSuccess"] == "2026-09-01T00:00:00Z" and pop["lastAttempt"] == "2026-09-03T00:00:00Z"
    assert pop["lastFailure"] == "Bad password" and pop["countries"] == ["GB", "US"]
    assert rows["unknown|Other clients"]["succeeded"] == 1


def test_inbox_rules_report(db, tenant):
    report, s = _sections(db, tenant, "inbox-rules")
    check = s["Rules to check"]
    assert [(r[0], r[1]) for r in check.rows] == [("ann@x.com", "."), ("ann@x.com", "Copy to home")]
    assert check.flags == ["high", "high"]
    assert [r[3] for r in s["Rules that send mail elsewhere"].rows] == ["ann.home@gmail.com", "boss@x.com"]
    assert len(s["All inbox rules"].rows) == 3
    assert report.highlights[0].value == 2


def test_mailbox_permissions_report(db, tenant):
    report, s = _sections(db, tenant, "mailbox-permissions")
    people = {r[0]: (r, f) for r, f in zip(s["By person"].rows, s["By person"].flags, strict=True)}
    ann, flag = people["ann@x.com"]
    # Send on behalf names "Ann" (her display name); it still resolves to her account.
    assert ann[1:] == ["Active", 1, "info@x.com", "info@x.com", "info@x.com"] and flag is None
    bob, flag = people["bob@x.com"]
    assert bob[1] == "Sign-in blocked" and flag == "medium"
    assert s["By mailbox"].rows == [["info@x.com", "Shared", "ann@x.com, bob@x.com", "ann@x.com", "ann@x.com"]]
    assert {r[3] for r in s["Every permission"].rows} == {"Full access", "Send as", "Send on behalf"}
    assert report.highlights[2].value == 1


def test_shared_mailboxes_report(db, tenant):
    report, s = _sections(db, tenant, "shared-mailboxes")
    section = s["Shared and resource mailboxes"]
    assert section.rows == [["info@x.com", "Shared", "Yes", "Never", "No", "ann@x.com, bob@x.com", "ann@x.com",
                             "586.6 KB"]]
    assert section.flags == ["medium"]
    assert section.intro.startswith("1 of 1 can be signed in to")
    assert report.highlights[0].value == "1 of 1"


def test_mailbox_size_report(db, tenant):
    report, s = _sections(db, tenant, "mailbox-size")
    section = s["Mailboxes by size"]
    assert [r[0] for r in section.rows] == ["ann@x.com", "info@x.com"]
    assert section.rows[0][2:5] == ["48.0 GB", "49.5 GB", "97%"] and section.flags == ["medium", None]
    assert section.rows[1][3] == "No limit"
    assert report.highlights[1].value == 1


def test_legacy_auth_report(db, tenant):
    report, s = _sections(db, tenant, "legacy-auth")
    log = s["Legacy sign-ins in the last 30 days"]
    assert [(r[0], r[1]) for r in log.rows] == [("info@x.com", "Authenticated SMTP"), ("ann@x.com", "IMAP4")]
    assert log.flags == ["medium", "low"]
    allows = s["What each mailbox allows"]
    # The test tenant leaves SMTP AUTH on for the organisation; Ann's mailbox turns it off.
    assert allows.rows == [["ann@x.com", "Off", "No", "No", "Yes"],
                           ["info@x.com", "On (organisation default)", "Yes", "Yes", "Yes"]]
    assert allows.flags == [None, None]
    assert [h.value for h in report.highlights] == [1, 1, "1 of 2"]


def test_legacy_auth_without_entra_p1(db, two_tenants):
    alpha, _ = two_tenants
    exchange_only = [{**SKUS[0], "servicePlans": [{"servicePlanName": "EXCHANGE_S_STANDARD",
                                                   "provisioningStatus": "Success"}]}]
    seed_tenant(db, alpha, skus=exchange_only, exchange=_exchange())
    report = build(db, store.get_tenant(db, alpha), "legacy-auth")
    assert not report.unavailable
    assert [s.title for s in report.sections] == ["What each mailbox allows"]
    assert any("Entra ID P1" in n for n in report.notes)


def test_mail_reports_unavailable_before_collection(db, two_tenants):
    tenant = store.get_tenant(db, two_tenants[1])
    for key in MAIL_REPORTS:
        assert build(db, tenant, key).unavailable, key


def test_mail_report_pages_tabs_and_xlsx(client, db, tenant):
    make_user(db, "admin")
    login(client, "admin")
    keys = list(REPORTS)
    assert keys[keys.index("mailboxes"):keys.index("mailboxes") + 6] == ["mailboxes", *MAIL_REPORTS]
    for key in MAIL_REPORTS:
        page = client.get(f"/tenants/{tenant['id']}/data/{key}")
        assert page.status_code == 200, key
        assert REPORTS[key][0] in page.get_data(as_text=True)
        resp = client.get(f"/tenants/{tenant['id']}/data/{key}.xlsx")
        assert resp.status_code == 200, key
        assert load_workbook(BytesIO(resp.data)).sheetnames
    assert "bob@x.com" in client.get(f"/tenants/{tenant['id']}/data/mailbox-permissions").get_data(as_text=True)


# -- Legacy sign-ins when Microsoft times out -------------------------------------------------

def _legacy_rows():
    rows = []
    now = datetime.now(UTC).replace(minute=0, second=0, microsecond=0)
    for i, client in enumerate(LEGACY_CLIENTS):
        for j in range(3):
            when = now - timedelta(days=2 + j * 9, hours=i)
            rows.append({"id": f"l{i}-{j}", "createdDateTime": when.strftime("%Y-%m-%dT%H:%M:%SZ"),
                         "userPrincipalName": ["ann@x.com", "info@x.com"][(i + j) % 2], "clientAppUsed": client,
                         "ipAddress": f"203.0.113.{i}", "appDisplayName": f"App {j}",
                         "location": {"countryOrRegion": ["GB", "RU", "FR"][j]},
                         "status": {"errorCode": 0 if j else 50126}})
    return rows


def _legacy_log(rows, fail):
    """A fake /v1.0/auditLogs/signIns that filters on time and clientAppUsed and fails when
    ``fail(clients, days)``, ``days`` being how long the query's time range is."""
    def handler(request):
        flt = parse_qs(urlsplit(request.url).query)["$filter"][0]
        times, _, ors = flt.partition(" and (")
        clients = [c.split("'")[1] for c in ors.split(" or ")]
        bounds = dict(c.split(" ")[1:] for c in times.split(" and "))
        start = datetime.fromisoformat(bounds["ge"].replace("Z", "+00:00"))
        end = datetime.fromisoformat(bounds["lt"].replace("Z", "+00:00")) if "lt" in bounds else datetime.now(UTC)
        if fail(clients, (end - start).total_seconds() / 86400):
            return 504, {}, json.dumps({"error": {"code": "GatewayTimeout", "message": "Request timed out"}})
        match = sorted((r for r in rows if r["clientAppUsed"] in clients and bounds["ge"] <= r["createdDateTime"]
                        and ("lt" not in bounds or r["createdDateTime"] < bounds["lt"])),
                       key=lambda r: r["createdDateTime"], reverse=True)
        return 200, {}, json.dumps({"value": match})
    responses.add_callback(responses.GET, f"{GRAPH_ROOT}/v1.0/auditLogs/signIns", callback=handler,
                           content_type="application/json")


def _legacy_run(db, tenant_id):
    store.enqueue_run(db, tenant_id, "legacy_signins")
    run = store.claim_next_run(db)
    status = execute_run(db, run, lambda t, c: GraphClient(lambda: "t", sleep=lambda s: None))
    return run["id"], status


@responses.activate
def test_legacy_signins_split_a_timed_out_filter_and_get_the_same_rows(db, two_tenants):
    alpha, bravo = two_tenants
    rows = _legacy_rows()
    _legacy_log(rows, fail=lambda clients, days: False)
    base_id, status = _legacy_run(db, alpha)
    assert status == "succeeded" and len(responses.calls) == 1
    expected = store.run_objects(db, base_id, "legacy_signin")

    # As on the test tenant: 30 days time out, a week answers. Split in time, the rows are the same.
    responses.reset()
    _legacy_log(rows, fail=lambda clients, days: days > 8)
    run_id, status = _legacy_run(db, bravo)
    assert status == "succeeded"
    assert store.run_objects(db, run_id, "legacy_signin") == expected
    assert len(responses.calls) == 2 + 2 * (2 + 1 + 1)  # 30 days, each 15 (newest first), their 7.5-day halves
    assert all(c.request.method == "GET" for c in responses.calls)

    # Even a day times out with all 13 clients: those days are split by client too.
    db.execute("DELETE FROM runs WHERE tenant_id = ?", (bravo,))
    responses.reset()
    _legacy_log(rows, fail=lambda clients, days: len(clients) > 4)
    run_id, status = _legacy_run(db, bravo)
    assert status == "succeeded"
    assert store.run_objects(db, run_id, "legacy_signin") == expected


@responses.activate
def test_legacy_signins_leave_out_a_client_that_keeps_failing(db, two_tenants):
    alpha, bravo = two_tenants
    rows = _legacy_rows()
    _legacy_log(rows, fail=lambda clients, days: len(clients) > 1 or clients == ["IMAP4"])
    run_id, status = _legacy_run(db, alpha)
    assert status == "partial"
    clients = {r["client"] for r in store.run_objects(db, run_id, "legacy_signin")}
    assert clients == set(LEGACY_CLIENTS) - {"IMAP4"}
    error = db.execute("SELECT error FROM runs WHERE id = ?", (run_id,)).fetchone()[0]
    assert "timed out for these clients" in error and "IMAP4" in error

    responses.reset()
    _legacy_log(rows, fail=lambda clients, days: True)
    run_id, status = _legacy_run(db, bravo)
    assert status == "failed" and store.run_objects(db, run_id, "legacy_signin") == []


@responses.activate
def test_legacy_signins_stop_splitting_when_the_log_is_too_slow(db, two_tenants, monkeypatch):
    from officesentry.collectors import sign_ins

    alpha, bravo = two_tenants
    rows = _legacy_rows()
    _legacy_log(rows, fail=lambda clients, days: True)
    run_id, status = _legacy_run(db, alpha)
    assert status == "failed"
    # Each query is tried twice (one retry), not seven times.
    assert len(responses.calls) % 2 == 0 and all(c.response.status_code == 504 for c in responses.calls)

    # The time limit starts at the first failure; once it's gone, the ranges left aren't asked for.
    now = [0.0]

    def clock():
        now[0] += sign_ins.FALLBACK_SECONDS / 4
        return now[0]
    monkeypatch.setattr(sign_ins, "clock", clock)
    responses.reset()
    _legacy_log(rows, fail=lambda clients, days: days > 8)
    run_id, status = _legacy_run(db, bravo)
    assert status == "partial"
    found = store.run_objects(db, run_id, "legacy_signin")
    # The newest 15 days were read (each client's sign-ins 2 and 11 days ago); the older 15 were given up on.
    assert sum(r["succeeded"] + r["failed"] for r in found) == 2 * len(LEGACY_CLIENTS)
    error = db.execute("SELECT error FROM runs WHERE id = ?", (run_id,)).fetchone()[0]
    assert "timed out for these clients" in error
