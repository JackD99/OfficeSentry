"""Over time, across clients: one measure for every client by month, the clients slipping first."""

from datetime import UTC, date, datetime
from io import BytesIO

import pytest
from openpyxl import load_workbook

from officesentry import store
from officesentry.reports import over_time
from tests.conftest import login, make_user
from tests.test_roles import role_user

NOW = datetime(2026, 10, 24, 9, 0, tzinfo=UTC)
SCORE = over_time.MEASURES["secure_score.pct"]
OPEN = over_time.MEASURES["open_severe"]


def metric(db, tenant_id, key, value, at):
    db.execute(
        "INSERT INTO metrics (tenant_id, key, value, captured_at) VALUES (?, ?, ?, ?)", (tenant_id, key, value, at)
    )


def scores(db, tenant_id, by_month: dict[str, float]):
    """A Secure Score at the start and (the one that counts) the end of each month named."""
    for month, value in by_month.items():
        metric(db, tenant_id, "secure_score.pct", 1.0, f"{month}-01T02:00:00Z")
        metric(db, tenant_id, "secure_score.pct", value, f"{month}-20T02:00:00Z")


@pytest.fixture
def three(db, connection_id):
    ids = [
        store.create_tenant(db, name, f"{name.lower()}.onmicrosoft.com", connection_id)
        for name in ("Alpha", "Bravo", "Charlie")
    ]
    alpha, bravo, charlie = ids
    scores(db, alpha, {"2026-08": 40, "2026-09": 45, "2026-10": 52})  # improving
    scores(db, bravo, {"2026-07": 70, "2026-09": 61, "2026-10": 58})  # slipping by 12
    scores(db, charlie, {"2026-10": 30})  # one month: nothing to compare
    return ids


def test_shades_say_how_good_a_value_is():
    assert [SCORE.level(v) for v in (95, 80, 79.6, 79.4, 65, 50, 35, 34, None)] == [0, 0, 0, 1, 1, 2, 3, 4, None]
    assert [OPEN.level(v) for v in (0, 1, 4, 5, 9, 10, 19, 20)] == [0, 1, 1, 2, 2, 3, 3, 4]
    assert (
        SCORE.legend() == "80% or more unshaded, 65% to 79% green, 50% to 64% amber, 35% to 49% orange, under 35% red"
    )
    assert OPEN.legend() == "0 unshaded, 1 to 4 green, 5 to 9 amber, 10 to 19 orange, over 19 red"
    assert (SCORE.worse(-3), SCORE.better_by(-3), OPEN.worse(2), OPEN.better_by(-2)) == (True, False, True, True)
    assert (SCORE.figure(54.4), SCORE.figure(54.4, unit=True), SCORE.figure(None)) == ("54", "54%", "–")
    assert (OPEN.change_text(5), OPEN.change_text(-12), SCORE.change_text(0.2), OPEN.change_text(None)) == (
        "+5",
        "−12",
        "0",
        "–",
    )


def test_each_month_end_and_the_slipping_first(db, three):
    alpha, bravo, charlie = three
    grid = over_time.build(db, [store.get_tenant(db, t) for t in three], SCORE, NOW)
    # Months before any client has a figure aren't shown.
    assert grid.months == [date(2026, 7, 1), date(2026, 8, 1), date(2026, 9, 1), date(2026, 10, 1)]
    rows = {r.tenant["name"]: r for r in grid.rows}
    assert rows["Alpha"].values == [None, 40, 45, 52]  # each month's last figure
    assert rows["Bravo"].values == [70, None, 61, 58]
    assert (rows["Alpha"].change, rows["Bravo"].change, rows["Charlie"].change) == (12, -12, None)
    assert [r.tenant["name"] for r in grid.rows] == ["Bravo", "Charlie", "Alpha"]  # slipped, then worst now
    assert [r.tenant["name"] for r in grid.slipped] == ["Bravo"]
    assert [r.tenant["name"] for r in grid.improved] == ["Alpha"]
    assert len(grid.counted) == 2
    # Percentages are averaged across clients; the change is over the clients with two months or more.
    assert grid.averages and grid.totals() == [70, 40, 53, (52 + 58 + 30) / 3]
    assert grid.now_and_change() == ((52 + 58 + 30) / 3, ((52 + 58) / 2) - ((40 + 70) / 2))

    table = grid.to_table()
    assert table.columns == ["Client", "Tenant ID", "Jul 2026 (%)", "Aug 2026 (%)", "Sep 2026 (%)", "Oct 2026 (%)"] + [
        "Change",
        "Note",
    ]
    assert table.rows[0][:1] + table.rows[0][2:] == ["Bravo", 70, "", 61, 58, -12, ""]


def test_a_new_install_still_shows_three_months(db, connection_id):
    t = store.create_tenant(db, "New", "new.onmicrosoft.com", connection_id)
    grid = over_time.build(db, [store.get_tenant(db, t)], SCORE, NOW)
    assert len(grid.months) == over_time.MIN_MONTHS
    assert grid.rows[0].note == "Nothing collected for this yet" and grid.rows[0].change is None


def test_open_critical_and_high_findings_as_they_stood(db, connection_id):
    t = store.create_tenant(db, "Alpha", "alpha.onmicrosoft.com", connection_id)
    db.execute(
        "INSERT INTO runs (tenant_id, collector, status, trigger, created_at, finished_at)"
        " VALUES (?, 'users', 'succeeded', 'manual', '2026-08-03T01:00:00Z', '2026-08-03T01:00:00Z')",
        (t,),
    )

    def finding(key, severity, first_seen, resolved=None, accepted=None, until=None):
        db.execute(
            "INSERT INTO findings (tenant_id, rule_key, object_id, severity, subject, detail, first_seen, last_seen,"
            " resolved_at, accepted_at, accepted_until) VALUES (?, ?, ?, ?, 's', 'd', ?, ?, ?, ?, ?)",
            (t, key, key, severity, first_seen, first_seen, resolved, accepted, until),
        )

    finding("a", "high", "2026-08-03T01:00:00Z")  # open throughout
    finding("b", "critical", "2026-08-03T01:00:00Z", resolved="2026-09-10T00:00:00Z")  # fixed in September
    finding("c", "high", "2026-09-20T00:00:00Z", accepted="2026-10-02T00:00:00Z")  # accepted in October
    finding("d", "high", "2026-08-03T01:00:00Z", accepted="2026-08-05T00:00:00Z", until="2026-09-15T00:00:00Z")
    finding("e", "medium", "2026-08-03T01:00:00Z")  # not critical or high
    grid = over_time.build(db, [store.get_tenant(db, t)], OPEN, NOW)
    assert grid.months == [date(2026, 8, 1), date(2026, 9, 1), date(2026, 10, 1)]
    # Before the first collection there's nothing; August: a, b (d accepted); September: a, c, d; October: a, d.
    assert grid.rows[0].values == [2, 3, 2]
    assert not grid.averages and grid.totals() == [2, 3, 2]


# -- The page and its downloads -------------------------------------------------------------------


def test_page_lists_every_client_slipping_first(client, db, three):
    make_user(db, "admin")
    login(client, "admin")
    html = client.get("/across/over-time?measure=secure_score.pct").get_data(as_text=True)
    assert "Secure Score, by month" in html
    assert html.index(">Bravo<") < html.index(">Charlie<") < html.index(">Alpha<")
    assert 'class="heat h' in html and "80% or more unshaded" in html
    assert 'aria-current="page">Secure Score<' in html
    # The default is open critical and high findings, and every summary's picker offers the page.
    assert "Open critical and high findings, by month" in client.get("/across/over-time").get_data(as_text=True)
    assert 'href="/across/over-time"' in client.get("/").get_data(as_text=True)
    assert client.get("/across?key=over-time").headers["Location"].endswith("/across/over-time")


def test_downloads_match_the_page(client, db, three):
    make_user(db, "admin")
    login(client, "admin")
    csv = client.get("/across/over-time.csv?measure=secure_score.pct")
    assert csv.status_code == 200 and "all-clients-over-time-secure-score" in csv.headers["Content-Disposition"]
    lines = csv.get_data(as_text=True).lstrip("﻿").splitlines()
    assert lines[0].startswith("Client,Tenant ID,") and lines[0].endswith("Oct 2026 (%),Change,Note")
    assert lines[1].startswith("Bravo,")
    book = load_workbook(BytesIO(client.get("/across/over-time.xlsx?measure=secure_score.pct").data))
    assert len(book.sheetnames) == 2
    assert db.execute("SELECT COUNT(*) FROM audit_log WHERE action = 'rollup_downloaded'").fetchone()[0] == 2


def test_only_the_clients_and_measures_someone_sees(client, db, three):
    alpha, bravo, _ = three
    make_user(db, "ana", role="analyst", tenants=[alpha])
    login(client, "ana")
    html = client.get("/across/over-time?measure=secure_score.pct").get_data(as_text=True)
    assert ">Alpha<" in html and ">Bravo<" not in html and ">Charlie<" not in html
    body = client.get("/across/over-time.csv?measure=secure_score.pct").get_data(as_text=True)
    assert "Alpha" in body and "Bravo" not in body


def test_a_role_that_does_not_see_findings(client, db, three):
    """First line doesn't see findings: the page opens on a measure it does see, and never offers findings."""
    role_user(db, "first", "first_line", all_tenants=True)
    login(client, "first")
    html = client.get("/across/over-time?measure=open_severe").get_data(as_text=True)
    assert "Open critical and high" not in html and "Users with MFA, by month" in html


def test_clients_cannot_open_it(client, db, three):
    make_user(db, "someone", role="client", tenants=[three[0]])
    login(client, "someone")
    assert client.get("/across/over-time").status_code == 403
    assert client.get("/across/over-time.csv").status_code == 403
