"""Site access: who owns, administers and can reach each SharePoint site (collector, check and report)."""

import io

from openpyxl import load_workbook

from officesentry import store
from officesentry.collectors.site_access import classify_login, describe_grant, guest_address
from officesentry.reports import data_reports
from officesentry.rules import evaluate
from officesentry.snapshot import TenantSnapshot
from tests import sharepoint_data as sd
from tests.conftest import login, make_user
from tests.seed import seed_tenant
from tests.test_sharepoint import _calls


def _snap(db, tenant_id):
    return TenantSnapshot(db, store.get_tenant(db, tenant_id))


def _report(db, tenant_id):
    return data_reports.build(db, store.get_tenant(db, tenant_id), "site-access")


def _section(report, title):
    return next(s for s in report.sections if s.title == title)


def test_classify_logins():
    assert classify_login("i:0#.f|membership|ann@x.com") == ("person", "ann@x.com")
    assert classify_login("i:0#.f|membership|pat_supplier.com#EXT#@x.onmicrosoft.com") == (
        "guest",
        "pat_supplier.com#ext#@x.onmicrosoft.com",
    )
    assert classify_login("c:0o.c|federateddirectoryclaimprovider|11111111-2222-3333-4444-555555555555") == (
        "group_members",
        "11111111-2222-3333-4444-555555555555",
    )
    assert classify_login("c:0o.c|federateddirectoryclaimprovider|11111111-2222-3333-4444-555555555555_o")[0] == (
        "group_owners"
    )
    assert classify_login("c:0t.c|tenant|11111111-2222-3333-4444-555555555555")[0] == "security_group"
    assert classify_login("c:0(.s|true")[0] == "everyone"
    assert classify_login(sd.EVERYONE_BUT_GUESTS)[0] == "everyone_but_guests"
    assert classify_login("i:0i.t|00000003-0000-0ff1-ce00-000000000000|app@sharepoint")[0] == "system"
    assert classify_login("SHAREPOINT\\system")[0] == "system"
    assert classify_login("Finance Owners") == ("sp_group", "Finance Owners")
    assert classify_login("") == ("unknown", "") and classify_login(None) == ("unknown", "")
    assert guest_address("pat_supplier.com#ext#@x.onmicrosoft.com", None) == "pat@supplier.com"
    assert guest_address("pat_supplier.com#ext#@x.onmicrosoft.com", "Pat@Supplier.com") == "pat@supplier.com"


def test_describe_grant_reads_every_identity_shape():
    assert describe_grant({"roles": ["owner"], "grantedToV2": {"siteGroup": {"displayName": "X Owners"}}}) == {
        "who": "X Owners",
        "how": "Site owners group",
        "can": "Full control",
        "outside": False,
    }
    assert describe_grant({"roles": ["read"], "grantedToV2": {"group": {"displayName": "Board", "id": "g4"}}})[
        "how"
    ] == ("Microsoft 365 or security group")
    everyone = describe_grant({"roles": ["read"], "grantedToV2": {"siteUser": {"loginName": "c:0(.s|true"}}})
    assert everyone["who"] == "Everyone, including guests" and everyone["outside"]
    link = describe_grant({"roles": ["read"], "link": {"scope": "anonymous", "type": "view"}})
    assert link["how"] == "Sharing link" and link["outside"]
    assert describe_grant({"roles": ["write"], "grantedToV2": {"user": {"email": "Bob@x.com"}}}) == {
        "who": "bob@x.com",
        "how": "Person",
        "can": "Edit",
        "outside": False,
    }
    app = {"grantedToV2": {"siteUser": {"loginName": "i:0i.t|00000003-0000-0ff1-ce00-000000000000|app@sharepoint"}}}
    assert describe_grant(app) is None and describe_grant({"roles": ["read"]}) is None
    assert describe_grant({"roles": "read", "grantedToV2": {"user": {"email": "a@x.com"}}})["can"] == "Unknown"


def test_collector_reads_owners_admins_guests_and_grants(db, two_tenants):
    alpha, _ = two_tenants
    seed_tenant(db, alpha)
    assert store.latest_run(db, alpha, "site_access")["status"] == "succeeded"
    snap = _snap(db, alpha)
    sites = {s["_id"]: s for s in snap.objects("site_access", "site")}
    assert set(sites) == {"site-fin", "site-old"}  # Ann's OneDrive is a personal site

    fin = sites["site-fin"]
    assert fin["group"] == {"id": "g1", "name": "Sales", "isTeam": True}  # from the group claim in the user list
    assert fin["ownerSource"] == "group" and fin["ownerCount"] == 2
    assert [o["name"] for o in fin["owners"]] == ["Ann", "Eve"]
    assert [a["upn"] for a in fin["admins"]] == ["ann@x.com"]
    assert fin["memberCount"] == 3 and fin["people"] == 2  # Ann and Eve recorded; the app and system rows skipped
    assert fin["guests"] == ["pat@supplier.com"] and fin["guestCount"] == 1  # once, from the list and the group
    assert fin["spGroups"] == ["Sales Owners", "Sales Members", "Sales Visitors"]
    # Everyone except external users is in Sales Visitors, which only SharePoint's own API shows.
    assert fin["everyone"] == "members" and fin["everyoneKnown"] and fin["sharePointRead"]
    assert [(a["who"], a["how"], a["can"], a["outside"]) for a in fin["access"]] == [
        ("Sales Owners", "Site owners group", "Full control", False),
        ("Sales Members", "Site members group", "Edit", False),
        ("Sales Visitors", "Site visitors group", "View", False),
        ("pat@supplier.com", "Guest", "View", True),
    ]
    assert fin["userList"] and fin["libraryAccess"] and fin["template"] == "Group" and fin["anyoneLinks"] == 2

    old = sites["site-old"]
    # Its administrator and the one person in its owners group are both Bob.
    assert old["group"] is None and old["ownerSource"] == "sharepoint"
    assert old["owners"] == [{"name": "Bob", "upn": "bob@x.com", "enabled": False}] and old["ownerCount"] == 0
    assert old["everyone"] == "members" and old["memberCount"] is None and old["guests"] == []
    assert old["access"][1] == {
        "who": "Everyone in the organisation",
        "how": "Everyone",
        "can": "View",
        "outside": False,
    }

    status = snap.objects("site_access", "access_status")[0]
    assert {k: v for k, v in status.items() if not k.startswith("_")} == {
        "available": True,
        "sites": 2,
        "unreadable": 0,
        "pages": 2,
        "complete": True,
    }
    assert snap.metric("sites.ownerless") == 1 and snap.metric("sites.single_owner") == 0
    assert snap.metric("sites.guests") == 1 and snap.metric("sites.open_to_everyone") == 2


def test_collector_reuses_the_site_list_and_asks_for_no_onedrives(db, two_tenants):
    alpha, _ = two_tenants
    seed_tenant(db, alpha)
    paths = _calls(db, alpha, "site_access")
    assert paths.count("/sites/getAllSites") == 0  # the other SharePoint collectors just listed the sites
    assert not any(p.startswith("/users/") for p in paths)
    assert paths.count("/sites/site-fin/lists/User%20Information%20List/items") == 1
    assert paths.count("/sites/site-old/drive/root/permissions") == 1
    assert not any("site-ann" in p for p in paths)


def test_collector_without_permission(db, two_tenants):
    alpha, _ = two_tenants
    seed_tenant(db, alpha, sharepoint_kw={"files": False})
    run = store.latest_run(db, alpha, "site_access")
    assert run["status"] == "partial" and "Sites.Read.All" in run["error"]
    assert "consent link" in _report(db, alpha).unavailable
    results = {r.key: r for r in evaluate(_snap(db, alpha))}
    assert results["site_owners"].status == "unavailable" and "Sites.Read.All" in results["site_owners"].summary


def test_site_owners_check(db, two_tenants):
    alpha, bravo = two_tenants
    seed_tenant(db, alpha)
    result = {r.key: r for r in evaluate(_snap(db, alpha))}["site_owners"]
    assert result.status == "fail" and result.summary == "1 site with no owner"
    assert [(f.severity, f.subject, f.detail) for f in result.findings] == [
        ("high", "Old Projects", "No enabled owner (site administrators and owners group)")
    ]
    assert {r.key: r for r in evaluate(_snap(db, bravo))}["site_owners"].status == "unavailable"


def test_report(db, two_tenants):
    alpha, _ = two_tenants
    seed_tenant(db, alpha)
    report = _report(db, alpha)
    tiles = {t.label: t for t in report.tiles}
    assert tiles["Sites with no owner"].value == 1 and tiles["Sites with no owner"].tone == "bad"
    assert tiles["One owner only"].value == 0 and tiles["Guests with access"].value == 1
    assert tiles["Open to everyone in the organisation"].value == 2 and tiles["Sites"].value == 2
    assert tiles["Open to everyone in the organisation"].foot == "Everyone or Everyone except external users"
    # Old Projects has no group, but SharePoint's own API has been read, so its groups' members are counted.
    assert not [n for n in report.notes if "Site permissions report" in n]

    needing = _section(report, "Sites needing an owner")
    assert [str(c) for c in needing.rows[0]] == [
        "Old Projects",
        "Communication site",
        "No owner",
        "Bob (disabled)",
        "Bob",
        needing.rows[0][5],
    ]
    assert needing.rows[0][0].safe_href == sd.OLD and needing.flags == ["high"]

    every = _section(report, "Who can access each site")
    assert [str(r[0]) for r in every.rows] == ["Finance", "Old Projects"]  # most guests first
    assert list(map(str, every.rows[0][2:8])) == ["Ann, Eve", "Ann", "3", "1", "Yes", "Yes"]
    assert list(map(str, every.rows[1][2:8])) == ["Bob (disabled)", "Bob", "1 person recorded", "0", "Yes", "No"]

    access = _section(report, "Access by site")
    fin = [r[1:] for r in access.rows if str(r[0]) == "Finance"]
    assert fin[0] == ["Owners of Sales: Ann, Eve", "Microsoft 365 group", "Full control", "No"]
    assert fin[1] == ["Members of Sales: 3 people, 1 guest", "Microsoft 365 group", "Edit", "Yes"]
    assert fin[-2:] == [
        ["pat@supplier.com", "Guest", "View", "Yes"],
        ["Everyone in the organisation", "Everyone", "Given access on the site", "No"],  # through Sales Visitors
    ]
    old = [r[1:] for r in access.rows if str(r[0]) == "Old Projects"]
    assert old == [
        ["Old Projects Owners", "Site owners group", "Full control", "No"],
        ["Everyone in the organisation", "Everyone", "View", "No"],
    ]
    guests = _section(report, "Guests and the sites they can reach")
    assert [str(c) for c in guests.rows[0]] == ["pat@supplier.com", "1", "Finance", "Member of the site's group"]
    admins = _section(report, "Site administrators")
    assert [(str(r[0]), r[1]) for r in admins.rows] == [("ann@x.com", 1), ("bob@x.com", 1)]
    assert "2 sites" in report.footnote


def test_sharepoint_groups_note_says_waiting_only_while_the_client_is_to_approve(db, two_tenants):
    alpha, _ = two_tenants
    seed_tenant(db, alpha, sharepoint_kw={"rest": "consent"}, our_app={"missing": ("SharePoint Sites.Read.All",)})
    (note,) = [n for n in _report(db, alpha).notes if "Site permissions report" in n]
    assert note.endswith("its members once the client approves the SharePoint permission.")


def test_without_sharepoints_own_api_it_says_what_it_cant_count(db, two_tenants):
    # Until the client approves the SharePoint permission, nobody can see who is in a site's SharePoint groups:
    # Old Projects' owners group may hold owners besides Bob, and Finance's Everyone claims (in every site's
    # user list) may or may not be in one of its groups. Neither is guessed.
    alpha, _ = two_tenants
    seed_tenant(db, alpha, sharepoint_kw={"rest": "consent"}, our_app={"missing": ("SharePoint Sites.Read.All",)})
    sites = {s["_id"]: s for s in _snap(db, alpha).objects("site_access", "site")}
    fin, old = sites["site-fin"], sites["site-old"]
    assert fin["everyone"] is None and fin["everyoneKnown"] is False and not fin["sharePointRead"]
    assert fin["ownerSource"] == "group" and fin["ownerCount"] == 2  # a group's owners come from Graph
    assert old["ownerSource"] is None and old["ownerCount"] is None
    assert [o["upn"] for o in old["owners"]] == ["bob@x.com"]
    assert old["everyone"] == "members" and old["everyoneKnown"]  # the default library grants it directly

    report = _report(db, alpha)
    tiles = {t.label: t for t in report.tiles}
    assert tiles["Sites with no owner"].value == 0 and tiles["Sites"].foot == "Owners unknown for 1 site"
    assert tiles["Open to everyone in the organisation"].value == 1
    assert tiles["Open to everyone in the organisation"].foot == "Not known for 1 site yet"
    every = {str(r[0]): list(map(str, r[2:8])) for r in _section(report, "Who can access each site").rows}
    assert every["Finance"][4] == "Unknown" and every["Old Projects"][4] == "Yes"
    assert every["Old Projects"][0] == "Bob (disabled)"

    result = {r.key: r for r in evaluate(_snap(db, alpha))}["site_owners"]
    assert result.status == "pass" and result.summary == "1 site checked, each with two or more owners"


def test_site_owners_check_waits_when_no_site_can_be_counted():
    from datetime import UTC, datetime

    from officesentry.checks.sharepoint import site_owners

    class Snap:
        def objects(self, collector, kind):
            if kind == "access_status":
                return [{"available": True}]
            return [{"_id": "s1", "name": "Projects", "ownerCount": None, "ownerSource": None}]

    status, summary, *_ = site_owners(Snap(), datetime.now(UTC))
    assert status == "unavailable" and "until the SharePoint permission is approved" in summary


def test_report_before_collection(db, two_tenants):
    _, bravo = two_tenants
    assert "yet" in _report(db, bravo).unavailable


def test_page_and_exports(client, db, two_tenants):
    alpha, bravo = two_tenants
    seed_tenant(db, alpha)
    make_user(db, "admin")
    login(client, "admin")
    page = client.get(f"/tenants/{alpha}/data/site-access")
    assert page.status_code == 200
    text = page.get_data(as_text=True)
    assert 'class="tiles"' in text and "Sites needing an owner" in text and "Old Projects" in text
    wb = load_workbook(io.BytesIO(client.get(f"/tenants/{alpha}/data/site-access.xlsx").data))
    assert "Access by site" in wb.sheetnames
    assert client.get(f"/tenants/{bravo}/data/site-access").status_code == 200  # the empty state renders
