from datetime import UTC, datetime, timedelta

from openpyxl import load_workbook

from officesentry import store
from officesentry.reports.export import pdf_available
from officesentry.reports.monthly_review import build
from tests.conftest import login, make_user
from tests.seed import seed_tenant


def test_build_report(db, two_tenants):
    alpha, _ = two_tenants
    seed_tenant(db, alpha)
    report = build(db, store.get_tenant(db, alpha))
    kpis = {k.key: k for k in report.kpis}
    assert kpis["secure_score.pct"].value == 42
    assert kpis["mfa.registered_pct"].value == 33.3
    assert report.failing[0].severity == "critical"  # most severe first
    assert report.counts["critical"] == 1
    assert "<svg" in report.secure_score_chart  # Microsoft's daily history gives 2 points
    assert all(not f.stale for f in report.freshness)


def test_kpi_delta_against_four_weeks_ago(db, two_tenants):
    alpha, _ = two_tenants
    old = (datetime.now(UTC) - timedelta(days=30)).strftime("%Y-%m-%dT%H:%M:%SZ")
    db.execute(
        "INSERT INTO metrics (run_id, tenant_id, key, value, captured_at) VALUES (NULL, ?, ?, ?, ?)",
        (alpha, "secure_score.pct", 35.0, old),
    )
    seed_tenant(db, alpha)
    k = {k.key: k for k in build(db, store.get_tenant(db, alpha)).kpis}["secure_score.pct"]
    assert (k.previous, k.delta, k.trend) == (35.0, 7.0, "better")


def test_report_pages(client, db, two_tenants):
    alpha, bravo = two_tenants
    seed_tenant(db, alpha)
    make_user(db, "admin")
    login(client, "admin")
    page = client.get(f"/tenants/{alpha}/report").get_data(as_text=True)
    assert "Monthly Security &amp; Licensing Review" in page and "Alpha Ltd" in page
    # A failing check leads with what's wrong, not its name (the goal), which would read as the opposite.
    assert "<h3>1 admin without MFA, and 1 whose registration isn&#39;t known</h3>" in page
    assert "Every admin has MFA registered" not in page
    xlsx = client.get(f"/tenants/{alpha}/report.xlsx")
    wb = load_workbook(__import__("io").BytesIO(xlsx.data))
    assert wb.sheetnames == ["Summary", "Checks", "Findings", "Licences", "Data freshness"]
    overview = client.get("/reports", follow_redirects=True).get_data(as_text=True)
    assert "Alpha Ltd" in overview and "Bravo Inc" in overview
    # A tenant with no data still renders (everything "not checked").
    assert client.get(f"/tenants/{bravo}/report").status_code == 200


def test_pdf_export(client, db, two_tenants):
    if not pdf_available():
        return
    alpha, _ = two_tenants
    seed_tenant(db, alpha)
    make_user(db, "admin")
    login(client, "admin")
    resp = client.get(f"/tenants/{alpha}/report.pdf")
    assert resp.status_code == 200 and resp.data[:4] == b"%PDF"


def test_client_sees_only_own_report(client, db, two_tenants):
    alpha, bravo = two_tenants
    make_user(db, "client", role="client", tenants=[alpha])
    login(client, "client")
    assert client.get(f"/tenants/{alpha}/report").status_code == 200
    for url in (f"/tenants/{bravo}/report", f"/tenants/{bravo}/report.xlsx", f"/tenants/{bravo}/report.pdf"):
        assert client.get(url).status_code == 404
    overview = client.get("/reports", follow_redirects=True).get_data(as_text=True)
    assert "Alpha Ltd" in overview and "Bravo Inc" not in overview
    assert client.get("/admin/prices").status_code == 403


def test_prices_page(client, db, two_tenants):
    alpha, _ = two_tenants
    seed_tenant(db, alpha)
    make_user(db, "admin")
    login(client, "admin")
    assert "Microsoft 365 Business Premium" in client.get("/admin/prices").get_data(as_text=True)
    client.post("/admin/prices", data={"currency": "USD", "price:SPB": "22", "price:OLD": "-5"})
    assert db.execute("SELECT monthly_price FROM sku_prices WHERE sku_part_number = 'SPB'").fetchone()[0] == 22
    assert db.execute("SELECT 1 FROM sku_prices WHERE sku_part_number = 'OLD'").fetchone() is None
    assert db.execute("SELECT value FROM settings WHERE key = 'currency'").fetchone()[0] == "USD"
    report = build(db, store.get_tenant(db, alpha))
    assert report.licences.monthly_waste == 22 * (3 + 1)  # 3 unassigned + Bob (disabled)


# -- The summary pages: found and fixed by month, what was fixed, the summary-only PDF --------------------


def _finding(db, tenant_id, key, first_seen, resolved_at=None, reopened_at=None, severity="high", subject="x"):
    db.execute(
        "INSERT INTO findings (tenant_id, rule_key, object_id, severity, subject, detail, first_seen, last_seen,"
        " resolved_at, reopened_at) VALUES (?, ?, ?, ?, ?, '', ?, ?, ?, ?)",
        (tenant_id, key, subject, severity, subject, first_seen, resolved_at or first_seen, resolved_at, reopened_at),
    )


def test_found_and_fixed_counts_each_month_from_the_first_collection(db, two_tenants):
    from officesentry.reports.monthly_review import ISO, _month_starts, found_and_fixed

    alpha, bravo = two_tenants
    now = datetime.now(UTC)
    starts = _month_starts(now, 3)
    last, this = (datetime(d.year, d.month, 1, 12, tzinfo=UTC).strftime(ISO) for d in starts[1:])
    _finding(db, alpha, "a", last, resolved_at=this)  # found last month, fixed this month
    _finding(db, alpha, "b", last)  # found last month, still open
    _finding(db, alpha, "c", "2020-01-01T00:00:00Z", reopened_at=this)  # old, came back this month
    _finding(db, bravo, "d", this)  # another client's
    months = found_and_fixed(db, alpha, now, "2000-01-02T00:00:00Z", months=3)
    assert [(m.month, m.found, m.fixed) for m in months] == [(starts[0], 0, 0), (starts[1], 2, 0), (starts[2], 1, 1)]
    # Before the first collection there's no data, and the first collection's own findings aren't "found".
    baseline = (datetime(starts[1].year, starts[1].month, 3, tzinfo=UTC)).strftime(ISO)
    months = found_and_fixed(db, alpha, now, baseline, months=3)
    assert [(m.found, m.fixed) for m in months] == [(None, None), (0, 0), (1, 1)]


def test_fixed_groups_by_check_most_fixed_first(db, two_tenants):
    from types import SimpleNamespace

    from officesentry.reports.monthly_review import fixed_groups

    rows = [
        {"rule_key": "mfa", "severity": "medium", "subject": "ann@x"},
        {"rule_key": "mfa", "severity": "high", "subject": "bob@x"},
        {"rule_key": "forwarding", "severity": "critical", "subject": "cat@x"},
        {"rule_key": "old_rule", "severity": "low", "subject": "dan@x"},
    ]
    results = [
        SimpleNamespace(key="mfa", title="Every user has MFA"),
        SimpleNamespace(key="forwarding", title="No forwards"),
    ]
    groups = fixed_groups(rows, results)
    assert [(g.title, g.severity, g.subjects) for g in groups] == [
        ("Every user has MFA", "high", ["ann@x", "bob@x"]),
        ("No forwards", "critical", ["cat@x"]),
        ("Old rule", "low", ["dan@x"]),
    ]


def test_months_chart_needs_two_months_of_data(db, two_tenants):
    from datetime import date

    from officesentry.reports.monthly_review import MonthCount

    alpha, _ = two_tenants
    seed_tenant(db, alpha)
    report = build(db, store.get_tenant(db, alpha))
    assert [m.found for m in report.months][:-1] == [None] * 11 and report.months[-1].found == 0
    assert report.months_chart == ""  # the first month alone isn't a trend
    report.months = [MonthCount(date(2026, m, 1), None if m < 5 else m, None if m < 5 else 1) for m in range(1, 13)]
    chart = report.months_chart
    assert "since May 2026" in chart and chart.count('class="nodata"') == 0 and chart.count("tone-2") == 8
    report.months[-1] = MonthCount(date(2026, 12, 1), 3, 0)
    report.months[:4] = [MonthCount(date(2026, m, 1), None, None) for m in range(1, 5)]
    report.months[4] = MonthCount(date(2026, 5, 1), 2, 2)
    chart = report.months_chart  # months before the first collection are left out; the first carries its year
    assert ">Jan" not in chart and ">May 2026<" in chart and ">Jun<" in chart


def test_summary_only_pdf_has_no_appendix(client, db, two_tenants, monkeypatch):
    from officesentry.reports import export

    alpha, _ = two_tenants
    seed_tenant(db, alpha)
    make_user(db, "admin")
    login(client, "admin")
    if pdf_available():
        resp = client.get(f"/tenants/{alpha}/report.pdf?appendix=0")
        assert resp.status_code == 200 and resp.data[:4] == b"%PDF"
        assert "security-review-summary-alpha-ltd.pdf" in resp.headers["Content-Disposition"]
        assert (
            "security-review-alpha-ltd.pdf" in client.get(f"/tenants/{alpha}/report.pdf").headers["Content-Disposition"]
        )
    monkeypatch.setattr(export, "pdf_available", lambda: False)
    full = client.get(f"/tenants/{alpha}/report.pdf").get_data(as_text=True)
    short = client.get(f"/tenants/{alpha}/report.pdf?appendix=0").get_data(as_text=True)
    for page in (full, short):
        assert "At a glance" in page and "This month" in page and "What we recommend next" in page
    assert 'id="appendix"' in full and "Fix first" in full and 'href="#check-' in full
    assert 'id="appendix"' not in short and "Fix first" not in short and 'href="#check-' not in short
    assert "Ask Harbour IT for the full report" in short or "for the full report" in short
    assert [r["detail"] for r in db.execute("SELECT detail FROM audit_log WHERE action = 'report_downloaded'")][
        -2:
    ] == [
        "pdf",
        "pdf summary",
    ]


def test_the_web_page_offers_both_pdfs_and_the_contents_split_summary_from_appendix(client, db, two_tenants):
    alpha, _ = two_tenants
    seed_tenant(db, alpha)
    make_user(db, "admin")
    login(client, "admin")
    page = client.get(f"/tenants/{alpha}/report").get_data(as_text=True)
    assert f'href="/tenants/{alpha}/report.pdf?appendix=0" title="The two summary pages without the appendix"' in page
    order = [page.index(s) for s in ("At a glance", "This month", 'id="appendix"', "<h2>Fix first")]
    assert order == sorted(order)


def test_summary_only_review_is_a_report_option(db, two_tenants, app, monkeypatch):
    from officesentry import delivery

    alpha, _ = two_tenants
    seed_tenant(db, alpha)
    tenant = store.get_tenant(db, alpha)
    option = delivery.catalog_map()["review-summary"]
    assert option.title == "Monthly review: summary only"
    monkeypatch.setattr(delivery.export, "pdf_available", lambda: True)  # Windows CI has no Pango
    assert delivery.attachment_plan(tenant, ["review", "review-summary"], True) == [
        ("review", "security-review-alpha-ltd.pdf"),
        ("review", "security-review-alpha-ltd.xlsx"),
        ("review-summary", "security-review-summary-alpha-ltd.pdf"),
    ]
    if not pdf_available():
        return
    files = delivery._attachments(app, db, tenant, ["review-summary"], True, datetime.now(UTC))
    assert [(f[0], f[1], f[2][:4]) for f in files] == [
        ("security-review-summary-alpha-ltd.pdf", "application/pdf", b"%PDF")
    ]
