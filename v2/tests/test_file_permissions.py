"""The File and folder permissions report and the outside-edit check."""

import io

from openpyxl import load_workbook

from officesentry import store
from officesentry.collectors.sharepoint import describe_permissions
from officesentry.reports import data_reports
from officesentry.rules import evaluate
from officesentry.snapshot import TenantSnapshot
from tests import sharepoint_data as sd
from tests.conftest import login, make_user
from tests.seed import seed_tenant


def _report(db, tenant_id):
    return data_reports.build(db, store.get_tenant(db, tenant_id), "file-permissions")


def _section(report, title):
    return next(s for s in report.sections if s.title == title)


def _column(section, name):
    i = section.columns.index(name)
    return [r[i] for r in section.rows]


def test_describe_permissions_lists_every_way_in():
    info = describe_permissions(
        [
            {
                "roles": ["read"],
                "link": {"scope": "anonymous", "type": "view"},
                "expirationDateTime": "2030-01-01T00:00:00Z",
            },
            {"roles": ["read"], "link": {"scope": "organization", "type": "view"}},
            {
                "roles": ["write"],
                "link": {"scope": "users", "type": "edit"},
                "grantedToIdentitiesV2": [{"user": {"email": "bob@x.com"}}, {"user": {"email": "p@other.com"}}],
            },
            {"roles": ["read"], "grantedToV2": {"group": {"displayName": "Sales"}}},
            {
                "roles": ["write"],
                "grantedToV2": {"siteUser": {"loginName": "i:0#.f|membership|g_other.com#ext#@x.onmicrosoft.com"}},
                "invitation": {"email": "g@other.com"},
                "expirationDateTime": "2026-01-01T00:00:00Z",
            },
            {"roles": ["owner"], "grantedToV2": {"siteGroup": {"displayName": "Owners"}}},
            {"roles": ["read"], "grantedToV2": {"user": {"email": "x@x.com"}}, "inheritedFrom": {"id": "root"}},
        ],
        {"x.com"},
    )
    assert [(g["who"], g["kind"], g["how"], g["can"], g["outside"]) for g in info["grants"]] == [
        ("Anyone with the link", "anyone", "Anyone link", "View", True),
        ("Everyone in the organisation", "organisation", "Organisation link", "View", False),
        ("bob@x.com", "person", "Link", "Edit", False),
        ("p@other.com", "guest", "Link", "Edit", True),
        ("Sales", "group", "Direct", "View", False),
        ("g@other.com", "guest", "Direct", "Edit", True),
    ]
    assert info["exposure"] == "anyone" and info["externalPeople"] == ["g@other.com", "p@other.com"]
    assert info["internalPeople"] == ["Sales", "bob@x.com"] and info["expires"] == "2026-01-01T00:00:00Z"


def test_expired_grants_are_listed_but_no_longer_count(db, two_tenants):
    from datetime import UTC, datetime

    now = datetime(2026, 6, 1, tzinfo=UTC)
    perms = [
        {
            "roles": ["read"],
            "link": {"scope": "users", "type": "view"},
            "grantedToIdentitiesV2": [{"user": {"email": "p@other.com"}}],
            "expirationDateTime": "2026-05-01T00:00:00Z",
        },
        {
            "roles": ["read"],
            "link": {"scope": "anonymous", "type": "view"},
            "expirationDateTime": "2026-05-01T00:00:00Z",
        },
    ]
    info = describe_permissions(perms, {"x.com"}, now)
    assert info["exposure"] is None and info["externalPeople"] == [] and info["anyoneLinks"] == 0
    assert [g["who"] for g in info["grants"]] == ["p@other.com", "Anyone with the link"]
    assert describe_permissions(perms, {"x.com"}, datetime(2026, 4, 1, tzinfo=UTC))["exposure"] == "anyone"


def test_report(db, two_tenants):
    alpha, _ = two_tenants
    seed_tenant(db, alpha)
    report = _report(db, alpha)
    assert report.unavailable is None and not report.notice
    tiles = {t.label: t for t in report.tiles}
    assert tiles["Files and folders with their own permissions"].value == 5
    assert tiles["People outside with access"].value == 1 and tiles["People outside with access"].tone == "warn"
    assert tiles["Newly shared in 30 days"].value == 1  # Budget.xlsx, shared 5 days ago
    assert tiles["Links past their expiry"].value == 1 and tiles["Expiring in 30 days"].value == 1

    fresh = _section(report, "Newly shared")
    assert [str(r[0]) for r in fresh.rows] == ["Budget.xlsx"]
    assert fresh.rows[0][2] == "partner@other.com (outside), old@other.com (outside)"
    assert fresh.rows[0][3] == "bob@x.com" and fresh.rows[0][4] == sd.days_ago(5)

    ending = _section(report, "Links expiring")
    assert _column(ending, "Who") == ["old@other.com", "partner@other.com"]  # expired first
    assert _column(ending, "Status") == ["Expired", "Expires soon"] and ending.flags == ["low", None]
    assert ending.rows[0][5] < 0 < ending.rows[1][5]  # days left

    who = _section(report, "Who can access what")
    assert who.group == "Location"
    rows = {(str(r[0]), r[2]): r for r in who.rows}
    assert rows[("Clients (folder)", "Anyone with the link")][3:6] == ["Anyone link", "Edit", "Yes"]
    assert rows[("Clients (folder)", "Anyone with the link")][6:9] == ["ann@x.com", "2026-08-20", "Never"]
    assert rows[("Internal.docx", "Sales")][3:6] == ["Direct", "View", "No"]
    assert rows[("Budget.xlsx", "old@other.com")][8] == sd.days_ago(3)
    flags = dict(zip([(str(r[0]), r[2]) for r in who.rows], who.flags, strict=True))
    assert flags[("Clients (folder)", "Anyone with the link")] == "high"
    assert flags[("Budget.xlsx", "partner@other.com")] == "low" and flags[("Budget.xlsx", "old@other.com")] is None
    assert str(rows[("Clients (folder)", "Anyone with the link")][0]) == "Clients (folder)"

    people = _section(report, "Access by person")
    assert _column(people, "Person or group")[0] == "partner@other.com"  # outside first; expired left out
    assert "old@other.com" not in _column(people, "Person or group")
    assert _column(people, "Type")[0] == "Guest or outside address"
    bob = next(r for r in people.rows if r[0] == "bob@x.com")
    assert bob[1] == "Member" and bob[2] == 1 and bob[3] == 1 and bob[4] == "Finance / Documents"
    assert "5 files and folders with permissions of their own" in report.footnote


def test_report_after_a_change_only_scan(db, two_tenants, monkeypatch):
    alpha, _ = two_tenants
    seed_tenant(db, alpha)
    monkeypatch.setattr(sd, "DELTA_CHANGES", {})
    from tests.test_sharepoint import _calls

    _calls(db, alpha)
    report = _report(db, alpha)
    assert "2 asked for changes only since the last scan, with 5 unchanged items kept" in report.footnote
    assert {t.label: t.value for t in report.tiles}["Files and folders with their own permissions"] == 5


def test_before_collection_and_old_scans(db, two_tenants):
    alpha, bravo = two_tenants
    assert "yet" in _report(db, bravo).unavailable
    seed_tenant(db, alpha, sharepoint_kw={"files": False}, our_app={"missing": ("Sites.Read.All",)})
    assert "consent link" in _report(db, alpha).unavailable
    # A scan stored by an earlier version has no grants: the report says so instead of showing nothing.
    seed_tenant(db, alpha)
    run = store.latest_run(db, alpha, "shared_files")
    from tests.stored import unpack

    unpack(db)
    db.execute(
        "UPDATE objects SET data = json_remove(data, '$.grants') WHERE run_id = ? AND kind = 'shared_item'",
        (run["id"],),
    )
    report = _report(db, alpha)
    assert "earlier version" in report.notice
    results = {r.key: r for r in evaluate(TenantSnapshot(db, store.get_tenant(db, alpha)))}
    assert results["outside_edit_access"].status == "unavailable"


def test_outside_edit_check(db, two_tenants, monkeypatch):
    alpha, _ = two_tenants
    monkeypatch.setitem(
        sd.PERMISSIONS,
        "it-policy",
        [
            {
                "id": "p3",
                "roles": ["write"],
                "link": {"scope": "users", "type": "edit"},
                "grantedToIdentitiesV2": [{"user": {"email": "editor@other.com"}}],
            },
            {
                "id": "p3b",
                "roles": ["write"],
                "grantedToV2": {"siteUser": {"loginName": "i:0#.f|membership|g_other.com#ext#@x.onmicrosoft.com"}},
                "invitation": {"email": "g@other.com"},
            },
        ],
    )
    seed_tenant(db, alpha)
    result = {r.key: r for r in evaluate(TenantSnapshot(db, store.get_tenant(db, alpha)))}["outside_edit_access"]
    assert result.status == "fail" and result.summary == "1 file or folder outside people can edit"
    assert result.findings[0].subject == "Policy.docx (Finance / Documents)"
    assert result.findings[0].detail == "Can be edited by editor@other.com, g@other.com"
    report = _report(db, alpha)
    people = {r[0]: r for r in _section(report, "Access by person").rows}
    assert people["g@other.com"][3] == 1 and _section(report, "Access by person").flags[0] == "medium"


def test_page_and_exports(client, db, two_tenants):
    alpha, bravo = two_tenants
    seed_tenant(db, alpha)
    make_user(db, "admin")
    login(client, "admin")
    page = client.get(f"/tenants/{alpha}/data/file-permissions").get_data(as_text=True)
    assert "Who can access what" in page and "Access by person" in page and 'class="tiles"' in page
    assert client.get(f"/tenants/{bravo}/data/file-permissions").status_code == 200  # empty state renders
    wb = load_workbook(io.BytesIO(client.get(f"/tenants/{alpha}/data/file-permissions.xlsx").data))
    assert wb.sheetnames[:3] == ["Summary", "Newly shared", "Links expiring"]
