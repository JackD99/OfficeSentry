"""Admin activity: directory audit collector, federation check and the Admin activity report."""

from datetime import UTC, datetime, timedelta
from io import BytesIO
from urllib.parse import parse_qs, urlparse

import pytest
import responses
from openpyxl import load_workbook

from officesentry import adminactivity, store
from officesentry.collectors import admin_activity as collector
from officesentry.devices import parse_ts
from officesentry.graph import GraphClient
from officesentry.reports import data_reports
from officesentry.rules import evaluate
from officesentry.snapshot import TenantSnapshot
from tests import adminactivitydata as ad
from tests.conftest import login, make_user
from tests.seed import seed_tenant
from tests.test_collectors import SKUS, V1, _metrics, _run

AUDITS = f"{V1}/auditLogs/directoryAudits"


@pytest.fixture
def factory():
    return lambda tenant, connection: GraphClient(lambda: "t", sleep=lambda s: None)


def _events(db, run_id):
    return {e["_id"]: e for e in store.run_objects(db, run_id, "admin_event")}


def _window(db, run_id):
    return store.run_objects(db, run_id, "audit_window")[0]


def _report(db, tenant_id):
    report = data_reports.build(db, store.get_tenant(db, tenant_id), "admin-activity", datetime.now(UTC))
    return report, {s.title: s for s in report.sections}


def _rule(db, tenant_id):
    return {r.key: r for r in evaluate(TenantSnapshot(db, store.get_tenant(db, tenant_id)))}[
        "domain_federation_changed"]


def _profile(db, tenant_id, factory, p1=True):
    skus = SKUS if p1 else [{**s, "servicePlans": [p for p in s["servicePlans"]
                                                    if not p["servicePlanName"].startswith("AAD_PREMIUM")]}
                            for s in SKUS]
    responses.get(f"{V1}/organization", json={"value": [{"id": "org1", "displayName": "Alpha"}]})
    responses.get(f"{V1}/subscribedSkus", json={"value": skus})
    assert _run(db, tenant_id, "tenant_profile", factory)[1] == "succeeded"


# -- Collector ---------------------------------------------------------------------------


@responses.activate
def test_keeps_only_allowlisted_activities(db, two_tenants, factory):
    alpha, _ = two_tenants
    responses.get(AUDITS, json={"value": ad.DIRECTORY_AUDITS})
    run_id, status = _run(db, alpha, "admin_activity", factory)
    assert status == "succeeded"
    query = parse_qs(urlparse(responses.calls[0].request.url).query)
    since = datetime.strptime(query["$filter"][0].removeprefix("activityDateTime ge "), "%Y-%m-%dT%H:%M:%SZ")
    assert abs(since.replace(tzinfo=UTC) - (datetime.now(UTC) - timedelta(days=30))) < timedelta(minutes=5)

    events = _events(db, run_id)
    assert set(events) == {e["id"] for e in ad.KEPT}
    assert not any(e["id"] in events for e in ad.NOISE)  # incl. a role name under the wrong category
    m = _metrics(db, run_id)
    assert m["admin_activity.events"] == len(ad.KEPT)
    assert m["admin_activity.role_changes"] == 3  # the failed grant doesn't count
    assert m["admin_activity.ca_changes"] == 3
    assert m["admin_activity.app_credential_changes"] == 3


@responses.activate
def test_events_are_normalised(db, two_tenants, factory):
    alpha, _ = two_tenants
    responses.get(AUDITS, json={"value": ad.DIRECTORY_AUDITS})
    run_id, _ = _run(db, alpha, "admin_activity", factory)
    e = _events(db, run_id)

    grant = e[ad.GLOBAL_ADMIN_GRANT["id"]]
    assert (grant["actor"], grant["actor_type"], grant["ip"]) == ("admin@x.com", "user", "203.0.113.40")
    assert (grant["kind"], grant["area"], grant["detail"], grant["targets"]) == (
        "role_grant", "Admin roles", "Global Administrator", ["bob@x.com"])
    pim = e[ad.PIM_ELIGIBLE["id"]]
    assert (pim["detail"], pim["targets"]) == ("Exchange Administrator", ["ann@x.com"])
    assert e[ad.ROLE_REMOVED["id"]]["detail"] == "Global Administrator"  # from the old value

    secret = e[ad.SECRET_ADDED["id"]]
    assert secret["activity"] == "Update application – Certificates and secrets management"  # trailing space
    assert (secret["detail"], secret["credential_added"]) == ("Added certificate", True)
    assert (e[ad.SECRET_REMOVED["id"]]["detail"], e[ad.SECRET_REMOVED["id"]]["credential_added"]) == (
        "Removed secret", False)
    sp = e[ad.SP_CREDENTIAL["id"]]
    assert (sp["actor"], sp["actor_type"], sp["ip"], sp["credential_added"]) == (
        "Contoso Automation", "app", "", True)

    ca = e[ad.CA_DISABLED["id"]]
    assert (ca["detail"], ca["ca_disabled"]) == ("On → Off", True)
    assert (e[ad.CA_ADDED["id"]]["detail"], e[ad.CA_ADDED["id"]]["ca_disabled"]) == ("On", False)
    assert e[ad.CONSENT["id"]]["detail"] == "Admin consent for everyone: Mail.Read User.Read"
    assert e[ad.CONSENT["id"]]["actor"] == "admin@x.com"  # UPN, not the "Azure ESTS Service" display name
    app_role = e[ad.APP_ROLE["id"]]
    assert (app_role["targets"], app_role["detail"]) == (["Mail Sync Pro"], "Mail.ReadWrite on Microsoft Graph")

    deleted = e[ad.USER_DELETED["id"]]
    assert len(deleted["targets"]) == adminactivity.MAX_TARGETS and deleted["more_targets"] == 2
    failed = e[ad.ROLE_GRANT_FAILED["id"]]
    assert (failed["result"], failed["reason"]) == ("failure", "Insufficient privileges")


@responses.activate
def test_follows_pages(db, two_tenants, factory):
    alpha, _ = two_tenants
    first, second = ad.DIRECTORY_AUDITS[:10], ad.DIRECTORY_AUDITS[10:]
    responses.get(AUDITS, json={"value": first, "@odata.nextLink": f"{AUDITS}?$skiptoken=page2"})
    responses.get(f"{AUDITS}?$skiptoken=page2", json={"value": second})
    run_id, status = _run(db, alpha, "admin_activity", factory)
    assert status == "succeeded"
    assert len(responses.calls) == 2
    assert set(_events(db, run_id)) == {e["id"] for e in ad.KEPT}
    window = _window(db, run_id)
    assert (window["read"], window["kept"], window["capped"]) == (len(ad.DIRECTORY_AUDITS), len(ad.KEPT), False)


@responses.activate
def test_cap_warns_instead_of_silently_truncating(db, two_tenants, factory, monkeypatch):
    alpha, _ = two_tenants
    monkeypatch.setattr(collector, "MAX_EVENTS", 5)
    monkeypatch.setattr(collector, "PAGE_SIZE", 3)
    pages = [ad.DIRECTORY_AUDITS[i:i + 3] for i in range(0, len(ad.DIRECTORY_AUDITS), 3)]
    for i, page in enumerate(pages):
        url = AUDITS if i == 0 else f"{AUDITS}?$skiptoken={i}"
        body = {"value": page}
        if i + 1 < len(pages):
            body["@odata.nextLink"] = f"{AUDITS}?$skiptoken={i + 1}"
        responses.get(url, json=body)
    run_id, status = _run(db, alpha, "admin_activity", factory)
    assert status == "partial"
    assert len(responses.calls) == 2  # stops paging at the cap
    run = db.execute("SELECT error FROM runs WHERE id = ?", (run_id,)).fetchone()
    assert "Only the newest 5 directory audit events were read" in run["error"]
    window = _window(db, run_id)
    assert (window["read"], window["capped"]) == (5, True)
    assert window["oldest"] == ad.DIRECTORY_AUDITS[4]["activityDateTime"]
    _, sections = _report(db, alpha)
    assert "Only the newest 5 audit events could be read" in sections["All admin changes"].intro


@responses.activate
def test_window_with_entra_p1(db, two_tenants, factory):
    alpha, _ = two_tenants
    _profile(db, alpha, factory, p1=True)
    responses.get(AUDITS, json={"value": ad.DIRECTORY_AUDITS})
    run_id, _ = _run(db, alpha, "admin_activity", factory)
    window = _window(db, run_id)
    assert (window["retention_days"], window["requested_days"]) == (30, 30)
    assert window["oldest"] == ad.DIRECTORY_AUDITS[-1]["activityDateTime"]
    assert window["newest"] == ad.DIRECTORY_AUDITS[0]["activityDateTime"]
    report, _ = _report(db, alpha)
    assert "in the last 30 days" in report.sections[0].intro and "P1" not in report.sections[0].intro


@responses.activate
def test_window_without_entra_p1(db, two_tenants, factory):
    alpha, _ = two_tenants
    _profile(db, alpha, factory, p1=False)
    recent = [e for e in ad.DIRECTORY_AUDITS if parse_ts(e["activityDateTime"])
              > datetime.now(UTC) - timedelta(days=7)]
    responses.get(AUDITS, json={"value": recent})
    run_id, _ = _run(db, alpha, "admin_activity", factory)
    assert _window(db, run_id)["retention_days"] == 7
    _, sections = _report(db, alpha)
    assert "the last 7 days (the tenant's audit log keeps only 7 days without Entra ID P1)" in \
        sections["Changes to review"].intro
    assert _rule(db, alpha).summary.endswith("in the last 7 days")


@responses.activate
def test_audit_log_forbidden_is_partial(db, two_tenants, factory):
    alpha, _ = two_tenants
    responses.get(AUDITS, status=403, json={"error": {"code": "Authorization_RequestDenied",
                                                      "message": "Insufficient privileges to complete the operation."}})
    run_id, status = _run(db, alpha, "admin_activity", factory)
    assert status == "partial"
    assert _window(db, run_id)["available"] is False
    report, _ = _report(db, alpha)
    assert "couldn't be read" in report.unavailable and not report.sections
    rule = _rule(db, alpha)
    assert rule.status == "unavailable" and "couldn't be read" in rule.summary


# -- Federation check --------------------------------------------------------------------


def test_federation_change_fails_check(db, two_tenants):
    alpha, _ = two_tenants
    seed_tenant(db, alpha)
    rule = _rule(db, alpha)
    assert rule.status == "fail" and rule.category == "Identity"
    assert [(f.severity, f.subject) for f in rule.findings] == [("high", "x.com")]
    assert "Set domain authentication by mallory@x.com" in rule.findings[0].detail


def test_failed_federation_change_passes(db, two_tenants, monkeypatch):
    alpha, _ = two_tenants
    failed = {**ad.FEDERATION, "result": "failure"}
    monkeypatch.setattr(ad, "DIRECTORY_AUDITS", [failed if e is ad.FEDERATION else e for e in ad.DIRECTORY_AUDITS])
    seed_tenant(db, alpha)
    rule = _rule(db, alpha)
    assert rule.status == "pass" and rule.summary == "No domain federation changes in the last 30 days"


def test_federation_check_unavailable_without_data(db, two_tenants):
    _, bravo = two_tenants
    assert _rule(db, bravo).status == "unavailable"


# -- Report ------------------------------------------------------------------------------


def test_admin_activity_report(db, two_tenants):
    alpha, _ = two_tenants
    seed_tenant(db, alpha)
    report, sections = _report(db, alpha)
    assert list(sections) == ["Changes to review", "All admin changes", "Who made changes"]

    review = sections["Changes to review"]
    got = [(row[1], flag) for row, flag in zip(review.rows, review.flags, strict=True)]
    assert got[0] == ("Set domain authentication", "high")
    assert sorted(got[1:]) == sorted([
        ("Add service principal credentials", "medium"),
        ("Update application – Certificates and secrets management", "medium"),
        ("Add member to role", "medium"),
        ("Add eligible member to role in PIM completed (permanent)", "medium"),
        ("Update conditional access policy", "medium"),
        ("Delete conditional access policy", "medium"),
    ])
    assert not any("mallory@x.com" == row[3] for row in review.rows)  # the failed grant is left out
    mediums = [row[0] for row, flag in zip(review.rows, review.flags, strict=True) if flag == "medium"]
    assert mediums == sorted(mediums, reverse=True)  # newest first within a severity

    every = sections["All admin changes"]
    assert len(every.rows) == len(ad.KEPT)
    assert [r[0] for r in every.rows] == sorted((r[0] for r in every.rows), reverse=True)
    failed = next(r for r in every.rows if r[3] == "mallory@x.com" and r[2] == "Add member to role")
    assert failed[6] == "Failed: Insufficient privileges"
    row = dict(zip(every.columns, next(r for r in every.rows if r[2] == "Delete user"), strict=True))
    assert row["Target"].endswith("and 2 more") and row["Area"] == "Accounts"

    who = sections["Who made changes"]
    assert who.rows[0][:3] == ["admin@x.com", "User", 14]
    assert ["Contoso Automation", "App"] in [r[:2] for r in who.rows]
    assert report.sources[0][0] == "Admin activity"


def test_admin_activity_report_without_data(db, two_tenants):
    _, bravo = two_tenants
    report, _ = _report(db, bravo)
    assert report.unavailable and not report.sections


def test_admin_activity_page_and_xlsx(client, db, two_tenants, monkeypatch):
    alpha, _ = two_tenants
    hostile = ad.event("x1", "ApplicationManagement", "Consent to application", ad.ago(1), ad.by_user("ann@x.com"),
                       [ad.target("ServicePrincipal", "<script>alert(1)</script>")])
    monkeypatch.setattr(ad, "DIRECTORY_AUDITS", [hostile] + ad.DIRECTORY_AUDITS)
    seed_tenant(db, alpha)
    make_user(db, "admin")
    login(client, "admin")
    page = client.get(f"/tenants/{alpha}/data/admin-activity")
    assert page.status_code == 200
    html = page.get_data(as_text=True)
    assert "Changes to review" in html and "Set domain authentication" in html
    assert "<script>alert(1)" not in html and "&lt;script&gt;alert(1)" in html  # names from Microsoft are escaped
    resp = client.get(f"/tenants/{alpha}/data/admin-activity.xlsx")
    assert resp.status_code == 200
    wb = load_workbook(BytesIO(resp.data))
    assert wb.sheetnames == ["Summary", "Changes to review", "All admin changes", "Who made changes"]
    assert wb["Changes to review"]["G2"].value == "high"
