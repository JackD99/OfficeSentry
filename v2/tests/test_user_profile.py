"""One page per account: what it brings together, links to it from reports, access and history."""

from datetime import UTC, datetime, timedelta
from io import BytesIO

from flask import render_template
from openpyxl import load_workbook

from officesentry import store
from officesentry.reports import data_reports, user_profile
from officesentry.reports.data_reports import UserLink
from officesentry.web.filters import day
from tests import graphdata as gd
from tests.conftest import login, make_user
from tests.seed import seed_tenant
from tests.stored import unpack
from tests.test_collectors import USERS

ANN = USERS[0]


def _seed(db, tenant_id, **kwargs):
    seed_tenant(db, tenant_id, users=[*USERS, gd.INFO_USER], **kwargs)


def _tenant(db, tenant_id):
    return store.get_tenant(db, tenant_id)


def test_profile_brings_the_account_together(db, two_tenants):
    alpha, _ = two_tenants
    _seed(db, alpha)
    profile = user_profile.build(db, _tenant(db, alpha), "u1")
    assert profile.name == "Ann" and profile.upn == "ann@x.com"
    parts = {p.key: p for p in profile.parts}
    assert list(parts) == ["licences", "usage", "mfa", "roles", "groups", "mailbox", "devices", "onedrive", "apps",
                           "signins", "signin_history", "activity", "timeline"]
    assert parts["licences"].tables[0].rows[0][0] == "Microsoft 365 Business Premium"
    groups = {row[0]: row[2] for row in parts["groups"].tables[0].rows}
    assert groups["Sales"] == "Owner"  # an owner of the team
    assert any(role == "Member" for role in groups.values())
    devices = [row[0] for row in parts["devices"].tables[0].rows]
    assert len(devices) == 2  # Ann's two managed devices, not Bob's
    # Ann has full access to the info@ mailbox: it shows under the mailboxes she can open.
    mailbox_tables = {t.title: t for t in parts["mailbox"].tables}
    can_open = mailbox_tables["Mailboxes this person can open or send from"]
    assert any(str(row[0]) == "info@x.com" for row in can_open.rows)
    # ...and links to that mailbox account's own page.
    link = next(row[0] for row in can_open.rows if str(row[0]) == "info@x.com")
    assert isinstance(link, UserLink) and link.page_href == f"/tenants/{alpha}/users/u7"
    assert profile.tiles[0].label == "Needs attention"
    assert profile.history_since


def test_disabled_account_with_licences_is_flagged(db, two_tenants):
    alpha, _ = two_tenants
    _seed(db, alpha)
    profile = user_profile.build(db, _tenant(db, alpha), "u2")
    texts = [a.text for a in profile.attention]
    assert any("Sign-in is blocked but the account still holds 1 paid licence" in t for t in texts)
    assert ("Sign-in blocked", "bad") in profile.badges


def test_unknown_account_is_none(db, two_tenants):
    alpha, _ = two_tenants
    _seed(db, alpha)
    assert user_profile.build(db, _tenant(db, alpha), "nobody") is None


def test_history_spots_licence_and_sign_in_changes(db, two_tenants):
    alpha, _ = two_tenants
    _seed(db, alpha)
    changed = [dict(USERS[0], accountEnabled=False, assignedLicenses=[]), *USERS[1:], gd.INFO_USER]
    seed_tenant(db, alpha, users=changed)
    profile = user_profile.build(db, _tenant(db, alpha), "u1")
    texts = [(e.area, e.text) for e in profile.history]
    assert ("Account", "Sign-in blocked") in texts
    assert ("Licences", "Microsoft 365 Business Premium removed") in texts


def test_report_cells_link_to_the_account_page(db, two_tenants):
    alpha, _ = two_tenants
    _seed(db, alpha)
    report = data_reports.build(db, _tenant(db, alpha), "mfa")
    cells = [v for s in report.sections for row in s.rows for v in row if isinstance(v, UserLink)]
    assert cells and all(c.page_href.startswith(f"/tenants/{alpha}/users/") for c in cells)
    # Downloads keep the plain name.
    assert all(c.safe_href is None for c in cells)
    assert cells[0] == str(cells[0])  # still the name for comparisons, sorting and exports


def test_pages_render_and_are_tenant_scoped(client, db, two_tenants):
    alpha, bravo = two_tenants
    _seed(db, alpha)
    make_user(db, "admin")
    make_user(db, "client", role="client", tenants=(bravo,))
    login(client, "admin")
    page = client.get(f"/tenants/{alpha}/users")
    assert page.status_code == 200
    html = page.get_data(as_text=True)
    assert f"/tenants/{alpha}/users/u1" in html and "Every account" in html
    page = client.get(f"/tenants/{alpha}/users/u1")
    assert page.status_code == 200
    html = page.get_data(as_text=True)
    for heading in ("Needs attention", "Licences", "MFA and sign-in protection", "Teams and groups", "Mailbox",
                    "Devices", "OneDrive", "History"):
        assert heading in html, heading
    assert client.get(f"/tenants/{alpha}/users/nobody").status_code == 404
    xlsx = client.get(f"/tenants/{alpha}/users/u1.xlsx")
    assert xlsx.status_code == 200
    sheets = load_workbook(BytesIO(xlsx.data)).sheetnames
    assert "Needs attention" in sheets and "Account" in sheets
    assert client.get(f"/tenants/{alpha}/users/u1.pdf").status_code == 200
    # Detail reports link account names to their page.
    assert f'href="/tenants/{alpha}/users/u1"' in client.get(f"/tenants/{alpha}/data/mfa").get_data(as_text=True)
    # Findings that name an account link to its page.
    assert f"/tenants/{alpha}/users/" in client.get(f"/tenants/{alpha}/findings").get_data(as_text=True)

    client.post("/logout")
    login(client, "client")
    assert client.get(f"/tenants/{alpha}/users/u1").status_code == 404
    assert client.get(f"/tenants/{alpha}/users").status_code == 404


def test_directory_counts(db, two_tenants):
    alpha, _ = two_tenants
    _seed(db, alpha)
    rows, tiles = user_profile.directory(db, _tenant(db, alpha), now=datetime(2026, 10, 2, tzinfo=UTC))
    assert {r.id for r in rows} == {"u1", "u2", "u3", "u7"}
    assert {t.label: t.value for t in tiles}["Accounts"] == 4
    bob = next(r for r in rows if r.id == "u2")
    assert bob.mfa == "–" and not bob.enabled


def test_usage_shows_on_the_account_page(db, two_tenants):
    alpha, _ = two_tenants
    _seed(db, alpha)
    profile = user_profile.build(db, _tenant(db, alpha), "u1")
    usage = {p.key: p for p in profile.parts}["usage"]
    assert not usage.unavailable
    labels = [label for label, _ in usage.facts]
    assert "Email" in labels and "Teams" in labels


# -- Leaver checklist and timeline -------------------------------------------------------------


def _reclaim(profile):
    part = {p.key: p for p in profile.parts}.get("reclaim")
    return {(row[0], str(row[1])): row[2] for row in part.tables[0].rows} if part and part.tables else None


def test_leaver_checklist_for_a_blocked_account(db, two_tenants):
    alpha, _ = two_tenants
    _seed(db, alpha)
    db.execute("INSERT INTO sku_prices (sku_part_number, monthly_price, updated_at) VALUES ('SPB', 18.1, '2026-10-01')")
    profile = user_profile.build(db, _tenant(db, alpha), "u2")  # Bob: blocked, Business Premium from a group
    assert profile.parts[0].key == "reclaim" and profile.parts[0].title == "Before you close this account"
    items = _reclaim(profile)
    licence = ("Licence", "Microsoft 365 Business Premium (£18.10 a month)")
    assert "take them out of that group" in items[licence]
    assert ("Group owner", "Only owner of Project Phoenix (team)") in items
    assert ("Admin role", "Directory Readers") in items
    assert any(area == "Device" and "LAPTOP-BOB" in found for area, found in items)
    assert any(area == "OneDrive" for area, _ in items)
    assert dict(profile.parts[0].facts)["Licence cost"] == "£18.10 a month, £217.20 a year"
    assert "£18.10 a month" in profile.parts[0].summary
    # An account whose sign-in is allowed only gets it when asked for.
    assert "reclaim" not in {p.key for p in user_profile.build(db, _tenant(db, alpha), "u1").parts}
    assert user_profile.build(db, _tenant(db, alpha), "u1", leaver=True).parts[0].key == "reclaim"


def _set_mailbox_size(db, tenant_id, uid, size):
    unpack(db)
    db.execute("UPDATE objects SET data = json_set(data, '$.totalBytes', ?) WHERE kind = 'mailbox_size'"
               " AND object_id = ? AND run_id IN (SELECT id FROM runs WHERE tenant_id = ?)", (size, uid, tenant_id))


def test_leaver_mailbox_advice(db, two_tenants):
    alpha, _ = two_tenants
    seed_tenant(db, alpha, users=[dict(USERS[0], accountEnabled=False), *USERS[1:], gd.INFO_USER])
    _set_mailbox_size(db, alpha, "u1", 5 * 1024 ** 3)
    items = _reclaim(user_profile.build(db, _tenant(db, alpha), "u1"))
    (found, mailbox), = [(found, todo) for (area, found), todo in items.items() if area == "Mailbox"]
    assert "Convert it to a shared mailbox" in mailbox and "under 50 GB need no licence" in mailbox
    assert "5.0 GB" in found and "forwards to ann@gmail.com" in found
    # Ann can open info@: that access should go.
    assert any(area == "Mailbox access" and "info@x.com" in found for area, found in items)

    _set_mailbox_size(db, alpha, "u1", 60 * 1024 ** 3)
    items = _reclaim(user_profile.build(db, _tenant(db, alpha), "u1"))
    mailbox = next(todo for (area, _), todo in items.items() if area == "Mailbox")
    assert "over the 50 GB" in mailbox and "Convert it" not in mailbox


def test_leaver_pdf_is_tenant_scoped(client, db, two_tenants):
    alpha, bravo = two_tenants
    _seed(db, alpha)
    make_user(db, "admin")
    make_user(db, "client", role="client", tenants=(bravo,))
    login(client, "admin")
    page = client.get(f"/tenants/{alpha}/users/u2").get_data(as_text=True)
    assert "Before you close this account" in page and f"/tenants/{alpha}/users/u2/leaver.pdf" in page
    page = client.get(f"/tenants/{alpha}/users/u1").get_data(as_text=True)
    assert "Before you close this account" not in page and "leaver=1" in page
    assert "Before you close this account" in client.get(f"/tenants/{alpha}/users/u1?leaver=1").get_data(as_text=True)
    sheet = client.get(f"/tenants/{alpha}/users/u2/leaver.pdf")
    assert sheet.status_code == 200
    if sheet.mimetype == "application/pdf":
        assert sheet.data.startswith(b"%PDF")
    else:  # no WeasyPrint: the print-ready page
        html = sheet.get_data(as_text=True)
        assert "Sign-off" in html and "Project Phoenix" in html
    assert client.get(f"/tenants/{alpha}/users/nobody/leaver.pdf").status_code == 404

    client.post("/logout")
    login(client, "client")
    assert client.get(f"/tenants/{alpha}/users/u2/leaver.pdf").status_code == 404


def test_leaver_sheet_lists_the_checklist(app, db, two_tenants):
    alpha, _ = two_tenants
    _seed(db, alpha)
    sheet = user_profile.leaver_sheet(db, _tenant(db, alpha), "u2")
    assert sheet.name == "Bob" and not sheet.enabled
    assert any(i.area == "Group owner" for i in sheet.checklist.items)
    with app.test_request_context():
        html = render_template("reports/leaver.html", sheet=sheet, pdf=False, print_view=True, print_name="x",
                               pdf_available=False, css_href="/static/report.css")
    assert "Leaver checklist" in html and "Sign-off" in html and "Only owner of Project Phoenix" in html


def test_timeline_dates_a_rule_by_the_snapshot_it_first_appeared_in(db, two_tenants):
    alpha, _ = two_tenants
    _seed(db, alpha, inbox_rules={})  # no inbox rules at first
    earlier = (datetime.now(UTC) - timedelta(days=10)).strftime("%Y-%m-%dT%H:%M:%SZ")
    db.execute("UPDATE runs SET finished_at = ? WHERE tenant_id = ?", (earlier, alpha))
    _seed(db, alpha)  # ...then Ann's rules appear
    later = db.execute("SELECT MAX(finished_at) FROM runs WHERE tenant_id = ? AND collector = 'inbox_rules'",
                       (alpha,)).fetchone()[0]
    profile = user_profile.build(db, _tenant(db, alpha), "u1")
    rows = {p.key: p for p in profile.parts}["timeline"].tables[0].rows
    rule = next(r for r in rows if r[2].startswith('Inbox rule "Copy to home" first seen'))
    assert rule[0] == day(later) and rule[1] == "Mailbox"
    # Present in the oldest retained snapshot: only "on or before".
    mfa = next(r for r in rows if r[1] == "MFA")
    assert mfa[0] == f"On or before {day(earlier)}"
    # Admin changes keep their own time.
    assert any(r[1] == "Admin change" for r in rows)
    assert len(rows) <= user_profile.TIMELINE_LIMIT
